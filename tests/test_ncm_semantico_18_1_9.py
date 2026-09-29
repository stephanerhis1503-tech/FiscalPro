from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.services.auditoria_cadastros_excel_service import AuditoriaCadastrosExcelService


def _indice(tmp_path: Path, linhas=None):
    p = tmp_path / "base_semantica.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Código", "Descrição", "NCM", "EAN"])
    for linha in (linhas or []):
        ws.append(linha)
    wb.save(p)
    rb = load_workbook(p, read_only=True, data_only=True)
    rws = rb["Produtos"]
    mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rws)
    indice = AuditoriaCadastrosExcelService._construir_indice_ncm_descricao(rws, mapa)
    return rb, indice


def _inferir(indice, descricao):
    return AuditoriaCadastrosExcelService._inferir_ncm_descricao(
        indice, codigo="", codigo_barras="", descricao=descricao
    )


def test_regra_semantica_prevalece_sobre_catalogo_ruidoso(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["A1", "MOTOR PARTIDA TITAN 160", "87141000", ""],
        ["A2", "ESCOVA ARRANQUE TITAN 160", "85119000", ""],
        ["A3", "FILTRO AR TITAN 160", "87141000", ""],
    ])
    try:
        assert _inferir(indice, "MOTOR PARTIDA TITAN 160").ncm == "85114000"
        assert _inferir(indice, "ESCOVA ARRANQUE TITAN 160").ncm == "85452000"
        assert _inferir(indice, "FILTRO AR TITAN 160").ncm == "84213100"
    finally:
        wb.close()


def test_regras_motor_e_eletrica_de_alta_confianca(tmp_path):
    wb, indice = _indice(tmp_path)
    try:
        casos = {
            "KIT PISTAO CG 160 0,50": "84099120",
            "VALVULA ADM CG 160": "84099114",
            "GUIA VALVULA TITAN 150": "84099117",
            "CDI BIZ 100": "85118030",
            "REGULADOR VOLTAGEM FAN 160": "85118020",
            "BOBINA IGNICAO CG 125": "85113020",
            "VELA IGNICAO NGK": "85111000",
            "BOMBA OLEO BIZ 125": "84133030",
            "BUZINA 12V UNIVERSAL": "85123000",
        }
        for descricao, ncm in casos.items():
            inf = _inferir(indice, descricao)
            assert inf is not None, descricao
            assert inf.ncm == ncm, descricao
            assert inf.confianca >= 90.0
    finally:
        wb.close()


def test_lampada_so_sugere_quando_tecnologia_ou_tensao_permite(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["L1", "LAMPADA FAROL UNIVERSAL", "87141000", ""],
        ["L2", "LAMPADA PISCA UNIVERSAL", "87141000", ""],
    ])
    try:
        led = _inferir(indice, "LAMPADA FAROL LED 12V")
        halogena = _inferir(indice, "LAMPADA FAROL H4 12V 35/35W")
        inc = _inferir(indice, "LAMPADA PAINEL 12V 2W")
        generica = _inferir(indice, "LAMPADA FAROL MODELO NOVO")
        assert led and led.ncm == "85395200"
        assert halogena and halogena.ncm == "85392110"
        assert inc and inc.ncm == "85392910"
        # A família LAMPADA é bloqueada para cópia aproximada. Sem tecnologia/tensão,
        # uma descrição apenas parecida no cadastro não vira verdade automática.
        assert generica is None
    finally:
        wb.close()


def test_corrente_com_medida_comercial_e_transmissao(tmp_path):
    wb, indice = _indice(tmp_path)
    try:
        for descricao in (
            "CORRENTE 520H 120L",
            "CORRENTE 428H COM RETENTOR",
            "CORRENTE TRANSMISSAO MOTO",
            "CORRENTE RELACAO TITAN 160",
        ):
            inf = _inferir(indice, descricao)
            assert inf is not None, descricao
            assert inf.ncm == "73151210", descricao
    finally:
        wb.close()


def test_familia_ambigua_continua_sem_palpite(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["C1", "CAPA BANCO TITAN PRETA", "39269090", ""],
        ["C2", "CAPA BANCO FAN PRETA", "87141000", ""],
        ["P1", "PAINEL COMPLETO TITAN", "87141000", ""],
    ])
    try:
        assert _inferir(indice, "CAPA BANCO BIZ PRETA") is None
        assert _inferir(indice, "PAINEL COMPLETO BIZ 125") is None
        assert _inferir(indice, "LAMPADA FAROL CG 160") is None
    finally:
        wb.close()


def test_consenso_muito_forte_de_familia_estavel(tmp_path):
    linhas = []
    for i in range(6):
        linhas.append([f"S{i}", f"SUPORTE PEDALEIRA MODELO {i} (MARCA X)", "87141000", ""])
    wb, indice = _indice(tmp_path, linhas)
    try:
        inf = _inferir(indice, "SUPORTE PEDALEIRA MODELO NOVO (MARCA X)")
        assert inf is not None
        assert inf.ncm == "87141000"
        assert "consenso" in inf.metodo.lower() or "mesma marca" in inf.metodo.lower()
    finally:
        wb.close()
