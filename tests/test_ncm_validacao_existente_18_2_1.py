from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.services.auditoria_cadastros_excel_service import AuditoriaCadastrosExcelService


def _indice(tmp_path: Path, linhas):
    p = tmp_path / "base_validacao_ncm.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Código", "Descrição", "NCM", "EAN"])
    for linha in linhas:
        ws.append(linha)
    wb.save(p)

    rb = load_workbook(p, read_only=True, data_only=True)
    rws = rb["Produtos"]
    mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rws)
    indice = AuditoriaCadastrosExcelService._construir_indice_ncm_descricao(rws, mapa)
    return rb, indice


def test_ncm_valido_mas_outlier_da_familia_sanfona_circuit(tmp_path):
    linhas = [
        ["S1", "SANFONA BENG. 11D PTO (CIRCUIT)", "87141000", ""],
        ["S2", "SANFONA BENG. 13D VERM (CIRCUIT)", "87141000", ""],
        ["S3", "SANFONA BENG. 18D PTO (CIRCUIT)", "87141000", ""],
        ["S4", "SANFONA BENG. 21D VERM (CIRCUIT)", "87141000", ""],
        ["S5", "SANFONA BENG. 24D PTO (CIRCUIT)", "87141000", ""],
        ["S6", "SANFONA BENG. 32D VERM (CIRCUIT)", "87141000", ""],
        ["S7", "SANFONA BENG. 35D PTO (CIRCUIT)", "87141000", ""],
        ["S8", "SANFONA BENG. 41D VERM (CIRCUIT)", "87141000", ""],
        ["S9", "SANFONA BENG. 45D PTO (CIRCUIT)", "87141000", ""],
        ["ERRO", "SANFONA BENG. 24D VERM (CIRCUIT)", "33030010", ""],
    ]
    wb, indice = _indice(tmp_path, linhas)
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_validacao_existente(
            indice,
            descricao="SANFONA BENG. 24D VERM (CIRCUIT)",
            ncm_atual="33030010",
        )
        assert inf is not None
        assert inf.ncm == "87141000"
        assert "consenso forte" in inf.metodo.lower()
        assert inf.confianca >= 90.0
    finally:
        wb.close()


def test_faixa_tanque_twister_nao_repete_ncm_de_hortalicas(tmp_path):
    wb, indice = _indice(tmp_path, [])
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_validacao_existente(
            indice,
            descricao="FAIXA TANQ TWISTER 16 L/D (HONDA)",
            ncm_atual="07109000",
        )
        assert inf is not None
        assert inf.ncm == "39199020"
        assert inf.confianca < 90.0
        assert "confirmar" in inf.referencia.lower()
    finally:
        wb.close()


def test_ncm_semantico_igual_ao_atual_nao_cria_falso_alerta(tmp_path):
    wb, indice = _indice(tmp_path, [])
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_validacao_existente(
            indice,
            descricao="MOTOR PARTIDA TITAN 160",
            ncm_atual="85114000",
        )
        assert inf is None
    finally:
        wb.close()
