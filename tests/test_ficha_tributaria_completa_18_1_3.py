from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.services.auditoria_cadastros_excel_service import (
    AuditoriaCadastrosExcelService,
    ExportadorFichaTributariaCompletaXLSX,
    ItemAuditoriaExcel,
    ResultadoAuditoriaExcel,
    _codigo_fiscal,
    _percentual_celula,
)


def test_cabecalhos_do_modelo_e_percentuais(tmp_path: Path):
    caminho = tmp_path / "modelo.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append([
        "COD. PRODUTO", "DESCRIÇÃO", "Código de Barras", "Cód. NCM", "EX", "CEST",
        "CST ICMS", "CFOP Saída", "ICMS Desonerado", "Red. BC ICMS(%)",
        "Aliq. ICMS(%)", "ICMS ST", "MVA ICMS-ST(%)", "Red. BC ICMS-ST(%)",
        "Aliq. ICMS-ST(%)", "Aliq. FEM(%)", "CST IPI", "Aliq. IPI(%)",
        "CST PIS/COFINS SAÍDA", "Aliq. PIS(%)", "Aliq. COFINS(%)", "Nat. de receita",
    ])
    ws.append([1, "Produto teste", "789", "40169300", "", "", 20, 5102, "NÃO", 0.60, 0.18, "NÃO", 0, 0, 0, 0, 51, 0, 1, 0.0065, 0.03, 213])
    for col in (10, 11, 13, 14, 15, 16, 18, 20, 21):
        ws.cell(2, col).number_format = "0.00%"
    wb.save(caminho)

    wb2 = load_workbook(caminho, read_only=True, data_only=True)
    ws2 = wb2.active
    mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(ws2)
    assert len(mapa) >= 20
    row = next(ws2.iter_rows(min_row=2, max_row=2))
    assert _percentual_celula(row, mapa, "aliq_icms") == 18.0
    assert _percentual_celula(row, mapa, "aliq_pis") == 0.65
    valores = tuple(c.value for c in row)
    assert _codigo_fiscal(AuditoriaCadastrosExcelService._valor(valores, mapa, "cst_icms"), 3) == "020"
    assert _codigo_fiscal(AuditoriaCadastrosExcelService._valor(valores, mapa, "cst_piscofins"), 2) == "01"
    wb2.close()


def test_exportador_ficha_tributaria(tmp_path: Path):
    resultado = ResultadoAuditoriaExcel(
        arquivo="entrada.xlsx",
        planilha="Produtos",
        contexto={"empresa": "Teste", "regime": "Lucro Real", "uf_origem": "MG", "uf_destino": "MG"},
        itens=[
            ItemAuditoriaExcel(
                linha_excel=2,
                codigo="ABC",
                descricao="Produto",
                ncm_atual="40169300",
                cest_atual="01.007.00",
                ncm_sugerido="40169300",
                cest_sugerido="01.007.00",
                status="OK",
                problema="Nenhuma divergência encontrada.",
                correcao_sugerida="Nenhuma divergência.",
                seguranca=100,
                codigo_barras="789",
                cst_icms_atual="060",
                cfop_atual="5405",
                aliquota_icms_atual=18.0,
                icms_st_atual="SIM",
                aliquota_icms_referencia=18.0,
                st_referencia="SIM",
                mva_referencia=71.78,
                cst_pis_atual="04",
                cst_pis_referencia="04",
                cst_cofins_atual="04",
                cst_cofins_referencia="04",
                pis_referencia=0.0,
                cofins_referencia=0.0,
                cst_ipi_atual="53",
                ipi_referencia=0.0,
            )
        ],
    )
    destino = tmp_path / "ficha.xlsx"
    ExportadorFichaTributariaCompletaXLSX.exportar(resultado, destino)
    wb = load_workbook(destino, read_only=True, data_only=True)
    ws = wb["Ficha Tributária"]
    assert ws.cell(1, 1).value == "COD. PRODUTO"
    assert ws.cell(1, 22).value == "Nat. de receita"
    assert ws.cell(1, 23).value == "Status Auditoria"
    assert ws.cell(2, 12).value == "SIM"
    assert ws.cell(2, 23).value == "OK"
    wb.close()
