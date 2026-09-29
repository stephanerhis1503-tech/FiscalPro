from pathlib import Path

from openpyxl import load_workbook

from src.financeiro.repositorio import ContasPagarRepositorio
from src.financeiro.servico import ContasPagarServico


def _servico(tmp_path: Path) -> ContasPagarServico:
    repo = ContasPagarRepositorio(tmp_path / "contas.db")
    return ContasPagarServico(repo)


def test_exportacao_usa_data_da_baixa(tmp_path):
    servico = _servico(tmp_path)
    empresa = servico.repositorio.listar_empresas()[0]
    servico.salvar_conta({
        "empresa": empresa,
        "fornecedor": "Fornecedor Teste",
        "valor": "100,00",
        "vencimento": "14/08/2026",
        "data_pagamento": "17/08/2026",
    })
    arquivo = servico.exportar_excel(tmp_path / "contas.xlsx")
    wb = load_workbook(arquivo, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    cabecalhos = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    assert "Vencimento" in cabecalhos
    assert "Data da Baixa" in cabecalhos
    assert "Data Pagamento" not in cabecalhos


def test_importador_mantem_compatibilidade_com_cabecalho_antigo_e_novo():
    fonte = Path('src/financeiro/servico.py').read_text(encoding='utf-8')
    assert 'indice("Data da Baixa", "Data da Baixa/Pagamento", "Data Pagamento", "Pagamento")' in fonte


def test_interface_deixa_semantica_da_baixa_explicita():
    fonte = Path('src/ui/painel_contas_pagar.py').read_text(encoding='utf-8')
    assert '"pagamento": "Data da baixa"' in fonte
    assert 'text="Dar baixa hoje"' in fonte
    assert 'text="Data da baixa:"' in fonte
