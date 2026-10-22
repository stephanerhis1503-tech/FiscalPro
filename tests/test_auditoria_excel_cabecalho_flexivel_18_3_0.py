from openpyxl import Workbook

from src.services.auditoria_cadastros_excel_service import AuditoriaCadastrosExcelService


def test_localiza_cabecalho_apos_linhas_de_titulo(tmp_path):
    p = tmp_path / "cadastro_com_titulo.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Cadastro"
    ws.append(["RELATORIO DE PRODUTOS"])
    ws.append(["EMPRESA TESTE"])
    ws.append([])
    ws.append(["Código", "Descrição", "NCM", "CEST"])
    ws.append(["A1", "PECA TESTE", "87141000", "0107600"])
    wb.save(p)

    rb = __import__("openpyxl").load_workbook(p, read_only=True, data_only=True)
    try:
        mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rb["Cadastro"])
        assert mapa["__linha_cabecalho__"] == 4
        assert mapa["codigo"] == 0
        assert mapa["descricao"] == 1
        assert mapa["ncm"] == 2
    finally:
        rb.close()


def test_codigo_e_opcional_na_estrutura_da_planilha(tmp_path):
    p = tmp_path / "cadastro_sem_codigo.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Descrição", "NCM", "CEST"])
    ws.append(["PECA TESTE", "87141000", "0107600"])
    wb.save(p)

    rb = __import__("openpyxl").load_workbook(p, read_only=True, data_only=True)
    try:
        ws2 = AuditoriaCadastrosExcelService._escolher_aba(rb)
        mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(ws2)
        assert "codigo" not in mapa
        assert mapa["descricao"] == 0
        assert mapa["ncm"] == 1
    finally:
        rb.close()
