from pathlib import Path

from openpyxl import load_workbook

from src.simulador.tabela_difal import aliquota_interna_destino
from src.sped.auditor_difal import AuditorDIFALSPED


def _sped_destino(cod_mun: str, c101: str | None = None, data: str = "01082026", aliq_inter: str = "12,00") -> list[str]:
    chave = "31260812345678000190550010000012341000012345"
    linhas = [
        "|0000|020|0|01082026|31082026|EMPRESA TESTE|12345678000190||MG|123|3106200|||B|1|\n",
        f"|0150|CLI|CLIENTE|1058||12345678909||{cod_mun}||RUA A|1|||\n",
        f"|C100|1|0|CLI|55|00|1|1234|{chave}|{data}|{data}|100,00|0|0,00|0,00|100,00|0|0,00|0,00|0,00|100,00|12,00|0,00|0,00|0,00|1,65|7,60|0,00|0,00|\n",
        f"|C190|000|6108|{aliq_inter}|100,00|100,00|12,00|0,00|0,00|0,00|0,00||\n",
    ]
    if c101:
        linhas.insert(3, c101 + "\n")
    return linhas


def test_tabela_modal_2026_e_historico():
    assert aliquota_interna_destino("AL", "01082026") == 20.5
    assert aliquota_interna_destino("AL", "31032026") == 19
    assert aliquota_interna_destino("MA", "01082026") == 23
    assert aliquota_interna_destino("RN", "01082026") == 20


def test_sem_xml_calcula_difal_devido_sp_para_base_c190():
    # MG -> SP = 12% interestadual; SP modal = 18%; base 100 => DIFAL 6.
    r = AuditorDIFALSPED().auditar(_sped_destino("3550308"), "EFD ICMS/IPI (Fiscal)")
    assert r.total_calculados_sem_xml == 1
    assert float(r.valor_difal_devido) == 6.0
    a = r.apontamentos[0]
    assert a.calculo_estimado is True
    assert a.origem_calculo == "SPED C190"
    assert float(a.difal_devido) == 6.0


def test_sem_xml_alagoas_usa_aliquota_20_5_e_fecoep_1():
    # MG -> AL = 7%; AL 20,5% em agosto/2026; base 100 => DIFAL 13,50 + FECOEP 1,00.
    r = AuditorDIFALSPED().auditar(_sped_destino("2704302", aliq_inter="7,00"), "EFD ICMS/IPI (Fiscal)")
    a = r.apontamentos[0]
    assert float(a.aliquota_interna) == 20.5
    assert float(a.aliquota_interestadual) == 7.0
    assert float(a.difal_devido) == 13.5
    assert a.fcp_calculado is True
    assert float(a.fcp_devido) == 1.0


def test_exportacao_excel_difal(tmp_path: Path):
    auditor = AuditorDIFALSPED()
    r = auditor.auditar(_sped_destino("3550308"), "EFD ICMS/IPI (Fiscal)", "EMPRESA", "08/2026")
    destino = auditor.salvar_excel(r, tmp_path / "difal.xlsx")
    assert destino.exists()
    wb = load_workbook(destino, read_only=True)
    assert set(["Resumo", "Detalhes DIFAL", "Sem XML"]).issubset(wb.sheetnames)
