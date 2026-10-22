from openpyxl import Workbook, load_workbook

from src.services.auditoria_cadastros_excel_service import (
    AuditoriaCadastrosExcelService,
    _cst_icms_cadastro,
)


def test_relatorio_mega_motos_cabecalho_mesclado_realinha_descricao(tmp_path):
    p = tmp_path / "mega_motos.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws["A1"] = "MEGA MOTOS TRILHA COMERCIO E SERVICOS LTDA"
    ws["A5"] = "Lista de produtos com tributos"
    ws["A6"] = "Situacao: Ativos. Tipo do calculo: Contribuinte. Unidade de federacao: MG."
    ws["A8"] = "Codigo"
    ws["C8"] = "Descricao"
    ws.merge_cells("C8:F8")
    ws["G8"] = "NCM"
    ws["H8"] = "CEST"
    ws["P8"] = "CST ICMS/IBS-CBS"
    ws.merge_cells("P8:Q8")
    ws["S8"] = "% ICMS"
    ws["A10"] = "SUB57"
    ws["D10"] = "SUPORT BAU TITAN150 09 PTO (TORK)"
    ws.merge_cells("D10:F10")
    ws["G10"] = "87141000"
    ws["H10"] = "01.076.00"
    ws["P10"] = "60/000"
    ws.merge_cells("P10:Q10")
    ws["S10"] = 0
    wb.save(p)

    rb = load_workbook(p, read_only=True, data_only=True)
    try:
        mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rb["Sheet"])
        assert mapa["__linha_cabecalho__"] == 8
        assert mapa["descricao"] == 3
        assert mapa["ncm"] == 6
        assert mapa["cest"] == 7
        assert mapa["cst_icms"] in {15, 16}
    finally:
        rb.close()


def test_cst_icms_separa_parte_ibs_cbs():
    assert _cst_icms_cadastro("60/000") == "060"
    assert _cst_icms_cadastro("00/200") == "000"
    assert _cst_icms_cadastro("060") == "060"
