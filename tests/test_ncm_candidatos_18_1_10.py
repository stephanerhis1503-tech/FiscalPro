from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.services.auditoria_cadastros_excel_service import (
    AuditoriaCadastrosExcelService,
    CandidatoNCMDescricao,
)


def _indice(tmp_path: Path, linhas):
    p = tmp_path / "base_candidatos.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Código", "Descrição", "NCM", "EAN"])
    for linha in linhas:
        ws.append(linha)
    wb.save(p)
    rb = load_workbook(p, read_only=True, data_only=True)
    rws = rb["Produtos"]
    mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rws)
    indice = AuditoriaCadastrosExcelService._construir_indice_ncm_descricao(rws, mapa)
    return rb, indice


def test_familia_ambigua_mostra_candidatos_sem_escolher(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["C1", "CAPA BANCO TITAN PRETA", "39269090", ""],
        ["C2", "CAPA BANCO FAN PRETA", "87141000", ""],
        ["C3", "CAPA BANCO BIZ PRETA", "39269090", ""],
    ])
    try:
        inferencia = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="CAPA BANCO POP PRETA"
        )
        assert inferencia is None
        candidatos = AuditoriaCadastrosExcelService._candidatos_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="CAPA BANCO POP PRETA", limite=3
        )
        ncms = {c.ncm for c in candidatos}
        assert "39269090" in ncms
        assert "87141000" in ncms
        assert len(candidatos) <= 3
    finally:
        wb.close()


def test_lampada_generica_recebe_opcoes_mas_nao_inferencia(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["L1", "LAMPADA FAROL TITAN 12V 35W", "85392910", ""],
        ["L2", "LAMPADA FAROL LED TITAN", "85395200", ""],
        ["L3", "LAMPADA FAROL H4 12V", "85392110", ""],
    ])
    try:
        # Sem tecnologia/tensão suficiente, continua sem classificação automática.
        inferencia = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="LAMPADA FAROL TITAN MODELO NOVO"
        )
        assert inferencia is None
        candidatos = AuditoriaCadastrosExcelService._candidatos_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="LAMPADA FAROL TITAN MODELO NOVO", limite=3
        )
        assert 1 <= len(candidatos) <= 3
        assert all(len(c.ncm) == 8 for c in candidatos)
    finally:
        wb.close()


def test_candidato_nao_e_usado_como_ncm_sugerido_automatico(tmp_path):
    wb, indice = _indice(tmp_path, [
        ["R1", "ROLAMENTO RODA DIANTEIRA MODELO A", "84821090", ""],
        ["R2", "ROLAMENTO AGULHA MODELO B", "84824000", ""],
    ])
    try:
        inferencia = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="ROLAMENTO RODA TRASEIRA MODELO C"
        )
        assert inferencia is None
        candidatos = AuditoriaCadastrosExcelService._candidatos_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="ROLAMENTO RODA TRASEIRA MODELO C", limite=3
        )
        assert candidatos
        texto = AuditoriaCadastrosExcelService._texto_candidatos_ncm(candidatos)
        assert any(c.ncm in texto for c in candidatos)
        assert "%" in texto
    finally:
        wb.close()


def test_lideranca_clara_vira_apenas_ncm_provavel_visual():
    candidatos = [
        CandidatoNCMDescricao(
            ncm="73151210",
            confianca=88.0,
            metodo="família",
            referencia="CORRENTE TRANSMISSAO",
            ocorrencias=14,
        ),
        CandidatoNCMDescricao(
            ncm="87141000",
            confianca=70.0,
            metodo="família",
            referencia="PEÇA MOTOCICLETA",
            ocorrencias=1,
        ),
        CandidatoNCMDescricao(
            ncm="73151290",
            confianca=62.0,
            metodo="família",
            referencia="OUTRA CORRENTE",
            ocorrencias=1,
        ),
    ]
    provavel = AuditoriaCadastrosExcelService._ncm_provavel_candidatos(candidatos)
    assert provavel is not None
    assert provavel.ncm == "73151210"


def test_empate_nao_elege_ncm_provavel():
    candidatos = [
        CandidatoNCMDescricao(
            ncm="73151210",
            confianca=77.0,
            metodo="família",
            referencia="CORRENTE",
            ocorrencias=3,
        ),
        CandidatoNCMDescricao(
            ncm="87141000",
            confianca=77.0,
            metodo="família",
            referencia="MOTOCICLETA",
            ocorrencias=2,
        ),
    ]
    assert AuditoriaCadastrosExcelService._ncm_provavel_candidatos(candidatos) is None


def test_uma_referencia_nao_basta_para_ncm_provavel():
    candidatos = [
        CandidatoNCMDescricao(
            ncm="87141000",
            confianca=91.0,
            metodo="semelhança",
            referencia="PEÇA",
            ocorrencias=1,
        ),
        CandidatoNCMDescricao(
            ncm="73151210",
            confianca=70.0,
            metodo="semelhança",
            referencia="CORRENTE",
            ocorrencias=1,
        ),
    ]
    assert AuditoriaCadastrosExcelService._ncm_provavel_candidatos(candidatos) is None
