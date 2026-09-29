from pathlib import Path

from openpyxl import load_workbook

from src.sped.importador_excel import (
    ImportadorExcelSPED,
    ItemConferenciaExcel,
    ResultadoValidacaoExcel,
)
from src.sped.motor_sped import MotorSPED


def test_classificacao_categorias_conferencia_excel():
    classificar = ImportadorExcelSPED.classificar_categoria_campo

    assert classificar("COD_ITEM") == "Identificação"
    assert classificar("CST_PIS") == "Tributário"
    assert classificar("VL_ITEM") == "Valor"
    assert classificar("Descrição do Produto/Serviços") == "Cadastral"
    assert classificar("CAMPO_PERSONALIZADO") == "Outro"


def test_filtros_e_relatorio_da_conferencia(tmp_path):
    itens = [
        ItemConferenciaExcel(
            tipo="Alteração",
            categoria="Tributário",
            sequencia=10,
            registro="C170",
            aba="C170",
            linha_excel=2,
            campo="CST_PIS",
            original="50",
            novo="01",
        ),
        ItemConferenciaExcel(
            tipo="Alteração",
            categoria="Cadastral",
            sequencia=20,
            registro="0200",
            aba="0200 - Cadastro de Itens",
            linha_excel=3,
            campo="Descrição do Produto/Serviços",
            original="ITEM A",
            novo="ITEM A TESTE",
        ),
    ]
    validacao = ResultadoValidacaoExcel(caminho=tmp_path / "ajustes.xlsx")
    validacao.conferencia = itens

    filtrados = ImportadorExcelSPED.filtrar_conferencia(
        itens,
        registro="0200",
        categoria="Cadastral",
        busca="TESTE",
    )
    assert filtrados == [itens[1]]

    destino = ImportadorExcelSPED.salvar_relatorio_conferencia(
        validacao,
        tmp_path / "conferencia.txt",
        itens=filtrados,
    )
    texto = destino.read_text(encoding="utf-8")
    assert "CONFERÊNCIA DE ALTERAÇÕES EXCEL → SPED" in texto
    assert "Registro 0200" in texto
    assert "Original: ITEM A" in texto
    assert "Novo: ITEM A TESTE" in texto
    assert "não aplica alterações automaticamente" in texto


def test_fluxo_real_0200_detecta_uma_alteracao_cadastral_sem_novo_erro(tmp_path):
    linhas = [
        "|0000|006|0|||01072026|31072026|EMPRESA|54304208000123|MG||||2|",
        "|0001|0|",
        "|0110|1||1|9|",
        "|0190|UN|Unidade|",
        "|0200|ITEM|ITEM TESTE|||UN|00|65061090||||0,00|",
        "|9999|6|",
    ]
    caminho_sped = tmp_path / "base.txt"
    caminho_sped.write_text("\r\n".join(linhas) + "\r\n", encoding="utf-8")

    motor = MotorSPED()
    motor.abrir(caminho_sped)
    caminho_excel = tmp_path / "base.xlsx"
    motor.exportar_excel(caminho_excel)

    workbook = load_workbook(caminho_excel)
    ws = workbook["0200 - Cadastro de Itens"]
    ws.cell(2, 3).value = "ITEM TESTE TESTE"
    workbook.save(caminho_excel)
    workbook.close()

    validacao = motor.validar_excel(caminho_excel)
    conferencia = motor.conferir_alteracoes_excel()

    assert validacao.valido is True
    assert validacao.total_alteracoes == 1
    assert validacao.pre_pva_novos_erros == 0
    assert len(conferencia) == 1
    assert conferencia[0].registro == "0200"
    assert conferencia[0].categoria == "Cadastral"
    assert conferencia[0].original == "ITEM TESTE"
    assert conferencia[0].novo == "ITEM TESTE TESTE"
