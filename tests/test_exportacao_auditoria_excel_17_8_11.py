from pathlib import Path

from openpyxl import load_workbook

from src.core.app_info import VERSAO_APP
from src.services.auditoria_cadastros_excel_service import (
    ExportadorAuditoriaCadastrosExcelXLSX,
    ItemAuditoriaExcel,
    ResultadoAuditoriaExcel,
    STATUS_CORRIGIR,
    STATUS_OK,
    STATUS_REVISAR,
)


def _resultado(qtd=1200):
    itens = []
    status = (STATUS_OK, STATUS_REVISAR, STATUS_CORRIGIR)
    for i in range(qtd):
        st = status[i % 3]
        itens.append(ItemAuditoriaExcel(
            linha_excel=i + 2,
            codigo=f"P{i:06d}",
            descricao=f"PRODUTO DE TESTE {i}",
            ncm_atual="87141000",
            cest_atual="01.076.00",
            ncm_sugerido="87141000",
            cest_sugerido="01.076.00",
            status=st,
            problema="" if st == STATUS_OK else "Revisar cadastro",
            correcao_sugerida="" if st == STATUS_OK else "Conferir tributação",
            seguranca=90.0,
            cst_pis_atual="01",
            pis_atual=1.65,
            cst_pis_referencia="01",
            pis_referencia=1.65,
            cst_cofins_atual="01",
            cofins_atual=7.6,
            cst_cofins_referencia="01",
            cofins_referencia=7.6,
        ))
    return ResultadoAuditoriaExcel(arquivo="cadastro.xlsx", planilha="Produtos", itens=itens)


def test_17811_exportacao_completa_preserva_linhas_e_recursos(tmp_path):
    destino = tmp_path / "completa.xlsx"
    progresso = []
    ExportadorAuditoriaCadastrosExcelXLSX.exportar(
        _resultado(), destino, progresso=lambda a, t, d: progresso.append((a, t, d))
    )
    assert destino.exists()
    wb = load_workbook(destino, read_only=False, data_only=False)
    ws = wb["Auditoria do cadastro"]
    assert ws.max_row == 1201
    assert ws.max_column == 18
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:R1201"
    assert len(ws.conditional_formatting) == 1
    assert wb["Resumo"]["B4"].value == "Auditoria completa"
    assert progresso and progresso[-1][0] == progresso[-1][1] == 1200
    wb.close()


def test_17811_exporta_somente_corrigir_e_revisar(tmp_path):
    destino = tmp_path / "pendencias.xlsx"
    ExportadorAuditoriaCadastrosExcelXLSX.exportar(
        _resultado(), destino, somente_pendencias=True
    )
    wb = load_workbook(destino, read_only=False, data_only=True)
    ws = wb["Auditoria do cadastro"]
    # 1200 itens / 3 status => 400 OK removidos + cabeçalho
    assert ws.max_row == 801
    assert all(row[0].value != STATUS_OK for row in ws.iter_rows(min_row=2, max_col=1))
    assert wb["Resumo"]["B4"].value == "Somente pendências (CORRIGIR + REVISAR)"
    wb.close()


def test_17811_interface_exporta_em_segundo_plano_e_tem_dois_modos():
    fonte = Path("src/ui/janela_auditoria_cadastros_excel.py").read_text(encoding="utf-8")
    assert "Exportar completa" in fonte
    assert "Exportar pendências" in fonte
    assert "export_progresso" in fonte
    assert "somente_pendencias=somente_pendencias" in fonte
    bloco = fonte.split("def _exportar", 1)[1]
    assert "threading.Thread(target=executar, daemon=True).start()" in bloco


def test_17811_exportador_nao_formata_celula_por_celula():
    fonte = Path("src/services/auditoria_cadastros_excel_service.py").read_text(encoding="utf-8")
    bloco = fonte.split("class ExportadorAuditoriaCadastrosExcelXLSX", 1)[1]
    assert "Workbook(write_only=True)" in bloco
    assert "FormulaRule" in bloco
    assert "for cel in ws[ws.max_row]" not in bloco


def test_versao_17_8_11_ou_superior():
    partes = tuple(int(p) for p in VERSAO_APP.split("."))
    assert partes >= (17, 8, 11)
