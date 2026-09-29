from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from src.services.analise_tributaria_lote_service import ResultadoAnaliseLote, ResultadoItemLote
from src.services.auditoria_cadastros_xml_service import (
    AuditoriaCadastrosXMLService,
    ExportadorAuditoriaCadastrosXMLXLSX,
    STATUS_CORRIGIR,
    STATUS_OK,
    STATUS_REVISAR,
)


def item(**kwargs):
    base = dict(
        fonte_tipo="XML NF-e",
        arquivo="nota.xml",
        documento="NF-e 1",
        chave="1" * 44,
        numero_item="1",
        codigo="ABC123",
        descricao="PNEU MOTO",
        ncm="40114000",
        ncm_oficial="Pneus novos de borracha, dos tipos utilizados em motocicletas",
        cfop="5405",
        uf_origem="MG",
        uf_destino="MG",
        data_operacao="2026-08-10",
        status="CONFIRMADO",
        confiabilidade=100.0,
        confirmado=True,
        exige_revisao=False,
        cst_pis_atual="04",
        aliquota_pis_atual=0.0,
        cst_pis_esperado="04",
        aliquota_pis_esperada=0.0,
        cst_cofins_atual="04",
        aliquota_cofins_atual=0.0,
        cst_cofins_esperado="04",
        aliquota_cofins_esperada=0.0,
        cst_icms_atual="60",
        aliquota_icms_atual=0.0,
        aliquota_icms_esperada=0.0,
        cest_atual="1600300",
        cest_esperado="1600300",
    )
    base.update(kwargs)
    return ResultadoItemLote(**base)


class TestAuditoriaCadastrosXML1770(unittest.TestCase):
    def analisar_fake(self, itens):
        fake = ResultadoAnaliseLote(itens=itens, arquivos_processados=2, fontes_solicitadas=1)
        with patch(
            "src.services.auditoria_cadastros_xml_service.AnaliseTributariaLoteService.analisar",
            return_value=fake,
        ):
            return AuditoriaCadastrosXMLService.analisar(["/tmp/xmls"], {"regime": "Lucro Real"})

    def test_agrupa_mesmo_codigo_em_um_cadastro(self):
        a = item(documento="NF-e 1", chave="1" * 44)
        b = item(documento="NF-e 2", chave="2" * 44)
        resultado = self.analisar_fake([a, b])
        self.assertEqual(1, len(resultado.produtos))
        produto = resultado.produtos[0]
        self.assertEqual(2, produto.ocorrencias)
        self.assertEqual(2, produto.documentos)
        self.assertEqual(STATUS_OK, produto.status)

    def test_divergencia_confirmada_marca_corrigir(self):
        a = item(
            status="DIVERGÊNCIA",
            confirmado=False,
            exige_revisao=True,
            cst_pis_atual="01",
            aliquota_pis_atual=1.65,
            divergencias=["CST PIS: documento 01; esperado 04."],
        )
        resultado = self.analisar_fake([a])
        produto = resultado.produtos[0]
        self.assertEqual(STATUS_CORRIGIR, produto.status)
        self.assertIn("PIS CST 04", produto.sugestao)

    def test_tratamentos_diferentes_mesmo_contexto_marca_revisar(self):
        a = item(documento="NF-e 1", chave="1" * 44)
        b = item(
            documento="NF-e 2", chave="2" * 44,
            cst_pis_atual="01", aliquota_pis_atual=1.65,
            status="CONFIRMADO",
        )
        resultado = self.analisar_fake([a, b])
        produto = resultado.produtos[0]
        self.assertTrue(produto.inconsistencia_interna)
        self.assertEqual(STATUS_REVISAR, produto.status)
        self.assertGreaterEqual(produto.tratamentos_distintos, 2)

    def test_diferenca_de_tratamento_em_ufs_distintas_nao_e_inconsistencia_sozinha(self):
        a = item(documento="NF-e 1", chave="1" * 44, uf_destino="MG", cfop="5405")
        b = item(
            documento="NF-e 2", chave="2" * 44, uf_destino="SP", cfop="6404",
            cst_icms_atual="60",
        )
        resultado = self.analisar_fake([a, b])
        produto = resultado.produtos[0]
        self.assertFalse(produto.inconsistencia_interna)
        self.assertEqual(STATUS_OK, produto.status)

    def test_ignora_item_de_entrada(self):
        saida = item(cfop="5405")
        entrada = item(codigo="ENTRADA", descricao="ITEM ENTRADA", cfop="1102")
        resultado = self.analisar_fake([saida, entrada])
        self.assertEqual(1, resultado.itens_saida)
        self.assertEqual(1, resultado.itens_ignorados_entrada)
        self.assertEqual(1, len(resultado.produtos))

    def test_exporta_excel_com_resumo_e_ocorrencias(self):
        resultado = self.analisar_fake([item()])
        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "auditoria.xlsx"
            arquivo = ExportadorAuditoriaCadastrosXMLXLSX.exportar(resultado, destino)
            self.assertTrue(Path(arquivo).exists())
            from openpyxl import load_workbook
            wb = load_workbook(arquivo, read_only=True)
            self.assertIn("Cadastros para revisar", wb.sheetnames)
            self.assertIn("Resumo", wb.sheetnames)
            self.assertIn("Ocorrências", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
