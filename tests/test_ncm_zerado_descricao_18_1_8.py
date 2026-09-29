from pathlib import Path
from openpyxl import Workbook

from src.services.auditoria_cadastros_excel_service import AuditoriaCadastrosExcelService


def _indice(tmp_path: Path):
    p = tmp_path / "base.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(["Código", "Descrição", "NCM", "EAN"])
    ws.append(["A1", "ACABAMENTO TRAVA PNEU VERDE (B BIKER)", "87141000", ""])
    ws.append(["A2", "ACABAMENTO GUIDAO PLASTICO PRETO (BIKER)", "87141000", ""])
    ws.append(["A3", "ACABAMENTO GUIDAO PLASTIC VERMELHO (BIKER)", "87141000", ""])
    ws.append(["B1", "PECA TESTE AZUL", "11111111", ""])
    ws.append(["B2", "PECA TESTE VERDE", "22222222", ""])
    wb.save(p)

    from openpyxl import load_workbook
    rb = load_workbook(p, read_only=True, data_only=True)
    rws = rb["Produtos"]
    mapa = AuditoriaCadastrosExcelService._mapear_cabecalhos(rws)
    indice = AuditoriaCadastrosExcelService._construir_indice_ncm_descricao(rws, mapa)
    return rb, indice


def test_inferencia_cor_da_mesma_familia(tmp_path):
    wb, indice = _indice(tmp_path)
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="",
            descricao="ACABAMENTO TRAVA PNEU VERMELHO (B BIKER)",
        )
        assert inf is not None
        assert inf.ncm == "87141000"
        assert inf.confianca >= 90
    finally:
        wb.close()


def test_inferencia_variante_mesma_marca(tmp_path):
    wb, indice = _indice(tmp_path)
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="",
            descricao="ACABAMENTO GUIDÃO PRETO (B BIKER)",
        )
        assert inf is not None
        assert inf.ncm == "87141000"
        assert "mesma marca" in inf.metodo
    finally:
        wb.close()


def test_nao_aprende_assinatura_ambigua(tmp_path):
    wb, indice = _indice(tmp_path)
    try:
        inf = AuditoriaCadastrosExcelService._inferir_ncm_descricao(
            indice, codigo="", codigo_barras="", descricao="PECA TESTE PRETA"
        )
        assert inf is None
    finally:
        wb.close()
