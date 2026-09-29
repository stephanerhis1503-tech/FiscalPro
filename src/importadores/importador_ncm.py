from pathlib import Path
from openpyxl import load_workbook

from src.repositorios.ncm_repository import NCMRepository


class ImportadorNCM:

    def __init__(self):

        self.inseridos = 0
        self.atualizados = 0
        self.sem_alteracao = 0
        self.erros = 0

    def importar(self, arquivo):

        arquivo = Path(arquivo)

        if not arquivo.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {arquivo}")

        print("=" * 60)
        print("IMPORTANDO TABELA OFICIAL NCM")
        print("=" * 60)

        wb = load_workbook(arquivo, read_only=True, data_only=True)

        ws = wb.active

        primeira = True

        for linha in ws.iter_rows(values_only=True):
            contador = 0

            contador += 1

            if contador % 500 == 0:
                print(f"{contador} linhas processadas...")
                
            if primeira:
                primeira = False
                continue

            try:

                ncm = str(linha[0]).replace(".", "").strip()

                descricao = linha[1] or ""

                data_inicio = linha[2]

                data_fim = linha[3]

                ato = linha[4] or ""

                numero = linha[5] or ""

                ano = linha[6] or ""

                resultado = NCMRepository.salvar(

                    ncm=ncm,

                    descricao=descricao,

                    data_inicio=data_inicio,

                    data_fim=data_fim,

                    ato_legal=ato,

                    numero_ato=numero,

                    ano_ato=ano

                )

                if resultado == "INSERIDO":
                    self.inseridos += 1

                elif resultado == "ATUALIZADO":
                    self.atualizados += 1

                else:
                    self.sem_alteracao += 1

            except Exception as erro:

                self.erros += 1

                print(f"Erro na linha: {erro}")

        print()

        print("=" * 60)

        print("IMPORTAÇÃO CONCLUÍDA")

        print("=" * 60)

        print(f"Inseridos.....: {self.inseridos}")

        print(f"Atualizados...: {self.atualizados}")

        print(f"Sem alteração.: {self.sem_alteracao}")

        print(f"Erros.........: {self.erros}")

        print("=" * 60)