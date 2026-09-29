"""Exportação do Simulador Tributário para Excel e PDF — Sprint 13.7."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, List, Sequence
from xml.sax.saxutils import escape

from src.simulador.motor_simulador import MotorSimuladorTributario, ResultadoSimulacao


def _brl(valor: Any) -> str:
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return str(valor or "")
    texto = f"{numero:,.2f}"
    return "R$ " + texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _percentual(valor: Any) -> str:
    try:
        return f"{float(valor):.4f}".rstrip("0").rstrip(".").replace(".", ",") + "%"
    except (TypeError, ValueError):
        return str(valor or "")


class ExportadorSimulacao:
    @classmethod
    def para_excel(cls, resultados: Sequence[ResultadoSimulacao], caminho: Any) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        if not resultados:
            raise ValueError("Não há simulações para exportar.")
        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Comparação"
        cabecalhos = [
            "Cenário", "NCM", "Descrição", "UF", "Regime", "Operação", "Finalidade",
            "Valor da operação", "Tributos atuais", "Carga atual", "IBS/CBS", "Carga reforma",
            "Diferença reforma - atual", "Valor estimado do documento", "Confiança",
        ]
        ws.append(cabecalhos)
        for item, resultado in zip(MotorSimuladorTributario.comparar(resultados), resultados):
            ws.append(
                [
                    item["cenario"], item["ncm"], resultado.descricao_ncm, item["uf"], item["regime"],
                    item["operacao"], item["finalidade"], float(item["valor_operacao"]),
                    float(item["tributos_atual"]), float(item["carga_atual"]) / 100,
                    float(item["reforma"]), float(item["carga_reforma"]) / 100,
                    float(item["diferenca"]), float(item["valor_documento"]),
                    float(item["confiabilidade"]) / 100,
                ]
            )
        for celula in ws[1]:
            celula.font = Font(bold=True)
            celula.fill = PatternFill("solid", fgColor="D9EAF7")
        for coluna in "HIJKL MNO".replace(" ", ""):
            if coluna in {"J", "L", "O"}:
                ws.column_dimensions[coluna].width = 16
            else:
                ws.column_dimensions[coluna].width = 21
        for coluna in ("H", "I", "K", "M", "N"):
            for celula in ws[coluna][1:]:
                celula.number_format = 'R$ #,##0.00'
        for coluna in ("J", "L", "O"):
            for celula in ws[coluna][1:]:
                celula.number_format = '0.00%'
        larguras = {"A": 24, "B": 12, "C": 44, "D": 14, "E": 20, "F": 14, "G": 20}
        for coluna, largura in larguras.items():
            ws.column_dimensions[coluna].width = largura
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for indice, resultado in enumerate(resultados, start=1):
            titulo = f"Cenario {indice}"[:31]
            detalhe = wb.create_sheet(titulo)
            entrada = resultado.entrada
            linhas_iniciais = [
                ("SIMULAÇÃO TRIBUTÁRIA", ""),
                ("Cenário", entrada.descricao_cenario),
                ("NCM", entrada.ncm),
                ("Descrição", resultado.descricao_ncm),
                ("Empresa", entrada.empresa),
                ("Regime", entrada.regime),
                ("Operação", entrada.operacao),
                ("Finalidade", entrada.finalidade),
                ("UF", f"{entrada.uf_origem} → {entrada.uf_destino}"),
                ("Destinatário", entrada.contribuinte),
                ("Data da operação", entrada.data_operacao),
                ("Regra atual", f"{resultado.regra_atual_id} | aderência {resultado.aderencia_regra_atual}"),
                ("Regra reforma", f"{resultado.regra_reforma_id} | aderência {resultado.aderencia_regra_reforma}"),
                ("Confiabilidade", f"{resultado.confiabilidade}% ({resultado.nivel_confiabilidade})"),
                ("Valor dos produtos", float(resultado.valor_produtos)),
                ("Valor da operação", float(resultado.valor_operacao)),
                ("Valor estimado do documento", float(resultado.valor_estimado_documento)),
                ("Total demonstrativo atual", float(resultado.total_tributos_atual)),
                ("Total IBS/CBS", float(resultado.total_reforma)),
                ("Diferença reforma - atual", float(resultado.diferenca_reforma_atual)),
            ]
            for linha in linhas_iniciais:
                detalhe.append(linha)
            detalhe.append([])
            detalhe.append(["Tributo", "Base", "Alíquota", "Valor", "Observação"])
            for linha in resultado.linhas:
                detalhe.append([
                    linha.tributo, float(linha.base), float(linha.aliquota) / 100,
                    float(linha.valor), linha.observacao,
                ])
            detalhe.append([])
            detalhe.append(["PREMISSAS"])
            for texto in resultado.premissas:
                detalhe.append(["•", texto])
            detalhe.append([])
            detalhe.append(["ALERTAS"])
            for texto in resultado.alertas:
                detalhe.append(["•", texto])

            for celula in detalhe[1]:
                celula.font = Font(bold=True, size=14)
            for linha in (22,):
                if linha <= detalhe.max_row:
                    for celula in detalhe[linha]:
                        celula.font = Font(bold=True)
                        celula.fill = PatternFill("solid", fgColor="D9EAF7")
            detalhe.column_dimensions["A"].width = 30
            detalhe.column_dimensions["B"].width = 70
            detalhe.column_dimensions["C"].width = 16
            detalhe.column_dimensions["D"].width = 18
            detalhe.column_dimensions["E"].width = 65
            for row in detalhe.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for linha in range(15, 21):
                detalhe.cell(linha, 2).number_format = 'R$ #,##0.00'
            for linha in range(23, 23 + len(resultado.linhas)):
                detalhe.cell(linha, 2).number_format = 'R$ #,##0.00'
                detalhe.cell(linha, 3).number_format = '0.0000%'
                detalhe.cell(linha, 4).number_format = 'R$ #,##0.00'
            detalhe.sheet_view.showGridLines = False

        wb.save(caminho)
        return caminho

    @classmethod
    def para_pdf(cls, resultados: Sequence[ResultadoSimulacao], caminho: Any) -> Path:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        if not resultados:
            raise ValueError("Não há simulações para exportar.")
        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(name="FiscalProTituloSim", parent=styles["Title"], fontSize=16, leading=19, textColor=colors.HexColor("#1F4E78")))
        styles.add(ParagraphStyle(name="FiscalProSecaoSim", parent=styles["Heading2"], fontSize=11, leading=14, textColor=colors.HexColor("#1F4E78"), spaceBefore=8, spaceAfter=5))
        styles.add(ParagraphStyle(name="FiscalProCorpoSim", parent=styles["BodyText"], fontSize=8, leading=10))
        doc = SimpleDocTemplate(
            str(caminho), pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=12 * mm, title="Simulação Tributária FiscalPro",
            author="FiscalPro - Stephane Rhis e ChatGPT",
        )
        elementos: List[Any] = [Paragraph("SIMULADOR TRIBUTÁRIO — FISCALPRO", styles["FiscalProTituloSim"]), Spacer(1, 4 * mm)]

        comparacao = [[
            "Cenário", "UF", "Valor operação", "Tributos atuais", "Carga atual",
            "IBS/CBS", "Carga reforma", "Diferença", "Confiança",
        ]]
        for item in MotorSimuladorTributario.comparar(resultados):
            comparacao.append([
                item["cenario"], item["uf"], _brl(item["valor_operacao"]), _brl(item["tributos_atual"]),
                _percentual(item["carga_atual"]), _brl(item["reforma"]), _percentual(item["carga_reforma"]),
                _brl(item["diferenca"]), _percentual(item["confiabilidade"]),
            ])
        tabela = Table(comparacao, repeatRows=1, colWidths=[42*mm, 24*mm, 30*mm, 30*mm, 23*mm, 27*mm, 24*mm, 28*mm, 22*mm])
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.extend([tabela, Spacer(1, 5 * mm)])

        for indice, resultado in enumerate(resultados, start=1):
            if indice > 1:
                elementos.append(PageBreak())
            entrada = resultado.entrada
            elementos.append(Paragraph(escape(entrada.descricao_cenario or f"Cenário {indice}"), styles["FiscalProSecaoSim"]))
            contexto = [
                ["NCM", entrada.ncm, "Descrição", escape(resultado.descricao_ncm)],
                ["Empresa", escape(entrada.empresa), "Regime", escape(entrada.regime)],
                ["Operação", escape(entrada.operacao), "Finalidade", escape(entrada.finalidade)],
                ["UF", f"{entrada.uf_origem} → {entrada.uf_destino}", "Destinatário", escape(entrada.contribuinte)],
                ["Regra atual", escape(f"{resultado.regra_atual_id} ({resultado.aderencia_regra_atual})"), "Regra reforma", escape(f"{resultado.regra_reforma_id} ({resultado.aderencia_regra_reforma})")],
            ]
            tctx = Table(contexto, colWidths=[27*mm, 70*mm, 28*mm, 128*mm])
            tctx.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF4F8")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EEF4F8")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            elementos.extend([tctx, Spacer(1, 4 * mm)])
            linhas = [["Tributo", "Base", "Alíquota", "Valor", "Observação"]]
            for linha in resultado.linhas:
                linhas.append([linha.tributo, _brl(linha.base), _percentual(linha.aliquota), _brl(linha.valor), escape(linha.observacao)])
            tlinhas = Table(linhas, repeatRows=1, colWidths=[38*mm, 30*mm, 25*mm, 30*mm, 130*mm])
            tlinhas.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            elementos.extend([tlinhas, Spacer(1, 4 * mm)])
            resumo = [[
                "Valor da operação", _brl(resultado.valor_operacao), "Tributos atuais", _brl(resultado.total_tributos_atual),
                "IBS/CBS", _brl(resultado.total_reforma), "Diferença", _brl(resultado.diferenca_reforma_atual),
            ]]
            tresumo = Table(resumo, colWidths=[31*mm, 30*mm, 31*mm, 30*mm, 25*mm, 30*mm, 25*mm, 30*mm])
            tresumo.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF4F8")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
            ]))
            elementos.append(tresumo)
            elementos.append(Spacer(1, 4 * mm))
            elementos.append(Paragraph("Premissas e alertas", styles["FiscalProSecaoSim"]))
            for texto in [*resultado.premissas, *resultado.alertas]:
                elementos.append(Paragraph("• " + escape(texto), styles["FiscalProCorpoSim"]))

        doc.build(elementos)
        return caminho


__all__ = ["ExportadorSimulacao"]
