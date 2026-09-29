"""Exportadores do Parecer Tributário Inteligente — Sprint 13.4."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, List, Tuple


class ExportadorParecer:
    @staticmethod
    def _texto(valor: Any) -> str:
        if valor is None:
            return ""
        if isinstance(valor, float):
            return f"{valor:.6f}".rstrip("0").rstrip(".").replace(".", ",")
        return str(valor)

    @classmethod
    def _linhas(cls, parecer) -> List[Tuple[str, str, str]]:
        linhas: List[Tuple[str, str, str]] = [
            ("titulo", "PARECER TRIBUTÁRIO INTELIGENTE", ""),
            ("campo", "NCM", parecer.ncm),
            ("campo", "Descrição", parecer.descricao),
            ("campo", "Confiabilidade", f"{parecer.confiabilidade:.0f}% ({parecer.nivel_confiabilidade})"),
            ("campo", "Gerado em", parecer.gerado_em),
            ("campo", "Versão do motor", parecer.versao_motor),
        ]
        for titulo, conteudo in parecer.secoes():
            linhas.append(("secao", titulo, ""))
            if isinstance(conteudo, dict):
                if conteudo:
                    linhas.extend(("campo", cls._texto(campo), cls._texto(valor)) for campo, valor in conteudo.items())
                else:
                    linhas.append(("item", "", "Não informado."))
            elif isinstance(conteudo, (list, tuple)):
                if conteudo:
                    linhas.extend(("item", "", cls._texto(item)) for item in conteudo)
                else:
                    linhas.append(("item", "", "Nenhum item."))
            else:
                linhas.append(("item", "", cls._texto(conteudo) or "Não informado."))
        linhas.extend(
            [
                ("secao", "IMPORTANTE", ""),
                (
                    "item",
                    "",
                    "Este parecer organiza e explica os dados cadastrados no FiscalPro. "
                    "A conclusão fiscal deve ser confirmada conforme a operação real e a fonte oficial vigente.",
                ),
            ]
        )
        return linhas

    @classmethod
    def para_excel(cls, parecer, caminho):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Parecer Tributário"

        linha_excel = 1
        for tipo, campo, valor in cls._linhas(parecer):
            if tipo in {"titulo", "secao"}:
                ws.merge_cells(start_row=linha_excel, start_column=1, end_row=linha_excel, end_column=2)
                celula = ws.cell(linha_excel, 1, campo)
                celula.font = Font(bold=True, size=14 if tipo == "titulo" else 11)
                celula.fill = PatternFill("solid", fgColor="D9EAF7" if tipo == "secao" else "B8D7EE")
                celula.alignment = Alignment(vertical="center", wrap_text=True)
                ws.row_dimensions[linha_excel].height = 24 if tipo == "titulo" else 20
            elif tipo == "campo":
                ws.cell(linha_excel, 1, campo).font = Font(bold=True)
                ws.cell(linha_excel, 2, valor)
            else:
                ws.cell(linha_excel, 1, "•")
                ws.cell(linha_excel, 2, valor)
            linha_excel += 1

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 100
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        wb.save(caminho)
        return caminho

    @classmethod
    def para_pdf(cls, parecer, caminho):
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        caminho = Path(caminho)
        caminho.parent.mkdir(parents=True, exist_ok=True)
        styles = getSampleStyleSheet()
        styles.add(
            ParagraphStyle(
                name="FiscalProTitulo",
                parent=styles["Title"],
                alignment=TA_CENTER,
                fontName="Helvetica-Bold",
                fontSize=16,
                leading=19,
                spaceAfter=12,
            )
        )
        styles.add(
            ParagraphStyle(
                name="FiscalProSecao",
                parent=styles["Heading2"],
                fontName="Helvetica-Bold",
                fontSize=11,
                leading=14,
                spaceBefore=8,
                spaceAfter=5,
                textColor=colors.HexColor("#1F4E78"),
            )
        )
        styles.add(
            ParagraphStyle(
                name="FiscalProCorpo",
                parent=styles["BodyText"],
                fontName="Helvetica",
                fontSize=8.5,
                leading=11,
                spaceAfter=1,
            )
        )
        styles.add(
            ParagraphStyle(
                name="FiscalProCampo",
                parent=styles["BodyText"],
                fontName="Helvetica-Bold",
                fontSize=8.3,
                leading=10.5,
            )
        )

        doc = SimpleDocTemplate(
            str(caminho),
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=14 * mm,
            bottomMargin=14 * mm,
            title=f"Parecer Tributário NCM {parecer.ncm}",
            author="FiscalPro - Stephane Rhis e ChatGPT",
        )
        elementos = []
        tabela: List[List[Any]] = []

        def descarregar_tabela() -> None:
            nonlocal tabela
            if not tabela:
                return
            objeto = Table(tabela, colWidths=[47 * mm, 128 * mm], repeatRows=0, hAlign="LEFT")
            objeto.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB7C4")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF4F8")),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            elementos.extend([objeto, Spacer(1, 4 * mm)])
            tabela = []

        for tipo, campo, valor in cls._linhas(parecer):
            if tipo == "titulo":
                descarregar_tabela()
                elementos.append(Paragraph(campo, styles["FiscalProTitulo"]))
            elif tipo == "secao":
                descarregar_tabela()
                elementos.append(Paragraph(campo, styles["FiscalProSecao"]))
            elif tipo == "campo":
                tabela.append(
                    [
                        Paragraph(cls._escapar(campo), styles["FiscalProCampo"]),
                        Paragraph(cls._escapar(valor), styles["FiscalProCorpo"]),
                    ]
                )
            else:
                descarregar_tabela()
                elementos.append(Paragraph("• " + cls._escapar(valor), styles["FiscalProCorpo"]))
        descarregar_tabela()
        doc.build(elementos)
        return caminho

    @staticmethod
    def _escapar(valor: Any) -> str:
        from xml.sax.saxutils import escape

        return escape(str(valor or "")).replace("\n", "<br/>")
