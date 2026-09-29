from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

from .layout_excel import LayoutRegistro, RESETAR_CONTEXTOS_AO_LER, obter_layout
from .campos_oficiais_excel import obter_cabecalhos_oficiais
from .motor_sped import ResultadoSPED

ProgressoCallback = Callable[[int, str], None]


@dataclass(frozen=True, slots=True)
class ResultadoExportacaoExcel:
    caminho: Path
    total_abas: int
    total_registros: int
    total_linhas: int


@dataclass(slots=True)
class ReferenciaTecnica:
    sequencia: int
    tipo: str
    id_tecnico: str
    registro: str
    aba: str
    linha_excel: int
    qtd_contexto: int
    qtd_campos: int
    linha_original: str


@dataclass(slots=True)
class LinhaExportacao:
    valores: list[str]
    referencia: ReferenciaTecnica


class ExportadorSPEDExcel:
    """Exporta o SPED em abas legíveis e reversíveis.

    As abas visíveis mantêm o formato organizado aprovado na Sprint 12.3. A
    Sprint 12.5 adiciona somente identificadores ocultos e uma aba técnica
    ``_FISCALPRO_ORDEM``. Essa estrutura permite editar campos existentes no
    Excel e reconstruir o TXT na mesma ordem do arquivo original.
    """

    LIMITE_LINHAS_EXCEL = 1_048_576
    LIMITE_DADOS_POR_ABA = LIMITE_LINHAS_EXCEL - 1
    LIMITE_COLUNAS_EXCEL = 16_384

    COR_CABECALHO = "4472C4"
    COR_TEXTO_CLARO = "FFFFFF"
    COR_BORDA = "B4C6E7"
    COR_LINHA_ALTERNADA = "EDF3F8"

    ABA_TECNICA = "_FISCALPRO_ORDEM"
    CABECALHO_ID = "__FP_ID__"
    ASSINATURA_FORMATO = "FISCALPRO_SPED_EXCEL"
    VERSAO_FORMATO = "1"

    def exportar(
        self,
        resultado: ResultadoSPED,
        caminho_saida: str | Path,
        progresso: ProgressoCallback | None = None,
    ) -> ResultadoExportacaoExcel:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as erro:
            raise RuntimeError(
                "A exportação para Excel precisa do pacote openpyxl. "
                "Instale com: python -m pip install openpyxl"
            ) from erro

        destino = Path(caminho_saida)
        if destino.suffix.lower() != ".xlsx":
            destino = destino.with_suffix(".xlsx")
        destino.parent.mkdir(parents=True, exist_ok=True)

        self._progresso(progresso, 3, "Organizando os registros do SPED...")
        ordem, dados, layouts, cabecalhos, referencias = self._preparar_dados(resultado)

        workbook = Workbook()
        workbook.properties.title = f"FiscalPro - {resultado.caminho.name}"
        workbook.properties.subject = "SPED organizado por registros"
        workbook.properties.creator = "FiscalPro"
        workbook.properties.description = (
            "Exportação reversível gerada pelo Motor SPED Inteligente. "
            "As abas visíveis são organizadas por registro e a estrutura "
            "técnica oculta permite converter novamente para TXT."
        )
        workbook.properties.created = datetime.now()

        estilos = {
            "Font": Font,
            "PatternFill": PatternFill,
            "Alignment": Alignment,
            "Border": Border,
            "Side": Side,
        }

        aba_padrao = workbook.active
        nomes_usados: set[str] = {self.ABA_TECNICA.casefold()}
        total_abas = 0
        total_codigos = max(len(ordem), 1)

        for posicao, codigo in enumerate(ordem, start=1):
            total_abas += self._montar_abas_registro(
                workbook=workbook,
                codigo=codigo,
                linhas=dados[codigo],
                layout=layouts[codigo],
                cabecalho=cabecalhos[codigo],
                estilos=estilos,
                nomes_usados=nomes_usados,
            )
            percentual = 8 + int((posicao / total_codigos) * 82)
            self._progresso(
                progresso,
                min(percentual, 90),
                f"Criando aba {codigo} ({posicao}/{len(ordem)})...",
            )

        if total_abas:
            workbook.remove(aba_padrao)
        else:
            aba_padrao.title = "Sem dados"
            aba_padrao["A1"] = "Nenhum registro SPED válido foi encontrado."
            total_abas = 1

        self._progresso(progresso, 92, "Criando estrutura técnica Excel → TXT...")
        self._criar_aba_tecnica(workbook, resultado, referencias)

        self._progresso(progresso, 96, "Salvando a planilha organizada...")
        try:
            workbook.save(destino)
        except PermissionError as erro:
            raise PermissionError(
                f"Não foi possível salvar {destino.name}. "
                "Feche o arquivo no Excel e tente novamente."
            ) from erro

        self._progresso(progresso, 100, "Excel organizado e reversível gerado com sucesso.")
        return ResultadoExportacaoExcel(
            caminho=destino,
            total_abas=total_abas,
            total_registros=resultado.estatisticas.total_registros_validos,
            total_linhas=resultado.estatisticas.total_linhas,
        )

    @classmethod
    def _mapear_participantes(cls, linhas: Iterable[str]) -> dict[str, tuple[str, str]]:
        """Retorna COD_PART -> (CNPJ/CPF, nome) a partir dos registros 0150."""
        mapa: dict[str, tuple[str, str]] = {}
        for linha in linhas:
            campos = cls._separar_campos(str(linha).rstrip("\r\n"))
            if not campos or campos[0].upper() != "0150":
                continue
            cod_part = campos[1].strip() if len(campos) > 1 else ""
            nome = campos[2].strip() if len(campos) > 2 else ""
            cnpj = campos[4].strip() if len(campos) > 4 else ""
            cpf = campos[5].strip() if len(campos) > 5 else ""
            if cod_part:
                mapa[cod_part] = (cnpj or cpf, nome)
        return mapa

    def _preparar_dados(self, resultado: ResultadoSPED):
        dados: dict[str, list[LinhaExportacao]] = defaultdict(list)
        contextos: dict[str, list[str]] = {}

        # Hotfix 17.8.98 — nas colunas de contexto “Participante(C100/D100)”
        # mostramos três informações legíveis: COD_PART, CNPJ/CPF e nome.
        # Os campos técnicos do registro pai continuam intactos, garantindo
        # Excel → TXT reversível.
        mapa_participantes = self._mapear_participantes(resultado.linhas)
        ordem: list[str] = []
        layouts: dict[str, LayoutRegistro | None] = {}
        maior_quantidade_propria: dict[str, int] = defaultdict(int)
        referencias: list[ReferenciaTecnica] = []

        for sequencia, linha in enumerate(resultado.linhas, start=1):
            texto = linha.rstrip("\r\n")
            campos = self._separar_campos(texto)
            if not campos or not campos[0]:
                referencias.append(
                    ReferenciaTecnica(
                        sequencia=sequencia,
                        tipo="RAW",
                        id_tecnico="",
                        registro="",
                        aba="",
                        linha_excel=0,
                        qtd_contexto=0,
                        qtd_campos=0,
                        linha_original=texto,
                    )
                )
                continue

            codigo = campos[0].upper()
            campos[0] = codigo

            for registro in RESETAR_CONTEXTOS_AO_LER.get(codigo, ()):
                contextos.pop(registro, None)

            if codigo not in dados:
                ordem.append(codigo)
                layouts[codigo] = obter_layout(codigo)

            layout = layouts[codigo]
            valores_contexto: list[str] = []
            if layout:
                for campo_contexto in layout.contexto:
                    registro_pai = contextos.get(campo_contexto.registro, [])
                    valor = (
                        registro_pai[campo_contexto.indice]
                        if campo_contexto.indice < len(registro_pai)
                        else ""
                    )
                    if campo_contexto.cabecalho.casefold().startswith("participante("):
                        cod_part = valor
                        documento, nome = mapa_participantes.get(cod_part, ("", ""))
                        valores_contexto.extend([cod_part, documento, nome])
                    else:
                        valores_contexto.append(valor)

            referencia = ReferenciaTecnica(
                sequencia=sequencia,
                tipo="REGISTRO",
                id_tecnico=f"FP{sequencia:09d}",
                registro=codigo,
                aba="",
                linha_excel=0,
                qtd_contexto=len(valores_contexto),
                qtd_campos=len(campos),
                linha_original=texto,
            )
            referencias.append(referencia)
            dados[codigo].append(
                LinhaExportacao(
                    valores=[*valores_contexto, *campos],
                    referencia=referencia,
                )
            )
            maior_quantidade_propria[codigo] = max(
                maior_quantidade_propria[codigo], len(campos)
            )
            contextos[codigo] = campos

        cabecalhos: dict[str, list[str]] = {}
        tipo_sped = resultado.estatisticas.tipo_sped
        eh_contribuicoes = str(tipo_sped or "").strip().casefold() == "efd contribuições".casefold()

        for codigo in ordem:
            layout = layouts[codigo]
            quantidade = maior_quantidade_propria[codigo]
            headers_contexto: list[str] = []
            if layout:
                for campo in layout.contexto:
                    if campo.cabecalho.casefold().startswith("participante("):
                        origem = campo.cabecalho[campo.cabecalho.find("("):]
                        headers_contexto.extend([
                            f"Código Participante{origem}",
                            f"CNPJ/CPF Participante{origem}",
                            f"Nome Participante{origem}",
                        ])
                    else:
                        headers_contexto.append(campo.cabecalho)

            oficiais = obter_cabecalhos_oficiais(tipo_sped, codigo)
            if oficiais:
                headers_proprios = list(oficiais)
            elif layout and not eh_contribuicoes:
                # No Fiscal, vários registros antigos já possuíam os nomes
                # técnicos corretos no layout legado. Na EFD-Contribuições
                # não reutilizamos um layout Fiscal de mesmo código, pois a
                # posição/semântica pode ser diferente (0000, 0200, C170,
                # D100, entre outros).
                headers_proprios = list(layout.cabecalhos)
            else:
                headers_proprios = ["REG"]

            # Campo adicional de um leiaute ainda não cadastrado: nunca
            # exibimos "CAMPO_01", que dá falsa impressão de nome oficial.
            # A posição fica explícita até o mapeamento técnico ser incluído.
            while len(headers_proprios) < quantidade:
                headers_proprios.append(
                    f"POSICAO_{len(headers_proprios) + 1:02d}_SEM_MAPEAMENTO"
                )

            cabecalhos[codigo] = [*headers_contexto, *headers_proprios]

        return ordem, dados, layouts, cabecalhos, referencias

    def _montar_abas_registro(
        self,
        workbook,
        codigo: str,
        linhas: list[LinhaExportacao],
        layout: LayoutRegistro | None,
        cabecalho: list[str],
        estilos: dict,
        nomes_usados: set[str],
    ) -> int:
        if len(cabecalho) + 1 > self.LIMITE_COLUNAS_EXCEL:
            raise RuntimeError(
                f"O registro {codigo} possui {len(cabecalho)} colunas, acima do limite do Excel."
            )

        nome_base = layout.nome_aba if layout else codigo
        total_abas = 0

        for numero_parte, lote in enumerate(
            self._em_lotes(linhas, self.LIMITE_DADOS_POR_ABA), start=1
        ):
            nome = self._nome_aba_unico(nome_base, numero_parte, nomes_usados)
            ws = workbook.create_sheet(nome)
            ws.append([*cabecalho, self.CABECALHO_ID])

            for linha in lote:
                valores = linha.valores
                valores_completos = [*valores, *([""] * (len(cabecalho) - len(valores)))]
                ws.append(
                    [
                        self._valor_excel(cabecalho[indice], valor)
                        for indice, valor in enumerate(valores_completos[: len(cabecalho)])
                    ]
                    + [linha.referencia.id_tecnico]
                )
                linha.referencia.aba = nome
                linha.referencia.linha_excel = ws.max_row

            self._finalizar_aba_dados(ws, cabecalho, estilos)
            total_abas += 1

        if not linhas:
            nome = self._nome_aba_unico(nome_base, 1, nomes_usados)
            ws = workbook.create_sheet(nome)
            ws.append([*cabecalho, self.CABECALHO_ID])
            self._finalizar_aba_dados(ws, cabecalho, estilos)
            total_abas = 1

        return total_abas

    def _criar_aba_tecnica(
        self,
        workbook,
        resultado: ResultadoSPED,
        referencias: list[ReferenciaTecnica],
    ) -> None:
        ws = workbook.create_sheet(self.ABA_TECNICA)
        quebra = self._detectar_quebra_linha(resultado.linhas)
        termina_com_quebra = bool(
            resultado.linhas
            and resultado.linhas[-1].endswith(("\r\n", "\n", "\r"))
        )

        bom_presente = self._possui_bom(resultado.caminho)
        metadados = (
            (self.ASSINATURA_FORMATO, self.VERSAO_FORMATO),
            ("ARQUIVO_ORIGEM", resultado.caminho.name),
            ("ENCODING", resultado.encoding),
            ("BOM_PRESENTE", "SIM" if bom_presente else "NAO"),
            ("QUEBRA_LINHA", quebra),
            ("FINAL_COM_QUEBRA", "SIM" if termina_com_quebra else "NAO"),
            ("TOTAL_LINHAS", len(resultado.linhas)),
            ("TOTAL_REGISTROS", resultado.estatisticas.total_registros_validos),
            ("GERADO_EM", datetime.now().isoformat(timespec="seconds")),
        )
        for chave, valor in metadados:
            ws.append([chave, valor])

        ws.append([])
        ws.append(
            [
                "SEQUENCIA",
                "TIPO",
                "FP_ID",
                "REGISTRO",
                "ABA_ORIGINAL",
                "LINHA_EXCEL",
                "QTD_CONTEXTO",
                "QTD_CAMPOS",
                "LINHA_ORIGINAL",
            ]
        )
        for ref in referencias:
            ws.append(
                [
                    ref.sequencia,
                    ref.tipo,
                    ref.id_tecnico,
                    ref.registro,
                    ref.aba,
                    ref.linha_excel,
                    ref.qtd_contexto,
                    ref.qtd_campos,
                    ref.linha_original,
                ]
            )

        ws.sheet_state = "veryHidden"

    def _finalizar_aba_dados(self, ws, cabecalho: list[str], estilos: dict) -> None:
        quantidade_colunas = len(cabecalho)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{self._letra_coluna(quantidade_colunas)}{max(ws.max_row, 1)}"
        self._estilizar_cabecalho(ws, quantidade_colunas, estilos)

        # Hotfix 17.7.7 — a precisão decimal precisa ser propriedade do
        # CAMPO SPED, e não um efeito do Excel/float. Além de deixar a planilha
        # legível (QTD com 5 casas, PIS/COFINS com 4, valores com 2 etc.),
        # isso evita que uma edição manual pareça ter perdido casas decimais.
        for indice, titulo in enumerate(cabecalho, start=1):
            letra = self._letra_coluna(indice)
            ws.column_dimensions[letra].width = self._largura_coluna(titulo)
            formato = self._formato_numero(titulo)
            if formato and ws.max_row >= 2:
                for (cell,) in ws.iter_rows(
                    min_row=2, max_row=ws.max_row, min_col=indice, max_col=indice
                ):
                    if cell.value not in (None, ""):
                        cell.number_format = formato

        coluna_id = self._letra_coluna(quantidade_colunas + 1)
        ws.column_dimensions[coluna_id].hidden = True
        ws.column_dimensions[coluna_id].width = 2

    def _estilizar_cabecalho(self, ws, quantidade_colunas: int, estilos: dict) -> None:
        Font = estilos["Font"]
        PatternFill = estilos["PatternFill"]
        Alignment = estilos["Alignment"]
        Border = estilos["Border"]
        Side = estilos["Side"]

        borda = Border(
            left=Side(style="thin", color=self.COR_BORDA),
            right=Side(style="thin", color=self.COR_BORDA),
            top=Side(style="thin", color=self.COR_BORDA),
            bottom=Side(style="thin", color=self.COR_BORDA),
        )
        for coluna in range(1, quantidade_colunas + 1):
            cell = ws.cell(1, coluna)
            cell.font = Font(bold=True, color=self.COR_TEXTO_CLARO)
            cell.fill = PatternFill("solid", fgColor=self.COR_CABECALHO)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = borda
        ws.row_dimensions[1].height = 34

    @classmethod
    def _valor_excel(cls, cabecalho: str, valor: str):
        texto = "" if valor is None else str(valor).strip()
        if not texto:
            return ""
        if not cls._cabecalho_numerico(cabecalho):
            return texto

        numero = texto.replace(" ", "").replace(",", ".")
        if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", numero):
            return texto
        try:
            if "." not in numero:
                return int(numero)
            return float(numero)
        except ValueError:
            return texto

    @classmethod
    def _cabecalho_numerico(cls, cabecalho: str) -> bool:
        return cls._casas_decimais(cabecalho) is not None

    @staticmethod
    def _casas_decimais(cabecalho: str) -> int | None:
        """Retorna a precisão técnica do campo SPED exibido no Excel.

        A regra prioriza os nomes técnicos oficiais introduzidos na 17.7.2.
        Na EFD-Contribuições, o Guia Prático define, entre outros: QTD do
        C170 com 5 casas; QUANT_BC_* com 3; ALIQ_PIS/COFINS e alíquotas por
        quantidade com 4; valores/bases monetárias com 2. Campos de ICMS/IPI
        permanecem com 2 casas. Cabeçalhos amigáveis legados continuam
        reconhecidos para preservar compatibilidade com planilhas anteriores.
        """

        titulo = str(cabecalho or "").strip().upper()
        if not titulo:
            return None

        # Totalizadores/contadores são inteiros.
        if (
            titulo.startswith("QTD_LIN")
            or titulo.startswith("QTD_REG")
            or titulo in {"QTD_CAMPOS", "QTD_CONTEXTO"}
        ):
            return 0

        # Quantidade comercial do item C170 (e equivalentes do leiaute
        # Fiscal) admite cinco casas decimais.
        if titulo == "QTD" or titulo.startswith("QTD(C"):
            return 5

        # Bases expressas em quantidade, conforme EFD-Contribuições.
        if titulo.startswith("QUANT_BC_"):
            return 3

        # Alíquotas de PIS/Cofins — percentuais ou específicas (R$/unidade).
        if titulo in {
            "ALIQ_PIS",
            "ALIQ_COFINS",
            "ALIQ_PIS_QUANT",
            "ALIQ_COFINS_QUANT",
        }:
            return 4

        # Mesma regra para cabeçalhos amigáveis de planilhas legadas.
        if ("ALÍQUOTA" in titulo or "ALIQUOTA" in titulo) and (
            "PIS" in titulo or "COFINS" in titulo
        ):
            return 4

        # ICMS/IPI e demais alíquotas percentuais tradicionais do leiaute.
        if titulo.startswith("ALIQ_") or "ALÍQUOTA" in titulo or "ALIQUOTA" in titulo:
            return 2

        # Fatores de conversão não devem ser arredondados para centavos.
        if "FAT_CONV" in titulo:
            return 6

        # Cabeçalhos legados de quantidade.
        if "QUANTIDADE" in titulo or titulo.startswith("QTD_") or titulo.startswith("QTDE_"):
            return 5

        # Valores, bases monetárias, receitas e descontos.
        marcadores_duas_casas = (
            "VALOR",
            "VL_",
            "BASE",
            "REC_BRU",
            "TOT_",
            "DESCONTO",
            "ABATIMENTO",
        )
        if any(marcador in titulo for marcador in marcadores_duas_casas):
            return 2

        return None

    @classmethod
    def _formato_numero(cls, cabecalho: str) -> str | None:
        casas = cls._casas_decimais(cabecalho)
        if casas is None:
            return None
        if casas == 0:
            return "#,##0"
        return "#,##0." + ("0" * casas)

    @staticmethod
    def _largura_coluna(cabecalho: str) -> int:
        titulo = cabecalho.upper()
        if "CHAVE" in titulo:
            return 48
        if any(termo in titulo for termo in ("DESCRI", "TXT", "INFORMAÇÃO", "COMPLEMENT")):
            return 38
        if any(termo in titulo for termo in ("CNPJ", "CPF", "ENDEREÇO", "RAZÃO", "NOME")):
            return 25
        if "DATA" in titulo or titulo.startswith("DT_"):
            return 15
        if any(termo in titulo for termo in ("VALOR", "VL_", "BASE", "ALÍQUOTA", "ALIQUOTA")):
            return 18
        return min(max(len(cabecalho) + 2, 12), 24)

    @staticmethod
    def _separar_campos(texto: str) -> list[str]:
        if not texto.startswith("|"):
            return []
        campos = texto.split("|")[1:]
        if texto.endswith("|") and campos:
            campos = campos[:-1]
        return campos

    @staticmethod
    def _possui_bom(caminho: Path) -> bool:
        try:
            with caminho.open("rb") as arquivo:
                inicio = arquivo.read(4)
        except OSError:
            return False
        return inicio.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff"))

    @staticmethod
    def _detectar_quebra_linha(linhas: list[str]) -> str:
        for linha in linhas:
            if linha.endswith("\r\n"):
                return "CRLF"
            if linha.endswith("\n"):
                return "LF"
            if linha.endswith("\r"):
                return "CR"
        return "CRLF"

    @staticmethod
    def _em_lotes(itens: Iterable, tamanho: int):
        lote = []
        for item in itens:
            lote.append(item)
            if len(lote) >= tamanho:
                yield lote
                lote = []
        if lote:
            yield lote

    @staticmethod
    def _nome_aba_unico(nome_base: str, numero_parte: int, usados: set[str]) -> str:
        seguro = re.sub(r"[\\/*?:\[\]]", "_", nome_base).strip() or "SEM_REG"
        sufixo = "" if numero_parte == 1 else f"_{numero_parte}"
        limite = 31 - len(sufixo)
        candidato = f"{seguro[:limite]}{sufixo}"

        contador = 2
        while candidato.casefold() in usados:
            extra = f"_{contador}"
            limite = 31 - len(extra)
            candidato = f"{seguro[:limite]}{extra}"
            contador += 1

        usados.add(candidato.casefold())
        return candidato

    @staticmethod
    def _letra_coluna(indice: int) -> str:
        letras = ""
        while indice:
            indice, resto = divmod(indice - 1, 26)
            letras = chr(65 + resto) + letras
        return letras

    @staticmethod
    def _progresso(callback: ProgressoCallback | None, percentual: int, mensagem: str) -> None:
        if callback:
            callback(percentual, mensagem)
