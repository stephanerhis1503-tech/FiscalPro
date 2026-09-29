from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Callable

from .exportador_excel import ExportadorSPEDExcel
from .estatisticas import CalculadorEstatisticasSPED
from .indice import IndiceSPED
from .pre_validador import PreValidadorPVA
from .sincronizador_bloco_m_excel import SincronizadorBlocoMExcel

ProgressoCallback = Callable[[int, str], None]


@dataclass(frozen=True, slots=True)
class ProblemaPlanilha:
    nivel: str
    mensagem: str
    sequencia: int | None = None
    registro: str = ""
    aba: str = ""
    linha_excel: int | None = None


@dataclass(frozen=True, slots=True)
class AlteracaoPlanilha:
    sequencia: int
    registro: str
    aba: str
    linha_excel: int
    campo: str
    original: str
    novo: str


@dataclass(frozen=True, slots=True)
class ItemConferenciaExcel:
    """Alteração encontrada entre o SPED original e a planilha validada.

    A Sprint 17.4.2 usa esta estrutura apenas para conferência visual e
    rastreabilidade. Ela não aplica nenhuma mudança por conta própria.
    """

    tipo: str
    categoria: str
    sequencia: int | None
    registro: str
    aba: str
    linha_excel: int | None
    campo: str
    original: str
    novo: str


@dataclass(slots=True)
class ResultadoValidacaoExcel:
    caminho: Path
    arquivo_origem: str = ""
    encoding: str = "utf-8"
    quebra_linha: str = "CRLF"
    final_com_quebra: bool = True
    bom_presente: bool = False
    total_linhas: int = 0
    total_registros: int = 0
    total_alteracoes: int = 0
    total_formulas: int = 0
    total_duplicidades_consolidadas: int = 0
    total_linhas_removidas: int = 0
    total_registros_opcionais_removidos: int = 0
    total_registros_bloco_m_removidos: int = 0
    modo_bloco_m_manual: bool = False
    total_ajustes_automaticos_c100: int = 0
    total_documentos_c100_retotalizados: int = 0
    total_ajustes_automaticos_c190: int = 0
    total_documentos_c190_reconstruidos: int = 0
    total_ajustes_automaticos_bloco_m: int = 0
    total_grupos_credito_bloco_m_consolidados: int = 0
    pre_pva_executado: bool = False
    pre_pva_tipo_sped: str = ""
    pre_pva_erros_origem: int = 0
    pre_pva_erros_reconstruido: int = 0
    pre_pva_novos_erros: int = 0
    pre_pva_avisos_origem: int = 0
    pre_pva_avisos_reconstruido: int = 0
    pre_pva_novos_avisos: int = 0
    problemas: list[ProblemaPlanilha] = field(default_factory=list)
    alteracoes: list[AlteracaoPlanilha] = field(default_factory=list)
    conferencia: list[ItemConferenciaExcel] = field(default_factory=list)
    duplicidades: list[DuplicidadeConsolidada] = field(default_factory=list)
    linhas_geradas: list[str] = field(default_factory=list, repr=False)

    @property
    def erros(self) -> list[ProblemaPlanilha]:
        return [p for p in self.problemas if p.nivel == "ERRO"]

    @property
    def avisos(self) -> list[ProblemaPlanilha]:
        return [p for p in self.problemas if p.nivel == "AVISO"]

    @property
    def valido(self) -> bool:
        return not self.erros


@dataclass(frozen=True, slots=True)
class DuplicidadeConsolidada:
    registro: str
    documento: str
    linhas_origem: tuple[int, ...]
    quantidade_linhas: int
    chave_tributaria: str


@dataclass(frozen=True, slots=True)
class ResultadoConversaoExcel:
    caminho_sped: Path
    caminho_relatorio: Path
    total_linhas: int
    total_registros: int
    total_alteracoes: int
    total_avisos: int
    total_duplicidades_consolidadas: int = 0
    total_linhas_removidas: int = 0
    total_registros_opcionais_removidos: int = 0
    total_registros_bloco_m_removidos: int = 0
    pre_pva_erros_reconstruido: int = 0
    pre_pva_novos_erros: int = 0
    pre_pva_avisos_reconstruido: int = 0
    pre_pva_novos_avisos: int = 0


@dataclass(frozen=True, slots=True)
class _ReferenciaLinha:
    sequencia: int
    tipo: str
    id_tecnico: str
    registro: str
    aba_original: str
    linha_excel_original: int
    qtd_contexto: int
    qtd_campos: int
    linha_original: str


class ImportadorExcelSPED:
    """Valida uma planilha reversível do FiscalPro e reconstrói o SPED TXT.

    A planilha altera registros existentes. Linhas novas continuam bloqueadas,
    mas registros analíticos C190 e D190 que se tornarem idênticos após ajustes
    de CFOP, CST ou alíquota são consolidados automaticamente, com soma dos
    valores e recálculo dos totalizadores do SPED. Os registros condicionais
    0205/0220 e, no modo de correção manual, os registros suportados do Bloco M
    podem ser removidos da planilha sem virar um falso erro de linha perdida.
    """

    ABA_TECNICA = ExportadorSPEDExcel.ABA_TECNICA
    CABECALHO_ID = ExportadorSPEDExcel.CABECALHO_ID
    ASSINATURA_FORMATO = ExportadorSPEDExcel.ASSINATURA_FORMATO
    VERSAO_FORMATO = ExportadorSPEDExcel.VERSAO_FORMATO
    LINHA_CABECALHO_TECNICO = 11
    REGISTROS_EXCLUSAO_OPCIONAL = frozenset({"0205", "0220"})
    # Hotfix 17.7.3 — o Excel também é uma ferramenta de CORREÇÃO. Se a
    # usuária apagar conscientemente uma linha do Bloco M, não devemos tratar
    # a ausência do __FP_ID__ como corrupção da planilha. Esses registros podem
    # ser removidos/alterados manualmente e o FiscalPro preserva a decisão,
    # recalcula apenas os totalizadores estruturais e deixa o Pré-PVA/PGE
    # apontarem eventual inconsistência fiscal remanescente.
    REGISTROS_EDICAO_MANUAL_BLOCO_M = frozenset(
        {
            "M100", "M105", "M200", "M210", "M400", "M410",
            "M500", "M505", "M600", "M610", "M800", "M810",
        }
    )

    def validar(
        self,
        caminho_excel: str | Path,
        progresso: ProgressoCallback | None = None,
    ) -> ResultadoValidacaoExcel:
        try:
            from openpyxl import load_workbook
        except ImportError as erro:
            raise RuntimeError(
                "A conversão Excel → TXT precisa do pacote openpyxl. "
                "Instale com: python -m pip install openpyxl"
            ) from erro

        caminho = Path(caminho_excel)
        if not caminho.exists():
            raise FileNotFoundError(f"Planilha não encontrada: {caminho}")
        if caminho.suffix.lower() != ".xlsx":
            raise RuntimeError("Selecione uma planilha .xlsx gerada pelo FiscalPro.")

        resultado = ResultadoValidacaoExcel(caminho=caminho)
        self._progresso(progresso, 5, "Abrindo a planilha do FiscalPro...")

        workbook = None
        workbook_valores = None
        try:
            workbook = load_workbook(caminho, data_only=False, read_only=False)
            workbook_valores = load_workbook(caminho, data_only=True, read_only=False)
        except Exception as erro:
            if workbook is not None:
                workbook.close()
            raise RuntimeError(f"Não foi possível abrir a planilha: {erro}") from erro

        try:
            if self.ABA_TECNICA not in workbook.sheetnames:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem=(
                            "A planilha não possui a estrutura reversível da Sprint 12.5. "
                            "Abra o SPED no FiscalPro atualizado e exporte novamente para Excel."
                        ),
                    )
                )
                return resultado

            tecnica = workbook[self.ABA_TECNICA]
            metadados = self._ler_metadados(tecnica)
            if metadados.get(self.ASSINATURA_FORMATO) != self.VERSAO_FORMATO:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem="Formato técnico da planilha ausente ou incompatível.",
                    )
                )
                return resultado

            resultado.arquivo_origem = metadados.get("ARQUIVO_ORIGEM", "")
            resultado.encoding = metadados.get("ENCODING", "utf-8") or "utf-8"
            resultado.bom_presente = metadados.get("BOM_PRESENTE", "NAO") == "SIM"
            resultado.quebra_linha = metadados.get("QUEBRA_LINHA", "CRLF") or "CRLF"
            resultado.final_com_quebra = metadados.get("FINAL_COM_QUEBRA", "SIM") == "SIM"

            referencias = self._ler_referencias(tecnica, resultado)
            if not referencias:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem="A estrutura técnica não contém as linhas do SPED.",
                    )
                )
                return resultado

            self._progresso(progresso, 20, "Localizando os registros nas abas do Excel...")
            mapa_ids = self._mapear_ids(workbook, resultado, referencias)
            ids_esperados = {r.id_tecnico for r in referencias if r.tipo == "REGISTRO"}
            ids_encontrados = set(mapa_ids)

            for id_extra in sorted(ids_encontrados - ids_esperados):
                ws, linha, coluna_id = mapa_ids[id_extra]
                registro_extra = self._registro_linha_extra(ws, linha, coluna_id)
                resultado.conferencia.append(
                    ItemConferenciaExcel(
                        tipo="Inclusão",
                        categoria="Outro",
                        sequencia=None,
                        registro=registro_extra,
                        aba=ws.title,
                        linha_excel=linha,
                        campo="REGISTRO",
                        original="(não existe no SPED original)",
                        novo=registro_extra or "Linha criada/copiada na planilha",
                    )
                )
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem=(
                            "Foi encontrada uma linha copiada ou criada sem vínculo com o SPED original. "
                            "A inclusão de registros será liberada em uma etapa futura."
                        ),
                        aba=ws.title,
                        linha_excel=linha,
                    )
                )

            total_referencias = max(len(referencias), 1)
            linhas_geradas: list[str] = []
            alteracoes: list[AlteracaoPlanilha] = []
            total_registros = 0
            total_formulas = 0
            formulas_sem_resultado = 0
            registros_opcionais_removidos: list[_ReferenciaLinha] = []
            registros_bloco_m_removidos: list[_ReferenciaLinha] = []
            registros_bloco_m_removidos_manuais: list[_ReferenciaLinha] = []
            registros_bloco_m_cascata: list[_ReferenciaLinha] = []
            posicao_gerada_por_sequencia: dict[int, int] = {}

            for posicao, referencia in enumerate(referencias, start=1):
                if referencia.tipo == "RAW":
                    posicao_gerada_por_sequencia[referencia.sequencia] = len(linhas_geradas)
                    linhas_geradas.append(referencia.linha_original)
                    continue

                total_registros += 1
                local = mapa_ids.get(referencia.id_tecnico)
                if local is None:
                    if referencia.registro.upper() in self.REGISTROS_EXCLUSAO_OPCIONAL:
                        registros_opcionais_removidos.append(referencia)
                        continue
                    if referencia.registro.upper() in self.REGISTROS_EDICAO_MANUAL_BLOCO_M:
                        registros_bloco_m_removidos.append(referencia)
                        continue
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="ERRO",
                            mensagem=(
                                "A linha original não foi encontrada. Ela pode ter sido apagada "
                                "ou o identificador oculto pode ter sido removido."
                            ),
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=referencia.aba_original,
                            linha_excel=referencia.linha_excel_original,
                        )
                    )
                    linhas_geradas.append(referencia.linha_original)
                    continue

                ws, linha_excel, coluna_id = local
                campos_originais = self._separar_campos(referencia.linha_original)
                if len(campos_originais) != referencia.qtd_campos:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="ERRO",
                            mensagem="Quantidade de campos da referência técnica está inconsistente.",
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=ws.title,
                            linha_excel=linha_excel,
                        )
                    )
                    linhas_geradas.append(referencia.linha_original)
                    continue

                coluna_inicial = referencia.qtd_contexto + 1
                coluna_final = coluna_inicial + referencia.qtd_campos - 1
                if coluna_final >= coluna_id:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="ERRO",
                            mensagem="As colunas da aba foram removidas ou deslocadas.",
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=ws.title,
                            linha_excel=linha_excel,
                        )
                    )
                    linhas_geradas.append(referencia.linha_original)
                    continue

                novos_campos: list[str] = []
                alteracoes_linha: list[AlteracaoPlanilha] = []
                linha_com_erro = False

                for indice, original in enumerate(campos_originais):
                    coluna = coluna_inicial + indice
                    cell = ws.cell(linha_excel, coluna)
                    cabecalho = str(ws.cell(1, coluna).value or f"CAMPO_{indice + 1:02d}")

                    valor_atual = cell.value
                    valor_veio_de_formula = False
                    if cell.data_type == "f" or (
                        isinstance(cell.value, str) and cell.value.startswith("=")
                    ):
                        valor_veio_de_formula = True
                        total_formulas += 1
                        ws_valores = (
                            workbook_valores[ws.title]
                            if ws.title in workbook_valores.sheetnames
                            else None
                        )
                        cell_valor = (
                            ws_valores.cell(linha_excel, coluna)
                            if ws_valores is not None
                            else None
                        )
                        valor_calculado = cell_valor.value if cell_valor is not None else None

                        if cell_valor is not None and cell_valor.data_type == "e":
                            resultado.problemas.append(
                                ProblemaPlanilha(
                                    nivel="ERRO",
                                    mensagem=(
                                        f"A fórmula do campo “{cabecalho}” resultou em erro "
                                        f"({valor_calculado}). Corrija a fórmula no Excel."
                                    ),
                                    sequencia=referencia.sequencia,
                                    registro=referencia.registro,
                                    aba=ws.title,
                                    linha_excel=linha_excel,
                                )
                            )
                            linha_com_erro = True
                            novos_campos.append(original)
                            continue

                        resultado_vazio_calculado = (
                            valor_calculado is None
                            and cell_valor is not None
                            and cell_valor.data_type in {"s", "str", "inlineStr"}
                        )
                        if valor_calculado is None and not resultado_vazio_calculado:
                            formulas_sem_resultado += 1
                            resultado.problemas.append(
                                ProblemaPlanilha(
                                    nivel="ERRO",
                                    mensagem=(
                                        f"A fórmula do campo “{cabecalho}” não possui resultado "
                                        "calculado salvo. Abra a planilha no Excel, pressione "
                                        "Ctrl+Alt+F9, salve, feche e valide novamente."
                                    ),
                                    sequencia=referencia.sequencia,
                                    registro=referencia.registro,
                                    aba=ws.title,
                                    linha_excel=linha_excel,
                                )
                            )
                            linha_com_erro = True
                            novos_campos.append(original)
                            continue

                        valor_atual = "" if resultado_vazio_calculado else valor_calculado

                    if cell.data_type == "e":
                        resultado.problemas.append(
                            ProblemaPlanilha(
                                nivel="ERRO",
                                mensagem=f"O campo “{cabecalho}” contém um erro do Excel.",
                                sequencia=referencia.sequencia,
                                registro=referencia.registro,
                                aba=ws.title,
                                linha_excel=linha_excel,
                            )
                        )
                        linha_com_erro = True
                        novos_campos.append(original)
                        continue

                    esperado_excel = ExportadorSPEDExcel._valor_excel(cabecalho, original)
                    if (
                        not valor_veio_de_formula
                        and self._valores_equivalentes(valor_atual, esperado_excel)
                    ):
                        novo = original
                    else:
                        try:
                            novo = self._serializar_valor(valor_atual, cabecalho)
                        except ValueError as erro:
                            resultado.problemas.append(
                                ProblemaPlanilha(
                                    nivel="ERRO",
                                    mensagem=f"Campo “{cabecalho}”: {erro}",
                                    sequencia=referencia.sequencia,
                                    registro=referencia.registro,
                                    aba=ws.title,
                                    linha_excel=linha_excel,
                                )
                            )
                            linha_com_erro = True
                            novo = original

                    novos_campos.append(novo)
                    if novo != original:
                        alteracoes_linha.append(
                            AlteracaoPlanilha(
                                sequencia=referencia.sequencia,
                                registro=referencia.registro,
                                aba=ws.title,
                                linha_excel=linha_excel,
                                campo=cabecalho,
                                original=original,
                                novo=novo,
                            )
                        )

                registro_removivel_esvaziado = (
                    referencia.registro.upper()
                    in (self.REGISTROS_EXCLUSAO_OPCIONAL | self.REGISTROS_EDICAO_MANUAL_BLOCO_M)
                    and novos_campos
                    and not any(str(campo or "").strip() for campo in novos_campos)
                )
                if registro_removivel_esvaziado:
                    if referencia.registro.upper() in self.REGISTROS_EXCLUSAO_OPCIONAL:
                        registros_opcionais_removidos.append(referencia)
                    else:
                        registros_bloco_m_removidos.append(referencia)
                    continue

                if novos_campos and novos_campos[0].upper() != referencia.registro:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="ERRO",
                            mensagem=(
                                f"O código do registro foi alterado de {referencia.registro} "
                                f"para {novos_campos[0] or '(vazio)'}."
                            ),
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=ws.title,
                            linha_excel=linha_excel,
                        )
                    )
                    linha_com_erro = True

                posicao_gerada_por_sequencia[referencia.sequencia] = len(linhas_geradas)
                if linha_com_erro:
                    linhas_geradas.append(referencia.linha_original)
                else:
                    linhas_geradas.append("|" + "|".join(novos_campos) + "|")
                    alteracoes.extend(alteracoes_linha)

                percentual = 20 + int((posicao / total_referencias) * 70)
                if posicao % 500 == 0 or posicao == total_referencias:
                    self._progresso(
                        progresso,
                        min(percentual, 90),
                        f"Validando linha {posicao:,} de {len(referencias):,}...",
                    )

            # Hotfix 17.7.4 — exclusão hierárquica determinística do Bloco M.
            # Se a usuária remove um M100/M500, os filhos de detalhamento que
            # pertenciam àquele pai não podem ficar órfãos no TXT. O PGE rejeita
            # M105 sem M100 e M505 sem M500. A cascata usa a ordem ORIGINAL da
            # aba técnica, por isso não tenta adivinhar a qual grupo o filho
            # pertencia, mesmo quando existem vários pais com a mesma alíquota.
            registros_bloco_m_removidos_manuais = list(registros_bloco_m_removidos)
            (
                linhas_geradas,
                registros_bloco_m_cascata,
            ) = self._aplicar_cascata_hierarquica_bloco_m(
                linhas_geradas,
                referencias,
                registros_bloco_m_removidos_manuais,
                posicao_gerada_por_sequencia,
            )
            if registros_bloco_m_cascata:
                registros_bloco_m_removidos.extend(registros_bloco_m_cascata)

            # O reposicionamento 17.7.5 é feito mais adiante, dentro do modo
            # manual e DEPOIS da assistência aritmética. Assim a correção de
            # ordem não altera valores que a usuária digitou conscientemente.
            reposicionamentos_bloco_m: list[tuple[str, int, int, str, int]] = []
            avisos_reposicionamento_bloco_m: list[str] = []

            resultado.alteracoes = alteracoes
            resultado.total_alteracoes = len(alteracoes)
            resultado.total_formulas = total_formulas

            # Hotfix 17.7.6 — saneia um erro estrutural típico do C175 antes
            # de qualquer apuração/Pré-PVA. Quando a tributação é ad valorem
            # (VL_BC + ALIQ percentual), ALIQ_*_QUANT só pode existir junto da
            # respectiva QUANT_BC. Fórmulas colocadas por engano nessa coluna
            # geravam valores como 0,00442 e o PGE rejeitava o registro.
            (
                linhas_geradas,
                ajustes_c175,
                avisos_c175,
            ) = self._normalizar_c175_aliquota_por_quantidade(linhas_geradas)
            if ajustes_c175:
                resultado.conferencia.extend(ajustes_c175)
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C175",
                        mensagem=(
                            f"O FiscalPro removeu automaticamente {len(ajustes_c175)} valor(es) "
                            "órfão(s) de ALIQ_PIS_QUANT/ALIQ_COFINS_QUANT em C175. "
                            "Esses campos só podem ser usados com base de cálculo em quantidade; "
                            "a base e a alíquota percentuais informadas foram preservadas."
                        ),
                    )
                )
            for mensagem in avisos_c175:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C175",
                        mensagem=mensagem,
                    )
                )

            # Hotfix 17.8.123 — no C175 o PGE exige coerência entre PIS e
            # COFINS. Quando a planilha corrige a base de apenas um dos dois
            # tributos (caso típico da exclusão do ICMS próprio), o outro lado
            # não pode permanecer com a base antiga. O FiscalPro espelha a base
            # que foi explicitamente alterada e recalcula somente o valor do
            # tributo espelhado, sem adivinhar quando ambos os lados foram
            # editados de forma divergente.
            (
                linhas_geradas,
                ajustes_bases_c175,
                avisos_bases_c175,
            ) = self._sincronizar_c175_bases_piscofins(
                linhas_geradas,
                alteracoes,
                posicao_gerada_por_sequencia,
            )
            if ajustes_bases_c175:
                resultado.conferencia.extend(ajustes_bases_c175)
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C175",
                        mensagem=(
                            f"O FiscalPro sincronizou automaticamente {len(ajustes_bases_c175)} "
                            "campo(s) de PIS/COFINS em C175 com bases divergentes. Quando um lado "
                            "já continha a base corrigida e o outro ainda estava na base bruta, "
                            "a base corrigida foi espelhada e o valor correspondente recalculado. "
                            "Os ajustes estão visíveis na Conferência Excel."
                        ),
                    )
                )
            for mensagem in avisos_bases_c175:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C175",
                        mensagem=mensagem,
                    )
                )

            resultado.conferencia.extend(
                self._item_conferencia_alteracao(alteracao)
                for alteracao in alteracoes
            )
            alteracoes_manuais_bloco_m = [
                alteracao
                for alteracao in alteracoes
                if alteracao.registro.upper() in self.REGISTROS_EDICAO_MANUAL_BLOCO_M
            ]
            resultado.modo_bloco_m_manual = bool(
                registros_bloco_m_removidos or alteracoes_manuais_bloco_m
            )

            # Hotfix 17.8.91: uma troca de CST de entrada para faixa sem direito
            # a crédito (70-75/98/99) precisa limpar os campos dependentes do
            # tributo no C170 antes de retotalizar o C100. Até a 17.8.90 a
            # alteração do CST era detectada, mas os valores antigos do item
            # permaneciam e, por consequência, o total do C100 não mudava.
            (
                linhas_geradas,
                ajustes_cst_c170,
                avisos_cst_c170,
            ) = self._normalizar_c170_por_cst_alterado(
                linhas_geradas,
                alteracoes,
                posicao_gerada_por_sequencia,
            )
            if ajustes_cst_c170:
                resultado.conferencia.extend(ajustes_cst_c170)
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C170",
                        mensagem=(
                            f"O FiscalPro ajustou automaticamente {len(ajustes_cst_c170)} campo(s) "
                            "dependente(s) de PIS/COFINS porque o CST de entrada foi alterado "
                            "para uma faixa sem direito a crédito. O C100 será retotalizado na sequência."
                        ),
                    )
                )
            for mensagem in avisos_cst_c170:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C170",
                        mensagem=mensagem,
                    )
                )

            # Hotfix 17.6.4: quando a planilha altera PIS/COFINS dos C170, os
            # totais VL_PIS/VL_COFINS do C100 pai precisam acompanhar a soma
            # dos itens. O ajuste é feito somente nos documentos afetados e
            # aparece na conferência como "Ajuste automático".
            (
                linhas_geradas,
                ajustes_c100,
                documentos_c100_retotalizados,
            ) = self._sincronizar_totais_c100_piscofins(
                linhas_geradas,
                referencias,
                alteracoes,
                posicao_gerada_por_sequencia,
            )
            if ajustes_c100:
                resultado.conferencia.extend(ajustes_c100)
                resultado.total_ajustes_automaticos_c100 = len(ajustes_c100)
                resultado.total_documentos_c100_retotalizados = documentos_c100_retotalizados
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        mensagem=(
                            f"O FiscalPro sincronizou automaticamente {len(ajustes_c100)} total(is) "
                            f"de PIS/COFINS em {documentos_c100_retotalizados} C100 afetado(s) "
                            "pelas alterações dos C170. Os ajustes estão visíveis na Conferência Excel."
                        ),
                    )
                )

            # Hotfix 17.6.5: reaplica no fluxo Excel -> TXT a mesma camada
            # segura de Bloco M que já existia nos fluxos de correção do TXT.
            # Hotfix 17.7.3: quando a própria usuária alterou/apagou registros
            # do Bloco M, a correção manual passa a prevalecer. O automático é
            # suspenso para não recriar registros removidos nem sobrescrever
            # alíquotas/bases digitadas conscientemente no Excel.
            linhas_originais = [referencia.linha_original for referencia in referencias]
            linhas_antes_bloco_m = len(linhas_geradas)

            # Hotfix 17.8.92 — o C190 é o resumo analítico do ICMS do C100.
            # Alterar CST_ICMS/CFOP/ALIQ_ICMS no C170 sem reagrupar o C190
            # deixa a escrituração inconsistente. A reconstrução abaixo atua
            # somente em documentos afetados, preserva o VL_OPR total já
            # escriturado e recalcula os valores somáveis a partir dos C170.
            (
                linhas_geradas,
                ajustes_c190,
                documentos_c190_reconstruidos,
                avisos_c190,
            ) = self._sincronizar_c190_icms_por_c170_alterado(
                linhas_geradas,
                referencias,
                alteracoes,
                posicao_gerada_por_sequencia,
            )
            if ajustes_c190:
                resultado.conferencia.extend(ajustes_c190)
                resultado.total_ajustes_automaticos_c190 = len(ajustes_c190)
                resultado.total_documentos_c190_reconstruidos = documentos_c190_reconstruidos
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="C190",
                        mensagem=(
                            f"O FiscalPro reconstruiu automaticamente o C190 de "
                            f"{documentos_c190_reconstruidos} documento(s) porque o resumo analítico "
                            "estava divergente dos C170 (CST/CFOP/alíquota). Confira o reagrupamento "
                            "na Conferência Excel."
                        ),
                    )
                )
            for mensagem in avisos_c190:
                resultado.problemas.append(
                    ProblemaPlanilha(nivel="AVISO", registro="C190", mensagem=mensagem)
                )

            tipo_sped_excel = self._identificar_tipo_sped_linhas(linhas_originais)
            if resultado.modo_bloco_m_manual:
                bloco_m_manual = SincronizadorBlocoMExcel().recalcular_modo_manual(
                    linhas_geradas,
                    tipo_sped_excel,
                )
                linhas_geradas = bloco_m_manual.linhas

                # Hotfix 17.7.5 — corrige apenas a ORDEM depois de preservar
                # integralmente a aritmética digitada. Ex.: um M105 que era
                # filho de um M100 excluído foi reaproveitado para o M100 de
                # 1,65%; a aba técnica ainda guarda a sequência antiga e o TXT
                # sairia M105 -> M100. Só movemos quando a base atual identifica
                # de forma única o novo pai.
                (
                    linhas_geradas,
                    reposicionamentos_bloco_m,
                    avisos_reposicionamento_bloco_m,
                ) = self._reordenar_filhos_reaproveitados_bloco_m(linhas_geradas)

                if reposicionamentos_bloco_m:
                    familias_movidas = {item[0] for item in reposicionamentos_bloco_m}
                    avisos_filtrados: list[str] = []
                    for aviso in bloco_m_manual.avisos:
                        if (
                            "M105" in familias_movidas
                            and "M100 linha" in aviso
                            and "ficou sem M105 filho" in aviso
                        ):
                            continue
                        if (
                            "M505" in familias_movidas
                            and "M500 linha" in aviso
                            and "ficou sem M505 filho" in aviso
                        ):
                            continue
                        avisos_filtrados.append(aviso)
                    bloco_m_manual.avisos = avisos_filtrados

                resultado.total_ajustes_automaticos_bloco_m = (
                    bloco_m_manual.campos_bloco_m_sincronizados
                )
                for registro, origem, destino, pai, linha_pai in reposicionamentos_bloco_m:
                    resultado.conferencia.append(
                        ItemConferenciaExcel(
                            tipo="Reposicionamento hierárquico",
                            categoria="Estrutural",
                            sequencia=None,
                            registro=registro,
                            aba="Bloco M",
                            linha_excel=None,
                            campo="ORDEM",
                            original=f"Antes do {pai} (posição {origem})",
                            novo=f"Após {pai} (posição {destino}; pai na posição {linha_pai})",
                        )
                    )
                for mensagem in avisos_reposicionamento_bloco_m:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="AVISO",
                            registro="Bloco M",
                            mensagem=mensagem,
                        )
                    )
                for ajuste in bloco_m_manual.ajustes:
                    resultado.conferencia.append(
                        ItemConferenciaExcel(
                            tipo="Assistência aritmética Bloco M",
                            categoria="Tributário",
                            sequencia=ajuste.numero_linha,
                            registro=ajuste.registro,
                            aba="Bloco M",
                            linha_excel=None,
                            campo=ajuste.campo,
                            original=ajuste.original,
                            novo=ajuste.novo,
                        )
                    )
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="Bloco M",
                        mensagem=(
                            "Modo manual do Bloco M ativado: o FiscalPro detectou "
                            f"{len(alteracoes_manuais_bloco_m)} campo(s) alterado(s), "
                            f"{len(registros_bloco_m_removidos_manuais)} registro(s) removido(s) "
                            "pela usuária"
                            + (
                                f" e {len(registros_bloco_m_cascata)} filho(s) dependente(s) "
                                "removido(s) automaticamente para preservar a hierarquia"
                                if registros_bloco_m_cascata
                                else ""
                            )
                            + ". As correções manuais foram preservadas e a "
                            "reclassificação automática M100/M105/M500/M505 foi suspensa "
                            "para não desfazer o que foi corrigido no Excel. O FiscalPro "
                            f"recalculou {bloco_m_manual.campos_bloco_m_sincronizados} campo(s) "
                            "aritmético(s) seguro(s) dos grupos simples existentes. Os "
                            "totalizadores estruturais serão recalculados e o resultado deve "
                            "ser conferido no Pré-PVA/PGE antes da transmissão."
                        ),
                    )
                )
                for mensagem in bloco_m_manual.avisos:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="AVISO",
                            registro="Bloco M",
                            mensagem=mensagem,
                        )
                    )
            else:
                bloco_m = SincronizadorBlocoMExcel().sincronizar(
                    linhas_originais,
                    linhas_geradas,
                    referencias,
                    posicao_gerada_por_sequencia,
                    tipo_sped_excel,
                )
                if bloco_m.erros:
                    for mensagem in bloco_m.erros:
                        resultado.problemas.append(
                            ProblemaPlanilha(
                                nivel="ERRO",
                                registro="Bloco M",
                                mensagem=mensagem,
                            )
                        )
                else:
                    linhas_geradas = bloco_m.linhas
                    resultado.total_ajustes_automaticos_bloco_m = (
                        bloco_m.campos_bloco_m_sincronizados
                    )
                    resultado.total_grupos_credito_bloco_m_consolidados = (
                        bloco_m.grupos_credito_consolidados
                    )
                    for ajuste in bloco_m.ajustes:
                        resultado.conferencia.append(
                            ItemConferenciaExcel(
                                tipo="Ajuste automático Bloco M",
                                categoria="Tributário",
                                sequencia=ajuste.numero_linha,
                                registro=ajuste.registro,
                                aba="Bloco M",
                                linha_excel=None,
                                campo=ajuste.campo,
                                original=ajuste.original,
                                novo=ajuste.novo,
                            )
                        )
                    if bloco_m.alterou:
                        resultado.problemas.append(
                            ProblemaPlanilha(
                                nivel="AVISO",
                                registro="Bloco M",
                                mensagem=(
                                    "O FiscalPro reaplicou automaticamente a correção segura do "
                                    f"Bloco M no fluxo Excel -> TXT: "
                                    f"{bloco_m.grupos_credito_consolidados} grupo(s) de crédito "
                                    f"consolidado(s) e {bloco_m.campos_bloco_m_sincronizados} "
                                    "ajuste(s)/totalização(ões) sincronizado(s). Os ajustes estão "
                                    "visíveis na Conferência Excel."
                                ),
                            )
                        )
                    for mensagem in bloco_m.avisos:
                        resultado.problemas.append(
                            ProblemaPlanilha(
                                nivel="AVISO", registro="Bloco M", mensagem=mensagem
                            )
                        )

            # Antes da consolidação, a planilha precisa reproduzir exatamente a
            # quantidade de linhas da estrutura técnica original. A redução de
            # linhas somente é permitida pelo tratamento seguro de C190/D190.
            total_esperado = self._inteiro_seguro(metadados.get("TOTAL_LINHAS"))
            esperado_apos_remocoes = (
                total_esperado
                - len(registros_opcionais_removidos)
                - len(registros_bloco_m_removidos)
            )
            if total_esperado and linhas_antes_bloco_m != esperado_apos_remocoes:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem=(
                            f"A planilha deveria reconstruir {esperado_apos_remocoes} linhas "
                            f"após as exclusões opcionais, mas foram obtidas "
                            f"{linhas_antes_bloco_m} antes das correções automáticas do Bloco M."
                        ),
                    )
                )

            linhas_geradas, duplicidades = self._consolidar_duplicidades_analiticas(
                linhas_geradas
            )
            linhas_duplicadas_removidas = sum(
                item.quantidade_linhas - 1 for item in duplicidades
            )

            if (
                registros_opcionais_removidos
                or registros_bloco_m_removidos
                or duplicidades
                or documentos_c190_reconstruidos
            ):
                linhas_geradas = self._recalcular_totalizadores(linhas_geradas)

            resultado.total_registros_opcionais_removidos = len(
                registros_opcionais_removidos
            )
            resultado.total_registros_bloco_m_removidos = len(
                registros_bloco_m_removidos
            )
            resultado.total_linhas_removidas = linhas_duplicadas_removidas

            if registros_opcionais_removidos:
                resultado.conferencia.extend(
                    ItemConferenciaExcel(
                        tipo="Exclusão",
                        categoria="Cadastral",
                        sequencia=referencia.sequencia,
                        registro=referencia.registro,
                        aba=referencia.aba_original,
                        linha_excel=referencia.linha_excel_original,
                        campo="REGISTRO",
                        original=referencia.linha_original,
                        novo="(registro removido da reconstrução)",
                    )
                    for referencia in registros_opcionais_removidos
                )
                contagem_opcionais = Counter(
                    referencia.registro.upper()
                    for referencia in registros_opcionais_removidos
                )
                resumo = ", ".join(
                    f"{codigo}: {quantidade}"
                    for codigo, quantidade in sorted(contagem_opcionais.items())
                )
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        mensagem=(
                            f"Foram removidos {len(registros_opcionais_removidos)} registro(s) "
                            f"condicional(is) ({resumo}). A geração do TXT foi liberada e os "
                            "totalizadores 0990, 9900, 9990 e 9999 foram recalculados."
                        ),
                    )
                )

            if registros_bloco_m_removidos:
                if registros_bloco_m_removidos_manuais:
                    resultado.conferencia.extend(
                        ItemConferenciaExcel(
                            tipo="Exclusão manual",
                            categoria="Tributário",
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=referencia.aba_original,
                            linha_excel=referencia.linha_excel_original,
                            campo="REGISTRO",
                            original=referencia.linha_original,
                            novo="(registro do Bloco M removido pela usuária)",
                        )
                        for referencia in registros_bloco_m_removidos_manuais
                    )
                if registros_bloco_m_cascata:
                    resultado.conferencia.extend(
                        ItemConferenciaExcel(
                            tipo="Exclusão hierárquica",
                            categoria="Tributário",
                            sequencia=referencia.sequencia,
                            registro=referencia.registro,
                            aba=referencia.aba_original,
                            linha_excel=referencia.linha_excel_original,
                            campo="REGISTRO",
                            original=referencia.linha_original,
                            novo=(
                                "(filho do Bloco M removido automaticamente porque o "
                                "registro-pai foi excluído pela usuária)"
                            ),
                        )
                        for referencia in registros_bloco_m_cascata
                    )
                contagem_m = Counter(
                    referencia.registro.upper()
                    for referencia in registros_bloco_m_removidos
                )
                resumo_m = ", ".join(
                    f"{codigo}: {quantidade}"
                    for codigo, quantidade in sorted(contagem_m.items())
                )
                complemento_cascata = (
                    f" Destas, {len(registros_bloco_m_cascata)} foram exclusões hierárquicas "
                    "de filhos dependentes (M105/M110/M115 ou M505/M510/M515) de um "
                    "M100/M500 removido."
                    if registros_bloco_m_cascata
                    else ""
                )
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        registro="Bloco M",
                        mensagem=(
                            f"Foram processadas {len(registros_bloco_m_removidos)} exclusão(ões) "
                            f"do Bloco M ({resumo_m}). "
                            f"{len(registros_bloco_m_removidos_manuais)} foram decisões manuais "
                            "da usuária."
                            + complemento_cascata
                            + " As exclusões não são tratadas como 'linha original não encontrada'. "
                            "O M990 e os totalizadores 0990/9900/9990/9999 foram "
                            "recalculados automaticamente."
                        ),
                    )
                )

            if duplicidades:
                resultado.duplicidades = duplicidades
                resultado.total_duplicidades_consolidadas = len(duplicidades)
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        mensagem=(
                            f"O FiscalPro consolidou automaticamente "
                            f"{len(duplicidades)} grupo(s) duplicado(s) de C190/D190 e "
                            f"removeu {resultado.total_linhas_removidas} linha(s), "
                            "somando os valores fiscais e recalculando os totalizadores."
                        ),
                    )
                )
                for item in duplicidades:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="AVISO",
                            registro=item.registro,
                            mensagem=(
                                f"Duplicidade {item.registro} consolidada no documento "
                                f"{item.documento}: {item.quantidade_linhas} linhas viraram 1. "
                                f"Chave: {item.chave_tributaria}."
                            ),
                        )
                    )

            # Hotfix 17.8.100 — linhas RAW vazias não fazem parte do leiaute da EFD.
            # Algumas planilhas reversíveis antigas carregaram uma linha RAW vazia
            # entre cada registro. Ao retornar Excel -> TXT isso virava uma linha física
            # em branco e o PVA acusava “Estrutura da linha inválida” nas posições
            # pares (2, 4, 6, ...), chegando rapidamente ao limite de 1000 erros.
            # Removemos somente linhas totalmente vazias/whitespace; nenhum registro
            # iniciado por |REG| é alterado.
            linhas_geradas = [
                str(linha).rstrip("\r\n")
                for linha in linhas_geradas
                if str(linha).strip()
            ]

            resultado.linhas_geradas = linhas_geradas
            resultado.total_linhas = len(linhas_geradas)
            resultado.total_registros = sum(
                1 for linha in linhas_geradas if self._separar_campos(linha)
            )

            self._detectar_c170_identicos(resultado)
            self._validar_estrutura_basica(resultado)

            # Sprint 17.4.1: além da integridade da planilha, simulamos o TXT em
            # memória e passamos o mesmo Pré-Validador PVA do FiscalPro antes de
            # liberar a gravação. Erros que já existiam no SPED de origem não
            # bloqueiam a conversão por si só; somente inconsistências NOVAS
            # criadas pela edição do Excel são bloqueantes nesta etapa.
            if not resultado.erros:
                self._pre_validar_reconstrucao(
                    resultado, linhas_originais, progresso
                )

            if total_formulas and not formulas_sem_resultado:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        mensagem=(
                            f"{total_formulas} fórmula(s) aceita(s). O FiscalPro usou os resultados "
                            "calculados que estavam salvos na planilha. Sempre recalcule e salve "
                            "o Excel antes de gerar o TXT."
                        ),
                    )
                )
            if resultado.total_registros and resultado.total_alteracoes > resultado.total_registros:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="AVISO",
                        mensagem=(
                            "Foram alterados muitos campos. Revise o relatório antes de importar "
                            "o TXT no PVA."
                        ),
                    )
                )

            self._progresso(
                progresso,
                100,
                "Planilha validada com sucesso." if resultado.valido else "Validação concluída com erros.",
            )
            return resultado
        finally:
            if workbook is not None:
                workbook.close()
            if workbook_valores is not None:
                workbook_valores.close()

    def gerar_txt(
        self,
        caminho_excel: str | Path,
        caminho_saida: str | Path,
        progresso: ProgressoCallback | None = None,
    ) -> ResultadoConversaoExcel:
        def progresso_validacao(percentual: int, mensagem: str) -> None:
            self._progresso(progresso, int(percentual * 0.85), mensagem)

        validacao = self.validar(caminho_excel, progresso_validacao)
        if not validacao.valido:
            detalhes = "\n".join(f"- {erro.mensagem}" for erro in validacao.erros[:5])
            if len(validacao.erros) > 5:
                detalhes += f"\n- ... e mais {len(validacao.erros) - 5} erro(s)."
            raise RuntimeError(
                "A planilha possui erros e o TXT não foi gerado.\n\n" + detalhes
            )

        destino = Path(caminho_saida)
        if destino.suffix.lower() != ".txt":
            destino = destino.with_suffix(".txt")
        destino.parent.mkdir(parents=True, exist_ok=True)

        # Hotfix 17.8.99 — o PVA precisa receber um registro SPED por linha física.
        # Alguns arquivos de origem podem chegar com terminador CR isolado. Embora
        # o FiscalPro consiga lê-los, o PVA do Windows pode interpretar esse retorno
        # como uma única linha gigantesca (ex.: registro 0000 com centenas de
        # milhares de campos). No retorno Excel -> TXT normalizamos CR para CRLF e
        # gravamos cada registro individualmente, sem depender de um join global.
        quebra_origem = (validacao.quebra_linha or "CRLF").upper()
        quebra = "\r\n" if quebra_origem == "CR" else {
            "CRLF": "\r\n",
            "LF": "\n",
        }.get(quebra_origem, "\r\n")

        # Segunda barreira: mesmo que uma planilha legada traga referências RAW
        # vazias, nunca gravamos linhas em branco no arquivo entregue ao PVA.
        linhas_saida = [
            str(linha).rstrip("\r\n")
            for linha in validacao.linhas_geradas
            if str(linha).strip()
        ]
        for numero, linha in enumerate(linhas_saida, start=1):
            if "\r" in linha or "\n" in linha:
                raise RuntimeError(
                    "O TXT não foi gerado porque um registro contém quebra de linha interna "
                    f"(linha lógica {numero}). Corrija a célula correspondente no Excel."
                )

        self._progresso(progresso, 94, "Gravando o arquivo SPED TXT...")
        encoding_saida = self._encoding_saida(validacao.encoding, validacao.bom_presente)
        temporario = destino.with_name(destino.name + ".fiscalpro_tmp")
        try:
            with temporario.open("w", encoding=encoding_saida, newline="") as stream:
                ultima = len(linhas_saida) - 1
                for indice, linha in enumerate(linhas_saida):
                    stream.write(linha)
                    if indice < ultima or validacao.final_com_quebra:
                        stream.write(quebra)

            # Validação física pós-gravação: antes de liberar o arquivo ao usuário,
            # garantimos que a quantidade de linhas no disco é exatamente a mesma
            # quantidade reconstruída pela planilha. Isso impede gerar silenciosamente
            # um SPED concatenado em uma única linha.
            with temporario.open("r", encoding=encoding_saida, newline="") as stream:
                linhas_fisicas = stream.readlines()
            if len(linhas_fisicas) != len(linhas_saida):
                raise RuntimeError(
                    "Falha de segurança ao gerar o TXT: a quantidade de linhas físicas "
                    f"({len(linhas_fisicas):,}) ficou diferente dos registros reconstruídos "
                    f"({len(linhas_saida):,}). O arquivo foi bloqueado para evitar erro no PVA."
                )

            temporario.replace(destino)
        except LookupError as erro:
            temporario.unlink(missing_ok=True)
            raise RuntimeError(
                f"A codificação original “{validacao.encoding}” não é reconhecida."
            ) from erro
        except UnicodeEncodeError as erro:
            temporario.unlink(missing_ok=True)
            raise RuntimeError(
                "Um texto digitado no Excel possui caracteres incompatíveis com a "
                f"codificação original ({validacao.encoding})."
            ) from erro
        except PermissionError as erro:
            temporario.unlink(missing_ok=True)
            raise PermissionError(
                f"Não foi possível salvar {destino.name}. Feche o arquivo e tente novamente."
            ) from erro
        except Exception:
            temporario.unlink(missing_ok=True)
            raise

        relatorio = destino.with_name(f"{destino.stem}_RELATORIO.txt")
        relatorio.write_text(
            self._montar_relatorio(validacao, destino),
            encoding="utf-8-sig",
        )
        self._progresso(progresso, 100, "SPED TXT e relatório gerados com sucesso.")

        return ResultadoConversaoExcel(
            caminho_sped=destino,
            caminho_relatorio=relatorio,
            total_linhas=validacao.total_linhas,
            total_registros=validacao.total_registros,
            total_alteracoes=validacao.total_alteracoes,
            total_avisos=len(validacao.avisos),
            total_duplicidades_consolidadas=validacao.total_duplicidades_consolidadas,
            total_linhas_removidas=validacao.total_linhas_removidas,
            total_registros_opcionais_removidos=(
                validacao.total_registros_opcionais_removidos
            ),
            total_registros_bloco_m_removidos=(
                validacao.total_registros_bloco_m_removidos
            ),
            pre_pva_erros_reconstruido=validacao.pre_pva_erros_reconstruido,
            pre_pva_novos_erros=validacao.pre_pva_novos_erros,
            pre_pva_avisos_reconstruido=validacao.pre_pva_avisos_reconstruido,
            pre_pva_novos_avisos=validacao.pre_pva_novos_avisos,
        )

    def _consolidar_duplicidades_analiticas(
        self, linhas: list[str]
    ) -> tuple[list[str], list[DuplicidadeConsolidada]]:
        """Agrupa C190/D190 equivalentes dentro do mesmo documento fiscal.

        Esses registros são totalizações analíticas. Depois de o usuário mudar
        CST, CFOP ou alíquota no Excel, duas totalizações antes diferentes podem
        passar a ter a mesma chave e precisam virar uma única linha.
        """

        configuracoes = {
            "C190": {
                "pai": "C100",
                "chave": (1, 2, 3, 11),
                "somar": (4, 5, 6, 7, 8, 9, 10),
                "min_campos": 12,
            },
            "D190": {
                "pai": "D100",
                "chave": (1, 2, 3, 8),
                "somar": (4, 5, 6, 7),
                "min_campos": 9,
            },
        }
        saida: list[str] = []
        grupos: dict[tuple, dict] = {}
        contador_pais = {"C100": 0, "D100": 0}
        pai_atual = {"C100": 0, "D100": 0}
        documento_atual = {"C100": "-", "D100": "-"}

        for numero_linha, linha in enumerate(linhas, start=1):
            campos = self._separar_campos(linha)
            if not campos:
                saida.append(linha)
                continue

            registro = campos[0].upper()
            if registro in contador_pais:
                contador_pais[registro] += 1
                pai_atual[registro] = contador_pais[registro]
                documento_atual[registro] = self._identificar_documento(registro, campos)
                saida.append(linha)
                continue

            config = configuracoes.get(registro)
            if config is None or len(campos) < config["min_campos"]:
                saida.append(linha)
                continue

            pai = config["pai"]
            identificador_pai = pai_atual[pai]
            if identificador_pai <= 0:
                saida.append(linha)
                continue

            chave_valores = tuple(
                self._normalizar_chave_analitica(campos[indice], indice == 3)
                for indice in config["chave"]
            )
            chave = (registro, identificador_pai, chave_valores)

            grupo = grupos.get(chave)
            if grupo is None:
                indice_saida = len(saida)
                saida.append(linha)
                grupos[chave] = {
                    "indice_saida": indice_saida,
                    "campos": campos,
                    "linhas": [numero_linha],
                    "documento": documento_atual[pai],
                    "config": config,
                }
                continue

            campos_base = grupo["campos"]
            for indice in config["somar"]:
                valor_base = self._decimal_seguro(campos_base[indice]) or Decimal("0")
                valor_novo = self._decimal_seguro(campos[indice]) or Decimal("0")
                campos_base[indice] = self._formatar_decimal_sped(
                    valor_base + valor_novo, "Valor da Operação"
                )
            grupo["linhas"].append(numero_linha)
            saida[grupo["indice_saida"]] = "|" + "|".join(campos_base) + "|"

        duplicidades: list[DuplicidadeConsolidada] = []
        for (registro, _pai, chave_valores), grupo in grupos.items():
            if len(grupo["linhas"]) <= 1:
                continue
            chave_texto = (
                f"CST {chave_valores[0] or '-'} | CFOP {chave_valores[1] or '-'} | "
                f"Alíquota {chave_valores[2] or '-'} | Obs. {chave_valores[3] or '-'}"
            )
            duplicidades.append(
                DuplicidadeConsolidada(
                    registro=registro,
                    documento=grupo["documento"],
                    linhas_origem=tuple(grupo["linhas"]),
                    quantidade_linhas=len(grupo["linhas"]),
                    chave_tributaria=chave_texto,
                )
            )
        return saida, duplicidades

    def _detectar_c170_identicos(self, resultado: ResultadoValidacaoExcel) -> None:
        """Avisa sobre itens idênticos, mas não os apaga automaticamente.

        C170 representa o item real da nota e duas linhas iguais podem ser
        legítimas. A consolidação automática fica restrita a C190/D190.
        """

        pai = 0
        documento = "-"
        vistos: dict[tuple, list[int]] = {}
        for numero_linha, linha in enumerate(resultado.linhas_geradas, start=1):
            campos = self._separar_campos(linha)
            if not campos:
                continue
            registro = campos[0].upper()
            if registro == "C100":
                pai += 1
                documento = self._identificar_documento("C100", campos)
                continue
            if registro != "C170" or len(campos) < 3:
                continue
            # Ignora NUM_ITEM (campo 1); os demais campos precisam ser idênticos.
            chave = (pai, tuple(campos[2:]))
            vistos.setdefault(chave, []).append(numero_linha)

        grupos = [(chave, linhas) for chave, linhas in vistos.items() if len(linhas) > 1]
        if grupos:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="AVISO",
                    registro="C170",
                    mensagem=(
                        f"Foram encontrados {len(grupos)} grupo(s) de itens C170 idênticos. "
                        "Eles foram mantidos porque podem representar itens legítimos da NF-e. "
                        "Confira a numeração dos itens antes de importar no PVA."
                    ),
                )
            )

    @classmethod
    def _aplicar_cascata_hierarquica_bloco_m(
        cls,
        linhas_geradas: list[str],
        referencias: list[_ReferenciaLinha],
        removidos_manuais: list[_ReferenciaLinha],
        posicoes_por_sequencia: dict[int, int],
    ) -> tuple[list[str], list[_ReferenciaLinha]]:
        """Remove filhos que ficaram órfãos após exclusão manual do pai no M.

        A decisão tributária continua sendo da usuária. A única automação aqui é
        estrutural: se um M100/M500 foi removido, os registros de detalhamento
        pertencentes àquele mesmo pai na estrutura ORIGINAL também precisam sair.
        A identificação é feita pela sequência técnica original, nunca apenas por
        alíquota ou por proximidade no TXT já reconstruído.
        """

        dependentes = {
            "M100": {"M105", "M110", "M115"},
            "M500": {"M505", "M510", "M515"},
        }
        pais_removidos = {
            referencia.sequencia: referencia.registro.upper()
            for referencia in removidos_manuais
            if referencia.registro.upper() in dependentes
        }
        if not pais_removidos:
            return list(linhas_geradas), []

        refs_registro = [r for r in referencias if r.tipo == "REGISTRO"]
        candidatos: list[_ReferenciaLinha] = []
        sequencias_candidatas: set[int] = set()

        for indice, referencia in enumerate(refs_registro):
            registro_pai = pais_removidos.get(referencia.sequencia)
            if not registro_pai:
                continue
            filhos_validos = dependentes[registro_pai]
            for proxima in refs_registro[indice + 1 :]:
                codigo = proxima.registro.upper()
                if codigo in filhos_validos:
                    # Filho apagado manualmente já não possui posição gerada e
                    # não precisa entrar novamente na cascata. Se ele ainda
                    # existe e foi EDITADO, porém, a usuária pode tê-lo
                    # reaproveitado para outro pai (caso real 17.7.5). Nesse
                    # cenário não o apagamos por causa da filiação original; a
                    # etapa seguinte tentará provar e corrigir a nova posição.
                    posicao_atual = posicoes_por_sequencia.get(proxima.sequencia)
                    if posicao_atual is not None and 0 <= posicao_atual < len(linhas_geradas):
                        linha_atual = str(linhas_geradas[posicao_atual]).rstrip("\r\n")
                        linha_original = str(proxima.linha_original).rstrip("\r\n")
                        if linha_atual != linha_original:
                            continue
                    if (
                        proxima.sequencia in posicoes_por_sequencia
                        and proxima.sequencia not in sequencias_candidatas
                    ):
                        candidatos.append(proxima)
                        sequencias_candidatas.add(proxima.sequencia)
                    continue
                # O subtree do crédito terminou ao aparecer o primeiro registro
                # que não é dependente daquele M100/M500.
                break

        if not candidatos:
            return list(linhas_geradas), []

        novas = list(linhas_geradas)
        removidos: list[_ReferenciaLinha] = []
        # As posições foram calculadas durante a reconstrução. Remover do fim
        # para o começo evita deslocar os índices ainda não processados.
        por_posicao = sorted(
            (
                (posicoes_por_sequencia[ref.sequencia], ref)
                for ref in candidatos
                if ref.sequencia in posicoes_por_sequencia
            ),
            key=lambda item: item[0],
            reverse=True,
        )
        for posicao, referencia in por_posicao:
            if posicao < 0 or posicao >= len(novas):
                continue
            if cls._codigo_linha(novas[posicao]) != referencia.registro.upper():
                # Guarda conservadora: se a posição não corresponder ao filho
                # esperado, não removemos uma linha diferente por engano.
                continue
            novas.pop(posicao)
            removidos.append(referencia)

        removidos.sort(key=lambda ref: ref.sequencia)
        return novas, removidos

    @classmethod
    def _reordenar_filhos_reaproveitados_bloco_m(
        cls,
        linhas: list[str],
    ) -> tuple[list[str], list[tuple[str, int, int, str, int]], list[str]]:
        """Reposiciona M105/M505 editados quando a nova filiação é comprovável.

        A planilha reversível guarda a sequência técnica ORIGINAL. Quando a
        usuária exclui um M100/M500 e reaproveita um filho existente para o
        grupo seguinte, o valor atual pode estar correto mas o TXT sair como
        ``M105 -> M100`` ou ``M505 -> M500``. O PGE rejeita essa hierarquia.

        Esta rotina não escolhe pai por proximidade e não altera valores. Ela
        só move um grupo de filhos órfãos se houver exatamente UM pai do mesmo
        tributo cuja base atual seja igual à soma das bases dos filhos já
        ligados a ele + as bases dos órfãos. Sem prova única, preserva a ordem e
        emite aviso para revisão no PGE.
        """

        novas = list(linhas)
        movimentos: list[tuple[str, int, int, str, int]] = []
        avisos: list[str] = []

        familias = (
            ("M100", "M105", {"M105", "M110", "M115"}),
            ("M500", "M505", {"M505", "M510", "M515"}),
        )

        def base_pai(indice: int) -> Decimal | None:
            campos = cls._separar_campos(novas[indice])
            return cls._decimal_seguro(campos[3] if len(campos) > 3 else None)

        def base_filho(indice: int) -> Decimal | None:
            campos = cls._separar_campos(novas[indice])
            # M105/M505: VL_BC_* efetiva do crédito é o campo 06.
            return cls._decimal_seguro(campos[6] if len(campos) > 6 else None)

        for pai, filho, dependentes in familias:
            # Pode haver mais de um trecho órfão. Recalcula os índices após
            # cada movimentação para não depender de posições antigas.
            for _ in range(50):
                pais: list[int] = []
                orfaos: list[int] = []
                pai_ativo: int | None = None

                for indice, linha in enumerate(novas):
                    codigo = cls._codigo_linha(linha)
                    if codigo == pai:
                        pais.append(indice)
                        pai_ativo = indice
                        continue
                    if codigo == filho:
                        if pai_ativo is None:
                            orfaos.append(indice)
                        continue
                    if codigo.startswith("M") and codigo not in dependentes:
                        pai_ativo = None

                if not orfaos:
                    break

                # Trata o primeiro trecho contíguo de filhos órfãos como uma
                # unidade; isso também cobre pais compostos por vários M105.
                grupo = [orfaos[0]]
                for indice in orfaos[1:]:
                    if indice == grupo[-1] + 1:
                        grupo.append(indice)
                    else:
                        break

                bases_orfas = [base_filho(i) for i in grupo]
                if any(valor is None for valor in bases_orfas):
                    avisos.append(
                        f"Bloco M manual: {filho} órfão não pôde ser reposicionado "
                        "porque a base atual não é numérica. A ordem foi preservada."
                    )
                    break
                total_orfao = sum((valor for valor in bases_orfas if valor is not None), Decimal("0"))

                candidatos: list[int] = []
                for indice_pai in pais:
                    base = base_pai(indice_pai)
                    if base is None:
                        continue

                    soma_existente = Decimal("0")
                    pos = indice_pai + 1
                    while pos < len(novas) and cls._codigo_linha(novas[pos]) in dependentes:
                        if cls._codigo_linha(novas[pos]) == filho:
                            valor = base_filho(pos)
                            if valor is None:
                                soma_existente = Decimal("NaN")
                                break
                            soma_existente += valor
                        pos += 1
                    if soma_existente.is_nan():
                        continue
                    if base == (soma_existente + total_orfao):
                        candidatos.append(indice_pai)

                if len(candidatos) != 1:
                    avisos.append(
                        f"Bloco M manual: encontrei {len(grupo)} {filho} órfão(s), "
                        f"mas a base não identificou um único {pai} de destino. "
                        "O FiscalPro não mudou a ordem automaticamente."
                    )
                    break

                indice_pai = candidatos[0]
                linhas_grupo = [novas[i] for i in grupo]
                origem_humana = grupo[0] + 1

                # Remove do fim para o começo.
                for indice in reversed(grupo):
                    novas.pop(indice)
                    if indice < indice_pai:
                        indice_pai -= 1

                # Insere logo após o pai e seus dependentes atuais.
                destino = indice_pai + 1
                while destino < len(novas) and cls._codigo_linha(novas[destino]) in dependentes:
                    destino += 1
                for deslocamento, linha in enumerate(linhas_grupo):
                    novas.insert(destino + deslocamento, linha)
                    movimentos.append(
                        (
                            filho,
                            origem_humana,
                            destino + deslocamento + 1,
                            pai,
                            indice_pai + 1,
                        )
                    )

            else:
                avisos.append(
                    f"Bloco M manual: limite de segurança atingido ao reorganizar {pai}/{filho}."
                )

        return novas, movimentos, avisos

    @classmethod
    def _recalcular_totalizadores(cls, linhas: list[str]) -> list[str]:
        novas = list(linhas)

        # Reconstrói o 9900 porque a quantidade de C190/D190 pode ter mudado.
        sem_9900 = [linha for linha in novas if cls._codigo_linha(linha) != "9900"]
        indice_9990 = next(
            (i for i, linha in enumerate(sem_9900) if cls._codigo_linha(linha) == "9990"),
            None,
        )
        if indice_9990 is not None:
            contagens = Counter(
                cls._codigo_linha(linha)
                for linha in sem_9900
                if cls._codigo_linha(linha)
            )
            codigos = sorted(set(contagens) | {"9900"})
            quantidade_9900 = len(codigos)
            registros_9900 = [
                f"|9900|{codigo}|"
                f"{quantidade_9900 if codigo == '9900' else contagens[codigo]}|"
                for codigo in codigos
            ]
            novas = sem_9900[:indice_9990] + registros_9900 + sem_9900[indice_9990:]

        blocos = {
            "0990": "0000",
            "1990": "1001",
            "B990": "B001",
            "C990": "C001",
            "D990": "D001",
            "E990": "E001",
            "G990": "G001",
            "H990": "H001",
            "K990": "K001",
            "M990": "M001",
            "9990": "9001",
        }
        posicoes: dict[str, list[int]] = {}
        for indice, linha in enumerate(novas):
            codigo = cls._codigo_linha(linha)
            if codigo:
                posicoes.setdefault(codigo, []).append(indice)

        for fechamento, abertura in blocos.items():
            for indice_fechamento in posicoes.get(fechamento, []):
                aberturas = [i for i in posicoes.get(abertura, []) if i <= indice_fechamento]
                if not aberturas:
                    continue
                indice_abertura = aberturas[-1]
                if fechamento == "9990" and posicoes.get("9999"):
                    quantidade = posicoes["9999"][0] - indice_abertura + 1
                else:
                    quantidade = indice_fechamento - indice_abertura + 1
                campos = cls._separar_campos(novas[indice_fechamento])
                if len(campos) >= 2:
                    campos[1] = str(quantidade)
                    novas[indice_fechamento] = "|" + "|".join(campos) + "|"

        indice_9999 = next(
            (i for i, linha in enumerate(novas) if cls._codigo_linha(linha) == "9999"),
            None,
        )
        if indice_9999 is not None:
            campos = cls._separar_campos(novas[indice_9999])
            if len(campos) >= 2:
                campos[1] = str(len(novas))
                novas[indice_9999] = "|" + "|".join(campos) + "|"
        return novas

    @staticmethod
    def _codigo_linha(linha: str) -> str:
        campos = ImportadorExcelSPED._separar_campos(linha)
        return campos[0].upper() if campos else ""

    @staticmethod
    def _normalizar_chave_analitica(valor: str, numerico: bool = False) -> str:
        texto = str(valor or "").strip()
        if numerico:
            numero = ImportadorExcelSPED._decimal_seguro(texto)
            if numero is not None:
                return format(numero.normalize(), "f")
        return texto.upper()

    @staticmethod
    def _identificar_documento(registro_pai: str, campos: list[str]) -> str:
        if registro_pai == "C100":
            numero = campos[7] if len(campos) > 7 else ""
            chave = campos[8] if len(campos) > 8 else ""
            return f"NF {numero or '-'}" + (f" / {chave}" if chave else "")
        numero = campos[8] if len(campos) > 8 else ""
        chave = campos[9] if len(campos) > 9 else ""
        return f"CT-e {numero or '-'}" + (f" / {chave}" if chave else "")

    def _mapear_ids(
        self,
        workbook,
        resultado: ResultadoValidacaoExcel,
        referencias: Iterable[_ReferenciaLinha] | None = None,
    ):
        mapa: dict[str, tuple[object, int, int]] = {}
        refs_por_aba: dict[str, list[_ReferenciaLinha]] = {}
        for ref in referencias or ():
            if ref.tipo != "REGISTRO" or not ref.id_tecnico or not ref.aba_original:
                continue
            refs_por_aba.setdefault(ref.aba_original, []).append(ref)

        for ws in workbook.worksheets:
            if ws.title == self.ABA_TECNICA:
                continue

            coluna_id = None
            for coluna in range(1, ws.max_column + 1):
                if ws.cell(1, coluna).value == self.CABECALHO_ID:
                    coluna_id = coluna
                    break

            if coluna_id is None:
                # Hotfix 17.8.121 — compatibilidade com planilhas já corrigidas
                # em que uma aba antiga perdeu somente a coluna oculta __FP_ID__.
                # Não adivinhamos vínculos: a recuperação por posição só é aceita
                # quando TODAS as linhas da aba continuam idênticas às referências
                # originais gravadas em _FISCALPRO_ORDEM. Assim a usuária não
                # precisa refazer centenas de correções de outras abas por causa
                # de um identificador técnico ausente em um registro não editado.
                refs_aba = refs_por_aba.get(ws.title, [])
                recuperados: list[tuple[str, tuple[object, int, int]]] = []
                seguro = bool(refs_aba) and self._aba_possui_dados(ws)
                if seguro:
                    for ref in refs_aba:
                        linha = int(ref.linha_excel_original or 0)
                        if linha < 2 or linha > ws.max_row:
                            seguro = False
                            break
                        campos_originais = self._separar_campos(ref.linha_original)
                        coluna_inicial = ref.qtd_contexto + 1
                        coluna_final = coluna_inicial + ref.qtd_campos - 1
                        if (
                            len(campos_originais) != ref.qtd_campos
                            or coluna_final > ws.max_column
                        ):
                            seguro = False
                            break

                        for indice, original in enumerate(campos_originais):
                            coluna = coluna_inicial + indice
                            cabecalho = str(
                                ws.cell(1, coluna).value or f"CAMPO_{indice + 1:02d}"
                            )
                            esperado = ExportadorSPEDExcel._valor_excel(cabecalho, original)
                            atual = ws.cell(linha, coluna).value
                            if not self._valores_equivalentes(atual, esperado):
                                seguro = False
                                break
                        if not seguro:
                            break
                        # Usa uma coluna virtual imediatamente após os dados. O
                        # laço de reconstrução só precisa dela como limite; não
                        # escreve nem altera a planilha do usuário.
                        recuperados.append(
                            (ref.id_tecnico, (ws, linha, ws.max_column + 1))
                        )

                if seguro and len(recuperados) == len(refs_aba):
                    for id_tecnico, local in recuperados:
                        mapa[id_tecnico] = local
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="AVISO",
                            mensagem=(
                                "A coluna técnica oculta __FP_ID__ não estava presente, "
                                "mas o FiscalPro recuperou com segurança os vínculos desta "
                                "aba usando a estrutura original gravada na planilha."
                            ),
                            aba=ws.title,
                        )
                    )
                    continue

                if self._aba_possui_dados(ws):
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="AVISO",
                            mensagem=(
                                "Aba adicional sem vínculo técnico; ela será ignorada na geração do TXT."
                            ),
                            aba=ws.title,
                        )
                    )
                continue

            for linha in range(2, ws.max_row + 1):
                valor_id = ws.cell(linha, coluna_id).value
                id_tecnico = "" if valor_id is None else str(valor_id).strip()

                if not id_tecnico:
                    possui_dados = any(
                        ws.cell(linha, coluna).value not in (None, "")
                        for coluna in range(1, coluna_id)
                    )
                    if possui_dados:
                        registro_extra = self._registro_linha_extra(ws, linha, coluna_id)
                        resultado.conferencia.append(
                            ItemConferenciaExcel(
                                tipo="Inclusão",
                                categoria="Outro",
                                sequencia=None,
                                registro=registro_extra,
                                aba=ws.title,
                                linha_excel=linha,
                                campo="REGISTRO",
                                original="(não existe no SPED original)",
                                novo=registro_extra or "Linha nova sem identificador técnico",
                            )
                        )
                        resultado.problemas.append(
                            ProblemaPlanilha(
                                nivel="ERRO",
                                mensagem=(
                                    "Linha nova sem identificador técnico. Nesta primeira versão, "
                                    "o Excel → TXT permite editar, mas ainda não incluir registros."
                                ),
                                aba=ws.title,
                                linha_excel=linha,
                            )
                        )
                    continue

                if id_tecnico in mapa:
                    resultado.problemas.append(
                        ProblemaPlanilha(
                            nivel="ERRO",
                            mensagem="Identificador técnico duplicado; a linha pode ter sido copiada.",
                            aba=ws.title,
                            linha_excel=linha,
                        )
                    )
                    continue
                mapa[id_tecnico] = (ws, linha, coluna_id)
        return mapa

    @staticmethod
    def _registro_linha_extra(ws, linha_excel: int, coluna_id: int) -> str:
        """Tenta identificar o código REG de uma linha sem referência técnica."""
        for coluna in range(1, coluna_id):
            cabecalho = str(ws.cell(1, coluna).value or "").strip().upper()
            if cabecalho == "REG":
                return str(ws.cell(linha_excel, coluna).value or "").strip().upper()
        return ""

    @classmethod
    def _normalizar_c170_por_cst_alterado(
        cls,
        linhas: list[str],
        alteracoes: list[AlteracaoPlanilha],
        posicao_gerada_por_sequencia: dict[int, int],
    ) -> tuple[list[str], list[ItemConferenciaExcel], list[str]]:
        """Normaliza PIS/COFINS quando o CST de uma entrada perde o direito a crédito.

        Regras conservadoras da 17.8.91:
        * atua somente em C170 de entrada (CFOP 1xxx/2xxx/3xxx);
        * CST 70-75/98/99: se a usuária alterou apenas o CST, zera base,
          alíquota, quantidade e valor do respectivo tributo;
        * se qualquer campo dependente também foi editado manualmente, a
          decisão da planilha é preservada e o Pré-PVA fica responsável por
          apontar eventual inconsistência;
        * CST 50-66: nunca inventa base/alíquota/valor. Se apenas o CST foi
          mudado para uma faixa de crédito e não há valor calculável, emite
          aviso para revisão.
        """
        sem_credito = {str(n) for n in range(70, 76)} | {"98", "99"}
        com_credito = {str(n) for n in range(50, 67)}
        cfg = {
            "PIS": {
                "cst": 24,
                "campos": {
                    25: ("VL_BC_PIS", "0,00"),
                    26: ("ALIQ_PIS", "0,0000"),
                    27: ("QUANT_BC_PIS", "0,00"),
                    28: ("ALIQ_PIS_QUANT", "0,0000"),
                    29: ("VL_PIS", "0,00"),
                },
            },
            "COFINS": {
                "cst": 30,
                "campos": {
                    31: ("VL_BC_COFINS", "0,00"),
                    32: ("ALIQ_COFINS", "0,0000"),
                    33: ("QUANT_BC_COFINS", "0,00"),
                    34: ("ALIQ_COFINS_QUANT", "0,0000"),
                    35: ("VL_COFINS", "0,00"),
                },
            },
        }

        def normalizar_nome(campo: str) -> str:
            return (
                str(campo or "").strip().upper()
                .replace("Á", "A").replace("À", "A").replace("Â", "A").replace("Ã", "A")
                .replace("É", "E").replace("Ê", "E").replace("Í", "I")
                .replace("Ó", "O").replace("Ô", "O").replace("Õ", "O")
                .replace("Ú", "U").replace("Ç", "C")
            )

        aliases = {
            "PIS": {
                "CST": {"CST_PIS", "CST DE PIS"},
                "DEPENDENTES": {
                    "VL_BC_PIS", "BASE DE PIS", "ALIQ_PIS", "ALIQUOTA DE PIS",
                    "QUANT_BC_PIS", "BASE DE PIS - QTDE",
                    "ALIQ_PIS_QUANT", "ALIQUOTA DE PIS QTDE",
                    "VL_PIS", "VALOR DE PIS",
                },
            },
            "COFINS": {
                "CST": {"CST_COFINS", "CST DE COFINS"},
                "DEPENDENTES": {
                    "VL_BC_COFINS", "BASE DE COFINS", "ALIQ_COFINS", "ALIQUOTA DE COFINS",
                    "QUANT_BC_COFINS", "BASE DE COFINS - QTDE",
                    "ALIQ_COFINS_QUANT", "ALIQUOTA DE COFINS QTDE",
                    "VL_COFINS", "VALOR DE COFINS",
                },
            },
        }

        cst_mudado: dict[tuple[int, str], AlteracaoPlanilha] = {}
        dependente_manual: set[tuple[int, str]] = set()
        for alt in alteracoes:
            if str(alt.registro or "").strip().upper() != "C170":
                continue
            nome = normalizar_nome(alt.campo)
            for tributo in ("PIS", "COFINS"):
                if nome in aliases[tributo]["CST"]:
                    cst_mudado[(alt.sequencia, tributo)] = alt
                elif nome in aliases[tributo]["DEPENDENTES"]:
                    dependente_manual.add((alt.sequencia, tributo))

        if not cst_mudado:
            return linhas, [], []

        novas = list(linhas)
        conferencia: list[ItemConferenciaExcel] = []
        avisos: list[str] = []

        for (sequencia, tributo), alt in sorted(cst_mudado.items()):
            posicao = posicao_gerada_por_sequencia.get(sequencia)
            if posicao is None or posicao >= len(novas):
                continue
            campos = cls._separar_campos(novas[posicao])
            if not campos or campos[0].strip().upper() != "C170":
                continue
            cfop = campos[10].strip() if len(campos) > 10 else ""
            if not cfop or cfop[0] not in {"1", "2", "3"}:
                continue

            conf = cfg[tributo]
            idx_cst = conf["cst"]
            if len(campos) <= idx_cst:
                continue
            cst_novo = campos[idx_cst].strip()

            if (sequencia, tributo) in dependente_manual:
                avisos.append(
                    f"C170 sequência {sequencia}: o CST de {tributo} foi alterado para {cst_novo or '-'}, "
                    "mas a planilha também contém edição manual em base/alíquota/valor. "
                    "O FiscalPro preservou a edição manual e não zerou esses campos automaticamente."
                )
                continue

            if cst_novo in sem_credito:
                mudou = False
                for indice, (nome_campo, zero) in conf["campos"].items():
                    if len(campos) <= indice:
                        continue
                    anterior = campos[indice]
                    valor = cls._decimal_seguro(anterior)
                    if valor is None and not str(anterior or "").strip():
                        continue
                    if valor == Decimal("0"):
                        continue
                    campos[indice] = zero
                    conferencia.append(
                        ItemConferenciaExcel(
                            tipo="Ajuste automático CST",
                            categoria="Tributário",
                            sequencia=sequencia,
                            registro="C170",
                            aba=alt.aba,
                            linha_excel=alt.linha_excel,
                            campo=nome_campo,
                            original=anterior,
                            novo=zero,
                        )
                    )
                    mudou = True
                if mudou:
                    novas[posicao] = "|" + "|".join(campos) + "|"
            elif cst_novo in com_credito:
                valor_idx = 29 if tributo == "PIS" else 35
                valor = cls._decimal_seguro(campos[valor_idx]) if len(campos) > valor_idx else None
                if valor in {None, Decimal("0")}:
                    avisos.append(
                        f"C170 sequência {sequencia}: o CST de {tributo} foi alterado para {cst_novo}, "
                        "que pode gerar crédito, mas não houve edição de base/alíquota/valor. "
                        "O FiscalPro não inventou o crédito; revise os campos do item."
                    )

        return novas, conferencia, avisos

    @classmethod
    def _sincronizar_c190_icms_por_c170_alterado(
        cls,
        linhas: list[str],
        referencias: list[_ReferenciaLinha],
        alteracoes: list[AlteracaoPlanilha],
        posicao_gerada_por_sequencia: dict[int, int],
    ) -> tuple[list[str], list[ItemConferenciaExcel], int, list[str]]:
        """Reconstrói C190 quando a chave analítica do ICMS muda no C170.

        O C190 consolida os itens pela combinação CST_ICMS + CFOP + ALIQ_ICMS.
        Portanto, uma alteração desses campos no C170 pode exigir criar, fundir
        ou eliminar linhas C190. A rotina é deliberadamente conservadora:

        * atua em C100 com mudança explícita no C170 ou divergência interna C170 x C190 comprovada por totais;
        * não sobrescreve C190 que a própria usuária também editou manualmente;
        * preserva o VL_OPR total já escriturado e o distribui proporcionalmente
          pelo VL_ITEM dos novos grupos, como já faz a rotina XML/DANFE;
        * soma diretamente dos C170: base/ICMS, base/ICMS-ST e IPI;
        * se existir redução de base (VL_RED_BC) no C190 ou COD_OBS distintos,
          não inventa rateio: preserva o documento e emite aviso para revisão.
        """

        def nome_normalizado(valor: str) -> str:
            return (
                str(valor or "").strip().upper()
                .replace("Á", "A").replace("À", "A").replace("Â", "A").replace("Ã", "A")
                .replace("É", "E").replace("Ê", "E").replace("Í", "I")
                .replace("Ó", "O").replace("Ô", "O").replace("Õ", "O")
                .replace("Ú", "U").replace("Ç", "C")
            )

        campos_gatilho = {
            "CST_ICMS", "CST ICMS", "CST DE ICMS",
            "CFOP",
            "ALIQ_ICMS", "ALIQUOTA ICMS", "ALIQUOTA DE ICMS",
        }

        referencia_por_sequencia = {ref.sequencia: ref for ref in referencias}
        pai_por_c170: dict[int, int] = {}
        pai_por_c190: dict[int, int] = {}
        filhos_c170: dict[int, list[int]] = {}
        filhos_c190: dict[int, list[int]] = {}
        c100_atual: int | None = None
        codigos_filho_c100 = {
            "C101", "C105", "C110", "C111", "C112", "C113", "C114",
            "C115", "C116", "C120", "C130", "C140", "C141", "C160",
            "C165", "C170", "C171", "C172", "C173", "C174", "C175",
            "C176", "C177", "C178", "C179", "C190", "C191", "C195",
            "C197", "C199",
        }

        for ref in referencias:
            codigo = str(ref.registro or "").strip().upper()
            if codigo == "C100":
                c100_atual = ref.sequencia
                filhos_c170.setdefault(c100_atual, [])
                filhos_c190.setdefault(c100_atual, [])
            elif codigo == "C170" and c100_atual is not None:
                pai_por_c170[ref.sequencia] = c100_atual
                filhos_c170.setdefault(c100_atual, []).append(ref.sequencia)
            elif codigo == "C190" and c100_atual is not None:
                pai_por_c190[ref.sequencia] = c100_atual
                filhos_c190.setdefault(c100_atual, []).append(ref.sequencia)
            elif codigo.startswith("C") and codigo in codigos_filho_c100:
                continue
            elif codigo:
                c100_atual = None

        afetados: set[int] = set()
        afetados_por_edicao: set[int] = set()
        c190_editado_manual: set[int] = set()
        for alt in alteracoes:
            registro = str(alt.registro or "").strip().upper()
            if registro == "C170" and nome_normalizado(alt.campo) in campos_gatilho:
                pai = pai_por_c170.get(alt.sequencia)
                if pai is not None:
                    afetados.add(pai)
                    afetados_por_edicao.add(pai)
            elif registro == "C190":
                pai = pai_por_c190.get(alt.sequencia)
                if pai is not None:
                    c190_editado_manual.add(pai)

        # Hotfix 17.8.93 — a divergência pode já existir no TXT que originou
        # a planilha. Isso acontece, por exemplo, quando uma rotina anterior
        # corrigiu CST/CFOP no C170 e o arquivo foi exportado para Excel antes
        # de o C190 ser reagrupado. Nesse cenário a comparação Excel x origem
        # mostra "0 campos alterados", embora o próprio SPED esteja incoerente.
        #
        # Para não reconstruir C190 por aproximação, a detecção interna só entra
        # no automático quando: (1) as chaves CST/CFOP/alíquota divergem e
        # (2) os totais monetários comprováveis em C170 e C190 (BC, ICMS,
        # BC-ST, ICMS-ST e IPI) fecham, admitindo apenas centavos de precisão.
        avisos_divergencia: list[str] = []
        tolerancia_totais = Decimal("0.02")
        for seq_c100, seqs_c170 in filhos_c170.items():
            seqs_c190 = filhos_c190.get(seq_c100, [])
            if not seqs_c170 or not seqs_c190 or seq_c100 in c190_editado_manual:
                continue

            linhas_c170: list[list[str]] = []
            linhas_c190: list[list[str]] = []
            for seq in seqs_c170:
                pos = posicao_gerada_por_sequencia.get(seq)
                if pos is not None and 0 <= pos < len(linhas):
                    campos = cls._separar_campos(linhas[pos])
                    if campos and campos[0].upper() == "C170" and len(campos) >= 24:
                        linhas_c170.append(campos)
            for seq in seqs_c190:
                pos = posicao_gerada_por_sequencia.get(seq)
                if pos is not None and 0 <= pos < len(linhas):
                    campos = cls._separar_campos(linhas[pos])
                    if campos and campos[0].upper() == "C190" and len(campos) >= 12:
                        linhas_c190.append(campos)
            if not linhas_c170 or not linhas_c190:
                continue

            chaves_c170: set[tuple[str, str, Decimal]] = set()
            chaves_c190_lista: list[tuple[str, str, Decimal]] = []
            estrutura_ok = True
            for c in linhas_c170:
                aliq = cls._decimal_seguro(c[13])
                if not c[9].strip() or not c[10].strip() or aliq is None:
                    estrutura_ok = False
                    break
                chaves_c170.add((c[9].strip(), c[10].strip(), aliq))
            if not estrutura_ok:
                continue
            for c in linhas_c190:
                aliq = cls._decimal_seguro(c[3])
                if not c[1].strip() or not c[2].strip() or aliq is None:
                    estrutura_ok = False
                    break
                chaves_c190_lista.append((c[1].strip(), c[2].strip(), aliq))
            if not estrutura_ok:
                continue

            chaves_c190 = set(chaves_c190_lista)
            diverge_chave = (
                chaves_c170 != chaves_c190
                or len(chaves_c190_lista) != len(chaves_c190)
            )
            if not diverge_chave:
                continue

            totais_c170 = (
                sum((cls._decimal_seguro(c[12]) or Decimal("0") for c in linhas_c170), Decimal("0")),
                sum((cls._decimal_seguro(c[14]) or Decimal("0") for c in linhas_c170), Decimal("0")),
                sum((cls._decimal_seguro(c[15]) or Decimal("0") for c in linhas_c170), Decimal("0")),
                sum((cls._decimal_seguro(c[17]) or Decimal("0") for c in linhas_c170), Decimal("0")),
                sum((cls._decimal_seguro(c[23]) or Decimal("0") for c in linhas_c170), Decimal("0")),
            )
            totais_c190 = (
                sum((cls._decimal_seguro(c[5]) or Decimal("0") for c in linhas_c190), Decimal("0")),
                sum((cls._decimal_seguro(c[6]) or Decimal("0") for c in linhas_c190), Decimal("0")),
                sum((cls._decimal_seguro(c[7]) or Decimal("0") for c in linhas_c190), Decimal("0")),
                sum((cls._decimal_seguro(c[8]) or Decimal("0") for c in linhas_c190), Decimal("0")),
                sum((cls._decimal_seguro(c[10]) or Decimal("0") for c in linhas_c190), Decimal("0")),
            )
            diferencas = tuple(a - b for a, b in zip(totais_c170, totais_c190))
            if all(d.copy_abs() <= tolerancia_totais for d in diferencas):
                afetados.add(seq_c100)
                continue

            pos_c100 = posicao_gerada_por_sequencia.get(seq_c100)
            campos_c100 = (
                cls._separar_campos(linhas[pos_c100])
                if pos_c100 is not None and 0 <= pos_c100 < len(linhas)
                else []
            )
            numero_doc = campos_c100[7].strip() if len(campos_c100) > 7 else str(seq_c100)
            nomes = ("BC ICMS", "ICMS", "BC ST", "ICMS-ST", "IPI")
            detalhe = ", ".join(
                f"{nome} {cls._formatar_decimal_sped(valor, 'VL_ICMS')}"
                for nome, valor in zip(nomes, diferencas)
                if valor.copy_abs() > tolerancia_totais
            )
            avisos_divergencia.append(
                f"NF {numero_doc}: C170 e C190 usam chaves analíticas diferentes, mas os totais "
                f"também divergem ({detalhe or 'diferença monetária'}). O C190 não foi reconstruído "
                "automaticamente porque os C170 podem não representar todos os grupos da nota; revise manualmente."
            )

        if not afetados:
            return linhas, [], 0, avisos_divergencia

        substituicoes: dict[int, tuple[set[int], list[str]]] = {}
        conferencia: list[ItemConferenciaExcel] = []
        avisos: list[str] = list(avisos_divergencia)
        documentos = 0

        for seq_c100 in sorted(afetados):
            pos_c100 = posicao_gerada_por_sequencia.get(seq_c100)
            campos_c100 = (
                cls._separar_campos(linhas[pos_c100])
                if pos_c100 is not None and 0 <= pos_c100 < len(linhas)
                else []
            )
            numero_doc = campos_c100[7].strip() if len(campos_c100) > 7 else str(seq_c100)

            if seq_c100 in c190_editado_manual:
                avisos.append(
                    f"NF {numero_doc}: houve edição manual no C190 e também alteração de chave "
                    "tributária no C170. O FiscalPro preservou o C190 manual para não sobrescrever "
                    "a decisão da planilha; revise o agrupamento no PVA."
                )
                continue

            seqs_c170 = filhos_c170.get(seq_c100, [])
            seqs_c190 = filhos_c190.get(seq_c100, [])
            posicoes_c190 = [
                posicao_gerada_por_sequencia[s]
                for s in seqs_c190
                if s in posicao_gerada_por_sequencia
                and posicao_gerada_por_sequencia[s] < len(linhas)
            ]
            if not seqs_c170 or not posicoes_c190:
                avisos.append(
                    f"NF {numero_doc}: não há C170/C190 suficientes para reconstruir o resumo do ICMS com segurança."
                )
                continue

            c190_atuais = [cls._separar_campos(linhas[p]) for p in posicoes_c190]
            if any(not c or c[0].upper() != "C190" or len(c) < 12 for c in c190_atuais):
                avisos.append(f"NF {numero_doc}: estrutura C190 inesperada; reagrupamento automático não aplicado.")
                continue

            reducao_total = sum(
                (cls._decimal_seguro(c[9]) or Decimal("0") for c in c190_atuais),
                Decimal("0"),
            )

            codigos_obs = {c[11].strip() for c in c190_atuais if len(c) > 11 and c[11].strip()}
            if len(codigos_obs) > 1:
                avisos.append(
                    f"NF {numero_doc}: existem COD_OBS distintos entre os C190. O vínculo da observação com os "
                    "novos grupos não pode ser provado pelo C170; reagrupamento automático não aplicado."
                )
                continue
            cod_obs = next(iter(codigos_obs), "")

            vl_opr_original = sum(
                (cls._decimal_seguro(c[4]) or Decimal("0") for c in c190_atuais),
                Decimal("0"),
            )

            grupos: dict[tuple[str, str, Decimal], dict[str, Decimal]] = {}
            falha = ""
            for seq_c170 in seqs_c170:
                pos = posicao_gerada_por_sequencia.get(seq_c170)
                if pos is None or pos >= len(linhas):
                    continue
                c = cls._separar_campos(linhas[pos])
                if not c or c[0].upper() != "C170" or len(c) < 24:
                    falha = "há C170 com estrutura incompleta"
                    break
                cst = c[9].strip()
                cfop = c[10].strip()
                aliq = cls._decimal_seguro(c[13])
                if not cst or not cfop or aliq is None:
                    falha = "há C170 sem CST_ICMS/CFOP/ALIQ_ICMS identificável"
                    break
                chave = (cst, cfop, aliq)
                valores = grupos.setdefault(
                    chave,
                    {
                        "vl_itens": Decimal("0"),
                        "vl_opr": Decimal("0"),
                        "base": Decimal("0"),
                        "icms": Decimal("0"),
                        "base_st": Decimal("0"),
                        "icms_st": Decimal("0"),
                        "ipi": Decimal("0"),
                    },
                )
                valores["vl_itens"] += cls._decimal_seguro(c[6]) or Decimal("0")
                valores["base"] += cls._decimal_seguro(c[12]) or Decimal("0")
                valores["icms"] += cls._decimal_seguro(c[14]) or Decimal("0")
                valores["base_st"] += cls._decimal_seguro(c[15]) or Decimal("0")
                valores["icms_st"] += cls._decimal_seguro(c[17]) or Decimal("0")
                valores["ipi"] += cls._decimal_seguro(c[23]) or Decimal("0")

            if falha or not grupos:
                avisos.append(f"NF {numero_doc}: {falha or 'não foi possível formar grupos C190 a partir dos C170'}.")
                continue

            soma_itens = sum((v["vl_itens"] for v in grupos.values()), Decimal("0"))
            chaves = sorted(grupos, key=lambda x: (x[0], x[1], x[2]))
            if reducao_total.copy_abs() > Decimal("0.005") and len(chaves) > 1:
                avisos.append(
                    f"NF {numero_doc}: o C190 possui VL_RED_BC diferente de zero e o novo C170 forma "
                    "mais de um grupo analítico. Como a redução não existe por item no C170, o FiscalPro "
                    "não inventou um rateio; revise o C190 manualmente."
                )
                continue
            restante = vl_opr_original
            for ordem, chave in enumerate(chaves, start=1):
                valores = grupos[chave]
                if len(chaves) == 1:
                    parcela = vl_opr_original
                elif ordem == len(chaves):
                    parcela = restante
                elif soma_itens > 0:
                    parcela = (vl_opr_original * valores["vl_itens"] / soma_itens).quantize(
                        Decimal("0.01"), rounding=ROUND_HALF_UP
                    )
                else:
                    parcela = valores["base"].quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                valores["vl_opr"] = parcela
                restante -= parcela

            novos_c190: list[str] = []
            for cst, cfop, aliq in chaves:
                v = grupos[(cst, cfop, aliq)]
                campos = [
                    "C190",
                    cst,
                    cfop,
                    cls._formatar_decimal_sped(aliq, "ALIQ_ICMS"),
                    cls._formatar_decimal_sped(v["vl_opr"], "VL_OPR"),
                    cls._formatar_decimal_sped(v["base"], "VL_BC_ICMS"),
                    cls._formatar_decimal_sped(v["icms"], "VL_ICMS"),
                    cls._formatar_decimal_sped(v["base_st"], "VL_BC_ICMS_ST"),
                    cls._formatar_decimal_sped(v["icms_st"], "VL_ICMS_ST"),
                    cls._formatar_decimal_sped(
                        reducao_total if len(chaves) == 1 else Decimal("0"),
                        "VL_RED_BC",
                    ),
                    cls._formatar_decimal_sped(v["ipi"], "VL_IPI"),
                    cod_obs,
                ]
                novos_c190.append("|" + "|".join(campos) + "|")

            resumo_antigo = "; ".join(
                f"{c[1]}/{c[2]}/{c[3]} = {c[4]}" for c in c190_atuais
            )
            resumo_novo = "; ".join(
                f"{cst}/{cfop}/{cls._formatar_decimal_sped(aliq, 'ALIQ_ICMS')} = "
                f"{cls._formatar_decimal_sped(grupos[(cst, cfop, aliq)]['vl_opr'], 'VL_OPR')}"
                for cst, cfop, aliq in chaves
            )

            primeiro = min(posicoes_c190)
            substituicoes[primeiro] = (set(posicoes_c190), novos_c190)
            ref_primeiro = referencia_por_sequencia.get(seqs_c190[0]) if seqs_c190 else None
            conferencia.append(
                ItemConferenciaExcel(
                    tipo="Ajuste automático C190",
                    categoria="Tributário",
                    sequencia=(seqs_c190[0] if seqs_c190 else seq_c100),
                    registro="C190",
                    aba=(ref_primeiro.aba_original if ref_primeiro else ""),
                    linha_excel=(ref_primeiro.linha_excel_original if ref_primeiro else None),
                    campo="REAGRUPAMENTO CST/CFOP/ALIQ_ICMS",
                    original=resumo_antigo,
                    novo=resumo_novo,
                )
            )
            documentos += 1

        if not substituicoes:
            return linhas, conferencia, documentos, avisos

        remover: set[int] = set()
        for indices, _ in substituicoes.values():
            remover.update(indices)

        novas: list[str] = []
        for indice, linha in enumerate(linhas):
            if indice in substituicoes:
                novas.extend(substituicoes[indice][1])
            if indice in remover:
                continue
            novas.append(linha)

        return novas, conferencia, documentos, avisos

    @classmethod
    def _sincronizar_totais_c100_piscofins(
        cls,
        linhas: list[str],
        referencias: list[_ReferenciaLinha],
        alteracoes: list[AlteracaoPlanilha],
        posicao_gerada_por_sequencia: dict[int, int],
    ) -> tuple[list[str], list[ItemConferenciaExcel], int]:
        """Retotaliza VL_PIS/VL_COFINS somente nos C100 afetados no Excel.

        A planilha pode alterar CST, base, alíquota ou valor nos C170. Depois
        dessas mudanças, o C100 precisa refletir a soma dos VL_PIS/VL_COFINS
        efetivamente reconstruídos nos filhos. O método não recalcula o valor
        do item: ele respeita o que está na planilha e sincroniza apenas o total
        documental correspondente.
        """
        campos_pis = {
            "CST_PIS", "VL_BC_PIS", "ALIQ_PIS", "QUANT_BC_PIS",
            "ALIQ_PIS_QUANT", "VL_PIS",
            "CST DE PIS", "BASE DE PIS", "ALIQUOTA DE PIS",
            "BASE DE PIS - QTDE", "ALIQUOTA DE PIS QTDE", "VALOR DE PIS",
        }
        campos_cofins = {
            "CST_COFINS", "VL_BC_COFINS", "ALIQ_COFINS",
            "QUANT_BC_COFINS", "ALIQ_COFINS_QUANT", "VL_COFINS",
            "CST DE COFINS", "BASE DE COFINS", "ALIQUOTA DE COFINS",
            "BASE DE COFINS - QTDE", "ALIQUOTA DE COFINS QTDE", "VALOR DE COFINS",
        }

        # Descobre o C100 pai de cada C170 pela estrutura original, sem depender
        # do número físico da linha depois de exclusões opcionais.
        pai_por_c170: dict[int, int] = {}
        filhos_por_c100: dict[int, list[int]] = {}
        c100_atual: int | None = None
        codigos_filho_c100 = {
            "C101", "C105", "C110", "C111", "C112", "C113", "C114",
            "C115", "C116", "C120", "C130", "C140", "C141", "C160",
            "C165", "C170", "C171", "C172", "C173", "C174", "C175",
            "C176", "C177", "C178", "C179", "C190", "C191", "C195",
            "C197", "C199",
        }
        referencia_por_sequencia = {ref.sequencia: ref for ref in referencias}
        for ref in referencias:
            codigo = str(ref.registro or "").strip().upper()
            if codigo == "C100":
                c100_atual = ref.sequencia
                filhos_por_c100.setdefault(c100_atual, [])
            elif codigo == "C170" and c100_atual is not None:
                pai_por_c170[ref.sequencia] = c100_atual
                filhos_por_c100.setdefault(c100_atual, []).append(ref.sequencia)
            elif codigo.startswith("C") and codigo in codigos_filho_c100:
                continue
            elif codigo:
                c100_atual = None

        afetados: dict[int, set[str]] = {}
        for alteracao in alteracoes:
            if str(alteracao.registro or "").upper() != "C170":
                continue
            campo = str(alteracao.campo or "").strip().upper()
            # Os cabeçalhos visíveis do Excel usam nomes amigáveis (ex.:
            # "Alíquota de Pis"). Normaliza acentos para aceitar tanto o nome
            # técnico quanto o nome mostrado na planilha.
            campo = (
                campo.replace("Á", "A").replace("À", "A").replace("Â", "A").replace("Ã", "A")
                .replace("É", "E").replace("Ê", "E").replace("Í", "I")
                .replace("Ó", "O").replace("Ô", "O").replace("Õ", "O")
                .replace("Ú", "U").replace("Ç", "C")
            )
            tributos: set[str] = set()
            if campo in campos_pis:
                tributos.add("PIS")
            if campo in campos_cofins:
                tributos.add("COFINS")
            if not tributos:
                continue
            pai = pai_por_c170.get(alteracao.sequencia)
            if pai is not None:
                afetados.setdefault(pai, set()).update(tributos)

        if not afetados:
            return linhas, [], 0

        novas = list(linhas)
        conferencia: list[ItemConferenciaExcel] = []
        documentos_retotalizados = 0
        indices_valor_c170 = {"PIS": 29, "COFINS": 35}
        indices_total_c100 = {"PIS": 25, "COFINS": 26}

        for sequencia_c100, tributos in sorted(afetados.items()):
            posicao_c100 = posicao_gerada_por_sequencia.get(sequencia_c100)
            if posicao_c100 is None or posicao_c100 >= len(novas):
                continue
            campos_c100 = cls._separar_campos(novas[posicao_c100])
            if not campos_c100 or campos_c100[0].upper() != "C100":
                continue

            houve_ajuste_documento = False
            for tributo in sorted(tributos):
                soma = Decimal("0")
                for sequencia_c170 in filhos_por_c100.get(sequencia_c100, []):
                    posicao_c170 = posicao_gerada_por_sequencia.get(sequencia_c170)
                    if posicao_c170 is None or posicao_c170 >= len(novas):
                        continue
                    campos_c170 = cls._separar_campos(novas[posicao_c170])
                    indice_valor = indices_valor_c170[tributo]
                    if len(campos_c170) <= indice_valor:
                        continue
                    valor = cls._decimal_seguro(campos_c170[indice_valor])
                    if valor is not None:
                        soma += valor
                soma = soma.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

                indice_total = indices_total_c100[tributo]
                if len(campos_c100) <= indice_total:
                    continue
                anterior = campos_c100[indice_total]
                anterior_decimal = cls._decimal_seguro(anterior)
                if anterior_decimal is not None:
                    anterior_decimal = anterior_decimal.quantize(
                        Decimal("0.01"), rounding=ROUND_HALF_UP
                    )
                if anterior_decimal == soma:
                    continue

                novo = format(soma, ".2f").replace(".", ",")
                campos_c100[indice_total] = novo
                ref_c100 = referencia_por_sequencia.get(sequencia_c100)
                conferencia.append(
                    ItemConferenciaExcel(
                        tipo="Ajuste automático",
                        categoria="Tributário",
                        sequencia=sequencia_c100,
                        registro="C100",
                        aba=(ref_c100.aba_original if ref_c100 else ""),
                        linha_excel=(ref_c100.linha_excel_original if ref_c100 else None),
                        campo=f"VL_{tributo}",
                        original=anterior,
                        novo=novo,
                    )
                )
                houve_ajuste_documento = True

            if houve_ajuste_documento:
                novas[posicao_c100] = "|" + "|".join(campos_c100) + "|"
                documentos_retotalizados += 1

        return novas, conferencia, documentos_retotalizados

    @classmethod
    def _item_conferencia_alteracao(
        cls,
        alteracao: AlteracaoPlanilha,
    ) -> ItemConferenciaExcel:
        return ItemConferenciaExcel(
            tipo="Alteração",
            categoria=cls.classificar_categoria_campo(alteracao.campo),
            sequencia=alteracao.sequencia,
            registro=alteracao.registro,
            aba=alteracao.aba,
            linha_excel=alteracao.linha_excel,
            campo=alteracao.campo,
            original=alteracao.original,
            novo=alteracao.novo,
        )

    @staticmethod
    def classificar_categoria_campo(campo: str) -> str:
        """Classifica o campo para facilitar a revisão na tela da 17.4.2."""
        nome = str(campo or "").strip().upper()
        if not nome:
            return "Outro"

        identificacao_exata = {
            "REG", "COD_PART", "COD_ITEM", "NUM_ITEM", "NUM_DOC", "SER",
            "SUB", "CHV_NFE", "CHV_CTE", "COD_MOD", "COD_CTA", "COD_CCUS",
            "COD_DOC", "COD_INF", "COD_OBS", "COD_AJ", "COD_LAN", "COD_VER",
        }
        if nome in identificacao_exata or nome.startswith(("COD_PART", "COD_ITEM", "CHV_")):
            return "Identificação"

        palavras_tributarias = (
            "CST", "CFOP", "NCM", "CEST", "ALIQ", "PIS", "COFINS", "ICMS",
            "IPI", "ISS", "FCP", "DIFAL", "MVA", "NAT_BC", "IND_ORIG_CRED",
            "COD_CRED", "COD_INC", "COD_TRIB", "COD_ENQ", "COD_REC",
        )
        if any(palavra in nome for palavra in palavras_tributarias):
            return "Tributário"

        if nome.startswith(("VL_", "QTD", "QTDE", "VALOR", "BASE")) or nome.endswith("_VL"):
            return "Valor"

        palavras_cadastrais = (
            "DESCR", "UNID", "UNIDADE", "TIPO_ITEM", "EX_IPI", "COD_GEN",
            "COD_LST", "DT_INI", "DT_FIM", "NOME", "END", "MUN", "SUFRAMA",
        )
        if any(palavra in nome for palavra in palavras_cadastrais):
            return "Cadastral"

        return "Outro"

    @classmethod
    def obter_conferencia(
        cls,
        validacao: ResultadoValidacaoExcel,
    ) -> list[ItemConferenciaExcel]:
        """Retorna a conferência da validação, com compatibilidade defensiva."""
        if validacao.conferencia:
            return list(validacao.conferencia)
        return [cls._item_conferencia_alteracao(item) for item in validacao.alteracoes]

    @staticmethod
    def filtrar_conferencia(
        itens: list[ItemConferenciaExcel],
        registro: str = "",
        tipo: str = "",
        categoria: str = "",
        busca: str = "",
    ) -> list[ItemConferenciaExcel]:
        registro_norm = str(registro or "").strip().casefold()
        tipo_norm = str(tipo or "").strip().casefold()
        categoria_norm = str(categoria or "").strip().casefold()
        busca_norm = str(busca or "").strip().casefold()

        if registro_norm in {"todos", "todas"}:
            registro_norm = ""
        if tipo_norm in {"todos", "todas"}:
            tipo_norm = ""
        if categoria_norm in {"todos", "todas"}:
            categoria_norm = ""

        filtrados: list[ItemConferenciaExcel] = []
        for item in itens:
            if registro_norm and item.registro.casefold() != registro_norm:
                continue
            if tipo_norm and item.tipo.casefold() != tipo_norm:
                continue
            if categoria_norm and item.categoria.casefold() != categoria_norm:
                continue
            if busca_norm:
                texto = " | ".join(
                    [
                        item.registro,
                        item.aba,
                        item.campo,
                        item.original,
                        item.novo,
                        item.tipo,
                        item.categoria,
                    ]
                ).casefold()
                if busca_norm not in texto:
                    continue
            filtrados.append(item)
        return filtrados

    @classmethod
    def salvar_relatorio_conferencia(
        cls,
        validacao: ResultadoValidacaoExcel,
        caminho_saida: str | Path,
        itens: list[ItemConferenciaExcel] | None = None,
    ) -> Path:
        destino = Path(caminho_saida)
        if destino.suffix.lower() != ".txt":
            destino = destino.with_suffix(".txt")
        destino.parent.mkdir(parents=True, exist_ok=True)

        conferencia = list(itens) if itens is not None else cls.obter_conferencia(validacao)
        tipos = Counter(item.tipo for item in conferencia)
        categorias = Counter(item.categoria for item in conferencia)

        linhas = [
            "FISCALPRO — CONFERÊNCIA DE ALTERAÇÕES EXCEL → SPED",
            "=" * 68,
            f"Gerado em: {datetime.now():%d/%m/%Y %H:%M:%S}",
            f"Planilha: {validacao.caminho}",
            f"Arquivo de origem: {validacao.arquivo_origem or '-'}",
            f"Alterações exibidas: {len(conferencia)}",
            (
                "Tipos: "
                + (", ".join(f"{nome}: {qtd}" for nome, qtd in sorted(tipos.items())) or "-")
            ),
            (
                "Categorias: "
                + (", ".join(f"{nome}: {qtd}" for nome, qtd in sorted(categorias.items())) or "-")
            ),
            f"Pré-PVA — novos erros pela planilha: {validacao.pre_pva_novos_erros}",
            f"Pré-PVA — novos avisos pela planilha: {validacao.pre_pva_novos_avisos}",
            "",
            "ALTERAÇÕES CONFERIDAS",
            "-" * 68,
        ]

        if not conferencia:
            linhas.append("Nenhuma alteração foi encontrada entre o SPED original e a planilha.")
        else:
            for indice, item in enumerate(conferencia, start=1):
                linhas.extend(
                    [
                        (
                            f"{indice}. {item.tipo} | {item.categoria} | "
                            f"Registro {item.registro or '-'} | Sequência {item.sequencia or '-'}"
                        ),
                        f"   Aba/Linha Excel: {item.aba or '-'} / {item.linha_excel or '-'}",
                        f"   Campo: {item.campo or '-'}",
                        f"   Original: {item.original}",
                        f"   Novo: {item.novo}",
                        "",
                    ]
                )

        linhas.extend(
            [
                "SEGURANÇA",
                "-" * 68,
                "Esta conferência é somente informativa e não aplica alterações automaticamente.",
                "Somente novos erros criados pela planilha bloqueiam o TXT no Pré-PVA comparativo.",
                "A validação oficial no PVA/PGE continua obrigatória antes da transmissão.",
                "",
            ]
        )
        destino.write_text("\n".join(linhas), encoding="utf-8")
        return destino

    def _ler_metadados(self, ws) -> dict[str, str]:
        metadados: dict[str, str] = {}
        for linha in range(1, self.LINHA_CABECALHO_TECNICO):
            chave = ws.cell(linha, 1).value
            valor = ws.cell(linha, 2).value
            if chave not in (None, ""):
                metadados[str(chave)] = "" if valor is None else str(valor)
        return metadados

    def _ler_referencias(
        self,
        ws,
        resultado: ResultadoValidacaoExcel,
    ) -> list[_ReferenciaLinha]:
        cabecalhos = {
            str(ws.cell(self.LINHA_CABECALHO_TECNICO, coluna).value): coluna
            for coluna in range(1, ws.max_column + 1)
            if ws.cell(self.LINHA_CABECALHO_TECNICO, coluna).value not in (None, "")
        }
        obrigatorios = {
            "SEQUENCIA",
            "TIPO",
            "FP_ID",
            "REGISTRO",
            "ABA_ORIGINAL",
            "LINHA_EXCEL",
            "QTD_CONTEXTO",
            "QTD_CAMPOS",
            "LINHA_ORIGINAL",
        }
        faltantes = obrigatorios - set(cabecalhos)
        if faltantes:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="ERRO",
                    mensagem="Estrutura técnica incompleta: " + ", ".join(sorted(faltantes)),
                )
            )
            return []

        referencias: list[_ReferenciaLinha] = []
        sequencias: set[int] = set()
        for linha in range(self.LINHA_CABECALHO_TECNICO + 1, ws.max_row + 1):
            sequencia = self._inteiro_seguro(ws.cell(linha, cabecalhos["SEQUENCIA"]).value)
            if not sequencia:
                continue
            if sequencia in sequencias:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem=f"Sequência técnica duplicada: {sequencia}.",
                        sequencia=sequencia,
                    )
                )
                continue
            sequencias.add(sequencia)

            referencias.append(
                _ReferenciaLinha(
                    sequencia=sequencia,
                    tipo=str(ws.cell(linha, cabecalhos["TIPO"]).value or ""),
                    id_tecnico=str(ws.cell(linha, cabecalhos["FP_ID"]).value or ""),
                    registro=str(ws.cell(linha, cabecalhos["REGISTRO"]).value or ""),
                    aba_original=str(ws.cell(linha, cabecalhos["ABA_ORIGINAL"]).value or ""),
                    linha_excel_original=self._inteiro_seguro(
                        ws.cell(linha, cabecalhos["LINHA_EXCEL"]).value
                    ),
                    qtd_contexto=self._inteiro_seguro(
                        ws.cell(linha, cabecalhos["QTD_CONTEXTO"]).value
                    ),
                    qtd_campos=self._inteiro_seguro(
                        ws.cell(linha, cabecalhos["QTD_CAMPOS"]).value
                    ),
                    linha_original=str(
                        ws.cell(linha, cabecalhos["LINHA_ORIGINAL"]).value or ""
                    ),
                )
            )

        referencias.sort(key=lambda r: r.sequencia)
        if referencias:
            esperado = list(range(1, len(referencias) + 1))
            obtido = [r.sequencia for r in referencias]
            if obtido != esperado:
                resultado.problemas.append(
                    ProblemaPlanilha(
                        nivel="ERRO",
                        mensagem="A ordem técnica das linhas foi alterada ou está incompleta.",
                    )
                )
        return referencias

    @staticmethod
    def _assinatura_apontamento_pva(item) -> tuple[str, ...]:
        """Chave estável para comparar causas antes/depois sem depender da linha."""
        return (
            str(getattr(item, "nivel", "") or ""),
            str(getattr(item, "categoria", "") or ""),
            str(getattr(item, "registro", "") or ""),
            str(getattr(item, "campo", "") or ""),
            str(getattr(item, "documento", "") or ""),
            str(getattr(item, "codigo_item", "") or ""),
            str(getattr(item, "mensagem", "") or ""),
        )

    @classmethod
    def _apontamentos_novos(cls, antes, depois):
        saldo = Counter(cls._assinatura_apontamento_pva(item) for item in antes)
        novos = []
        for item in depois:
            chave = cls._assinatura_apontamento_pva(item)
            if saldo[chave] > 0:
                saldo[chave] -= 1
            else:
                novos.append(item)
        return novos

    @staticmethod
    def _identificar_tipo_sped_linhas(linhas: list[str]) -> str:
        indice = IndiceSPED().construir(linhas)
        return CalculadorEstatisticasSPED().calcular(linhas, indice).tipo_sped

    def _pre_validar_reconstrucao(
        self,
        resultado: ResultadoValidacaoExcel,
        linhas_originais: list[str],
        progresso: ProgressoCallback | None = None,
    ) -> None:
        """Compara o Pré-PVA do SPED original com o TXT reconstruído em memória.

        A conversão Excel → TXT não deve transformar um erro que já existia no
        arquivo de origem em bloqueio artificial. Por isso o FiscalPro trabalha
        por diferença: bloqueia somente causas novas introduzidas pela planilha.
        """
        if not resultado.linhas_geradas or not linhas_originais:
            return

        self._progresso(progresso, 92, "Executando Pré-PVA no TXT reconstruído...")
        tipo_sped = self._identificar_tipo_sped_linhas(linhas_originais)
        if tipo_sped not in {"EFD Contribuições", "EFD ICMS/IPI (Fiscal)"}:
            tipo_sped = self._identificar_tipo_sped_linhas(resultado.linhas_geradas)

        resultado.pre_pva_tipo_sped = tipo_sped
        if tipo_sped not in {"EFD Contribuições", "EFD ICMS/IPI (Fiscal)"}:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="AVISO",
                    mensagem=(
                        "O tipo do SPED não pôde ser identificado com segurança; "
                        "a validação estrutural do Excel foi concluída, mas o Pré-PVA "
                        "comparativo não foi executado."
                    ),
                )
            )
            return

        validador = PreValidadorPVA()
        antes = validador.validar(linhas_originais, tipo_sped)
        depois = PreValidadorPVA().validar(resultado.linhas_geradas, tipo_sped)

        novos_erros = self._apontamentos_novos(antes.erros, depois.erros)
        novos_avisos = self._apontamentos_novos(antes.avisos, depois.avisos)

        resultado.pre_pva_executado = True
        resultado.pre_pva_erros_origem = len(antes.erros)
        resultado.pre_pva_erros_reconstruido = len(depois.erros)
        resultado.pre_pva_novos_erros = len(novos_erros)
        resultado.pre_pva_avisos_origem = len(antes.avisos)
        resultado.pre_pva_avisos_reconstruido = len(depois.avisos)
        resultado.pre_pva_novos_avisos = len(novos_avisos)

        for item in novos_erros:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="ERRO",
                    mensagem=(
                        f"Pré-PVA criou nova pendência [{item.categoria}]: {item.mensagem}"
                    ),
                    sequencia=item.numero_linha,
                    registro=item.registro,
                )
            )

        for item in novos_avisos:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="AVISO",
                    mensagem=(
                        f"Pré-PVA encontrou novo aviso [{item.categoria}]: {item.mensagem}"
                    ),
                    sequencia=item.numero_linha,
                    registro=item.registro,
                )
            )

        if antes.erros and not novos_erros:
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="AVISO",
                    mensagem=(
                        f"O SPED de origem já possuía {len(antes.erros)} erro(s) nas regras "
                        "locais do Pré-PVA. A planilha não criou novos erros; essas pendências "
                        "originais continuam para revisão no fluxo normal do FiscalPro/PVA."
                    ),
                )
            )

        self._progresso(
            progresso,
            97,
            (
                f"Pré-PVA comparativo: {len(depois.erros)} erro(s) no reconstruído, "
                f"{len(novos_erros)} novo(s)."
            ),
        )

    def _validar_estrutura_basica(self, resultado: ResultadoValidacaoExcel) -> None:
        registros = []
        for linha in resultado.linhas_geradas:
            campos = self._separar_campos(linha)
            if campos:
                registros.append(campos[0].upper())

        if not registros:
            resultado.problemas.append(
                ProblemaPlanilha(nivel="ERRO", mensagem="Nenhum registro SPED foi reconstruído.")
            )
            return
        if registros[0] != "0000":
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="ERRO",
                    mensagem=f"O primeiro registro deveria ser 0000, mas foi encontrado {registros[0]}.",
                )
            )
        if registros[-1] != "9999":
            resultado.problemas.append(
                ProblemaPlanilha(
                    nivel="ERRO",
                    mensagem=f"O último registro deveria ser 9999, mas foi encontrado {registros[-1]}.",
                )
            )

    @classmethod
    def _serializar_valor(cls, valor, cabecalho: str) -> str:
        if valor is None:
            return ""
        if isinstance(valor, datetime):
            return valor.strftime("%d%m%Y")
        if isinstance(valor, date):
            return valor.strftime("%d%m%Y")
        if isinstance(valor, bool):
            return "1" if valor else "0"
        if isinstance(valor, (int, float, Decimal)):
            try:
                numero = Decimal(str(valor))
            except InvalidOperation as erro:
                raise ValueError("valor numérico inválido") from erro
            return cls._formatar_decimal_sped(numero, cabecalho)

        texto = str(valor)
        if "|" in texto:
            raise ValueError("o caractere | não pode ser usado dentro de um campo SPED")
        if "\r" in texto or "\n" in texto:
            raise ValueError("quebras de linha não podem ser usadas dentro de um campo SPED")

        texto = texto.strip()
        if ExportadorSPEDExcel._cabecalho_numerico(cabecalho):
            numero_texto = texto.replace(" ", "")
            if re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", numero_texto):
                try:
                    numero = Decimal(numero_texto.replace(",", "."))
                except InvalidOperation as erro:
                    raise ValueError("valor numérico inválido") from erro
                return cls._formatar_decimal_sped(numero, cabecalho)
        return texto

    @classmethod
    def _formatar_decimal_sped(cls, numero: Decimal, cabecalho: str) -> str:
        """Serializa a célula usando a precisão do campo SPED.

        Hotfix 17.7.7: a quantidade de casas não é mais inferida por uma
        regra genérica de ``VL_/ALIQ``. Usamos a mesma precisão técnica do
        exportador. Assim, ao editar 0,65 no Excel, ALIQ_PIS volta como
        ``0,6500``; uma QTD igual a 1 volta como ``1,00000``; e valores
        monetários continuam com duas casas.
        """

        casas = ExportadorSPEDExcel._casas_decimais(cabecalho)
        if casas is not None:
            escala = Decimal("1").scaleb(-casas)
            arredondado = numero.quantize(escala, rounding=ROUND_HALF_UP)
            return format(arredondado, f".{casas}f").replace(".", ",")

        texto = format(numero, "f")
        texto = texto.rstrip("0").rstrip(".") if "." in texto else texto
        return texto.replace(".", ",")

    @staticmethod
    def _valores_equivalentes(atual, esperado) -> bool:
        if atual in (None, "") and esperado in (None, ""):
            return True
        if isinstance(atual, bool) or isinstance(esperado, bool):
            return atual == esperado

        numero_atual = ImportadorExcelSPED._decimal_seguro(atual)
        numero_esperado = ImportadorExcelSPED._decimal_seguro(esperado)
        if numero_atual is not None and numero_esperado is not None:
            return numero_atual == numero_esperado
        return str(atual) == str(esperado)

    @staticmethod
    def _decimal_seguro(valor) -> Decimal | None:
        if isinstance(valor, bool) or valor in (None, ""):
            return None
        if isinstance(valor, (int, float, Decimal)):
            try:
                return Decimal(str(valor))
            except InvalidOperation:
                return None
        if isinstance(valor, str):
            texto = valor.strip().replace(" ", "").replace(",", ".")
            if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", texto):
                return None
            try:
                return Decimal(texto)
            except InvalidOperation:
                return None
        return None

    @classmethod
    def _normalizar_c175_aliquota_por_quantidade(
        cls, linhas: list[str]
    ) -> tuple[list[str], list[ItemConferenciaExcel], list[str]]:
        """Remove alíquota por quantidade órfã quando o C175 é ad valorem.

        No leiaute do C175 há duas modalidades alternativas de cálculo:
        - ad valorem: VL_BC + ALIQ percentual;
        - por unidade: QUANT_BC + ALIQ_*_QUANT (R$/unidade).

        O caso seguro para correção automática é ALIQ_*_QUANT preenchida sem
        QUANT_BC, enquanto VL_BC e ALIQ percentuais estão informadas. Nessa
        situação a alíquota específica não possui base em quantidade e não
        participa do cálculo válido do registro, portanto é removida sem
        alterar base, alíquota percentual ou valor da contribuição.
        """
        novas = list(linhas)
        ajustes: list[ItemConferenciaExcel] = []
        avisos: list[str] = []
        modalidades = (
            ("PIS", 5, 6, 7, 8),
            ("COFINS", 11, 12, 13, 14),
        )

        for numero_linha, linha in enumerate(novas, start=1):
            campos = cls._separar_campos(linha)
            if not campos or str(campos[0]).strip().upper() != "C175":
                continue
            alterou_linha = False

            for tributo, idx_base, idx_aliq, idx_qtd, idx_aliq_qtd in modalidades:
                if max(idx_base, idx_aliq, idx_qtd, idx_aliq_qtd) >= len(campos):
                    continue
                base = str(campos[idx_base] or "").strip()
                aliquota = str(campos[idx_aliq] or "").strip()
                quantidade = str(campos[idx_qtd] or "").strip()
                aliquota_quantidade = str(campos[idx_aliq_qtd] or "").strip()

                if aliquota_quantidade and not quantidade and base and aliquota:
                    campos[idx_aliq_qtd] = ""
                    alterou_linha = True
                    ajustes.append(
                        ItemConferenciaExcel(
                            tipo="Ajuste automático C175",
                            categoria="Estrutural",
                            sequencia=numero_linha,
                            registro="C175",
                            aba="C175",
                            linha_excel=None,
                            campo=f"ALIQ_{tributo}_QUANT",
                            original=aliquota_quantidade,
                            novo="",
                        )
                    )
                elif aliquota_quantidade and not quantidade and not (base and aliquota):
                    avisos.append(
                        f"C175 linha {numero_linha}: ALIQ_{tributo}_QUANT está preenchida "
                        f"({aliquota_quantidade}) sem QUANT_BC_{tributo}. O FiscalPro não "
                        "apagou porque também não há uma modalidade ad valorem completa; "
                        "revise esse registro."
                    )

            if alterou_linha:
                novas[numero_linha - 1] = "|" + "|".join(campos) + "|"

        return novas, ajustes, avisos

    @classmethod
    def _sincronizar_c175_bases_piscofins(
        cls,
        linhas: list[str],
        alteracoes: list[AlteracaoPlanilha],
        posicao_gerada_por_sequencia: dict[int, int],
    ) -> tuple[list[str], list[ItemConferenciaExcel], list[str]]:
        """Sincroniza bases PIS/COFINS do C175 quando a intenção é inequívoca.

        O PGE valida o C175 de forma pareada: para a mesma tributação ad valorem,
        CST e base de PIS/COFINS precisam permanecer coerentes. Um caso recorrente
        no fluxo Excel é a usuária corrigir a exclusão do ICMS em apenas uma das
        bases. Como a outra continua com o valor bruto, o TXT fica com pares como
        5,00 x 4,10 e o PGE rejeita o registro.

        A correção automática é deliberadamente conservadora:
        * atua somente em C175 ad valorem (sem modalidade por quantidade);
        * exige CST de PIS e COFINS iguais e alíquotas percentuais válidas;
        * se exatamente uma base foi alterada no Excel, essa base prevalece;
        * se nenhuma base foi editada, só repara o padrão inequívoco herdado em
          que um lado ainda é igual ao VL_OPR e o outro já contém base reduzida;
        * recalcula o valor do tributo espelhado apenas quando esse valor não foi
          editado manualmente;
        * se as duas bases foram editadas e ficaram diferentes, não escolhe uma
          delas: preserva a planilha e deixa a divergência para revisão/Pré-PVA.
        """

        def nome_normalizado(valor: str) -> str:
            return (
                str(valor or "").strip().upper()
                .replace("Á", "A").replace("À", "A").replace("Â", "A").replace("Ã", "A")
                .replace("É", "E").replace("Ê", "E").replace("Í", "I")
                .replace("Ó", "O").replace("Ô", "O").replace("Õ", "O")
                .replace("Ú", "U").replace("Ç", "C")
            )

        aliases_base = {
            "PIS": {"VL_BC_PIS", "BASE DE PIS", "BASE PIS"},
            "COFINS": {"VL_BC_COFINS", "BASE DE COFINS", "BASE COFINS"},
        }
        aliases_valor = {
            "PIS": {"VL_PIS", "VALOR DE PIS", "VALOR PIS"},
            "COFINS": {"VL_COFINS", "VALOR DE COFINS", "VALOR COFINS"},
        }

        base_manual: dict[tuple[int, str], AlteracaoPlanilha] = {}
        valor_manual: set[tuple[int, str]] = set()
        for alt in alteracoes:
            if str(alt.registro or "").strip().upper() != "C175":
                continue
            campo = nome_normalizado(alt.campo)
            for tributo in ("PIS", "COFINS"):
                if campo in aliases_base[tributo]:
                    base_manual[(alt.sequencia, tributo)] = alt
                elif campo in aliases_valor[tributo]:
                    valor_manual.add((alt.sequencia, tributo))

        # Varre todos os C175 reconstruídos. Além das edições explícitas de
        # base, a 17.8.123 também repara uma divergência herdada de versões
        # anteriores: uma contribuição ficava com VL_BC = VL_OPR (base bruta)
        # e a outra já trazia a base líquida após exclusão do ICMS.
        sequencias = sorted(posicao_gerada_por_sequencia)

        novas = list(linhas)
        ajustes: list[ItemConferenciaExcel] = []
        avisos: list[str] = []

        # Índices reais do C175 sem os pipes externos.
        idx = {
            "CST_PIS": 4,
            "VL_BC_PIS": 5,
            "ALIQ_PIS": 6,
            "QUANT_BC_PIS": 7,
            "ALIQ_PIS_QUANT": 8,
            "VL_PIS": 9,
            "CST_COFINS": 10,
            "VL_BC_COFINS": 11,
            "ALIQ_COFINS": 12,
            "QUANT_BC_COFINS": 13,
            "ALIQ_COFINS_QUANT": 14,
            "VL_COFINS": 15,
        }

        for sequencia in sequencias:
            posicao = posicao_gerada_por_sequencia.get(sequencia)
            if posicao is None or posicao >= len(novas):
                continue
            campos = cls._separar_campos(novas[posicao])
            if not campos or campos[0].strip().upper() != "C175":
                continue
            if len(campos) <= idx["VL_COFINS"]:
                continue

            base_pis = cls._decimal_seguro(campos[idx["VL_BC_PIS"]])
            base_cofins = cls._decimal_seguro(campos[idx["VL_BC_COFINS"]])
            if base_pis is None or base_cofins is None or base_pis == base_cofins:
                continue

            cst_pis = str(campos[idx["CST_PIS"]] or "").strip()
            cst_cofins = str(campos[idx["CST_COFINS"]] or "").strip()
            if not cst_pis or not cst_cofins or cst_pis != cst_cofins:
                avisos.append(
                    f"C175 sequência {sequencia}: as bases diferem, mas PIS e COFINS não têm "
                    f"o mesmo CST ({cst_pis or '-'} x {cst_cofins or '-'}). "
                    "A base não foi sincronizada automaticamente."
                )
                continue

            # Modalidade por quantidade não pode ser misturada com a lógica
            # ad valorem abaixo.
            usa_quantidade = any(
                str(campos[i] or "").strip()
                for i in (
                    idx["QUANT_BC_PIS"], idx["ALIQ_PIS_QUANT"],
                    idx["QUANT_BC_COFINS"], idx["ALIQ_COFINS_QUANT"],
                )
            )
            if usa_quantidade:
                avisos.append(
                    f"C175 sequência {sequencia}: as bases diferem e existe cálculo por quantidade; "
                    "o FiscalPro preservou o registro sem espelhamento automático."
                )
                continue

            aliq_pis = cls._decimal_seguro(campos[idx["ALIQ_PIS"]])
            aliq_cofins = cls._decimal_seguro(campos[idx["ALIQ_COFINS"]])
            if aliq_pis is None or aliq_cofins is None:
                avisos.append(
                    f"C175 sequência {sequencia}: as bases diferem, mas falta alíquota percentual "
                    "para recalcular com segurança. O FiscalPro não alterou o registro."
                )
                continue

            pis_manual = (sequencia, "PIS") in base_manual
            cofins_manual = (sequencia, "COFINS") in base_manual
            if pis_manual and cofins_manual:
                avisos.append(
                    f"C175 sequência {sequencia}: VL_BC_PIS e VL_BC_COFINS foram editadas "
                    "manualmente e ficaram diferentes. O FiscalPro preservou as duas decisões "
                    "e não escolheu uma base automaticamente."
                )
                continue

            aba_conferencia = "C175"
            linha_excel_conferencia = None
            if pis_manual ^ cofins_manual:
                fonte = "PIS" if pis_manual else "COFINS"
                destino = "COFINS" if pis_manual else "PIS"
                alt_fonte = base_manual[(sequencia, fonte)]
                aba_conferencia = alt_fonte.aba
                linha_excel_conferencia = alt_fonte.linha_excel
            else:
                # Reparo seguro de divergência herdada: um lado ainda está na
                # base bruta (igual ao VL_OPR) e o outro já está reduzido.
                # Nesse desenho, a base reduzida é a única que contém a
                # correção tributária; espelhá-la elimina o falso par
                # 5,00 x 4,10 sem inventar percentual de exclusão.
                vl_opr = cls._decimal_seguro(campos[2]) if len(campos) > 2 else None
                if vl_opr is None:
                    continue
                tol = Decimal("0.01")
                pis_bruta = abs(base_pis - vl_opr) <= tol
                cofins_bruta = abs(base_cofins - vl_opr) <= tol
                if pis_bruta and base_cofins < base_pis:
                    fonte, destino = "COFINS", "PIS"
                elif cofins_bruta and base_pis < base_cofins:
                    fonte, destino = "PIS", "COFINS"
                else:
                    continue

            idx_base_fonte = idx[f"VL_BC_{fonte}"]
            idx_base_destino = idx[f"VL_BC_{destino}"]
            idx_aliq_destino = idx[f"ALIQ_{destino}"]
            idx_valor_destino = idx[f"VL_{destino}"]

            base_fonte = cls._decimal_seguro(campos[idx_base_fonte])
            aliq_destino = cls._decimal_seguro(campos[idx_aliq_destino])
            if base_fonte is None or aliq_destino is None:
                continue

            base_anterior = campos[idx_base_destino]
            base_nova = cls._formatar_decimal_sped(base_fonte, f"VL_BC_{destino}")
            campos[idx_base_destino] = base_nova
            ajustes.append(
                ItemConferenciaExcel(
                    tipo="Sincronização automática C175",
                    categoria="Tributário",
                    sequencia=sequencia,
                    registro="C175",
                    aba=aba_conferencia,
                    linha_excel=linha_excel_conferencia,
                    campo=f"VL_BC_{destino}",
                    original=base_anterior,
                    novo=base_nova,
                )
            )

            # Se o valor do lado espelhado também foi editado manualmente, a
            # planilha prevalece. A base é sincronizada porque essa intenção é
            # inequívoca, mas o valor fica explícito para revisão.
            if (sequencia, destino) in valor_manual:
                avisos.append(
                    f"C175 sequência {sequencia}: a base de {destino} foi sincronizada com {fonte}, "
                    f"mas VL_{destino} também foi editado manualmente; o valor foi preservado para revisão."
                )
            else:
                valor_anterior = campos[idx_valor_destino]
                valor_novo_dec = (base_fonte * aliq_destino / Decimal("100")).quantize(
                    Decimal("0.01"), rounding=ROUND_HALF_UP
                )
                valor_novo = cls._formatar_decimal_sped(valor_novo_dec, f"VL_{destino}")
                campos[idx_valor_destino] = valor_novo
                if str(valor_anterior or "").strip() != valor_novo:
                    ajustes.append(
                        ItemConferenciaExcel(
                            tipo="Recálculo automático C175",
                            categoria="Tributário",
                            sequencia=sequencia,
                            registro="C175",
                            aba=aba_conferencia,
                            linha_excel=linha_excel_conferencia,
                            campo=f"VL_{destino}",
                            original=valor_anterior,
                            novo=valor_novo,
                        )
                    )

            novas[posicao] = "|" + "|".join(campos) + "|"

        return novas, ajustes, avisos

    @staticmethod
    def _separar_campos(texto: str) -> list[str]:
        if not texto.startswith("|"):
            return []
        campos = texto.split("|")[1:]
        if texto.endswith("|") and campos:
            campos = campos[:-1]
        return campos

    @staticmethod
    def _aba_possui_dados(ws) -> bool:
        for linha in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
            if any(valor not in (None, "") for valor in linha):
                return True
        return False

    @staticmethod
    def _inteiro_seguro(valor) -> int:
        try:
            return int(valor or 0)
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _encoding_saida(encoding: str, bom_presente: bool) -> str:
        # O PVA interpreta o BOM UTF-8 como caracteres antes do registro |0000|.
        # Arquivos SPED devem ser gravados sem BOM, mesmo quando a planilha técnica
        # informa que o arquivo de origem possuía a assinatura UTF-8.
        normalizado = encoding.lower().replace("_", "-")
        if normalizado in {"utf-8-sig", "utf8-sig", "utf-8", "utf8"}:
            return "utf-8"
        return encoding

    @staticmethod
    def _montar_relatorio(validacao: ResultadoValidacaoExcel, destino: Path) -> str:
        linhas = [
            "FISCALPRO — RELATÓRIO EXCEL → SPED TXT",
            "=" * 58,
            f"Gerado em: {datetime.now():%d/%m/%Y %H:%M:%S}",
            f"Planilha: {validacao.caminho}",
            f"Arquivo de origem: {validacao.arquivo_origem or '-'}",
            f"TXT gerado: {destino}",
            f"Codificação: {validacao.encoding}",
            f"Linhas reconstruídas: {validacao.total_linhas}",
            f"Registros reconstruídos: {validacao.total_registros}",
            f"Campos alterados: {validacao.total_alteracoes}",
            f"Ajustes automáticos C100: {validacao.total_ajustes_automaticos_c100}",
            f"Documentos C100 retotalizados: {validacao.total_documentos_c100_retotalizados}",
            f"Ajustes automáticos C190: {validacao.total_ajustes_automaticos_c190}",
            f"Documentos C190 reconstruídos: {validacao.total_documentos_c190_reconstruidos}",
            f"Ajustes automáticos Bloco M: {validacao.total_ajustes_automaticos_bloco_m}",
            f"Grupos M100/M500 consolidados: {validacao.total_grupos_credito_bloco_m_consolidados}",
            f"Duplicidades consolidadas: {validacao.total_duplicidades_consolidadas}",
            f"Linhas duplicadas removidas: {validacao.total_linhas_removidas}",
            (
                "Registros condicionais 0205/0220 removidos: "
                f"{validacao.total_registros_opcionais_removidos}"
            ),
            (
                "Registros do Bloco M removidos manualmente: "
                f"{validacao.total_registros_bloco_m_removidos}"
            ),
            f"Modo manual do Bloco M: {'SIM' if validacao.modo_bloco_m_manual else 'NÃO'}",
            f"Avisos: {len(validacao.avisos)}",
            "",
            "PRÉ-PVA COMPARATIVO",
            "-" * 58,
            (
                f"Executado: {'SIM' if validacao.pre_pva_executado else 'NÃO'}"
            ),
            f"Tipo SPED: {validacao.pre_pva_tipo_sped or '-'}",
            f"Erros no arquivo de origem: {validacao.pre_pva_erros_origem}",
            f"Erros no TXT reconstruído: {validacao.pre_pva_erros_reconstruido}",
            f"Novos erros criados pela planilha: {validacao.pre_pva_novos_erros}",
            f"Avisos no arquivo de origem: {validacao.pre_pva_avisos_origem}",
            f"Avisos no TXT reconstruído: {validacao.pre_pva_avisos_reconstruido}",
            f"Novos avisos criados pela planilha: {validacao.pre_pva_novos_avisos}",
            "",
            "ALTERAÇÕES IDENTIFICADAS",
            "-" * 58,
        ]

        if validacao.alteracoes:
            for alteracao in validacao.alteracoes:
                linhas.extend(
                    [
                        (
                            f"Sequência {alteracao.sequencia} | Registro {alteracao.registro} | "
                            f"Aba {alteracao.aba} | Linha Excel {alteracao.linha_excel}"
                        ),
                        f"Campo: {alteracao.campo}",
                        f"Antes: {alteracao.original}",
                        f"Depois: {alteracao.novo}",
                        "",
                    ]
                )
        else:
            linhas.append("Nenhum campo foi alterado. O TXT reproduz o arquivo original.")
            linhas.append("")

        ajustes_automaticos = [
            item for item in validacao.conferencia
            if str(item.tipo or "").startswith("Ajuste automático")
        ]
        if ajustes_automaticos:
            linhas.extend(["AJUSTES AUTOMÁTICOS", "-" * 58])
            for item in ajustes_automaticos:
                linhas.extend(
                    [
                        f"{item.tipo} | Registro {item.registro or '-'} | Campo {item.campo or '-'}",
                        f"Antes: {item.original}",
                        f"Depois: {item.novo}",
                        "",
                    ]
                )

        if validacao.duplicidades:
            linhas.extend(["DUPLICIDADES CONSOLIDADAS", "-" * 58])
            for item in validacao.duplicidades:
                linhas.extend(
                    [
                        (
                            f"{item.registro} | {item.documento} | "
                            f"{item.quantidade_linhas} linhas → 1"
                        ),
                        f"Linhas reconstruídas de origem: {', '.join(map(str, item.linhas_origem))}",
                        f"Chave tributária: {item.chave_tributaria}",
                        "Valores fiscais somados automaticamente.",
                        "",
                    ]
                )

        if validacao.avisos:
            linhas.extend(["AVISOS", "-" * 58])
            for aviso in validacao.avisos:
                local = ""
                if aviso.aba:
                    local = f" | Aba {aviso.aba}"
                if aviso.linha_excel:
                    local += f" | Linha {aviso.linha_excel}"
                linhas.append(f"- {aviso.mensagem}{local}")

        linhas.extend(
            [
                "",
                "SEGURANÇA",
                "-" * 58,
                "O arquivo Excel original não foi alterado.",
                "A planilha passa por validação estrutural e Pré-PVA comparativo antes da gravação.",
                "Somente novos erros criados pela edição do Excel bloqueiam o TXT nesta etapa.",
                "A validação oficial no PVA continua obrigatória antes da transmissão.",
                "Linhas novas continuam bloqueadas. C190/D190 duplicados são consolidados com segurança.",
            ]
        )
        return "\n".join(linhas) + "\n"

    @staticmethod
    def _progresso(callback: ProgressoCallback | None, percentual: int, mensagem: str) -> None:
        if callback:
            callback(percentual, mensagem)
