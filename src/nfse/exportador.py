"""Exportações Excel e XML do módulo NFS-e Nacional."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def _seguro_nome(texto: str) -> str:
    texto = re.sub(r"[^A-Za-z0-9._-]+", "_", str(texto or "").strip())
    return texto.strip("._") or "NFSe"


def exportar_excel(registros: list[dict], caminho: str | Path) -> Path:
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "NFS-e"
    colunas = [
        ("Empresa", "empresa_nome"), ("CNPJ Empresa", "empresa_cnpj"), ("Direção", "direcao"),
        ("Nº NFS-e", "numero_nfse"), ("Data emissão", "data_emissao"), ("Competência", "competencia"),
        ("Prestador CNPJ/CPF", "prestador_doc"), ("Prestador", "prestador_nome"),
        ("Tomador CNPJ/CPF", "tomador_doc"), ("Tomador", "tomador_nome"),
        ("Valor serviço", "valor_servico"), ("Valor líquido", "valor_liquido"),
        ("ISS", "iss"), ("PIS", "pis"), ("COFINS", "cofins"), ("CSLL", "csll"), ("IR", "ir"), ("INSS", "inss"),
        ("Situação", "situacao"), ("NSU", "nsu"), ("Chave de acesso", "chave_acesso"), ("Tipo documento", "tipo_documento"),
    ]
    ws.append([c[0] for c in colunas])
    for cel in ws[1]:
        cel.font = Font(bold=True)
        cel.alignment = Alignment(horizontal="center")
    for registro in registros:
        ws.append([registro.get(chave, "") for _, chave in colunas])
    moedas = {11, 12, 13, 14, 15, 16, 17, 18}
    for linha in range(2, ws.max_row + 1):
        for col in moedas:
            ws.cell(linha, col).number_format = 'R$ #,##0.00'
    larguras = [24, 18, 14, 14, 13, 13, 18, 34, 18, 34, 16, 16, 14, 14, 14, 14, 14, 14, 14, 12, 52, 18]
    for idx, largura in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(idx)].width = largura
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Indicador", "Valor"])
    resumo["A1"].font = resumo["B1"].font = Font(bold=True)
    emitidas = [r for r in registros if r.get("direcao") == "EMITIDA" and not r.get("tipo_evento")]
    recebidas = [r for r in registros if r.get("direcao") == "RECEBIDA" and not r.get("tipo_evento")]
    resumo.append(["Quantidade total", len(registros)])
    resumo.append(["NFS-e emitidas", len(emitidas)])
    resumo.append(["Valor emitidas", sum(float(r.get("valor_servico") or 0) for r in emitidas)])
    resumo.append(["NFS-e recebidas", len(recebidas)])
    resumo.append(["Valor recebidas", sum(float(r.get("valor_servico") or 0) for r in recebidas)])
    resumo.column_dimensions["A"].width = 28
    resumo.column_dimensions["B"].width = 20
    resumo["B4"].number_format = resumo["B6"].number_format = 'R$ #,##0.00'
    wb.save(destino)
    return destino


def exportar_xmls(registros: list[dict], pasta: str | Path) -> tuple[Path, int]:
    destino = Path(pasta)
    destino.mkdir(parents=True, exist_ok=True)
    total = 0
    for reg in registros:
        xml = str(reg.get("xml") or "").strip()
        if not xml:
            continue
        empresa = _seguro_nome(reg.get("empresa_nome") or reg.get("empresa_cnpj") or "Empresa")
        sub = destino / empresa
        sub.mkdir(parents=True, exist_ok=True)
        nome = reg.get("numero_nfse") or reg.get("chave_acesso") or f"NSU_{reg.get('nsu', '')}"
        if reg.get("tipo_evento"):
            nome = f"EVENTO_{reg.get('tipo_evento')}_{nome}"
        arquivo = sub / f"{_seguro_nome(nome)}.xml"
        arquivo.write_text(xml, encoding="utf-8")
        total += 1
    return destino, total
