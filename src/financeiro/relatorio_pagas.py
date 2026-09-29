"""Relatório mensal de contas pagas do FiscalPro.

A referência do relatório é a DATA DA BAIXA. Desde a 17.8.114 o relatório
prioriza o CNPJ gravado na própria conta, permitindo vários CNPJs por fornecedor.
"""

from __future__ import annotations

import re
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence


@dataclass(frozen=True)
class ResultadoRelatorioContasPagas:
    caminho: Path
    quantidade_contas: int
    quantidade_empresas: int
    total_geral: float
    totais_empresas: tuple[tuple[str, float], ...]


def _campo(registro: Any, chave: str, padrao: object = ""):
    try:
        return registro[chave]
    except Exception:
        return padrao


def _data_excel(valor_iso: object):
    texto = str(valor_iso or "").strip()
    try:
        return datetime.strptime(texto, "%Y-%m-%d")
    except ValueError:
        return texto


def _nome_arquivo(valor: object) -> str:
    texto = str(valor or "").strip()
    if not texto:
        return ""
    try:
        return Path(texto).name
    except Exception:
        return texto


def _tem_anexo(valor: object) -> bool:
    texto = str(valor or "").strip()
    return bool(texto and Path(texto).is_file())


def _formatar_cnpj(valor: object) -> str:
    digitos = "".join(ch for ch in str(valor or "") if ch.isdigit())
    if len(digitos) == 14:
        return (
            f"{digitos[:2]}.{digitos[2:5]}.{digitos[5:8]}/"
            f"{digitos[8:12]}-{digitos[12:]}"
        )
    return str(valor or "").strip()


class GeradorRelatorioContasPagas:
    COR_PRIMARIA = "1F4E78"
    COR_CABECALHO = "D9EAF7"
    COR_TOTAL = "E2F0D9"
    COR_BORDA = "AAB7C4"

    def gerar(
        self,
        contas: Sequence[Any],
        destino: str | Path,
        *,
        competencia_iso: str,
        empresas: Sequence[str] = (),
    ) -> ResultadoRelatorioContasPagas:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as erro:
            raise RuntimeError(
                "O componente de Excel do FiscalPro não está disponível. "
                "Reinstale a versão completa do sistema."
            ) from erro

        pasta = Path(destino).expanduser()
        pasta.mkdir(parents=True, exist_ok=True)
        mes = competencia_iso[5:7]
        ano = competencia_iso[:4]
        caminho = pasta / f"CONTAS_PAGAS_{mes}-{ano}.xlsx"

        selecionadas = tuple(dict.fromkeys(str(e).strip() for e in empresas if str(e).strip()))
        grupos: "OrderedDict[str, list[Any]]" = OrderedDict((empresa, []) for empresa in selecionadas)
        for conta in contas:
            empresa = str(_campo(conta, "empresa") or "Empresa não identificada")
            grupos.setdefault(empresa, []).append(conta)

        totais = OrderedDict(
            (empresa, round(sum(float(_campo(c, "valor", 0) or 0) for c in itens), 2))
            for empresa, itens in grupos.items()
        )
        total_geral = round(sum(totais.values()), 2)

        preenchimento_cab = PatternFill("solid", fgColor=self.COR_PRIMARIA)
        fonte_cab = Font(color="FFFFFF", bold=True)
        borda = Border(
            left=Side(style="thin", color=self.COR_BORDA),
            right=Side(style="thin", color=self.COR_BORDA),
            top=Side(style="thin", color=self.COR_BORDA),
            bottom=Side(style="thin", color=self.COR_BORDA),
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:D1")
        ws["A1"] = "RELATÓRIO MENSAL DE CONTAS PAGAS"
        ws["A1"].font = Font(bold=True, size=15, color=self.COR_PRIMARIA)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("A2:D2")
        ws["A2"] = f"Data do pagamento: {mes}/{ano}"
        ws["A2"].alignment = Alignment(horizontal="center")
        ws.append([])
        ws.append(["Empresa", "Quantidade paga", "Total pago", "% do total"])
        for cel in ws[4]:
            cel.fill = preenchimento_cab
            cel.font = fonte_cab
            cel.alignment = Alignment(horizontal="center", vertical="center")
            cel.border = borda
        for empresa, itens in grupos.items():
            total = totais[empresa]
            ws.append([empresa, len(itens), total, (total / total_geral if total_geral else 0)])
        if not grupos:
            ws.append(["Nenhuma empresa selecionada", 0, 0, 0])
        ws.append(["TOTAL GERAL", len(contas), total_geral, (1 if total_geral else 0)])
        linha_total = ws.max_row
        for cel in ws[linha_total]:
            cel.fill = PatternFill("solid", fgColor=self.COR_TOTAL)
            cel.font = Font(bold=True)
            cel.border = borda
        for linha in range(5, linha_total + 1):
            ws.cell(linha, 3).number_format = 'R$ #,##0.00'
            ws.cell(linha, 4).number_format = '0.00%'
            for col in range(1, 5):
                ws.cell(linha, col).border = borda
        for col, width in {"A":35,"B":18,"C":18,"D":14}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:D{linha_total}"

        # Aba principal já ordenada no formato mais útil para a contabilidade.
        det = wb.create_sheet("Pagamentos")
        det.sheet_view.showGridLines = False
        det.merge_cells("A1:M1")
        det["A1"] = f"CONTAS PAGAS EM {mes}/{ano} — REFERÊNCIA: DATA DO PAGAMENTO"
        det["A1"].font = Font(bold=True, size=14, color=self.COR_PRIMARIA)
        det["A1"].alignment = Alignment(horizontal="center")
        det.append([])
        det.append([
            "Data de pagamento", "Vencimento", "Fornecedor", "CNPJ do fornecedor",
            "Nº Documento", "Empresa", "Valor", "Descrição", "Categoria",
            "Competência", "Origem", "Anexo", "Arquivo",
        ])
        for cel in det[3]:
            cel.fill = preenchimento_cab
            cel.font = fonte_cab
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cel.border = borda

        ordenadas = sorted(
            contas,
            key=lambda c: (
                str(_campo(c, "data_pagamento") or ""),
                str(_campo(c, "empresa") or "").casefold(),
                str(_campo(c, "fornecedor") or "").casefold(),
                int(_campo(c, "id", 0) or 0),
            ),
        )
        for conta in ordenadas:
            caminho_doc = str(_campo(conta, "caminho_documento") or "")
            det.append([
                _data_excel(_campo(conta, "data_pagamento")),
                _data_excel(_campo(conta, "vencimento")),
                _campo(conta, "fornecedor"),
                _formatar_cnpj(_campo(conta, "fornecedor_cnpj_relatorio") or _campo(conta, "fornecedor_cnpj")),
                str(_campo(conta, "numero_documento") or ""),
                _campo(conta, "empresa"),
                float(_campo(conta, "valor", 0) or 0),
                _campo(conta, "descricao"),
                _campo(conta, "categoria"),
                str(_campo(conta, "competencia") or ""),
                _campo(conta, "origem"),
                "SIM" if _tem_anexo(caminho_doc) else "NÃO",
                _nome_arquivo(caminho_doc),
            ])

        primeira_dados = 4
        ultima_dados = det.max_row
        for linha in range(primeira_dados, ultima_dados + 1):
            det.cell(linha, 1).number_format = "dd/mm/yyyy"
            det.cell(linha, 2).number_format = "dd/mm/yyyy"
            det.cell(linha, 7).number_format = 'R$ #,##0.00'
            for col in range(1, 14):
                det.cell(linha, col).alignment = Alignment(vertical="top", wrap_text=True)
                det.cell(linha, col).border = borda
        total_row = det.max_row + 1
        det.cell(total_row, 6, "TOTAL PAGO")
        det.cell(total_row, 7, total_geral)
        det.cell(total_row, 7).number_format = 'R$ #,##0.00'
        for col in range(1, 14):
            det.cell(total_row, col).fill = PatternFill("solid", fgColor=self.COR_TOTAL)
            det.cell(total_row, col).font = Font(bold=True)
            det.cell(total_row, col).border = borda
        larguras = {
            "A":18, "B":14, "C":32, "D":21, "E":18, "F":26, "G":16,
            "H":34, "I":23, "J":13, "K":14, "L":10, "M":36,
        }
        for coluna, largura in larguras.items():
            det.column_dimensions[coluna].width = largura
        det.freeze_panes = "A4"
        det.auto_filter.ref = f"A3:M{max(3, ultima_dados)}"

        # Uma aba por empresa, também com CNPJ do fornecedor.
        for empresa, itens in grupos.items():
            if not itens:
                continue
            titulo = re.sub(r"[\\/*?:\[\]]", "-", empresa)[:31] or "Empresa"
            base = titulo
            indice = 2
            while titulo in wb.sheetnames:
                sufixo = f" {indice}"
                titulo = base[:31-len(sufixo)] + sufixo
                indice += 1
            emp = wb.create_sheet(titulo)
            emp.sheet_view.showGridLines = False
            emp.merge_cells("A1:I1")
            emp["A1"] = f"CONTAS PAGAS — {empresa} — {mes}/{ano}"
            emp["A1"].font = Font(bold=True, size=13, color=self.COR_PRIMARIA)
            emp["A1"].alignment = Alignment(horizontal="center")
            emp.append([])
            emp.append([
                "Data de pagamento", "Vencimento", "Fornecedor", "CNPJ do fornecedor",
                "Nº Documento", "Valor", "Descrição", "Categoria", "Anexo",
            ])
            for cel in emp[3]:
                cel.fill = preenchimento_cab
                cel.font = fonte_cab
                cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cel.border = borda
            itens_ordenados = sorted(
                itens,
                key=lambda c: (
                    str(_campo(c, "data_pagamento") or ""),
                    str(_campo(c, "fornecedor") or "").casefold(),
                ),
            )
            for conta in itens_ordenados:
                emp.append([
                    _data_excel(_campo(conta, "data_pagamento")),
                    _data_excel(_campo(conta, "vencimento")),
                    _campo(conta, "fornecedor"),
                    _formatar_cnpj(_campo(conta, "fornecedor_cnpj_relatorio") or _campo(conta, "fornecedor_cnpj")),
                    str(_campo(conta, "numero_documento") or ""),
                    float(_campo(conta, "valor", 0) or 0),
                    _campo(conta, "descricao"),
                    _campo(conta, "categoria"),
                    "SIM" if _tem_anexo(_campo(conta, "caminho_documento")) else "NÃO",
                ])
            last = emp.max_row
            for linha in range(4, last + 1):
                emp.cell(linha, 1).number_format = "dd/mm/yyyy"
                emp.cell(linha, 2).number_format = "dd/mm/yyyy"
                emp.cell(linha, 6).number_format = 'R$ #,##0.00'
                for col in range(1, 10):
                    emp.cell(linha, col).alignment = Alignment(vertical="top", wrap_text=True)
                    emp.cell(linha, col).border = borda
            tr = emp.max_row + 1
            emp.cell(tr, 5, f"TOTAL {empresa}")
            emp.cell(tr, 6, totais[empresa])
            emp.cell(tr, 6).number_format = 'R$ #,##0.00'
            for col in range(1, 10):
                emp.cell(tr, col).fill = PatternFill("solid", fgColor=self.COR_TOTAL)
                emp.cell(tr, col).font = Font(bold=True)
                emp.cell(tr, col).border = borda
            for col, width in {
                "A":18,"B":14,"C":32,"D":21,"E":18,"F":16,"G":34,"H":23,"I":10,
            }.items():
                emp.column_dimensions[col].width = width
            emp.freeze_panes = "A4"
            emp.auto_filter.ref = f"A3:I{max(3, last)}"

        wb.save(caminho)
        return ResultadoRelatorioContasPagas(
            caminho=caminho,
            quantidade_contas=len(contas),
            quantidade_empresas=sum(1 for itens in grupos.values() if itens),
            total_geral=total_geral,
            totais_empresas=tuple((empresa, totais[empresa]) for empresa in grupos if grupos[empresa]),
        )


__all__ = ["GeradorRelatorioContasPagas", "ResultadoRelatorioContasPagas"]
