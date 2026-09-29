"""Relatórios semanais do módulo Contas a Pagar.

Sprint 16.2.1 — relatórios consolidados ou separados por empresa, em PDF e Excel, com dependências PDF locais.
"""

from __future__ import annotations

import os
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Sequence
from xml.sax.saxutils import escape


def _preparar_dependencias_pdf() -> None:
    """Disponibiliza as bibliotecas PDF empacotadas com o FiscalPro.

    A Sprint 16.2.1 leva ReportLab, Pillow e charset-normalizer em
    ``C:\\FiscalPro\\vendor_pdf``. O caminho só é acrescentado quando o
    ReportLab não está instalado no Python usado para abrir o FiscalPro.
    """
    try:
        import reportlab  # noqa: F401
        return
    except ImportError:
        pass

    raiz = Path(__file__).resolve().parents[2]
    vendor = raiz / "vendor_pdf"
    if vendor.is_dir():
        caminho = str(vendor)
        if caminho not in sys.path:
            sys.path.insert(0, caminho)


_preparar_dependencias_pdf()


@dataclass(frozen=True)
class OpcoesRelatorioContas:
    inicio_iso: str
    fim_iso: str
    empresa: str = ""
    incluir_vencidas: bool = False
    detalhado: bool = True
    separado: bool = False
    gerar_pdf: bool = True
    gerar_excel: bool = True
    empresas: tuple[str, ...] = ()


@dataclass(frozen=True)
class ResultadoRelatorioContas:
    arquivos: tuple[Path, ...]
    quantidade_contas: int
    quantidade_empresas: int
    total_geral: float
    totais_empresas: tuple[tuple[str, float], ...]


def _moeda(valor: Any) -> str:
    numero = Decimal(str(valor or 0)).quantize(Decimal("0.01"))
    return f"R$ {numero:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data_br(valor_iso: str) -> str:
    try:
        return datetime.strptime(valor_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return str(valor_iso or "")


def _nome_seguro(texto: str, limite: int = 70) -> str:
    texto = re.sub(r"[^A-Za-z0-9À-ÿ._ -]+", "_", str(texto or "").strip())
    texto = re.sub(r"\s+", "_", texto).strip("._")
    return (texto or "RELATORIO")[:limite]


def _agrupar(
    contas: Sequence[Any], empresas: Sequence[str] = ()
) -> "OrderedDict[str, list[Any]]":
    grupos: "OrderedDict[str, list[Any]]" = OrderedDict(
        (str(empresa), []) for empresa in empresas
    )
    for conta in contas:
        empresa = str(conta["empresa"] or "Empresa não identificada")
        grupos.setdefault(empresa, []).append(conta)
    return grupos


def _totais(grupos: "OrderedDict[str, list[Any]]") -> "OrderedDict[str, float]":
    return OrderedDict(
        (empresa, round(sum(float(conta["valor"] or 0) for conta in contas), 2))
        for empresa, contas in grupos.items()
    )


class GeradorRelatoriosContasPagar:
    COR_PRIMARIA = "1F4E78"
    COR_CABECALHO = "D9EAF7"
    COR_TOTAL = "E2F0D9"
    COR_ALERTA = "FFF2CC"
    COR_BORDA = "AAB7C4"

    def gerar(
        self,
        contas: Sequence[Any],
        destino: str | Path,
        opcoes: OpcoesRelatorioContas,
    ) -> ResultadoRelatorioContas:
        pasta = Path(destino)
        pasta.mkdir(parents=True, exist_ok=True)
        grupos = _agrupar(contas, opcoes.empresas)
        totais = _totais(grupos)
        total_geral = round(sum(totais.values()), 2)
        arquivos: list[Path] = []

        sufixo = f"{opcoes.inicio_iso.replace('-', '')}_A_{opcoes.fim_iso.replace('-', '')}"
        if opcoes.separado:
            pasta_saida = pasta / f"CONTAS_A_PAGAR_{sufixo}_SEPARADO"
            pasta_saida.mkdir(parents=True, exist_ok=True)
            if opcoes.gerar_pdf:
                arquivos.extend(self._pdf_separado(grupos, totais, total_geral, pasta_saida, opcoes))
            if opcoes.gerar_excel:
                arquivos.extend(self._excel_separado(grupos, totais, total_geral, pasta_saida, opcoes))
        else:
            base = pasta / f"CONTAS_A_PAGAR_{sufixo}_CONSOLIDADO"
            if opcoes.gerar_pdf:
                arquivos.append(self._pdf_consolidado(grupos, totais, total_geral, base.with_suffix(".pdf"), opcoes))
            if opcoes.gerar_excel:
                arquivos.append(self._excel_consolidado(grupos, totais, total_geral, base.with_suffix(".xlsx"), opcoes))

        return ResultadoRelatorioContas(
            arquivos=tuple(arquivos),
            quantidade_contas=len(contas),
            quantidade_empresas=len(grupos),
            total_geral=total_geral,
            totais_empresas=tuple(totais.items()),
        )

    @staticmethod
    def _periodo_texto(opcoes: OpcoesRelatorioContas) -> str:
        texto = f"Período: {_data_br(opcoes.inicio_iso)} a {_data_br(opcoes.fim_iso)}"
        if opcoes.incluir_vencidas:
            texto += " | Inclui contas vencidas anteriores e ainda não pagas"
        else:
            texto += " | Somente contas não pagas com vencimento no período"
        return texto

    def _resumo_pdf(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        opcoes: OpcoesRelatorioContas,
        *,
        titulo: str,
    ) -> list[Any]:
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        titulo_style = ParagraphStyle(
            "FiscalProFinanceiroTitulo",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#" + self.COR_PRIMARIA),
            spaceAfter=5,
        )
        corpo = ParagraphStyle(
            "FiscalProFinanceiroCorpo",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#404040"),
        )
        elementos: list[Any] = [
            Paragraph(escape(titulo), titulo_style),
            Paragraph(escape(self._periodo_texto(opcoes)), corpo),
            Spacer(1, 4 * mm),
        ]

        linhas = [["Empresa", "Quantidade", "Total"]]
        for empresa, contas in grupos.items():
            linhas.append([escape(empresa), str(len(contas)), _moeda(totais[empresa])])
        if not grupos:
            linhas.append(["Nenhuma conta encontrada", "0", _moeda(0)])
        linhas.append(["TOTAL GERAL DO GRUPO", str(sum(len(v) for v in grupos.values())), _moeda(total_geral)])
        tabela = Table(linhas, repeatRows=1, colWidths=[150 * mm, 35 * mm, 45 * mm])
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + self.COR_CABECALHO)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + self.COR_PRIMARIA)),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + self.COR_TOTAL)),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#" + self.COR_BORDA)),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos.extend([tabela, Spacer(1, 5 * mm)])
        return elementos

    def _detalhe_empresa_pdf(
        self,
        empresa: str,
        contas: Sequence[Any],
        total: float,
        *,
        incluir_titulo: bool = True,
    ) -> list[Any]:
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        secao = ParagraphStyle(
            "FiscalProFinanceiroSecao",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=13,
            textColor=colors.HexColor("#" + self.COR_PRIMARIA),
            spaceBefore=3,
            spaceAfter=5,
        )
        celula = ParagraphStyle(
            "FiscalProFinanceiroCelula",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7,
            leading=8.5,
        )
        elementos: list[Any] = []
        if incluir_titulo:
            elementos.append(Paragraph(escape(empresa), secao))

        linhas: list[list[Any]] = [[
            "Vencimento", "Fornecedor", "Descrição", "Categoria", "Documento", "Valor"
        ]]
        for conta in contas:
            linhas.append([
                _data_br(str(conta["vencimento"] or "")),
                Paragraph(escape(str(conta["fornecedor"] or "")), celula),
                Paragraph(escape(str(conta["descricao"] or "")), celula),
                Paragraph(escape(str(conta["categoria"] or "")), celula),
                Paragraph(escape(str(conta["numero_documento"] or conta["observacoes"] or "")), celula),
                _moeda(conta["valor"]),
            ])
        if not contas:
            linhas.append(["-", "Nenhuma conta encontrada", "", "", "", _moeda(0)])
        linhas.append(["", "", "", "", f"TOTAL {empresa}", _moeda(total)])
        tabela = Table(
            linhas,
            repeatRows=1,
            colWidths=[26 * mm, 52 * mm, 62 * mm, 42 * mm, 43 * mm, 30 * mm],
        )
        tabela.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + self.COR_CABECALHO)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + self.COR_PRIMARIA)),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#" + self.COR_TOTAL)),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (0, 1), (0, -1), "CENTER"),
            ("ALIGN", (-1, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#" + self.COR_BORDA)),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.extend([tabela, Spacer(1, 5 * mm)])
        return elementos

    @staticmethod
    def _rodape_pdf(canvas, doc) -> None:
        from reportlab.lib import colors
        from reportlab.lib.units import mm

        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D9E2F3"))
        canvas.line(12 * mm, 10 * mm, doc.pagesize[0] - 12 * mm, 10 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(12 * mm, 6 * mm, "FiscalPro - Contas a Pagar")
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 6 * mm, f"Página {doc.page}")
        canvas.restoreState()

    def _pdf_consolidado(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        caminho: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> Path:
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.units import mm
            from reportlab.platypus import PageBreak, SimpleDocTemplate
        except ImportError as erro:
            raise RuntimeError("Instale o pacote reportlab para gerar o relatório em PDF.") from erro

        caminho.parent.mkdir(parents=True, exist_ok=True)
        doc = SimpleDocTemplate(
            str(caminho),
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=11 * mm,
            bottomMargin=14 * mm,
            title="Relatório de Contas a Pagar - FiscalPro",
            author="FiscalPro - Stephane Rhis + ChatGPT",
        )
        elementos = self._resumo_pdf(
            grupos, totais, total_geral, opcoes,
            titulo="RELATÓRIO DE VENCIMENTOS - CONTAS A PAGAR",
        )
        if opcoes.detalhado:
            detalhados = [(empresa, contas) for empresa, contas in grupos.items() if contas]
            for indice, (empresa, contas) in enumerate(detalhados):
                if indice > 0:
                    elementos.append(PageBreak())
                elementos.extend(self._detalhe_empresa_pdf(empresa, contas, totais[empresa]))
        doc.build(elementos, onFirstPage=self._rodape_pdf, onLaterPages=self._rodape_pdf)
        return caminho

    def _pdf_empresa(
        self,
        empresa: str,
        contas: Sequence[Any],
        total: float,
        caminho: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> Path:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate

        caminho.parent.mkdir(parents=True, exist_ok=True)
        grupos = OrderedDict(((empresa, list(contas)),))
        totais = OrderedDict(((empresa, total),))
        doc = SimpleDocTemplate(
            str(caminho), pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=11 * mm, bottomMargin=14 * mm,
            title=f"Contas a Pagar - {empresa}", author="FiscalPro - Stephane Rhis + ChatGPT",
        )
        elementos = self._resumo_pdf(
            grupos, totais, total, opcoes,
            titulo=f"CONTAS A PAGAR - {empresa}",
        )
        if opcoes.detalhado:
            elementos.extend(self._detalhe_empresa_pdf(empresa, contas, total, incluir_titulo=False))
        doc.build(elementos, onFirstPage=self._rodape_pdf, onLaterPages=self._rodape_pdf)
        return caminho

    def _pdf_resumo_geral(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        caminho: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> Path:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate

        doc = SimpleDocTemplate(
            str(caminho), pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=11 * mm, bottomMargin=14 * mm,
            title="Resumo Geral de Contas a Pagar", author="FiscalPro - Stephane Rhis + ChatGPT",
        )
        elementos = self._resumo_pdf(
            grupos, totais, total_geral, opcoes,
            titulo="RESUMO GERAL - CONTAS A PAGAR",
        )
        doc.build(elementos, onFirstPage=self._rodape_pdf, onLaterPages=self._rodape_pdf)
        return caminho

    def _pdf_separado(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        pasta: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> list[Path]:
        arquivos = [
            self._pdf_resumo_geral(
                grupos, totais, total_geral, pasta / "00_RESUMO_GERAL.pdf", opcoes
            )
        ]
        for empresa, contas in grupos.items():
            if not contas:
                continue
            arquivos.append(
                self._pdf_empresa(
                    empresa, contas, totais[empresa],
                    pasta / f"CONTAS_A_PAGAR_{_nome_seguro(empresa)}.pdf", opcoes,
                )
            )
        return arquivos

    @staticmethod
    def _estilizar_cabecalho(celulas: Iterable[Any], *, cor: str = "1F4E78") -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        for celula in celulas:
            celula.fill = PatternFill("solid", fgColor=cor)
            celula.font = Font(color="FFFFFF", bold=True)
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _preencher_resumo_excel(
        self,
        ws,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        opcoes: OpcoesRelatorioContas,
    ) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:D1")
        ws["A1"] = "RELATÓRIO DE VENCIMENTOS - CONTAS A PAGAR"
        ws["A1"].font = Font(bold=True, size=15, color=self.COR_PRIMARIA)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("A2:D2")
        ws["A2"] = self._periodo_texto(opcoes)
        ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
        ws.append([])
        ws.append(["Empresa", "Quantidade", "Total", "% do total"])
        self._estilizar_cabecalho(ws[4])
        for empresa, contas in grupos.items():
            percentual = (totais[empresa] / total_geral) if total_geral else 0
            ws.append([empresa, len(contas), totais[empresa], percentual])
        if not grupos:
            ws.append(["Nenhuma conta encontrada", 0, 0, 0])
        ws.append(["TOTAL GERAL DO GRUPO", sum(len(v) for v in grupos.values()), total_geral, 1 if total_geral else 0])
        ultima = ws.max_row
        for celula in ws[ultima]:
            celula.fill = PatternFill("solid", fgColor=self.COR_TOTAL)
            celula.font = Font(bold=True)
        for linha in range(5, ultima + 1):
            ws.cell(linha, 3).number_format = 'R$ #,##0.00'
            ws.cell(linha, 4).number_format = '0.00%'
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 14
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A4:D{ultima}"

    def _preencher_empresa_excel(self, ws, empresa: str, contas: Sequence[Any], total: float) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:G1")
        ws["A1"] = f"CONTAS A PAGAR - {empresa}"
        ws["A1"].font = Font(bold=True, size=14, color=self.COR_PRIMARIA)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.append([])
        ws.append(["Vencimento", "Fornecedor", "Descrição", "Categoria", "Documento", "Origem", "Valor"])
        self._estilizar_cabecalho(ws[3])
        for conta in contas:
            ws.append([
                _data_br(str(conta["vencimento"] or "")),
                conta["fornecedor"], conta["descricao"], conta["categoria"],
                conta["numero_documento"] or conta["observacoes"], conta["origem"],
                float(conta["valor"] or 0),
            ])
        if not contas:
            ws.append(["", "Nenhuma conta encontrada", "", "", "", "", 0])
        ws.append(["", "", "", "", "", f"TOTAL {empresa}", total])
        ultima = ws.max_row
        for celula in ws[ultima]:
            celula.fill = PatternFill("solid", fgColor=self.COR_TOTAL)
            celula.font = Font(bold=True)
        for linha in range(4, ultima + 1):
            ws.cell(linha, 7).number_format = 'R$ #,##0.00'
            for coluna in range(1, 8):
                ws.cell(linha, coluna).alignment = Alignment(vertical="top", wrap_text=True)
        larguras = {"A": 14, "B": 30, "C": 34, "D": 24, "E": 25, "F": 16, "G": 16}
        for coluna, largura in larguras.items():
            ws.column_dimensions[coluna].width = largura
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:G{max(3, ultima - 1)}"

    def _excel_consolidado(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        caminho: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> Path:
        try:
            from openpyxl import Workbook
        except ImportError as erro:
            raise RuntimeError("Instale o pacote openpyxl para gerar o relatório em Excel.") from erro

        caminho.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        resumo = wb.active
        resumo.title = "Resumo"
        self._preencher_resumo_excel(resumo, grupos, totais, total_geral, opcoes)
        if opcoes.detalhado:
            for empresa, contas in grupos.items():
                if not contas:
                    continue
                titulo = re.sub(r"[\\/*?:\[\]]", "-", empresa)[:31]
                ws = wb.create_sheet(titulo)
                self._preencher_empresa_excel(ws, empresa, contas, totais[empresa])
        wb.save(caminho)
        return caminho

    def _excel_resumo_geral(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        caminho: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> Path:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo Geral"
        self._preencher_resumo_excel(ws, grupos, totais, total_geral, opcoes)
        wb.save(caminho)
        return caminho

    def _excel_empresa(
        self,
        empresa: str,
        contas: Sequence[Any],
        total: float,
        caminho: Path,
    ) -> Path:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Contas a Pagar"
        self._preencher_empresa_excel(ws, empresa, contas, total)
        wb.save(caminho)
        return caminho

    def _excel_separado(
        self,
        grupos: "OrderedDict[str, list[Any]]",
        totais: "OrderedDict[str, float]",
        total_geral: float,
        pasta: Path,
        opcoes: OpcoesRelatorioContas,
    ) -> list[Path]:
        arquivos = [
            self._excel_resumo_geral(
                grupos, totais, total_geral, pasta / "00_RESUMO_GERAL.xlsx", opcoes
            )
        ]
        for empresa, contas in grupos.items():
            if not contas:
                continue
            arquivos.append(
                self._excel_empresa(
                    empresa, contas, totais[empresa],
                    pasta / f"CONTAS_A_PAGAR_{_nome_seguro(empresa)}.xlsx",
                )
            )
        return arquivos


__all__ = [
    "GeradorRelatoriosContasPagar",
    "OpcoesRelatorioContas",
    "ResultadoRelatorioContas",
]
