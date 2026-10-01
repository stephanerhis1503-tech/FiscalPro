"""Regras de negócio do Contas a Pagar."""

from __future__ import annotations

import re
import shutil
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from email.utils import parseaddr
from pathlib import Path

from .repositorio import ContasPagarRepositorio
from .agenda_faturas import AgendaFaturasServico
from .documentos_contabilidade import GestorDocumentosFinanceiros
from .download_documentos import BaixadorDocumentosAdobe


def _normalizar_texto(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto or "")
    texto = "".join(letra for letra in texto if not unicodedata.combining(letra))
    texto = re.sub(r"\s+", " ", texto).strip().upper()
    return texto


class ContasPagarServico:
    STATUS = ("A VENCER", "VENCIDO", "PAGO")

    def __init__(self, repositorio: ContasPagarRepositorio | None = None):
        self.repositorio = repositorio or ContasPagarRepositorio()
        self.agenda = AgendaFaturasServico(self.repositorio)
        self.documentos = GestorDocumentosFinanceiros(self.repositorio.banco.parent)
        self.baixador_documentos = BaixadorDocumentosAdobe(self.repositorio.banco.parent)
        self._arquivo_pasta_download_scans = self.repositorio.banco.parent / "pasta_download_scans.txt"

    @staticmethod
    def data_iso(valor: object, *, obrigatoria: bool = False) -> str:
        if valor in (None, ""):
            if obrigatoria:
                raise ValueError("Informe a data de vencimento.")
            return ""
        if isinstance(valor, datetime):
            return valor.date().isoformat()
        if isinstance(valor, date):
            return valor.isoformat()
        texto = str(valor).strip()
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                return datetime.strptime(texto, formato).date().isoformat()
            except ValueError:
                continue
        raise ValueError(f"Data inválida: {texto}. Use dd/mm/aaaa.")

    @staticmethod
    def data_br(valor_iso: str) -> str:
        if not valor_iso:
            return ""
        try:
            return datetime.strptime(valor_iso, "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            return valor_iso

    @staticmethod
    def valor_decimal(valor: object) -> Decimal:
        if isinstance(valor, Decimal):
            numero = valor
        elif isinstance(valor, (int, float)):
            numero = Decimal(str(valor))
        else:
            texto = str(valor or "").strip().replace("R$", "").replace(" ", "")
            if not texto:
                raise ValueError("Informe o valor da conta.")
            if "," in texto:
                texto = texto.replace(".", "").replace(",", ".")
            try:
                numero = Decimal(texto)
            except InvalidOperation as erro:
                raise ValueError("Valor inválido. Exemplo: 1250,90.") from erro
        if numero <= 0:
            raise ValueError("O valor deve ser maior que zero.")
        return numero.quantize(Decimal("0.01"))

    @staticmethod
    def texto_documento(valor: object) -> str:
        if valor in (None, ""):
            return ""
        if isinstance(valor, bool):
            return str(valor)
        if isinstance(valor, int):
            return str(valor)
        if isinstance(valor, float):
            if valor.is_integer():
                return format(valor, ".0f")
            return format(valor, "f").rstrip("0").rstrip(".")
        return str(valor).strip()

    @staticmethod
    def valor_br(valor: object) -> str:
        numero = Decimal(str(valor or 0)).quantize(Decimal("0.01"))
        return f"R$ {numero:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

    @staticmethod
    def status(vencimento_iso: str, data_pagamento_iso: str = "") -> str:
        if data_pagamento_iso:
            return "PAGO"
        return "VENCIDO" if vencimento_iso < date.today().isoformat() else "A VENCER"

    @staticmethod
    def competencia(vencimento_iso: str, informada: str = "") -> str:
        texto = (informada or "").strip()
        if texto:
            match = re.fullmatch(r"(\d{2})/(\d{4})", texto)
            if match:
                mes = int(match.group(1))
                if 1 <= mes <= 12:
                    return f"{match.group(2)}-{mes:02d}"
            match_iso = re.fullmatch(r"(\d{4})-(\d{2})", texto)
            if match_iso and 1 <= int(match_iso.group(2)) <= 12:
                return texto
            raise ValueError("Competência inválida. Escolha um mês entre 01 e 12.")
        return vencimento_iso[:7]

    @staticmethod
    def opcoes_competencia(
        *,
        ano_inicial: int = 2026,
        ano_final: int | None = None,
        adicionais: object = (),
    ) -> tuple[str, ...]:
        """Gera competências mensais de 01/2026 até o último ano necessário.

        O ano corrente é acrescentado automaticamente. Competências já existentes
        e datas futuras informadas pela tela também ampliam a lista sem excluir os
        anos anteriores.
        """

        anos = {int(ano_inicial), int(ano_final or date.today().year)}
        if isinstance(adicionais, str):
            adicionais = (adicionais,)
        for valor in adicionais or ():
            texto = str(valor or "").strip()
            match_br = re.fullmatch(r"(\d{2})/(\d{4})", texto)
            match_iso = re.fullmatch(r"(\d{4})-(\d{2})", texto)
            if match_br and 1 <= int(match_br.group(1)) <= 12:
                anos.add(int(match_br.group(2)))
            elif match_iso and 1 <= int(match_iso.group(2)) <= 12:
                anos.add(int(match_iso.group(1)))
        ultimo_ano = max(max(anos), int(ano_inicial))
        return tuple(
            f"{mes:02d}/{ano:04d}"
            for ano in range(int(ano_inicial), ultimo_ano + 1)
            for mes in range(1, 13)
        )

    @staticmethod
    def competencia_br(competencia_iso: str) -> str:
        if re.fullmatch(r"\d{4}-\d{2}", competencia_iso or ""):
            return f"{competencia_iso[5:7]}/{competencia_iso[:4]}"
        return competencia_iso or ""

    @staticmethod
    def chave_duplicidade(
        empresa: str,
        fornecedor: str,
        vencimento_iso: str,
        valor: Decimal,
        numero_documento: str = "",
        fornecedor_cnpj: str = "",
    ) -> str:
        return "|".join(
            (
                _normalizar_texto(empresa),
                _normalizar_texto(fornecedor),
                vencimento_iso,
                f"{valor:.2f}",
                _normalizar_texto(numero_documento),
                _normalizar_texto(fornecedor_cnpj),
            )
        )

    def preparar_dados(self, dados: dict[str, object]) -> dict[str, object]:
        empresa = str(dados.get("empresa") or "").strip()
        fornecedor = str(dados.get("fornecedor") or "").strip()
        if not empresa:
            raise ValueError("Selecione a empresa.")
        if empresa not in self.repositorio.listar_empresas():
            raise ValueError("A empresa selecionada não está cadastrada.")
        if not fornecedor:
            raise ValueError("Informe o fornecedor ou beneficiário.")

        vencimento = self.data_iso(dados.get("vencimento"), obrigatoria=True)
        pagamento = self.data_iso(dados.get("data_pagamento"))
        valor = self.valor_decimal(dados.get("valor"))
        categoria = str(dados.get("categoria") or "OUTROS").strip().upper() or "OUTROS"
        numero_documento = str(dados.get("numero_documento") or "").strip()
        competencia = self.competencia(vencimento, str(dados.get("competencia") or ""))
        fornecedor = self.repositorio.salvar_fornecedor(fornecedor)
        cnpjs_cadastrados = self.repositorio.listar_cnpjs_fornecedor(fornecedor)
        fornecedor_cnpj = self.repositorio.validar_cnpj(dados.get("fornecedor_cnpj"))
        if fornecedor_cnpj:
            fornecedor_cnpj = self.repositorio.salvar_cnpj_fornecedor(
                fornecedor, fornecedor_cnpj
            )
        elif len(cnpjs_cadastrados) == 1:
            fornecedor_cnpj = cnpjs_cadastrados[0]
        elif len(cnpjs_cadastrados) > 1:
            raise ValueError(
                "Este fornecedor possui mais de um CNPJ cadastrado. "
                "Selecione o CNPJ desta conta."
            )

        categoria = self.repositorio.salvar_categoria(categoria)
        chave = self.chave_duplicidade(
            empresa, fornecedor, vencimento, valor, numero_documento, fornecedor_cnpj
        )
        robo_id = dados.get("robo_documento_id")
        if robo_id in ("", None):
            robo_id = None
        else:
            robo_id = int(robo_id)

        return {
            "empresa": empresa,
            "fornecedor": fornecedor,
            "fornecedor_cnpj": fornecedor_cnpj,
            "descricao": str(dados.get("descricao") or "").strip(),
            "categoria": categoria,
            "valor": float(valor),
            "vencimento": vencimento,
            "data_pagamento": pagamento,
            "status": self.status(vencimento, pagamento),
            "numero_documento": numero_documento,
            "competencia": competencia,
            "origem": str(dados.get("origem") or "MANUAL").strip().upper(),
            "caminho_documento": str(dados.get("caminho_documento") or "").strip(),
            "robo_documento_id": robo_id,
            "observacoes": str(dados.get("observacoes") or "").strip(),
            "chave_duplicidade": chave,
        }

    def salvar_conta(
        self,
        dados: dict[str, object],
        *,
        conta_id: int | None = None,
        permitir_duplicidade: bool = False,
    ) -> tuple[int, list[object]]:
        preparado = self.preparar_dados(dados)
        duplicados = self.repositorio.localizar_duplicidades(
            str(preparado["chave_duplicidade"]), ignorar_id=conta_id
        )
        if duplicados and not permitir_duplicidade:
            return 0, list(duplicados)
        if conta_id is None:
            conta_id = self.repositorio.inserir_conta(preparado)
        else:
            self.repositorio.atualizar_conta(conta_id, preparado)

        # Hotfix 17.8.102: arquivos locais selecionados pela usuária passam a
        # ser copiados para a área persistente do FiscalPro. Links e caminhos
        # antigos inválidos são preservados para não quebrar bases anteriores.
        caminho = str(preparado.get("caminho_documento") or "").strip()
        if caminho and Path(caminho).expanduser().is_file():
            interno = self.documentos.internalizar(
                caminho,
                conta_id=int(conta_id),
                empresa=str(preparado["empresa"]),
                competencia=str(preparado["competencia"]),
                fornecedor=str(preparado["fornecedor"]),
                vencimento=str(preparado["vencimento"]),
                numero_documento=preparado.get("numero_documento", ""),
            )
            if str(interno) != caminho:
                preparado["caminho_documento"] = str(interno)
                self.repositorio.atualizar_conta(int(conta_id), preparado)

        self.agenda.sincronizar_conta(int(conta_id))
        return int(conta_id), list(duplicados)

    def dar_baixa_conta(self, conta_id: int, data_pagamento: object | None = None) -> int:
        """Marca uma conta existente como paga e sincroniza a Agenda de Faturas.

        Se a conta já possuir data de baixa, preserva a data original e apenas
        refaz a sincronização da agenda.
        """
        conta = self.repositorio.obter_conta(int(conta_id))
        if not conta:
            raise ValueError("A conta selecionada não foi encontrada.")

        if str(conta["data_pagamento"] or "").strip():
            self.agenda.sincronizar_conta(int(conta_id))
            return int(conta_id)

        dados = dict(conta)
        dados["data_pagamento"] = self.data_iso(data_pagamento or date.today())
        salvo_id, _ = self.salvar_conta(
            dados, conta_id=int(conta_id), permitir_duplicidade=True
        )
        return int(salvo_id)

    def atualizar_status(self) -> None:
        self.repositorio.atualizar_status_em_lote(date.today().isoformat())

    def resumo(self, empresa: str = "", competencia: str = "") -> dict[str, float | int]:
        self.atualizar_status()
        resumo = self.repositorio.resumo(empresa=empresa, competencia=competencia)
        hoje = date.today()
        resumo["proximos_7_dias"] = self.repositorio.resumo_proximos_dias(
            hoje.isoformat(), (hoje + timedelta(days=7)).isoformat(), empresa
        )
        return resumo

    @staticmethod
    def periodo_proxima_semana(referencia: date | None = None) -> tuple[date, date]:
        """Retorna a próxima semana completa, de sábado a sexta-feira.

        O início é sempre o próximo sábado posterior à data de referência.
        Assim, quando a consulta é feita em um sábado, o período começa no
        sábado seguinte, preservando o sentido de "próxima semana".
        """

        hoje = referencia or date.today()
        dias_ate_sabado = (5 - hoje.weekday()) % 7
        if dias_ate_sabado == 0:
            dias_ate_sabado = 7
        inicio = hoje + timedelta(days=dias_ate_sabado)
        return inicio, inicio + timedelta(days=6)

    def contas_para_relatorio(
        self,
        inicio: object,
        fim: object,
        *,
        empresa: str = "",
        empresas: tuple[str, ...] = (),
        incluir_vencidas: bool = False,
    ) -> list[object]:
        inicio_iso = self.data_iso(inicio, obrigatoria=True)
        fim_iso = self.data_iso(fim, obrigatoria=True)
        if fim_iso < inicio_iso:
            raise ValueError("A data final não pode ser anterior à data inicial.")

        cadastradas = tuple(self.repositorio.listar_empresas())
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if empresa and not selecionadas:
            selecionadas = (empresa,)
        invalidas = [nome for nome in selecionadas if nome not in cadastradas]
        if invalidas:
            raise ValueError(f"Empresa não cadastrada: {invalidas[0]}.")
        if empresa and empresa not in cadastradas:
            raise ValueError("A empresa selecionada não está cadastrada.")

        self.atualizar_status()
        if selecionadas:
            contas: list[object] = []
            for nome in selecionadas:
                contas.extend(
                    self.repositorio.listar_contas_periodo(
                        inicio_iso,
                        fim_iso,
                        empresa=nome,
                        incluir_vencidas=incluir_vencidas,
                    )
                )
            return contas

        return list(
            self.repositorio.listar_contas_periodo(
                inicio_iso,
                fim_iso,
                empresa="",
                incluir_vencidas=incluir_vencidas,
            )
        )

    def gerar_relatorio_vencimentos(
        self,
        destino: str | Path,
        *,
        inicio: object,
        fim: object,
        empresa: str = "",
        empresas: tuple[str, ...] = (),
        incluir_vencidas: bool = False,
        detalhado: bool = True,
        separado: bool = False,
        gerar_pdf: bool = True,
        gerar_excel: bool = True,
    ):
        if not gerar_pdf and not gerar_excel:
            raise ValueError("Selecione pelo menos um formato: PDF ou Excel.")
        inicio_iso = self.data_iso(inicio, obrigatoria=True)
        fim_iso = self.data_iso(fim, obrigatoria=True)
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if empresa and not selecionadas:
            selecionadas = (empresa,)
        if not selecionadas:
            selecionadas = tuple(self.repositorio.listar_empresas())

        contas = self.contas_para_relatorio(
            inicio_iso,
            fim_iso,
            empresas=selecionadas,
            incluir_vencidas=incluir_vencidas,
        )
        from .relatorios import GeradorRelatoriosContasPagar, OpcoesRelatorioContas

        opcoes = OpcoesRelatorioContas(
            inicio_iso=inicio_iso,
            fim_iso=fim_iso,
            empresa=empresa if len(selecionadas) == 1 else "",
            incluir_vencidas=incluir_vencidas,
            detalhado=detalhado,
            separado=separado,
            gerar_pdf=gerar_pdf,
            gerar_excel=gerar_excel,
            empresas=selecionadas,
        )
        return GeradorRelatoriosContasPagar().gerar(contas, destino, opcoes)

    def contas_competencia(
        self, competencia: str, *, empresas: tuple[str, ...] = ()
    ) -> list[object]:
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if not selecionadas:
            selecionadas = tuple(self.repositorio.listar_empresas())

        # Para os relatórios da contabilidade, prioriza o CNPJ gravado na conta.
        # Em lançamentos antigos sem CNPJ próprio, usa automaticamente o cadastro
        # do fornecedor somente quando existir um único CNPJ ativo, evitando
        # escolher o estabelecimento errado quando houver matriz/filiais.
        contas: list[object] = []
        for empresa in selecionadas:
            for conta in self.repositorio.listar_contas(
                empresa=empresa, competencia=competencia_iso, limite=100000
            ):
                dados = dict(conta)
                cnpj_relatorio = str(dados.get("fornecedor_cnpj") or "").strip()
                if not cnpj_relatorio:
                    fornecedor = str(dados.get("fornecedor") or "").strip()
                    cnpjs = self.repositorio.listar_cnpjs_fornecedor(fornecedor) if fornecedor else []
                    if len(cnpjs) == 1:
                        cnpj_relatorio = str(cnpjs[0] or "").strip()
                dados["fornecedor_cnpj_relatorio"] = cnpj_relatorio
                contas.append(dados)
        return contas

    def resumo_documentos_competencia(
        self, competencia: str, *, empresas: tuple[str, ...] = ()
    ):
        contas = self.contas_competencia(competencia, empresas=empresas)
        return self.documentos.resumo(contas)

    def anexar_documento(self, conta_id: int, caminho: str | Path) -> Path:
        conta = self.repositorio.obter_conta(int(conta_id))
        if not conta:
            raise ValueError("A conta selecionada não foi encontrada.")
        dados = dict(conta)
        dados["caminho_documento"] = str(caminho)
        salvo_id, _ = self.salvar_conta(
            dados, conta_id=int(conta_id), permitir_duplicidade=True
        )
        atualizada = self.repositorio.obter_conta(int(salvo_id))
        if not atualizada:
            raise ValueError("Não foi possível atualizar o documento da conta.")
        return Path(str(atualizada["caminho_documento"]))

    def vincular_scan_recebido(
        self, conta_id: int, caminho: str | Path
    ) -> tuple[Path, Path | None]:
        """Vincula um scan baixado e o retira da fila principal de recebidos.

        O documento permanente continua em dados/financeiro/documentos. O arquivo
        bruto vindo do Adobe é movido para a subpasta VINCULADOS apenas depois de
        a cópia permanente ter sido criada com sucesso.
        """

        fonte = Path(caminho).expanduser()
        if not fonte.is_file():
            raise FileNotFoundError("O scan selecionado não foi encontrado.")

        conta = self.repositorio.obter_conta(int(conta_id))
        if not conta:
            raise ValueError("A conta selecionada não foi encontrada.")

        permanente = self.anexar_documento(int(conta_id), fonte)
        if not permanente.is_file():
            raise ValueError("O FiscalPro não conseguiu criar a cópia permanente do anexo.")

        arquivado: Path | None = None
        try:
            competencia = str(conta["competencia"] or "").strip()
            pasta_recebidos = self.pasta_documentos_recebidos(competencia or date.today().strftime("%m/%Y"))
            fonte_resolvida = fonte.resolve()
            dentro_recebidos = False
            try:
                fonte_resolvida.relative_to(pasta_recebidos.resolve())
                dentro_recebidos = True
            except (OSError, ValueError):
                dentro_recebidos = False

            if dentro_recebidos and fonte_resolvida != permanente.resolve() and fonte_resolvida.exists():
                pasta_vinculados = pasta_recebidos / "VINCULADOS"
                pasta_vinculados.mkdir(parents=True, exist_ok=True)
                destino = pasta_vinculados / permanente.name
                if destino.exists():
                    indice = 2
                    while True:
                        candidato = destino.with_name(f"{destino.stem}_{indice}{destino.suffix}")
                        if not candidato.exists():
                            destino = candidato
                            break
                        indice += 1
                shutil.move(str(fonte_resolvida), str(destino))
                arquivado = destino.resolve()
        except OSError:
            # O vínculo principal já está seguro na pasta documentos; falha ao
            # arquivar o bruto não desfaz a associação da conta.
            arquivado = None

        return permanente, arquivado

    def extrair_links_documentos_adobe(self, texto: str) -> tuple[str, ...]:
        return self.baixador_documentos.extrair_links(texto)

    def pasta_documentos_recebidos(self, competencia: str) -> Path:
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        return self.baixador_documentos.pasta_competencia(competencia_iso)

    def pasta_download_scans_preferida(self, competencia: str) -> Path:
        """Retorna a pasta escolhida pela usuária ou a pasta interna da competência."""
        try:
            if self._arquivo_pasta_download_scans.exists():
                texto = self._arquivo_pasta_download_scans.read_text(encoding="utf-8").strip()
                if texto:
                    pasta = Path(texto).expanduser()
                    if pasta.exists() and pasta.is_dir():
                        return pasta.resolve()
        except Exception:
            pass
        return self.pasta_documentos_recebidos(competencia)

    def salvar_pasta_download_scans(self, pasta: str | Path) -> Path:
        destino = Path(pasta).expanduser().resolve()
        destino.mkdir(parents=True, exist_ok=True)
        self._arquivo_pasta_download_scans.write_text(str(destino), encoding="utf-8")
        return destino

    def usar_pasta_interna_download_scans(self, competencia: str) -> Path:
        try:
            if self._arquivo_pasta_download_scans.exists():
                self._arquivo_pasta_download_scans.unlink()
        except Exception:
            pass
        return self.pasta_documentos_recebidos(competencia)

    def baixar_documentos_adobe(
        self,
        texto: str,
        *,
        competencia: str,
        pasta_destino: str | Path | None = None,
        progresso=None,
    ):
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        destino = None
        if pasta_destino:
            destino = self.salvar_pasta_download_scans(pasta_destino)
        return self.baixador_documentos.baixar_lote(
            texto, competencia=competencia_iso, pasta_destino=destino, progresso=progresso
        )

    def gerar_pacote_contabilidade(
        self,
        destino: str | Path,
        *,
        competencia: str,
        empresas: tuple[str, ...] = (),
    ):
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if not selecionadas:
            selecionadas = tuple(self.repositorio.listar_empresas())
        contas = self.contas_competencia(competencia_iso, empresas=selecionadas)
        return self.documentos.gerar_pacote(
            contas, destino, competencia=competencia_iso, empresas=selecionadas
        )

    def contas_pagas_mes(
        self, competencia: str, *, empresas: tuple[str, ...] = ()
    ) -> list[object]:
        """Retorna contas cuja DATA DA BAIXA pertence ao mês informado."""
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        ano = int(competencia_iso[:4])
        mes = int(competencia_iso[5:7])
        inicio = date(ano, mes, 1)
        if mes == 12:
            proximo = date(ano + 1, 1, 1)
        else:
            proximo = date(ano, mes + 1, 1)
        fim = proximo - timedelta(days=1)

        cadastradas = tuple(self.repositorio.listar_empresas())
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if not selecionadas:
            selecionadas = cadastradas
        invalidas = [nome for nome in selecionadas if nome not in cadastradas]
        if invalidas:
            raise ValueError(f"Empresa não cadastrada: {invalidas[0]}.")

        return list(
            self.repositorio.listar_contas_pagas_periodo(
                inicio.isoformat(), fim.isoformat(), empresas=selecionadas
            )
        )

    def resumo_contas_pagas_mes(
        self, competencia: str, *, empresas: tuple[str, ...] = ()
    ) -> dict[str, object]:
        contas = self.contas_pagas_mes(competencia, empresas=empresas)
        totais: dict[str, float] = {}
        for conta in contas:
            empresa = str(conta["empresa"] or "Empresa não identificada")
            totais[empresa] = round(totais.get(empresa, 0.0) + float(conta["valor"] or 0), 2)
        fornecedores_sem_cnpj = sorted({
            str(conta["fornecedor"] or "").strip()
            for conta in contas
            if str(conta["fornecedor"] or "").strip()
            and not str(
                (
                    conta["fornecedor_cnpj_relatorio"]
                    if "fornecedor_cnpj_relatorio" in conta.keys()
                    else conta["fornecedor_cnpj"]
                )
                or ""
            ).strip()
        }, key=str.casefold)
        return {
            "quantidade": len(contas),
            "total": round(sum(float(conta["valor"] or 0) for conta in contas), 2),
            "totais_empresas": tuple(sorted(totais.items())),
            "fornecedores_sem_cnpj": tuple(fornecedores_sem_cnpj),
            "quantidade_fornecedores_sem_cnpj": len(fornecedores_sem_cnpj),
        }

    def gerar_relatorio_contas_pagas_mes(
        self,
        destino: str | Path,
        *,
        competencia: str,
        empresas: tuple[str, ...] = (),
    ):
        competencia_iso = self.competencia(date.today().isoformat(), competencia)
        selecionadas = tuple(dict.fromkeys(str(nome).strip() for nome in empresas if str(nome).strip()))
        if not selecionadas:
            selecionadas = tuple(self.repositorio.listar_empresas())
        contas = self.contas_pagas_mes(competencia_iso, empresas=selecionadas)
        if not contas:
            raise ValueError(
                f"Nenhuma conta com data de baixa em {self.competencia_br(competencia_iso)} "
                "foi encontrada para as empresas selecionadas."
            )
        from .relatorio_pagas import GeradorRelatorioContasPagas

        return GeradorRelatorioContasPagas().gerar(
            contas, destino, competencia_iso=competencia_iso, empresas=selecionadas
        )

    def sincronizar_fila_robo(self) -> dict[str, int]:
        """Copia documentos lidos pelo robô para a fila financeira.

        Nenhum lançamento é criado automaticamente. Arquivos incompletos ficam
        marcados para revisão e só entram no contas a pagar após aprovação.
        """

        from src.robo_email.repositorio import RoboEmailRepositorio

        robo = RoboEmailRepositorio()
        documentos = robo.listar_documentos(limite=5000)
        resultado = {"analisados": 0, "novos": 0, "revisar": 0, "duplicados": 0, "ignorados": 0}
        for documento in documentos:
            resultado["analisados"] += 1
            valor = documento["valor_documento"]
            vencimento = str(documento["vencimento_documento"] or "")
            empresa = str(documento["empresa"] or "")
            categoria = str(documento["categoria"] or "")
            tipo = str(documento["documento_tipo"] or categoria or "Documento")

            # Documentos sem valor e sem vencimento não ajudam no Contas a Pagar.
            if valor in (None, 0, 0.0) and not vencimento:
                resultado["ignorados"] += 1
                continue

            remetente = (
                str(documento["remetente"] or "")
                if "remetente" in documento.keys()
                else ""
            )
            nome_remetente, email_remetente = parseaddr(remetente)
            assunto = (
                str(documento["assunto"] or "")
                if "assunto" in documento.keys()
                else ""
            )
            fornecedor = (
                nome_remetente
                or email_remetente
                or str(documento["tributo_documento"] or "")
                or assunto
                or tipo
            ).strip() or "Fornecedor a conferir"
            numero = str(documento["documento_numero"] or "")
            competencia = str(
                documento["competencia_documento"] or documento["competencia"] or ""
            )
            caminho = str(documento["caminho_salvo"] or "")
            alertas: list[str] = []
            if not empresa or empresa == "Não Identificada":
                alertas.append("Empresa não identificada")
            if valor in (None, 0, 0.0):
                alertas.append("Valor não localizado")
            if not vencimento:
                alertas.append("Vencimento não localizado")
            status = "REVISAR" if alertas else "NOVO"
            chave = ""
            if empresa and valor not in (None, 0, 0.0) and vencimento:
                chave = self.chave_duplicidade(
                    empresa,
                    fornecedor,
                    vencimento,
                    Decimal(str(valor)),
                    numero,
                )
                if self.repositorio.localizar_duplicidades(chave):
                    status = "DUPLICADO"
                    alertas.append("Possível conta já lançada")

            self.repositorio.inserir_ou_atualizar_fila(
                {
                    "robo_documento_id": int(documento["id"]),
                    "empresa": empresa,
                    "fornecedor": fornecedor,
                    "tipo_documento": tipo,
                    "numero_documento": numero,
                    "vencimento": vencimento,
                    "valor": float(valor) if valor not in (None, "") else None,
                    "competencia": competencia,
                    "caminho_documento": caminho,
                    "status": status,
                    "motivo": "; ".join(alertas),
                    "chave_duplicidade": chave,
                }
            )
            if status == "REVISAR":
                resultado["revisar"] += 1
            elif status == "DUPLICADO":
                resultado["duplicados"] += 1
            else:
                resultado["novos"] += 1
        return resultado

    def importar_excel(self, caminho: str | Path) -> dict[str, int]:
        """Importa a planilha histórica do Grupo Mega Motos.

        Reconhece tanto a versão antiga quanto a versão preparada para o Robô.
        Linhas duplicadas não são importadas novamente.
        """

        try:
            from openpyxl import load_workbook
        except ImportError as erro:
            raise RuntimeError("Instale o pacote openpyxl para importar a planilha.") from erro

        arquivo = Path(caminho)
        if not arquivo.is_file():
            raise ValueError("A planilha selecionada não foi encontrada.")
        wb = load_workbook(arquivo, data_only=True, read_only=True)
        aliases = {
            "MEGA S. PROF": "Mega Profissional",
            "MEGA M. SERV": "Mega Serviços",
            "MEGA MIX E-COM": "Mega Mix E-commerce",
            "MEGA TO E-COM": "Mega T.O. E-commerce",
            "MEGA M. COMER": "Mega Motos Comércio",
            "MEGA M. TRILHA": "Mega Motos Trilha",
        }
        resultado = {"lidas": 0, "importadas": 0, "duplicadas": 0, "ignoradas": 0}
        for planilha in wb.worksheets:
            empresa = ""
            nome_norm = _normalizar_texto(planilha.title)
            for trecho, empresa_nome in aliases.items():
                if _normalizar_texto(trecho) in nome_norm:
                    empresa = empresa_nome
                    break
            if not empresa:
                continue

            cabecalhos: dict[str, int] = {}
            primeira = next(planilha.iter_rows(min_row=1, max_row=1, values_only=True), ())
            for indice, valor in enumerate(primeira):
                cabecalhos[_normalizar_texto(str(valor or ""))] = indice

            def indice(*nomes: str) -> int | None:
                for nome in nomes:
                    alvo = _normalizar_texto(nome)
                    if alvo in cabecalhos:
                        return cabecalhos[alvo]
                return None

            idx_venc = indice("Vencimento")
            idx_fornecedor = indice("Fornecedor")
            idx_desc = indice("Descrição", "Descricao")
            idx_cat = indice("Categoria")
            idx_valor = indice("Valor")
            idx_pag = indice("Data da Baixa", "Data da Baixa/Pagamento", "Data Pagamento", "Pagamento")
            idx_doc = indice("Documento / Observações", "Observações", "Documento")
            idx_comp = indice("Competência", "Competencia")
            idx_origem = indice("Origem")
            idx_arquivo = indice("Arquivo / Link")
            if idx_venc is None or idx_fornecedor is None or idx_valor is None:
                continue

            for linha in planilha.iter_rows(min_row=2, values_only=True):
                resultado["lidas"] += 1
                try:
                    vencimento = linha[idx_venc] if idx_venc < len(linha) else None
                    fornecedor = linha[idx_fornecedor] if idx_fornecedor < len(linha) else None
                    valor = linha[idx_valor] if idx_valor < len(linha) else None
                    if vencimento in (None, "") or fornecedor in (None, "") or valor in (None, "", "-"):
                        resultado["ignoradas"] += 1
                        continue
                    dados = {
                        "empresa": empresa,
                        "fornecedor": fornecedor,
                        "descricao": linha[idx_desc] if idx_desc is not None and idx_desc < len(linha) else "",
                        "categoria": linha[idx_cat] if idx_cat is not None and idx_cat < len(linha) else "OUTROS",
                        "valor": valor,
                        "vencimento": vencimento,
                        "data_pagamento": linha[idx_pag] if idx_pag is not None and idx_pag < len(linha) else "",
                        "numero_documento": self.texto_documento(linha[idx_doc]) if idx_doc is not None and idx_doc < len(linha) else "",
                        "observacoes": self.texto_documento(linha[idx_doc]) if idx_doc is not None and idx_doc < len(linha) else "",
                        "competencia": linha[idx_comp] if idx_comp is not None and idx_comp < len(linha) else "",
                        "origem": linha[idx_origem] if idx_origem is not None and idx_origem < len(linha) else "IMPORTAÇÃO",
                        "caminho_documento": linha[idx_arquivo] if idx_arquivo is not None and idx_arquivo < len(linha) else "",
                    }
                    conta_id, duplicados = self.salvar_conta(dados)
                    if duplicados:
                        resultado["duplicadas"] += 1
                    elif conta_id:
                        resultado["importadas"] += 1
                except (ValueError, TypeError, IndexError):
                    resultado["ignoradas"] += 1
        wb.close()
        return resultado

    def exportar_excel(self, caminho: str | Path) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as erro:
            raise RuntimeError("Instale o pacote openpyxl para exportar a planilha.") from erro

        destino = Path(caminho)
        wb = Workbook()
        wb.remove(wb.active)
        cabecalhos = [
            "Vencimento", "Fornecedor", "Descrição", "Categoria", "Valor",
            "Data da Baixa", "Status", "Documento / Observações", "Competência",
            "Origem", "Arquivo / Link",
        ]
        azul = "1F4E78"
        for empresa in self.repositorio.listar_empresas():
            titulo = re.sub(r"[\\/*?:\[\]]", "-", empresa)[:31]
            ws = wb.create_sheet(titulo)
            ws.append(cabecalhos)
            for celula in ws[1]:
                celula.fill = PatternFill("solid", fgColor=azul)
                celula.font = Font(color="FFFFFF", bold=True)
                celula.alignment = Alignment(horizontal="center")
            for conta in self.repositorio.listar_contas(empresa=empresa):
                ws.append([
                    self.data_br(conta["vencimento"]), conta["fornecedor"], conta["descricao"],
                    conta["categoria"], float(conta["valor"]),
                    self.data_br(conta["data_pagamento"]), conta["status"],
                    conta["observacoes"] or conta["numero_documento"],
                    self.competencia_br(conta["competencia"]), conta["origem"],
                    conta["caminho_documento"],
                ])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.column_dimensions["A"].width = 13
            ws.column_dimensions["B"].width = 28
            ws.column_dimensions["C"].width = 28
            ws.column_dimensions["D"].width = 22
            ws.column_dimensions["E"].width = 15
            ws.column_dimensions["F"].width = 15
            ws.column_dimensions["G"].width = 13
            ws.column_dimensions["H"].width = 30
            ws.column_dimensions["I"].width = 12
            ws.column_dimensions["J"].width = 16
            ws.column_dimensions["K"].width = 42
            for celula in ws["E"][1:]:
                celula.number_format = 'R$ #,##0.00'
            for linha in ws.iter_rows(min_row=2):
                for celula in linha:
                    celula.alignment = Alignment(vertical="top", wrap_text=True)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)
        return destino
