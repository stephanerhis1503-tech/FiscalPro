"""Auditoria preventiva de cadastros tributários a partir de XMLs de saída.

Sprint 17.7.0
--------------
Reaproveita o motor de análise tributária em lote aprovado e consolida as
ocorrências por código de produto. O módulo é somente analítico: não altera
XML, cadastro, banco tributário ou SPED.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
import re
import unicodedata

from src.services.analise_tributaria_lote_service import (
    AnaliseTributariaLoteService,
    ResultadoAnaliseLote,
    ResultadoItemLote,
)


STATUS_CORRIGIR = "CORRIGIR"
STATUS_REVISAR = "REVISAR"
STATUS_OK = "OK"


def _normalizar_texto(valor: Any) -> str:
    texto = str(valor or "").strip().upper()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(texto.split())


def _digitos(valor: Any) -> str:
    return "".join(c for c in str(valor or "") if c.isdigit())


def _fmt_numero(valor: Optional[float]) -> str:
    if valor is None:
        return "-"
    return f"{float(valor):.4f}".replace(".", ",")


def _modo(valores: Iterable[Any], ignorar_vazio: bool = True) -> Any:
    lista = []
    for valor in valores:
        if ignorar_vazio and valor in (None, ""):
            continue
        lista.append(valor)
    if not lista:
        return ""
    return Counter(lista).most_common(1)[0][0]


@dataclass(slots=True)
class CadastroAuditado:
    chave_produto: str
    codigo: str
    descricao: str
    ncm: str
    status: str
    confiabilidade: float
    ocorrencias: int
    documentos: int
    tratamentos_distintos: int
    inconsistencia_interna: bool
    cfops_observados: str
    cst_icms_observados: str
    pis_observado: str
    cofins_observado: str
    cest_observado: str
    sugestao: str
    divergencias: List[str] = field(default_factory=list)
    pendencias: List[str] = field(default_factory=list)
    exemplos_documentos: List[str] = field(default_factory=list)
    ocorrencias_detalhadas: List[ResultadoItemLote] = field(default_factory=list, repr=False)

    def para_dict(self) -> Dict[str, Any]:
        dados = asdict(self)
        dados.pop("ocorrencias_detalhadas", None)
        return dados


@dataclass(slots=True)
class ResultadoAuditoriaCadastrosXML:
    produtos: List[CadastroAuditado] = field(default_factory=list)
    analise_base: ResultadoAnaliseLote | None = None
    itens_saida: int = 0
    itens_ignorados_entrada: int = 0
    erros: List[str] = field(default_factory=list)

    def resumo(self) -> Dict[str, int]:
        return {
            "produtos": len(self.produtos),
            "corrigir": sum(1 for p in self.produtos if p.status == STATUS_CORRIGIR),
            "revisar": sum(1 for p in self.produtos if p.status == STATUS_REVISAR),
            "ok": sum(1 for p in self.produtos if p.status == STATUS_OK),
            "inconsistentes": sum(1 for p in self.produtos if p.inconsistencia_interna),
            "itens_saida": self.itens_saida,
            "itens_ignorados_entrada": self.itens_ignorados_entrada,
        }


class AuditoriaCadastrosXMLService:
    """Agrupa a análise tributária dos XMLs por cadastro de produto."""

    @classmethod
    def analisar(
        cls,
        caminhos: Sequence[str | Path],
        contexto: Optional[Dict[str, Any]] = None,
    ) -> ResultadoAuditoriaCadastrosXML:
        contexto_final = dict(contexto or {})
        contexto_final.setdefault("operacao", "VENDA")
        contexto_final.setdefault("finalidade", "REVENDA")
        contexto_final.setdefault("usar_uf_data_documento", True)

        analise = AnaliseTributariaLoteService().analisar(caminhos, contexto_final)
        itens_saida: List[ResultadoItemLote] = []
        ignorados = 0
        for item in analise.itens:
            # XMLs de saída usam CFOP iniciado por 5, 6 ou 7. Mantemos itens sem
            # CFOP para revisão, pois sua ausência já é um problema de cadastro.
            cfop = _digitos(item.cfop)
            if cfop and not cfop.startswith(("5", "6", "7")):
                ignorados += 1
                continue
            itens_saida.append(item)

        grupos: Dict[str, List[ResultadoItemLote]] = defaultdict(list)
        for item in itens_saida:
            grupos[cls._chave_produto(item)].append(item)

        produtos = [cls._consolidar(chave, ocorrencias) for chave, ocorrencias in grupos.items()]
        ordem = {STATUS_CORRIGIR: 0, STATUS_REVISAR: 1, STATUS_OK: 2}
        produtos.sort(key=lambda p: (ordem.get(p.status, 9), -p.ocorrencias, p.descricao, p.codigo))

        erros = list(analise.erros)
        if analise.itens and not itens_saida:
            erros.append(
                "Nenhum item de saída foi identificado. Confira se a pasta contém NF-e de saída "
                "(CFOP iniciado por 5, 6 ou 7)."
            )

        return ResultadoAuditoriaCadastrosXML(
            produtos=produtos,
            analise_base=analise,
            itens_saida=len(itens_saida),
            itens_ignorados_entrada=ignorados,
            erros=erros,
        )

    @staticmethod
    def _chave_produto(item: ResultadoItemLote) -> str:
        codigo = _normalizar_texto(item.codigo)
        if codigo:
            return f"COD:{codigo}"
        ncm = _digitos(item.ncm)
        descricao = _normalizar_texto(item.descricao)
        return f"SEM_COD:{ncm}:{descricao}"

    @classmethod
    def _consolidar(cls, chave: str, itens: List[ResultadoItemLote]) -> CadastroAuditado:
        codigo = str(_modo(item.codigo for item in itens) or "")
        descricao = str(_modo(item.descricao for item in itens) or "")
        ncm = _digitos(_modo(item.ncm for item in itens))

        divergencias = cls._unicos(
            texto
            for item in itens
            for texto in item.divergencias
            if str(texto or "").strip()
        )
        pendencias = cls._unicos(
            texto
            for item in itens
            for texto in item.pendencias
            if str(texto or "").strip()
        )

        tratamentos = {cls._assinatura_tratamento(item) for item in itens}
        inconsistencia = cls._tem_inconsistencia_no_mesmo_contexto(itens)
        if inconsistencia:
            pendencias.insert(
                0,
                "O mesmo código de produto apareceu com tratamentos tributários diferentes "
                "em operações equivalentes. Revise o cadastro/regra do ERP.",
            )

        if any(item.status == "DIVERGÊNCIA" for item in itens):
            status = STATUS_CORRIGIR
        elif inconsistencia or any(item.status in {"REVISAR", "SEM NCM", "ERRO"} for item in itens):
            status = STATUS_REVISAR
        else:
            status = STATUS_OK

        confiancas = [float(item.confiabilidade or 0.0) for item in itens if float(item.confiabilidade or 0.0) > 0]
        confiabilidade = min(confiancas) if confiancas else 0.0

        cfops = cls._juntar(item.cfop for item in itens)
        csts_icms = cls._juntar(item.cst_icms_atual for item in itens)
        pis = cls._juntar(
            f"CST {item.cst_pis_atual or '-'} / {_fmt_numero(item.aliquota_pis_atual)}%"
            for item in itens
        )
        cofins = cls._juntar(
            f"CST {item.cst_cofins_atual or '-'} / {_fmt_numero(item.aliquota_cofins_atual)}%"
            for item in itens
        )
        cest = cls._juntar(_digitos(item.cest_atual) for item in itens if _digitos(item.cest_atual)) or "-"

        sugestao = cls._montar_sugestao(itens, divergencias, pendencias)
        docs = cls._unicos(item.documento for item in itens if item.documento)[:5]
        documentos = len({item.chave or item.documento or item.arquivo for item in itens})

        return CadastroAuditado(
            chave_produto=chave,
            codigo=codigo,
            descricao=descricao,
            ncm=ncm,
            status=status,
            confiabilidade=round(confiabilidade, 2),
            ocorrencias=len(itens),
            documentos=documentos,
            tratamentos_distintos=len(tratamentos),
            inconsistencia_interna=inconsistencia,
            cfops_observados=cfops or "-",
            cst_icms_observados=csts_icms or "-",
            pis_observado=pis or "-",
            cofins_observado=cofins or "-",
            cest_observado=cest,
            sugestao=sugestao,
            divergencias=divergencias,
            pendencias=pendencias,
            exemplos_documentos=docs,
            ocorrencias_detalhadas=list(itens),
        )

    @staticmethod
    def _assinatura_tratamento(item: ResultadoItemLote) -> Tuple[Any, ...]:
        return (
            _digitos(item.cfop),
            str(item.cst_icms_atual or "").strip(),
            round(float(item.aliquota_icms_atual or 0.0), 4),
            str(item.cst_pis_atual or "").strip(),
            round(float(item.aliquota_pis_atual or 0.0), 4),
            str(item.cst_cofins_atual or "").strip(),
            round(float(item.aliquota_cofins_atual or 0.0), 4),
            _digitos(item.cest_atual),
        )

    @classmethod
    def _tem_inconsistencia_no_mesmo_contexto(cls, itens: List[ResultadoItemLote]) -> bool:
        por_contexto: Dict[Tuple[str, str], set[Tuple[Any, ...]]] = defaultdict(set)
        for item in itens:
            contexto = (str(item.uf_origem or "").upper(), str(item.uf_destino or "").upper())
            por_contexto[contexto].add(cls._assinatura_tratamento(item))
        return any(len(assinaturas) > 1 for assinaturas in por_contexto.values())

    @classmethod
    def _montar_sugestao(
        cls,
        itens: List[ResultadoItemLote],
        divergencias: List[str],
        pendencias: List[str],
    ) -> str:
        partes: List[str] = []

        cst_pis = _modo(item.cst_pis_esperado for item in itens)
        aliq_pis = _modo(item.aliquota_pis_esperada for item in itens if item.aliquota_pis_esperada is not None)
        cst_cof = _modo(item.cst_cofins_esperado for item in itens)
        aliq_cof = _modo(item.aliquota_cofins_esperada for item in itens if item.aliquota_cofins_esperada is not None)
        aliq_icms = _modo(item.aliquota_icms_esperada for item in itens if item.aliquota_icms_esperada is not None)
        cest = _digitos(_modo(item.cest_esperado for item in itens if item.cest_esperado))

        if cst_pis:
            partes.append(f"PIS CST {cst_pis}" + (f" / {_fmt_numero(aliq_pis)}%" if aliq_pis != "" else ""))
        if cst_cof:
            partes.append(f"COFINS CST {cst_cof}" + (f" / {_fmt_numero(aliq_cof)}%" if aliq_cof != "" else ""))
        if aliq_icms != "":
            partes.append(f"ICMS {_fmt_numero(aliq_icms)}%")
        if cest:
            partes.append(f"CEST {cest}")

        if not partes:
            if divergencias:
                return "Revisar os campos divergentes antes de alterar o cadastro."
            if pendencias:
                return "Revisão fiscal necessária; o motor não possui segurança suficiente para sugerir todos os campos."
            return "Cadastro compatível com as regras confirmadas pelo motor."
        return " | ".join(partes)

    @staticmethod
    def _juntar(valores: Iterable[Any], limite: int = 6) -> str:
        unicos = AuditoriaCadastrosXMLService._unicos(
            str(valor).strip() for valor in valores if str(valor or "").strip()
        )
        if len(unicos) > limite:
            return " | ".join(unicos[:limite]) + f" | +{len(unicos) - limite}"
        return " | ".join(unicos)

    @staticmethod
    def _unicos(valores: Iterable[Any]) -> List[str]:
        saida: List[str] = []
        vistos: set[str] = set()
        for valor in valores:
            texto = str(valor or "").strip()
            if not texto:
                continue
            chave = _normalizar_texto(texto)
            if chave in vistos:
                continue
            vistos.add(chave)
            saida.append(texto)
        return saida


class ExportadorAuditoriaCadastrosXMLXLSX:
    """Gera planilha de trabalho para correção do cadastro no ERP."""

    @staticmethod
    def exportar(resultado: ResultadoAuditoriaCadastrosXML, destino: str | Path) -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
            from openpyxl.utils import get_column_letter
        except ImportError as erro:  # pragma: no cover - dependência do projeto
            raise RuntimeError("Instale o pacote openpyxl para exportar o relatório.") from erro

        destino = Path(destino)
        wb = Workbook()
        ws = wb.active
        ws.title = "Cadastros para revisar"

        cabecalhos = [
            "Status", "Código", "Descrição", "NCM", "Ocorrências", "Documentos",
            "Tratamentos distintos", "Inconsistência interna", "CFOPs observados",
            "CST ICMS observados", "PIS observado", "COFINS observado", "CEST observado",
            "Segurança %", "Correção sugerida", "Divergências confirmadas", "Pendências",
            "Exemplos de documentos",
        ]
        ws.append(cabecalhos)
        for celula in ws[1]:
            celula.font = Font(bold=True)
            celula.alignment = Alignment(vertical="top", wrap_text=True)

        for produto in resultado.produtos:
            ws.append([
                produto.status,
                produto.codigo,
                produto.descricao,
                produto.ncm,
                produto.ocorrencias,
                produto.documentos,
                produto.tratamentos_distintos,
                "SIM" if produto.inconsistencia_interna else "NÃO",
                produto.cfops_observados,
                produto.cst_icms_observados,
                produto.pis_observado,
                produto.cofins_observado,
                produto.cest_observado,
                produto.confiabilidade,
                produto.sugestao,
                " | ".join(produto.divergencias),
                " | ".join(produto.pendencias),
                " | ".join(produto.exemplos_documentos),
            ])

        larguras = [12, 18, 42, 12, 12, 12, 18, 18, 26, 22, 28, 28, 16, 12, 55, 75, 75, 42]
        for indice, largura in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(indice)].width = largura
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for linha in ws.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)

        ws_resumo = wb.create_sheet("Resumo")
        resumo = resultado.resumo()
        linhas_resumo = [
            ("Produtos únicos", resumo["produtos"]),
            ("Corrigir", resumo["corrigir"]),
            ("Revisar", resumo["revisar"]),
            ("OK", resumo["ok"]),
            ("Tratamentos inconsistentes", resumo["inconsistentes"]),
            ("Itens de saída analisados", resumo["itens_saida"]),
            ("Itens de entrada ignorados", resumo["itens_ignorados_entrada"]),
            ("Erros de leitura", len(resultado.erros)),
        ]
        ws_resumo.append(["Indicador", "Quantidade"])
        for celula in ws_resumo[1]:
            celula.font = Font(bold=True)
        for linha in linhas_resumo:
            ws_resumo.append(linha)
        ws_resumo.column_dimensions["A"].width = 34
        ws_resumo.column_dimensions["B"].width = 16

        ws_oc = wb.create_sheet("Ocorrências")
        cab_oc = [
            "Código", "Descrição", "NCM", "Documento", "Chave", "Item", "CFOP",
            "UF origem", "UF destino", "Status item", "CST ICMS", "Alíquota ICMS",
            "CST PIS", "Alíquota PIS", "CST COFINS", "Alíquota COFINS", "CEST",
            "Divergências", "Pendências",
        ]
        ws_oc.append(cab_oc)
        for celula in ws_oc[1]:
            celula.font = Font(bold=True)
        for produto in resultado.produtos:
            for item in produto.ocorrencias_detalhadas:
                ws_oc.append([
                    item.codigo, item.descricao, item.ncm, item.documento, item.chave,
                    item.numero_item, item.cfop, item.uf_origem, item.uf_destino, item.status,
                    item.cst_icms_atual, item.aliquota_icms_atual, item.cst_pis_atual,
                    item.aliquota_pis_atual, item.cst_cofins_atual, item.aliquota_cofins_atual,
                    item.cest_atual, " | ".join(item.divergencias), " | ".join(item.pendencias),
                ])
        ws_oc.freeze_panes = "A2"
        ws_oc.auto_filter.ref = ws_oc.dimensions
        for coluna, largura in enumerate([18, 42, 12, 24, 48, 8, 10, 10, 10, 14, 12, 14, 10, 14, 12, 14, 14, 70, 70], start=1):
            ws_oc.column_dimensions[get_column_letter(coluna)].width = largura
        for linha in ws_oc.iter_rows():
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)

        wb.save(destino)
        return str(destino.resolve())


__all__ = [
    "AuditoriaCadastrosXMLService",
    "CadastroAuditado",
    "ExportadorAuditoriaCadastrosXMLXLSX",
    "ResultadoAuditoriaCadastrosXML",
    "STATUS_CORRIGIR",
    "STATUS_REVISAR",
    "STATUS_OK",
]
