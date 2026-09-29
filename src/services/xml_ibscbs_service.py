"""Leitura e conferência de IBS/CBS em XMLs fiscais.

O serviço foi desenhado para não depender de uma lista fechada de tags:
- identifica blocos ``IBSCBS`` e ``IBSCBSTot`` em XMLs;
- preserva o caminho completo de todas as tags folha e seus valores;
- quando o documento é NF-e, também monta um resumo amigável por item;
- confere os totais informados no XML contra a soma dos itens, quando possível;
- aceita XML avulso, ZIP com XMLs e pastas.

Nenhum XML é alterado.
"""

from __future__ import annotations

from dataclasses import dataclass, field, asdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Iterable
import xml.etree.ElementTree as ET
import zipfile


TOLERANCIA = Decimal("0.02")
TOLERANCIA_ALIQUOTA = Decimal("0.0001")

# Regra determinística usada pela auditoria automática em 2026.
# Fonte técnica: Informe Técnico 2025.002 (Portal Nacional da NF-e).
# Para 2026, a regra geral usa pIBSUF=0,1%, pIBSMun=0% e pCBS=0,9%.
ALIQUOTAS_PADRAO_2026 = {
    "pIBSUF": Decimal("0.1"),
    "pIBSMun": Decimal("0"),
    "pCBS": Decimal("0.9"),
}
CST_PADRAO = "000"
CCLASSTRIB_PADRAO = "000001"


def _local(tag: str) -> str:
    return str(tag or "").split("}", 1)[-1]


def _texto(elem: ET.Element | None) -> str:
    return (elem.text or "").strip() if elem is not None else ""


def _filho(elem: ET.Element | None, nome: str) -> ET.Element | None:
    if elem is None:
        return None
    for filho in list(elem):
        if _local(filho.tag) == nome:
            return filho
    return None


def _desc(elem: ET.Element | None, nome: str) -> ET.Element | None:
    if elem is None:
        return None
    for no in elem.iter():
        if _local(no.tag) == nome:
            return no
    return None


def _valor_caminho(elem: ET.Element | None, caminho: str) -> str:
    atual = elem
    for parte in caminho.split("/"):
        atual = _filho(atual, parte)
        if atual is None:
            return ""
    return _texto(atual)


def _dec(valor: str | int | float | Decimal | None) -> Decimal:
    try:
        return Decimal(str(valor or "0").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _fmt_decimal(valor: Decimal | str | int | float | None, casas: int = 2) -> str:
    n = _dec(valor)
    return f"{n:.{casas}f}"


def _doc_sem_mascara(valor: str) -> str:
    return "".join(ch for ch in str(valor or "") if ch.isdigit())


@dataclass
class TagIBSCBS:
    arquivo: str
    chave: str = ""
    numero_nf: str = ""
    emissao: str = ""
    escopo: str = ""
    item: str = ""
    codigo_produto: str = ""
    descricao: str = ""
    ncm: str = ""
    cfop: str = ""
    bloco: str = ""
    caminho: str = ""
    tag: str = ""
    valor: str = ""


@dataclass
class ItemIBSCBS:
    arquivo: str
    chave: str = ""
    numero_nf: str = ""
    emissao: str = ""
    emitente: str = ""
    emitente_cnpj: str = ""
    item: str = ""
    codigo_produto: str = ""
    descricao: str = ""
    ncm: str = ""
    cfop: str = ""
    cst: str = ""
    cclass_trib: str = ""
    vbc: str = ""
    p_ibs_uf: str = ""
    v_ibs_uf: str = ""
    p_ibs_mun: str = ""
    v_ibs_mun: str = ""
    v_ibs: str = ""
    p_cbs: str = ""
    v_cbs: str = ""
    qtd_tags: int = 0


@dataclass
class AuditoriaItemIBSCBS:
    arquivo: str
    chave: str = ""
    numero_nf: str = ""
    emissao: str = ""
    emitente: str = ""
    emitente_cnpj: str = ""
    crt: str = ""
    item: str = ""
    codigo_produto: str = ""
    descricao: str = ""
    ncm: str = ""
    cfop: str = ""
    status: str = ""
    perfil: str = ""
    regra: str = ""
    cst_encontrado: str = ""
    cst_esperado: str = ""
    cclass_encontrado: str = ""
    cclass_esperado: str = ""
    vbc: str = ""
    p_ibs_uf_encontrado: str = ""
    p_ibs_uf_esperado: str = ""
    v_ibs_uf_encontrado: str = ""
    v_ibs_uf_calculado: str = ""
    p_ibs_mun_encontrado: str = ""
    p_ibs_mun_esperado: str = ""
    v_ibs_mun_encontrado: str = ""
    v_ibs_mun_calculado: str = ""
    v_ibs_encontrado: str = ""
    v_ibs_calculado: str = ""
    p_cbs_encontrado: str = ""
    p_cbs_esperado: str = ""
    v_cbs_encontrado: str = ""
    v_cbs_calculado: str = ""
    motivos: str = ""


@dataclass
class ResumoXMLIBSCBS:
    arquivo: str
    chave: str = ""
    numero_nf: str = ""
    emissao: str = ""
    emitente: str = ""
    emitente_cnpj: str = ""
    destinatario: str = ""
    destinatario_cnpj: str = ""
    itens_xml: int = 0
    itens_com_ibscbs: int = 0
    qtd_tags: int = 0
    total_base: str = ""
    total_ibs_uf: str = ""
    total_ibs_mun: str = ""
    total_ibs: str = ""
    total_cbs: str = ""
    soma_base_itens: str = ""
    soma_ibs_itens: str = ""
    soma_cbs_itens: str = ""
    status: str = ""
    detalhe: str = ""


@dataclass
class ResultadoLeituraIBSCBS:
    resumos: list[ResumoXMLIBSCBS] = field(default_factory=list)
    itens: list[ItemIBSCBS] = field(default_factory=list)
    tags: list[TagIBSCBS] = field(default_factory=list)
    auditorias: list[AuditoriaItemIBSCBS] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)

    @property
    def arquivos(self) -> int:
        return len(self.resumos)

    @property
    def arquivos_com_ibscbs(self) -> int:
        return sum(1 for r in self.resumos if r.qtd_tags > 0)

    @property
    def auditoria_ok(self) -> int:
        return sum(1 for a in self.auditorias if a.status == "OK")

    @property
    def auditoria_revisar(self) -> int:
        return sum(1 for a in self.auditorias if a.status == "REVISAR")

    @property
    def auditoria_especial(self) -> int:
        return sum(1 for a in self.auditorias if a.status == "REGRA ESPECIAL")


@dataclass
class _Documento:
    nome: str
    dados: bytes


class ServicoXMLIBSCBS:
    """Extrai IBS/CBS de XMLs sem modificar os documentos."""

    @staticmethod
    def listar_arquivos_pasta(pasta: str | Path) -> list[Path]:
        raiz = Path(pasta)
        return sorted([p for p in raiz.rglob("*") if p.is_file() and p.suffix.lower() in {".xml", ".zip"}])

    @classmethod
    def analisar(cls, caminhos: Iterable[str | Path]) -> ResultadoLeituraIBSCBS:
        resultado = ResultadoLeituraIBSCBS()
        for caminho in caminhos:
            p = Path(caminho)
            try:
                docs = cls._abrir_arquivo(p)
            except Exception as exc:
                resultado.erros.append(f"{p.name}: {exc}")
                continue
            for doc in docs:
                try:
                    resumo, itens, tags, auditorias = cls._analisar_documento(doc)
                    resultado.resumos.append(resumo)
                    resultado.itens.extend(itens)
                    resultado.tags.extend(tags)
                    resultado.auditorias.extend(auditorias)
                except Exception as exc:
                    resultado.erros.append(f"{doc.nome}: {exc}")
        return resultado

    @staticmethod
    def _abrir_arquivo(caminho: Path) -> list[_Documento]:
        ext = caminho.suffix.lower()
        if ext == ".xml":
            return [_Documento(caminho.name, caminho.read_bytes())]
        if ext == ".zip":
            documentos: list[_Documento] = []
            with zipfile.ZipFile(caminho, "r") as zf:
                for nome in sorted(zf.namelist()):
                    if nome.lower().endswith(".xml") and not nome.endswith("/"):
                        documentos.append(_Documento(f"{caminho.name}::{nome}", zf.read(nome)))
            return documentos
        raise ValueError("Formato não suportado. Use XML ou ZIP.")

    @classmethod
    def _analisar_documento(
        cls, doc: _Documento
    ) -> tuple[ResumoXMLIBSCBS, list[ItemIBSCBS], list[TagIBSCBS], list[AuditoriaItemIBSCBS]]:
        raiz = ET.parse(BytesIO(doc.dados)).getroot()
        inf = _desc(raiz, "infNFe")
        ide = _filho(inf, "ide") if inf is not None else _desc(raiz, "ide")
        emit = _filho(inf, "emit") if inf is not None else _desc(raiz, "emit")
        dest = _filho(inf, "dest") if inf is not None else _desc(raiz, "dest")

        chave = ""
        if inf is not None:
            chave = str(inf.attrib.get("Id") or "")
            if chave.startswith("NFe"):
                chave = chave[3:]
        if not chave:
            chave = _texto(_desc(raiz, "chNFe"))

        numero_nf = _valor_caminho(ide, "nNF") if ide is not None else ""
        emissao = (_valor_caminho(ide, "dhEmi") or _valor_caminho(ide, "dEmi")) if ide is not None else ""
        emitente = _valor_caminho(emit, "xNome") if emit is not None else ""
        emit_cnpj = _doc_sem_mascara(_valor_caminho(emit, "CNPJ") or _valor_caminho(emit, "CPF")) if emit is not None else ""
        crt = _valor_caminho(emit, "CRT") if emit is not None else ""
        destinatario = _valor_caminho(dest, "xNome") if dest is not None else ""
        dest_cnpj = _doc_sem_mascara(_valor_caminho(dest, "CNPJ") or _valor_caminho(dest, "CPF")) if dest is not None else ""

        tags: list[TagIBSCBS] = []
        itens: list[ItemIBSCBS] = []
        auditorias: list[AuditoriaItemIBSCBS] = []
        todos_det = [n for n in (inf.iter() if inf is not None else raiz.iter()) if _local(n.tag) == "det"]

        for det in todos_det:
            prod = _filho(det, "prod")
            imposto = _filho(det, "imposto")
            bloco = _filho(imposto, "IBSCBS")
            item_num = str(det.attrib.get("nItem") or "")
            cod = _valor_caminho(prod, "cProd")
            desc = _valor_caminho(prod, "xProd")
            ncm = _valor_caminho(prod, "NCM")
            cfop = _valor_caminho(prod, "CFOP")

            auditorias.append(
                cls._auditar_item(
                    arquivo=doc.nome,
                    chave=chave,
                    numero_nf=numero_nf,
                    emissao=emissao,
                    emitente=emitente,
                    emitente_cnpj=emit_cnpj,
                    crt=crt,
                    item=item_num,
                    codigo_produto=cod,
                    descricao=desc,
                    ncm=ncm,
                    cfop=cfop,
                    bloco=bloco,
                )
            )
            if bloco is None:
                continue
            tags_item = cls._flatten_tags(
                bloco,
                arquivo=doc.nome,
                chave=chave,
                numero_nf=numero_nf,
                emissao=emissao,
                escopo="ITEM",
                item=item_num,
                codigo_produto=cod,
                descricao=desc,
                ncm=ncm,
                cfop=cfop,
                bloco="IBSCBS",
            )
            tags.extend(tags_item)
            itens.append(
                ItemIBSCBS(
                    arquivo=doc.nome,
                    chave=chave,
                    numero_nf=numero_nf,
                    emissao=emissao,
                    emitente=emitente,
                    emitente_cnpj=emit_cnpj,
                    item=item_num,
                    codigo_produto=cod,
                    descricao=desc,
                    ncm=ncm,
                    cfop=cfop,
                    cst=_valor_caminho(bloco, "CST"),
                    cclass_trib=_valor_caminho(bloco, "cClassTrib"),
                    vbc=_valor_caminho(bloco, "gIBSCBS/vBC"),
                    p_ibs_uf=_valor_caminho(bloco, "gIBSCBS/gIBSUF/pIBSUF"),
                    v_ibs_uf=_valor_caminho(bloco, "gIBSCBS/gIBSUF/vIBSUF"),
                    p_ibs_mun=_valor_caminho(bloco, "gIBSCBS/gIBSMun/pIBSMun"),
                    v_ibs_mun=_valor_caminho(bloco, "gIBSCBS/gIBSMun/vIBSMun"),
                    v_ibs=_valor_caminho(bloco, "gIBSCBS/vIBS"),
                    p_cbs=_valor_caminho(bloco, "gIBSCBS/gCBS/pCBS"),
                    v_cbs=_valor_caminho(bloco, "gIBSCBS/gCBS/vCBS"),
                    qtd_tags=len(tags_item),
                )
            )

        # Totais oficiais da NF-e (quando presentes).
        total = _filho(inf, "total") if inf is not None else _desc(raiz, "total")
        tot_ibscbs = _filho(total, "IBSCBSTot") if total is not None else _desc(raiz, "IBSCBSTot")
        if tot_ibscbs is not None:
            tags.extend(
                cls._flatten_tags(
                    tot_ibscbs,
                    arquivo=doc.nome,
                    chave=chave,
                    numero_nf=numero_nf,
                    emissao=emissao,
                    escopo="TOTAL",
                    item="",
                    codigo_produto="",
                    descricao="",
                    ncm="",
                    cfop="",
                    bloco="IBSCBSTot",
                )
            )

        # Captura genérica: se não for NF-e, ainda identifica blocos IBSCBS/IBSCBSTot
        # existentes em outros XMLs sem depender do namespace.
        if not tags:
            visitados: set[int] = set()
            for no in raiz.iter():
                nome = _local(no.tag)
                if nome not in {"IBSCBS", "IBSCBSTot"} or id(no) in visitados:
                    continue
                visitados.add(id(no))
                escopo = "TOTAL" if nome == "IBSCBSTot" else "DOCUMENTO"
                tags.extend(
                    cls._flatten_tags(
                        no, arquivo=doc.nome, chave=chave, numero_nf=numero_nf, emissao=emissao,
                        escopo=escopo, item="", codigo_produto="", descricao="", ncm="", cfop="", bloco=nome,
                    )
                )

        soma_base = sum((_dec(i.vbc) for i in itens), Decimal("0"))
        soma_ibs = sum((_dec(i.v_ibs) for i in itens), Decimal("0"))
        soma_cbs = sum((_dec(i.v_cbs) for i in itens), Decimal("0"))

        total_base = _valor_caminho(tot_ibscbs, "vBCIBSCBS") if tot_ibscbs is not None else ""
        total_ibs_uf = _valor_caminho(tot_ibscbs, "gIBS/gIBSUF/vIBSUF") if tot_ibscbs is not None else ""
        total_ibs_mun = _valor_caminho(tot_ibscbs, "gIBS/gIBSMun/vIBSMun") if tot_ibscbs is not None else ""
        total_ibs = _valor_caminho(tot_ibscbs, "gIBS/vIBS") if tot_ibscbs is not None else ""
        total_cbs = _valor_caminho(tot_ibscbs, "gCBS/vCBS") if tot_ibscbs is not None else ""

        status, detalhe = cls._conferir_totais(itens, total_base, total_ibs, total_cbs, bool(tags))
        resumo = ResumoXMLIBSCBS(
            arquivo=doc.nome,
            chave=chave,
            numero_nf=numero_nf,
            emissao=emissao,
            emitente=emitente,
            emitente_cnpj=emit_cnpj,
            destinatario=destinatario,
            destinatario_cnpj=dest_cnpj,
            itens_xml=len(todos_det),
            itens_com_ibscbs=len(itens),
            qtd_tags=len(tags),
            total_base=total_base,
            total_ibs_uf=total_ibs_uf,
            total_ibs_mun=total_ibs_mun,
            total_ibs=total_ibs,
            total_cbs=total_cbs,
            soma_base_itens=_fmt_decimal(soma_base) if itens else "",
            soma_ibs_itens=_fmt_decimal(soma_ibs) if itens else "",
            soma_cbs_itens=_fmt_decimal(soma_cbs) if itens else "",
            status=status,
            detalhe=detalhe,
        )
        return resumo, itens, tags, auditorias

    @staticmethod
    def _ano_emissao(emissao: str) -> int | None:
        valor = str(emissao or "").strip()
        if len(valor) >= 4 and valor[:4].isdigit():
            return int(valor[:4])
        return None

    @staticmethod
    def _centavos(valor: Decimal) -> Decimal:
        return valor.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    @classmethod
    def _auditar_item(
        cls,
        *,
        arquivo: str,
        chave: str,
        numero_nf: str,
        emissao: str,
        emitente: str,
        emitente_cnpj: str,
        crt: str,
        item: str,
        codigo_produto: str,
        descricao: str,
        ncm: str,
        cfop: str,
        bloco: ET.Element | None,
    ) -> AuditoriaItemIBSCBS:
        base = AuditoriaItemIBSCBS(
            arquivo=arquivo,
            chave=chave,
            numero_nf=numero_nf,
            emissao=emissao,
            emitente=emitente,
            emitente_cnpj=emitente_cnpj,
            crt=crt,
            item=item,
            codigo_produto=codigo_produto,
            descricao=descricao,
            ncm=ncm,
            cfop=cfop,
        )

        if bloco is None:
            base.status = "REVISAR"
            base.perfil = "GRUPO AUSENTE"
            base.regra = "Obrigação/enquadramento depende do regime, data e operação"
            base.motivos = (
                "Item sem grupo IBSCBS. Conferir se a operação está alcançada pela obrigação "
                "e qual CST/cClassTrib deve ser usado."
            )
            return base

        cst = _valor_caminho(bloco, "CST")
        cclass = _valor_caminho(bloco, "cClassTrib")
        vbc = _valor_caminho(bloco, "gIBSCBS/vBC")
        p_ibs_uf = _valor_caminho(bloco, "gIBSCBS/gIBSUF/pIBSUF")
        v_ibs_uf = _valor_caminho(bloco, "gIBSCBS/gIBSUF/vIBSUF")
        p_ibs_mun = _valor_caminho(bloco, "gIBSCBS/gIBSMun/pIBSMun")
        v_ibs_mun = _valor_caminho(bloco, "gIBSCBS/gIBSMun/vIBSMun")
        v_ibs = _valor_caminho(bloco, "gIBSCBS/vIBS")
        p_cbs = _valor_caminho(bloco, "gIBSCBS/gCBS/pCBS")
        v_cbs = _valor_caminho(bloco, "gIBSCBS/gCBS/vCBS")

        base.cst_encontrado = cst
        base.cclass_encontrado = cclass
        base.vbc = vbc
        base.p_ibs_uf_encontrado = p_ibs_uf
        base.v_ibs_uf_encontrado = v_ibs_uf
        base.p_ibs_mun_encontrado = p_ibs_mun
        base.v_ibs_mun_encontrado = v_ibs_mun
        base.v_ibs_encontrado = v_ibs
        base.p_cbs_encontrado = p_cbs
        base.v_cbs_encontrado = v_cbs

        problemas: list[str] = []

        if not cst:
            problemas.append("CST IBS/CBS ausente")
        if not cclass:
            problemas.append("cClassTrib ausente")
        if cst and cclass and not cclass.startswith(cst):
            problemas.append(f"CST {cst} incompatível com cClassTrib {cclass}")

        # Conferência estrutural universal: quando os três valores existem,
        # o IBS total do item deve corresponder à soma UF + Município.
        if v_ibs and v_ibs_uf and v_ibs_mun:
            soma_componentes = cls._centavos(_dec(v_ibs_uf) + _dec(v_ibs_mun))
            if (soma_componentes - _dec(v_ibs)).copy_abs() > TOLERANCIA:
                problemas.append(
                    f"vIBS {_fmt_decimal(v_ibs)} difere de IBS UF + Município {_fmt_decimal(soma_componentes)}"
                )

        ano = cls._ano_emissao(emissao)
        perfil_padrao_2026 = (
            ano == 2026
            and crt == "3"
            and cst == CST_PADRAO
            and cclass == CCLASSTRIB_PADRAO
        )

        if perfil_padrao_2026:
            base.perfil = "REGRA GERAL 2026"
            base.regra = "CST 000 / cClassTrib 000001 • CRT 3 • alíquotas padrão 2026"
            base.cst_esperado = CST_PADRAO
            base.cclass_esperado = CCLASSTRIB_PADRAO
            base.p_ibs_uf_esperado = _fmt_decimal(ALIQUOTAS_PADRAO_2026["pIBSUF"], 4)
            base.p_ibs_mun_esperado = _fmt_decimal(ALIQUOTAS_PADRAO_2026["pIBSMun"], 4)
            base.p_cbs_esperado = _fmt_decimal(ALIQUOTAS_PADRAO_2026["pCBS"], 4)

            if not vbc:
                problemas.append("vBC ausente para tributação integral")
            if not p_ibs_uf:
                problemas.append("pIBSUF ausente")
            elif (_dec(p_ibs_uf) - ALIQUOTAS_PADRAO_2026["pIBSUF"]).copy_abs() > TOLERANCIA_ALIQUOTA:
                problemas.append(f"pIBSUF {p_ibs_uf}% × esperado 0,1000%")
            if not p_ibs_mun:
                problemas.append("pIBSMun ausente")
            elif (_dec(p_ibs_mun) - ALIQUOTAS_PADRAO_2026["pIBSMun"]).copy_abs() > TOLERANCIA_ALIQUOTA:
                problemas.append(f"pIBSMun {p_ibs_mun}% × esperado 0,0000%")
            if not p_cbs:
                problemas.append("pCBS ausente")
            elif (_dec(p_cbs) - ALIQUOTAS_PADRAO_2026["pCBS"]).copy_abs() > TOLERANCIA_ALIQUOTA:
                problemas.append(f"pCBS {p_cbs}% × esperado 0,9000%")

            if vbc:
                bc = _dec(vbc)
                calc_ibs_uf = cls._centavos(bc * ALIQUOTAS_PADRAO_2026["pIBSUF"] / Decimal("100"))
                calc_ibs_mun = cls._centavos(bc * ALIQUOTAS_PADRAO_2026["pIBSMun"] / Decimal("100"))
                calc_ibs = cls._centavos(calc_ibs_uf + calc_ibs_mun)
                calc_cbs = cls._centavos(bc * ALIQUOTAS_PADRAO_2026["pCBS"] / Decimal("100"))

                base.v_ibs_uf_calculado = _fmt_decimal(calc_ibs_uf)
                base.v_ibs_mun_calculado = _fmt_decimal(calc_ibs_mun)
                base.v_ibs_calculado = _fmt_decimal(calc_ibs)
                base.v_cbs_calculado = _fmt_decimal(calc_cbs)

                checks = (
                    ("vIBSUF", v_ibs_uf, calc_ibs_uf),
                    ("vIBSMun", v_ibs_mun, calc_ibs_mun),
                    ("vIBS", v_ibs, calc_ibs),
                    ("vCBS", v_cbs, calc_cbs),
                )
                for nome, encontrado, calculado in checks:
                    if encontrado == "":
                        problemas.append(f"{nome} ausente")
                    elif (_dec(encontrado) - calculado).copy_abs() > TOLERANCIA:
                        problemas.append(
                            f"{nome} {_fmt_decimal(encontrado)} × calculado {_fmt_decimal(calculado)}"
                        )

            base.status = "REVISAR" if problemas else "OK"
            base.motivos = "; ".join(problemas) if problemas else (
                "CST, cClassTrib, alíquotas e valores conferem com a regra geral de 2026."
            )
            return base

        # Fora da regra geral, o FiscalPro não inventa um enquadramento legal.
        # Ainda sinaliza incoerências objetivas de estrutura, mas deixa a escolha
        # da tributação específica para conferência.
        base.perfil = "REGRA ESPECIAL"
        if ano != 2026:
            base.regra = f"Emissão {ano or 'sem ano'}: regra automática de alíquotas não aplicada"
        elif crt != "3":
            base.regra = f"CRT {crt or 'não informado'}: regra geral do regime normal não aplicada"
        elif cst != CST_PADRAO or cclass != CCLASSTRIB_PADRAO:
            base.regra = "CST/cClassTrib específico: conferir tabela oficial e condições da operação"
        else:
            base.regra = "Enquadramento fora do perfil automático"

        if problemas:
            base.status = "REVISAR"
            base.motivos = "; ".join(problemas)
        else:
            base.status = "REGRA ESPECIAL"
            base.motivos = (
                "Estrutura básica coerente, mas o FiscalPro não substitui o enquadramento específico "
                "por produto/operação. Conferir cClassTrib e benefícios/reduções aplicáveis."
            )
        return base

    @classmethod
    def _flatten_tags(
        cls,
        raiz: ET.Element,
        *,
        arquivo: str,
        chave: str,
        numero_nf: str,
        emissao: str,
        escopo: str,
        item: str,
        codigo_produto: str,
        descricao: str,
        ncm: str,
        cfop: str,
        bloco: str,
    ) -> list[TagIBSCBS]:
        saida: list[TagIBSCBS] = []

        def visitar(no: ET.Element, partes: list[str]):
            nome = _local(no.tag)
            caminho = partes + [nome]
            filhos = list(no)
            valor = _texto(no)
            if not filhos:
                saida.append(
                    TagIBSCBS(
                        arquivo=arquivo,
                        chave=chave,
                        numero_nf=numero_nf,
                        emissao=emissao,
                        escopo=escopo,
                        item=item,
                        codigo_produto=codigo_produto,
                        descricao=descricao,
                        ncm=ncm,
                        cfop=cfop,
                        bloco=bloco,
                        caminho="/".join(caminho),
                        tag=nome,
                        valor=valor,
                    )
                )
                return
            for filho in filhos:
                visitar(filho, caminho)

        # O nome do bloco já será a primeira parte do caminho.
        visitar(raiz, [])
        return saida

    @staticmethod
    def _conferir_totais(
        itens: list[ItemIBSCBS], total_base: str, total_ibs: str, total_cbs: str, tem_tags: bool
    ) -> tuple[str, str]:
        if not tem_tags:
            return "SEM IBS/CBS", "Nenhum bloco IBSCBS/IBSCBSTot encontrado."
        if not itens:
            return "LOCALIZADO", "Tags IBS/CBS localizadas; documento sem itens NF-e resumíveis."
        divergencias: list[str] = []
        pares = (
            ("Base", sum((_dec(i.vbc) for i in itens), Decimal("0")), total_base),
            ("IBS", sum((_dec(i.v_ibs) for i in itens), Decimal("0")), total_ibs),
            ("CBS", sum((_dec(i.v_cbs) for i in itens), Decimal("0")), total_cbs),
        )
        houve_total = False
        for nome, soma, total in pares:
            if str(total or "").strip() == "":
                continue
            houve_total = True
            dif = (soma - _dec(total)).copy_abs()
            if dif > TOLERANCIA:
                divergencias.append(f"{nome}: itens {_fmt_decimal(soma)} × total {_fmt_decimal(total)}")
        if divergencias:
            return "REVISAR", "Divergência de totalização: " + "; ".join(divergencias)
        if houve_total:
            return "OK", "Somas dos itens conferem com os totais IBS/CBS do XML."
        return "LOCALIZADO", "Itens IBS/CBS encontrados; bloco IBSCBSTot não localizado para conferência."

    @staticmethod
    def exportar_excel(resultado: ResultadoLeituraIBSCBS, destino: str | Path) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        destino = Path(destino)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo_XML"

        def escrever_planilha(planilha, cabecalhos: list[str], linhas: list[list]):
            planilha.append(cabecalhos)
            for c in planilha[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="D9E7F5")
                c.alignment = Alignment(horizontal="center")
            for linha in linhas:
                planilha.append(linha)
            planilha.freeze_panes = "A2"
            planilha.auto_filter.ref = planilha.dimensions
            for idx, col in enumerate(cabecalhos, 1):
                max_len = len(col)
                for cell in planilha[get_column_letter(idx)][: min(planilha.max_row, 500)]:
                    max_len = max(max_len, len(str(cell.value or "")))
                planilha.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 42)

        escrever_planilha(
            ws,
            ["Arquivo", "Chave NF-e", "NF", "Emissão", "Emitente", "CNPJ emitente", "Destinatário", "CNPJ destinatário",
             "Itens XML", "Itens com IBS/CBS", "Tags", "Base total", "IBS UF total", "IBS Mun total", "IBS total", "CBS total",
             "Soma base itens", "Soma IBS itens", "Soma CBS itens", "Status", "Detalhe"],
            [[r.arquivo, r.chave, r.numero_nf, r.emissao, r.emitente, r.emitente_cnpj, r.destinatario, r.destinatario_cnpj,
              r.itens_xml, r.itens_com_ibscbs, r.qtd_tags, r.total_base, r.total_ibs_uf, r.total_ibs_mun, r.total_ibs, r.total_cbs,
              r.soma_base_itens, r.soma_ibs_itens, r.soma_cbs_itens, r.status, r.detalhe] for r in resultado.resumos],
        )

        wi = wb.create_sheet("Itens_IBS_CBS")
        escrever_planilha(
            wi,
            ["Arquivo", "Chave NF-e", "NF", "Emissão", "Emitente", "CNPJ emitente", "Item", "Código", "Descrição", "NCM", "CFOP",
             "CST IBS/CBS", "cClassTrib", "vBC", "pIBSUF", "vIBSUF", "pIBSMun", "vIBSMun", "vIBS", "pCBS", "vCBS", "Qtd tags"],
            [[i.arquivo, i.chave, i.numero_nf, i.emissao, i.emitente, i.emitente_cnpj, i.item, i.codigo_produto, i.descricao, i.ncm,
              i.cfop, i.cst, i.cclass_trib, i.vbc, i.p_ibs_uf, i.v_ibs_uf, i.p_ibs_mun, i.v_ibs_mun, i.v_ibs, i.p_cbs, i.v_cbs,
              i.qtd_tags] for i in resultado.itens],
        )

        wt = wb.create_sheet("Todas_Tags_IBS_CBS")
        escrever_planilha(
            wt,
            ["Arquivo", "Chave NF-e", "NF", "Emissão", "Escopo", "Item", "Código", "Descrição", "NCM", "CFOP", "Bloco",
             "Caminho da tag", "Tag", "Valor"],
            [[t.arquivo, t.chave, t.numero_nf, t.emissao, t.escopo, t.item, t.codigo_produto, t.descricao, t.ncm, t.cfop, t.bloco,
              t.caminho, t.tag, t.valor] for t in resultado.tags],
        )


        wa = wb.create_sheet("Auditoria_IBS_CBS")
        escrever_planilha(
            wa,
            [
                "Status", "Perfil", "Regra", "Arquivo", "Chave NF-e", "NF", "Emissão", "Emitente", "CNPJ emitente", "CRT",
                "Item", "Código", "Descrição", "NCM", "CFOP",
                "CST encontrado", "CST esperado", "cClassTrib encontrado", "cClassTrib esperado", "vBC",
                "pIBSUF encontrado", "pIBSUF esperado", "vIBSUF encontrado", "vIBSUF calculado",
                "pIBSMun encontrado", "pIBSMun esperado", "vIBSMun encontrado", "vIBSMun calculado",
                "vIBS encontrado", "vIBS calculado",
                "pCBS encontrado", "pCBS esperado", "vCBS encontrado", "vCBS calculado", "Motivos",
            ],
            [[
                a.status, a.perfil, a.regra, a.arquivo, a.chave, a.numero_nf, a.emissao, a.emitente, a.emitente_cnpj, a.crt,
                a.item, a.codigo_produto, a.descricao, a.ncm, a.cfop,
                a.cst_encontrado, a.cst_esperado, a.cclass_encontrado, a.cclass_esperado, a.vbc,
                a.p_ibs_uf_encontrado, a.p_ibs_uf_esperado, a.v_ibs_uf_encontrado, a.v_ibs_uf_calculado,
                a.p_ibs_mun_encontrado, a.p_ibs_mun_esperado, a.v_ibs_mun_encontrado, a.v_ibs_mun_calculado,
                a.v_ibs_encontrado, a.v_ibs_calculado,
                a.p_cbs_encontrado, a.p_cbs_esperado, a.v_cbs_encontrado, a.v_cbs_calculado, a.motivos,
            ] for a in resultado.auditorias],
        )

        if resultado.erros:
            we = wb.create_sheet("Erros_Leitura")
            escrever_planilha(we, ["Erro"], [[e] for e in resultado.erros])

        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)
        return destino
