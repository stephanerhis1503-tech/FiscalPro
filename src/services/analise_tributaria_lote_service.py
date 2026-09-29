"""Análise tributária em lote de NF-e XML, ZIP e SPED.

Sprint 17.2.0
----------------
A rotina é somente analítica: ela não altera XML, SPED, banco tributário ou
cadastros do usuário. Cada item é confrontado com os motores aprovados de NCM,
TIPI, PIS/Cofins, ICMS próprio, ICMS-ST e FCP para produzir um relatório
rastreável, com divergências confirmadas separadas de pontos condicionais.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Sequence, Tuple
import io
import os
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

from src.inteligencia.base_oficial.repositorio import BaseOficialRepository
from src.modelos.registro_0200 import Registro0200
from src.modelos.registro_c170 import RegistroC170
from src.services.base_ncm_nacional_service import BaseNCMNacionalService
from src.services.icms_uf_service import ICMSUFService
from src.services.icms_mg_lote_integracao_service import ICMSMGLoteIntegracaoService
from src.services.piscofins_nacional_service import PISCOFINSNacionalService
from src.services.empresas_regimes_service import EmpresasRegimesService


MAX_ARQUIVOS_LOTE = 5000
MAX_BYTES_POR_ARQUIVO = 25 * 1024 * 1024
MAX_BYTES_ZIP = 350 * 1024 * 1024
ORIGENS_IMPORTADAS = {"1", "2", "6", "7", "8"}

_IBGE_UF = {
    "11": "RO", "12": "AC", "13": "AM", "14": "RR", "15": "PA", "16": "AP", "17": "TO",
    "21": "MA", "22": "PI", "23": "CE", "24": "RN", "25": "PB", "26": "PE", "27": "AL",
    "28": "SE", "29": "BA", "31": "MG", "32": "ES", "33": "RJ", "35": "SP", "41": "PR",
    "42": "SC", "43": "RS", "50": "MS", "51": "MT", "52": "GO", "53": "DF",
}


def _local(tag: str) -> str:
    return str(tag or "").split("}")[-1]


def _filho(no: Optional[ET.Element], nome: str) -> Optional[ET.Element]:
    if no is None:
        return None
    for item in list(no):
        if _local(item.tag) == nome:
            return item
    return None


def _descendente(no: Optional[ET.Element], nome: str) -> Optional[ET.Element]:
    if no is None:
        return None
    for item in no.iter():
        if _local(item.tag) == nome:
            return item
    return None


def _texto(no: Optional[ET.Element], nome: str, padrao: str = "") -> str:
    alvo = _filho(no, nome)
    return (alvo.text or "").strip() if alvo is not None else padrao


def _numero(valor: Any) -> float:
    texto = str(valor if valor not in (None, "") else "0").strip().replace("%", "")
    if not texto:
        return 0.0
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except (TypeError, ValueError):
        return 0.0


def _so_digitos(valor: Any) -> str:
    return "".join(c for c in str(valor or "") if c.isdigit())


def _normalizar_texto(valor: Any) -> str:
    texto = str(valor or "").strip().upper()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(texto.split())


def _booleano(valor: Any) -> bool:
    if isinstance(valor, bool):
        return valor
    return str(valor or "").strip().upper() in {"1", "S", "SIM", "TRUE", "VERDADEIRO", "YES"}


def _data_iso(valor: Any, padrao: str = "") -> str:
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d%m%Y"):
        try:
            return datetime.strptime(texto[:10] if formato != "%d%m%Y" else texto[:8], formato).date().isoformat()
        except ValueError:
            continue
    return padrao or date.today().isoformat()


def _quase_igual(a: Any, b: Any, tolerancia: float = 0.01) -> bool:
    return abs(_numero(a) - _numero(b)) <= tolerancia


@dataclass
class ItemBrutoLote:
    fonte_tipo: str
    arquivo: str
    documento: str = ""
    chave: str = ""
    numero_item: str = ""
    codigo: str = ""
    descricao: str = ""
    ncm: str = ""
    ex_tipi: str = ""
    cest_atual: str = ""
    cfop: str = ""
    uf_origem: str = ""
    uf_destino: str = ""
    data_operacao: str = ""
    operacao_sugerida: str = ""
    origem_mercadoria: str = ""
    cst_icms_atual: str = ""
    aliquota_icms_atual: float = 0.0
    valor_icms_st_atual: float = 0.0
    mva_st_atual: float = 0.0
    aliquota_fcp_st_atual: float = 0.0
    cst_pis_atual: str = ""
    aliquota_pis_atual: float = 0.0
    cst_cofins_atual: str = ""
    aliquota_cofins_atual: float = 0.0
    cst_ipi_atual: str = ""
    aliquota_ipi_atual: float = 0.0
    numero_linha_fonte: int = 0


@dataclass
class ResultadoItemLote:
    fonte_tipo: str
    arquivo: str
    documento: str
    chave: str
    numero_item: str
    codigo: str
    descricao: str
    ncm: str
    ncm_oficial: str
    cfop: str
    uf_origem: str
    uf_destino: str
    data_operacao: str
    status: str
    confiabilidade: float
    confirmado: bool
    exige_revisao: bool
    numero_linha_fonte: int = 0
    divergencias: List[str] = field(default_factory=list)
    pendencias: List[str] = field(default_factory=list)
    cst_pis_atual: str = ""
    aliquota_pis_atual: float = 0.0
    cst_pis_esperado: str = ""
    aliquota_pis_esperada: Optional[float] = None
    cst_cofins_atual: str = ""
    aliquota_cofins_atual: float = 0.0
    cst_cofins_esperado: str = ""
    aliquota_cofins_esperada: Optional[float] = None
    piscofins_status: str = ""
    piscofins_confirmado: bool = False
    confiabilidade_piscofins: float = 0.0
    cst_icms_atual: str = ""
    aliquota_icms_atual: float = 0.0
    aliquota_icms_esperada: Optional[float] = None
    icms_status: str = ""
    icms_confirmado: bool = False
    aliquota_icms_confirmada: bool = False
    confiabilidade_icms: float = 0.0
    cest_atual: str = ""
    cest_esperado: str = ""
    st_status: str = ""
    st_confirmado: bool = False
    confiabilidade_st: float = 0.0
    mva_original: Optional[float] = None
    mva_ajustada: Optional[float] = None
    mva_aplicada: Optional[float] = None
    mva_tipo: str = ""
    mva_st_atual: float = 0.0
    st_vigencia_referencia: str = ""
    fcp_esperado: Optional[float] = None
    fcp_status: str = ""
    fcp_confirmado: bool = False
    confiabilidade_fcp: float = 0.0
    aliquota_fcp_st_atual: float = 0.0
    valor_icms_st_atual: float = 0.0
    cst_ipi_atual: str = ""
    aliquota_ipi_atual: float = 0.0
    aliquota_ipi_referencia: Optional[float] = None
    tipi_status: str = ""
    fundamento_piscofins: str = ""
    fundamento_icms: str = ""
    fonte_piscofins: str = ""
    fonte_icms: str = ""
    observacao: str = ""

    def para_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class ResultadoAnaliseLote:
    itens: List[ResultadoItemLote] = field(default_factory=list)
    erros: List[str] = field(default_factory=list)
    arquivos_processados: int = 0
    fontes_solicitadas: int = 0
    iniciado_em: str = ""
    finalizado_em: str = ""
    contexto: Dict[str, Any] = field(default_factory=dict)

    def resumo(self) -> Dict[str, Any]:
        contagem = {
            "CONFIRMADO": 0,
            "DIVERGÊNCIA": 0,
            "REVISAR": 0,
            "SEM NCM": 0,
            "ERRO": 0,
        }
        for item in self.itens:
            chave = item.status if item.status in contagem else "REVISAR"
            contagem[chave] += 1
        return {
            "arquivos_processados": self.arquivos_processados,
            "fontes_solicitadas": self.fontes_solicitadas,
            "itens": len(self.itens),
            "ncm_unicos": len({item.ncm for item in self.itens if len(item.ncm) == 8}),
            "confirmados": contagem["CONFIRMADO"],
            "divergencias": contagem["DIVERGÊNCIA"],
            "revisar": contagem["REVISAR"],
            "sem_ncm": contagem["SEM NCM"],
            "erros_itens": contagem["ERRO"],
            "erros_arquivos": len(self.erros),
        }

    def para_dict(self) -> Dict[str, Any]:
        return {
            "itens": [item.para_dict() for item in self.itens],
            "erros": list(self.erros),
            "arquivos_processados": self.arquivos_processados,
            "fontes_solicitadas": self.fontes_solicitadas,
            "iniciado_em": self.iniciado_em,
            "finalizado_em": self.finalizado_em,
            "contexto": dict(self.contexto),
            "resumo": self.resumo(),
        }


class AnaliseTributariaLoteService:
    """Orquestra a análise tributária de múltiplos documentos sem corrigi-los."""

    def __init__(self) -> None:
        self._cache_ncm: Dict[str, Optional[Dict[str, Any]]] = {}
        self._cache_tipi: Dict[Tuple[str, str], Dict[str, Any]] = {}
        self._cache_pis: Dict[Tuple[Any, ...], Dict[str, Any]] = {}
        self._cache_icms: Dict[Tuple[Any, ...], Dict[str, Any]] = {}

    def analisar_item(
        self,
        item: ItemBrutoLote,
        contexto: Optional[Dict[str, Any]] = None,
    ) -> ResultadoItemLote:
        """Analisa um item avulso reutilizando os mesmos motores e caches do lote."""
        contexto_final = self._normalizar_contexto(contexto or {})
        try:
            BaseNCMNacionalService.garantir_instalada()
        except Exception:
            pass
        return self._analisar_item(item, contexto_final)

    def analisar(
        self,
        caminhos: Sequence[str | os.PathLike[str]],
        contexto: Optional[Dict[str, Any]] = None,
    ) -> ResultadoAnaliseLote:
        contexto_final = self._normalizar_contexto(contexto or {})
        resultado = ResultadoAnaliseLote(
            fontes_solicitadas=len(caminhos),
            iniciado_em=datetime.now().isoformat(timespec="seconds"),
            contexto=contexto_final,
        )
        try:
            BaseNCMNacionalService.garantir_instalada()
        except Exception as erro:  # A análise continua com os motores locais.
            resultado.erros.append(f"Base nacional: {erro}")

        total_fontes = 0
        for nome, extensao, conteudo in self._iterar_arquivos(caminhos, resultado.erros):
            if total_fontes >= MAX_ARQUIVOS_LOTE:
                resultado.erros.append(
                    f"Limite de {MAX_ARQUIVOS_LOTE} arquivos atingido; os demais não foram processados."
                )
                break
            total_fontes += 1
            try:
                if extensao == ".xml":
                    brutos = self._ler_xml(conteudo, nome, contexto_final)
                elif extensao == ".txt":
                    brutos = self._ler_sped(conteudo, nome, contexto_final)
                else:
                    continue
                resultado.arquivos_processados += 1
                for bruto in brutos:
                    resultado.itens.append(self._analisar_item(bruto, contexto_final))
            except Exception as erro:
                resultado.erros.append(f"{nome}: {erro}")

        resultado.finalizado_em = datetime.now().isoformat(timespec="seconds")
        return resultado

    @staticmethod
    def _normalizar_contexto(contexto: Dict[str, Any]) -> Dict[str, Any]:
        contexto = EmpresasRegimesService.aplicar_contexto(contexto)
        empresa = str(contexto.get("empresa") or "").strip()
        return {
            "empresa": empresa,
            "regime": str(contexto.get("regime") or "LUCRO REAL").strip().upper(),
            "operacao": str(contexto.get("operacao") or "VENDA").strip().upper(),
            "finalidade": str(contexto.get("finalidade") or "REVENDA").strip().upper(),
            "uf_origem": str(contexto.get("uf_origem") or "MG").strip().upper(),
            "uf_destino": str(contexto.get("uf_destino") or "MG").strip().upper(),
            "data_operacao": _data_iso(contexto.get("data_operacao")),
            "consumidor_final": _booleano(contexto.get("consumidor_final")),
            "contrato_fidelidade": _booleano(contexto.get("contrato_fidelidade")),
            "usar_uf_data_documento": contexto.get("usar_uf_data_documento", True) is not False,
            "usar_operacao_sped": contexto.get("usar_operacao_sped", True) is not False,
        }

    def _iterar_arquivos(
        self,
        caminhos: Sequence[str | os.PathLike[str]],
        erros: List[str],
    ) -> Iterator[Tuple[str, str, bytes]]:
        vistos: set[str] = set()
        for caminho_bruto in caminhos:
            caminho = Path(caminho_bruto)
            if not caminho.exists():
                erros.append(f"Fonte não encontrada: {caminho}")
                continue
            candidatos: Iterable[Path]
            if caminho.is_dir():
                candidatos = sorted(
                    item for item in caminho.rglob("*")
                    if item.is_file() and item.suffix.lower() in {".xml", ".txt", ".zip"}
                )
            else:
                candidatos = (caminho,)

            for arquivo in candidatos:
                chave = str(arquivo.resolve()).lower()
                if chave in vistos:
                    continue
                vistos.add(chave)
                extensao = arquivo.suffix.lower()
                try:
                    tamanho = arquivo.stat().st_size
                    if tamanho > MAX_BYTES_ZIP and extensao == ".zip":
                        erros.append(f"ZIP ignorado por exceder o limite de tamanho: {arquivo.name}")
                        continue
                    if tamanho > MAX_BYTES_POR_ARQUIVO and extensao != ".zip":
                        erros.append(f"Arquivo ignorado por exceder 25 MB: {arquivo.name}")
                        continue
                    if extensao == ".zip":
                        yield from self._iterar_zip(arquivo, erros)
                    elif extensao in {".xml", ".txt"}:
                        yield str(arquivo), extensao, arquivo.read_bytes()
                except OSError as erro:
                    erros.append(f"{arquivo}: {erro}")

    @staticmethod
    def _iterar_zip(arquivo: Path, erros: List[str]) -> Iterator[Tuple[str, str, bytes]]:
        try:
            with zipfile.ZipFile(arquivo, "r") as pacote:
                total_descompactado = 0
                for info in pacote.infolist():
                    if info.is_dir():
                        continue
                    extensao = Path(info.filename).suffix.lower()
                    if extensao not in {".xml", ".txt"}:
                        continue
                    if info.file_size > MAX_BYTES_POR_ARQUIVO:
                        erros.append(f"{arquivo.name}/{info.filename}: item maior que 25 MB ignorado.")
                        continue
                    total_descompactado += info.file_size
                    if total_descompactado > MAX_BYTES_ZIP:
                        erros.append(f"{arquivo.name}: limite descompactado atingido; membros restantes ignorados.")
                        break
                    nome_seguro = Path(info.filename).name
                    if not nome_seguro:
                        continue
                    try:
                        conteudo = pacote.read(info)
                    except (OSError, RuntimeError, zipfile.BadZipFile) as erro:
                        erros.append(f"{arquivo.name}/{info.filename}: {erro}")
                        continue
                    yield f"{arquivo.name}::{nome_seguro}", extensao, conteudo
        except zipfile.BadZipFile as erro:
            erros.append(f"ZIP inválido {arquivo.name}: {erro}")

    def _ler_xml(
        self,
        conteudo: bytes,
        nome: str,
        contexto: Dict[str, Any],
    ) -> List[ItemBrutoLote]:
        try:
            raiz = ET.fromstring(conteudo)
        except ET.ParseError as erro:
            raise ValueError(f"XML inválido: {erro}") from erro
        inf_nfe = _descendente(raiz, "infNFe")
        if inf_nfe is None:
            raise ValueError("O XML não contém uma NF-e reconhecida (infNFe).")

        ide = _filho(inf_nfe, "ide")
        emit = _filho(inf_nfe, "emit")
        dest = _filho(inf_nfe, "dest")
        ender_emit = _filho(emit, "enderEmit")
        ender_dest = _filho(dest, "enderDest")
        numero = _texto(ide, "nNF")
        serie = _texto(ide, "serie")
        documento = f"NF-e {numero}" + (f" série {serie}" if serie else "")
        chave = str(inf_nfe.attrib.get("Id") or "").removeprefix("NFe")
        usar_doc = bool(contexto.get("usar_uf_data_documento", True))
        uf_origem = _texto(ender_emit, "UF").upper() if usar_doc else contexto["uf_origem"]
        uf_destino = _texto(ender_dest, "UF").upper() if usar_doc else contexto["uf_destino"]
        data_doc = _texto(ide, "dhEmi") or _texto(ide, "dEmi")
        data_operacao = _data_iso(data_doc, contexto["data_operacao"]) if usar_doc else contexto["data_operacao"]

        itens: List[ItemBrutoLote] = []
        for det in list(inf_nfe):
            if _local(det.tag) != "det":
                continue
            prod = _filho(det, "prod")
            imposto = _filho(det, "imposto")
            if prod is None:
                continue

            icms_tipo = self._grupo_tributo(imposto, "ICMS", "ICMS")
            pis_tipo = self._grupo_tributo(imposto, "PIS", "PIS")
            cofins_tipo = self._grupo_tributo(imposto, "COFINS", "COFINS")
            ipi_tipo = self._grupo_tributo(imposto, "IPI", "IPI")
            origem = _texto(icms_tipo, "orig")
            cst_icms = _texto(icms_tipo, "CST") or _texto(icms_tipo, "CSOSN")
            numero_item = str(det.attrib.get("nItem") or (len(itens) + 1))
            itens.append(ItemBrutoLote(
                fonte_tipo="XML NF-e",
                arquivo=nome,
                documento=documento,
                chave=chave,
                numero_item=numero_item,
                codigo=_texto(prod, "cProd"),
                descricao=_texto(prod, "xProd"),
                ncm=_so_digitos(_texto(prod, "NCM")),
                ex_tipi=_so_digitos(_texto(prod, "EXTIPI")),
                cest_atual=_so_digitos(_texto(prod, "CEST")),
                cfop=_texto(prod, "CFOP"),
                uf_origem=uf_origem or contexto["uf_origem"],
                uf_destino=uf_destino or contexto["uf_destino"],
                data_operacao=data_operacao,
                operacao_sugerida=contexto["operacao"],
                origem_mercadoria=origem,
                cst_icms_atual=cst_icms,
                aliquota_icms_atual=_numero(_texto(icms_tipo, "pICMS")),
                valor_icms_st_atual=max(
                    _numero(_texto(icms_tipo, "vICMSST")),
                    _numero(_texto(icms_tipo, "vICMSSTRet")),
                ),
                mva_st_atual=_numero(_texto(icms_tipo, "pMVAST")),
                aliquota_fcp_st_atual=max(
                    _numero(_texto(icms_tipo, "pFCPST")),
                    _numero(_texto(icms_tipo, "pFCP")),
                ),
                cst_pis_atual=_texto(pis_tipo, "CST"),
                aliquota_pis_atual=_numero(_texto(pis_tipo, "pPIS")),
                cst_cofins_atual=_texto(cofins_tipo, "CST"),
                aliquota_cofins_atual=_numero(_texto(cofins_tipo, "pCOFINS")),
                cst_ipi_atual=_texto(ipi_tipo, "CST"),
                aliquota_ipi_atual=_numero(_texto(ipi_tipo, "pIPI")),
            ))
        if not itens:
            raise ValueError("A NF-e não possui itens de produto reconhecidos.")
        return itens

    @staticmethod
    def _grupo_tributo(
        imposto: Optional[ET.Element],
        grupo_nome: str,
        prefixo_filho: str,
    ) -> Optional[ET.Element]:
        grupo = _filho(imposto, grupo_nome)
        if grupo is None:
            return None
        for filho in list(grupo):
            if _local(filho.tag).upper().startswith(prefixo_filho.upper()):
                return filho
        return grupo

    def _ler_sped(
        self,
        conteudo: bytes,
        nome: str,
        contexto: Dict[str, Any],
    ) -> List[ItemBrutoLote]:
        texto = self._decodificar_sped(conteudo)
        linhas = texto.splitlines()
        cadastros: Dict[str, Registro0200] = {}
        participantes_uf: Dict[str, str] = {}
        empresa_uf = contexto["uf_origem"]
        data_periodo = contexto["data_operacao"]

        for linha in linhas:
            if linha.startswith("|0000|"):
                campos = linha.split("|")
                if len(campos) > 9 and str(campos[9]).strip().upper():
                    empresa_uf = str(campos[9]).strip().upper()
                if len(campos) > 4 and str(campos[4]).strip():
                    data_periodo = _data_iso(campos[4], data_periodo)
            elif linha.startswith("|0150|"):
                campos = linha.split("|")
                cod_part = campos[2].strip() if len(campos) > 2 else ""
                cod_mun = _so_digitos(campos[8] if len(campos) > 8 else "")
                uf = _IBGE_UF.get(cod_mun[:2], "")
                if cod_part and uf:
                    participantes_uf[cod_part] = uf
            elif linha.startswith("|0200|"):
                registro = Registro0200(linha)
                if registro.codigo:
                    cadastros[registro.codigo] = registro

        itens: List[ItemBrutoLote] = []
        c100: Dict[str, Any] = {}
        for numero_linha, linha in enumerate(linhas, start=1):
            if linha.startswith("|C100|"):
                campos = linha.split("|")
                ind_oper = campos[2].strip() if len(campos) > 2 else ""
                cod_part = campos[4].strip() if len(campos) > 4 else ""
                participante_uf = participantes_uf.get(cod_part, "")
                entrada = ind_oper == "0"
                origem = participante_uf if entrada else empresa_uf
                destino = empresa_uf if entrada else participante_uf
                c100 = {
                    "operacao": "ENTRADA" if entrada else "VENDA",
                    "origem": origem or contexto["uf_origem"],
                    "destino": destino or contexto["uf_destino"],
                    "documento": campos[8].strip() if len(campos) > 8 else "",
                    "chave": campos[9].strip() if len(campos) > 9 else "",
                    "data": _data_iso(campos[10] if len(campos) > 10 else "", data_periodo),
                }
                continue
            if not linha.startswith("|C170|"):
                continue
            registro = RegistroC170(linha)
            cadastro = cadastros.get(registro.codigo)
            registro.vincular_cadastro(cadastro)
            operacao = (
                c100.get("operacao")
                if contexto.get("usar_operacao_sped", True) and c100.get("operacao")
                else contexto["operacao"]
            )
            itens.append(ItemBrutoLote(
                fonte_tipo="SPED C170",
                arquivo=nome,
                documento=(f"Documento {c100.get('documento')}" if c100.get("documento") else "SPED"),
                chave=str(c100.get("chave") or ""),
                numero_item=registro.numero_item or str(numero_linha),
                codigo=registro.codigo,
                descricao=registro.descricao_0200 or registro.descricao_complementar,
                ncm=registro.ncm,
                ex_tipi=str(cadastro.ex_ipi if cadastro is not None else ""),
                cest_atual=registro.cest,
                cfop=registro.cfop,
                uf_origem=str(c100.get("origem") or contexto["uf_origem"]),
                uf_destino=str(c100.get("destino") or contexto["uf_destino"]),
                data_operacao=str(c100.get("data") or data_periodo),
                operacao_sugerida=str(operacao or contexto["operacao"]),
                cst_icms_atual=registro.cst_icms,
                aliquota_icms_atual=registro.aliquota_icms,
                valor_icms_st_atual=registro.valor_st,
                mva_st_atual=0.0,
                aliquota_fcp_st_atual=0.0,
                cst_pis_atual=registro.cst_pis,
                aliquota_pis_atual=registro.aliquota_pis,
                cst_cofins_atual=registro.cst_cofins,
                aliquota_cofins_atual=registro.aliquota_cofins,
                cst_ipi_atual=registro.cst_ipi,
                aliquota_ipi_atual=registro.aliquota_ipi,
                numero_linha_fonte=numero_linha,
            ))
        if not itens:
            raise ValueError("Nenhum registro C170 foi encontrado no SPED.")
        return itens

    @staticmethod
    def _decodificar_sped(conteudo: bytes) -> str:
        ultimo: Optional[Exception] = None
        for encoding in ("utf-8-sig", "latin-1", "cp1252"):
            try:
                return conteudo.decode(encoding)
            except UnicodeDecodeError as erro:
                ultimo = erro
        raise ValueError("Não foi possível identificar a codificação do SPED.") from ultimo

    def _analisar_item(
        self,
        item: ItemBrutoLote,
        contexto_padrao: Dict[str, Any],
    ) -> ResultadoItemLote:
        ncm = _so_digitos(item.ncm)
        if len(ncm) != 8:
            return ResultadoItemLote(
                fonte_tipo=item.fonte_tipo,
                arquivo=item.arquivo,
                documento=item.documento,
                chave=item.chave,
                numero_item=item.numero_item,
                codigo=item.codigo,
                descricao=item.descricao,
                ncm=ncm,
                ncm_oficial="",
                cfop=item.cfop,
                uf_origem=item.uf_origem,
                uf_destino=item.uf_destino,
                data_operacao=item.data_operacao,
                status="SEM NCM",
                confiabilidade=0.0,
                confirmado=False,
                exige_revisao=True,
                numero_linha_fonte=item.numero_linha_fonte,
                pendencias=["NCM ausente ou diferente de 8 dígitos."],
                cest_atual=item.cest_atual,
                cst_pis_atual=item.cst_pis_atual,
                aliquota_pis_atual=item.aliquota_pis_atual,
                cst_cofins_atual=item.cst_cofins_atual,
                aliquota_cofins_atual=item.aliquota_cofins_atual,
                cst_icms_atual=item.cst_icms_atual,
                aliquota_icms_atual=item.aliquota_icms_atual,
                valor_icms_st_atual=item.valor_icms_st_atual,
                mva_st_atual=item.mva_st_atual,
                aliquota_fcp_st_atual=item.aliquota_fcp_st_atual,
                cst_ipi_atual=item.cst_ipi_atual,
                aliquota_ipi_atual=item.aliquota_ipi_atual,
            )

        contexto = dict(contexto_padrao)
        contexto.update({
            "operacao": item.operacao_sugerida or contexto_padrao["operacao"],
            "uf_origem": item.uf_origem or contexto_padrao["uf_origem"],
            "uf_destino": item.uf_destino or contexto_padrao["uf_destino"],
            "data_operacao": _data_iso(item.data_operacao, contexto_padrao["data_operacao"]),
            "mercadoria_importada": item.origem_mercadoria in ORIGENS_IMPORTADAS,
        })

        oficial = self._buscar_ncm(ncm)
        tipi = self._buscar_tipi(ncm, item.ex_tipi)
        pis = self._analisar_pis(ncm, item.descricao, item.ex_tipi, contexto)
        icms = self._analisar_icms(ncm, item.descricao, contexto)

        divergencias: List[str] = []
        pendencias: List[str] = []
        if oficial is None:
            pendencias.append("NCM não localizado na base nacional oficial instalada.")

        pis_confirmado = bool(pis.get("confirmado"))
        cst_pis_esperado = str(pis.get("cst_pis") or pis.get("sugestao_cst_pis") or "")
        aliq_pis_esperada = self._opcional_numero(
            pis.get("aliquota_pis") if pis_confirmado else pis.get("sugestao_aliquota_pis")
        )
        cst_cofins_esperado = str(pis.get("cst_cofins") or pis.get("sugestao_cst_cofins") or "")
        aliq_cofins_esperada = self._opcional_numero(
            pis.get("aliquota_cofins") if pis_confirmado else pis.get("sugestao_aliquota_cofins")
        )
        if pis_confirmado:
            self._comparar_texto(divergencias, "CST PIS", item.cst_pis_atual, cst_pis_esperado)
            self._comparar_numero(divergencias, "Alíquota PIS", item.aliquota_pis_atual, aliq_pis_esperada)
            self._comparar_texto(divergencias, "CST Cofins", item.cst_cofins_atual, cst_cofins_esperado)
            self._comparar_numero(divergencias, "Alíquota Cofins", item.aliquota_cofins_atual, aliq_cofins_esperada)
        else:
            pendencias.append(str(pis.get("status") or "PIS/Cofins exige revisão."))

        icms_confirmado = bool(icms.get("confirmado"))
        aliq_icms_esperada = self._opcional_numero(icms.get("aliquota_operacao"))
        if bool(icms.get("aliquota_operacao_confirmada")) and aliq_icms_esperada is not None:
            self._comparar_numero(
                divergencias, "Alíquota ICMS", item.aliquota_icms_atual, aliq_icms_esperada
            )
        elif not bool(icms.get("aliquota_operacao_confirmada")):
            pendencias.append(str(icms.get("aliquota_operacao_status") or "Alíquota ICMS exige revisão."))

        cest_esperado = _so_digitos(icms.get("cest"))
        st_confirmado = bool(icms.get("st_confirmado"))
        decisao_st = bool(icms.get("st_decisao_confirmada"))
        if st_confirmado and cest_esperado:
            self._comparar_texto(divergencias, "CEST", _so_digitos(item.cest_atual), cest_esperado)
        elif decisao_st and not st_confirmado:
            if item.valor_icms_st_atual > 0.01:
                divergencias.append(
                    f"ICMS-ST informado ({item.valor_icms_st_atual:.2f}), mas a regra estadual instalada não aplica ST na data."
                )
        else:
            pendencias.append(str(icms.get("st_status") or "ICMS-ST/MVA exige revisão."))

        if icms.get("st_responsabilidade_confirmada") is False:
            pendencias.append(str(
                icms.get("st_responsabilidade_status")
                or "Responsabilidade pelo recolhimento do ICMS-ST exige revisão."
            ))

        mva_aplicada = self._opcional_numero(icms.get("mva_aplicada"))
        if st_confirmado and item.mva_st_atual > 0 and mva_aplicada is not None:
            self._comparar_numero(divergencias, "MVA-ST", item.mva_st_atual, mva_aplicada)

        fcp_esperado = self._opcional_numero(icms.get("fcp"))
        fcp_confirmado = bool(icms.get("fcp_confirmado"))
        if fcp_confirmado and fcp_esperado is not None:
            if item.aliquota_fcp_st_atual > 0 or fcp_esperado > 0:
                self._comparar_numero(
                    divergencias, "Alíquota FCP-ST",
                    item.aliquota_fcp_st_atual, fcp_esperado
                )

        if not bool(icms.get("aliquota_interna_confirmada")):
            pendencias.append(str(icms.get("aliquota_interna_status") or "Alíquota interna exige revisão."))
        if not bool(icms.get("fcp_confirmado")):
            pendencias.append(str(icms.get("fcp_status") or "FCP exige revisão."))
        if icms.get("beneficio_status") and "NÃO AUTOMATIZADO" in _normalizar_texto(icms.get("beneficio_status")):
            pendencias.append("Benefícios, reduções, isenções e diferimentos não foram automatizados para este item.")

        tipi_status = str(tipi.get("status") or "")
        if tipi.get("exige_revisao"):
            pendencias.append(tipi_status or "TIPI/EX exige revisão.")

        confiabilidade_piscofins = float(pis.get("confiabilidade") or 0.0)
        confiabilidade_icms = float(icms.get("confiabilidade_aliquota") or 0.0)
        confiabilidade_st = float(icms.get("confiabilidade_st") or 0.0)
        confiabilidade_fcp = float(icms.get("confiabilidade_fcp") or 0.0)
        aliquota_icms_confirmada = bool(icms.get("aliquota_operacao_confirmada"))

        confiancas = [
            confiabilidade_piscofins,
            confiabilidade_icms,
            confiabilidade_st,
            confiabilidade_fcp,
        ]
        confiancas = [valor for valor in confiancas if valor > 0]
        confiabilidade = min(confiancas) if confiancas else 0.0

        if divergencias:
            status = "DIVERGÊNCIA"
        elif pendencias:
            status = "REVISAR"
        else:
            status = "CONFIRMADO"
        confirmado = status == "CONFIRMADO"

        observacoes = [
            str(pis.get("observacao") or "").strip(),
            str(icms.get("observacao") or "").strip(),
            str(tipi.get("observacao") or "").strip(),
        ]
        return ResultadoItemLote(
            fonte_tipo=item.fonte_tipo,
            arquivo=item.arquivo,
            documento=item.documento,
            chave=item.chave,
            numero_item=item.numero_item,
            codigo=item.codigo,
            descricao=item.descricao,
            ncm=ncm,
            ncm_oficial=str((oficial or {}).get("descricao_completa") or (oficial or {}).get("descricao") or ""),
            cfop=item.cfop,
            uf_origem=str(contexto["uf_origem"]),
            uf_destino=str(contexto["uf_destino"]),
            data_operacao=str(contexto["data_operacao"]),
            status=status,
            confiabilidade=round(confiabilidade, 2),
            confirmado=confirmado,
            exige_revisao=not confirmado,
            numero_linha_fonte=item.numero_linha_fonte,
            divergencias=divergencias,
            pendencias=list(dict.fromkeys(texto for texto in pendencias if texto)),
            cst_pis_atual=item.cst_pis_atual,
            aliquota_pis_atual=item.aliquota_pis_atual,
            cst_pis_esperado=cst_pis_esperado,
            aliquota_pis_esperada=aliq_pis_esperada,
            cst_cofins_atual=item.cst_cofins_atual,
            aliquota_cofins_atual=item.aliquota_cofins_atual,
            cst_cofins_esperado=cst_cofins_esperado,
            aliquota_cofins_esperada=aliq_cofins_esperada,
            piscofins_status=str(pis.get("status") or ""),
            piscofins_confirmado=pis_confirmado,
            confiabilidade_piscofins=round(confiabilidade_piscofins, 2),
            cst_icms_atual=item.cst_icms_atual,
            aliquota_icms_atual=item.aliquota_icms_atual,
            aliquota_icms_esperada=aliq_icms_esperada,
            icms_status=str(icms.get("status") or ""),
            icms_confirmado=icms_confirmado,
            aliquota_icms_confirmada=aliquota_icms_confirmada,
            confiabilidade_icms=round(confiabilidade_icms, 2),
            cest_atual=_so_digitos(item.cest_atual),
            cest_esperado=cest_esperado,
            st_status=str(icms.get("st_status") or ""),
            st_confirmado=st_confirmado,
            confiabilidade_st=round(confiabilidade_st, 2),
            mva_original=self._opcional_numero(icms.get("mva_original")),
            mva_ajustada=self._opcional_numero(icms.get("mva_ajustada")),
            mva_aplicada=mva_aplicada,
            mva_tipo=str(icms.get("mva_tipo") or ""),
            mva_st_atual=item.mva_st_atual,
            st_vigencia_referencia=str(icms.get("st_vigencia_referencia") or ""),
            fcp_esperado=fcp_esperado,
            fcp_status=str(icms.get("fcp_status") or ""),
            fcp_confirmado=fcp_confirmado,
            confiabilidade_fcp=round(confiabilidade_fcp, 2),
            aliquota_fcp_st_atual=item.aliquota_fcp_st_atual,
            valor_icms_st_atual=item.valor_icms_st_atual,
            cst_ipi_atual=item.cst_ipi_atual,
            aliquota_ipi_atual=item.aliquota_ipi_atual,
            aliquota_ipi_referencia=self._opcional_numero(tipi.get("aliquota")),
            tipi_status=tipi_status,
            fundamento_piscofins=(
                f"{pis.get('fundamento') or ''} {pis.get('artigo') or ''}"
            ).strip(),
            fundamento_icms=(
                f"{icms.get('fundamento_aliquota') or ''} | "
                f"{icms.get('fundamento_st') or ''} | "
                f"{icms.get('fundamento_fcp') or ''}"
            ).strip(" |"),
            fonte_piscofins=str(pis.get("fonte_url") or ""),
            fonte_icms=(
                f"{icms.get('fonte_aliquota') or ''} | "
                f"{icms.get('fonte_st') or ''} | "
                f"{icms.get('fonte_fcp') or ''}"
            ).strip(" |"),
            observacao=" ".join(texto for texto in observacoes if texto),
        )

    def _buscar_ncm(self, ncm: str) -> Optional[Dict[str, Any]]:
        if ncm not in self._cache_ncm:
            self._cache_ncm[ncm] = BaseOficialRepository.buscar_ncm_oficial(ncm)
        return self._cache_ncm[ncm]

    def _buscar_tipi(self, ncm: str, ex_tipi: str) -> Dict[str, Any]:
        chave = (ncm, _so_digitos(ex_tipi))
        if chave in self._cache_tipi:
            return self._cache_tipi[chave]
        linhas = BaseOficialRepository.buscar_tipi(ncm)
        ex = chave[1]
        escolhida: Optional[Dict[str, Any]] = None
        if ex:
            escolhida = next((linha for linha in linhas if _so_digitos(linha.get("ex_tipi")) == ex), None)
        if escolhida is None:
            escolhida = next((linha for linha in linhas if not str(linha.get("ex_tipi") or "").strip()), None)
        if escolhida is None and len(linhas) == 1:
            escolhida = linhas[0]
        possui_ex = any(str(linha.get("ex_tipi") or "").strip() for linha in linhas)
        if escolhida is None:
            resultado = {
                "status": "TIPI NÃO LOCALIZADA",
                "aliquota": None,
                "exige_revisao": True,
                "observacao": "Não foi encontrada referência TIPI para o NCM.",
            }
        else:
            exige = bool(possui_ex and not ex and str(escolhida.get("ex_tipi") or "").strip())
            resultado = {
                "status": "TIPI CONFIRMADA" if not exige else "TIPI DEPENDE DO EX",
                "aliquota": escolhida.get("aliquota"),
                "exige_revisao": exige,
                "observacao": (
                    "Alíquota apresentada como referência TIPI; a incidência na operação concreta depende do contribuinte e da operação."
                ),
            }
        self._cache_tipi[chave] = resultado
        return resultado

    def _analisar_pis(
        self, ncm: str, descricao: str, ex_tipi: str, contexto: Dict[str, Any]
    ) -> Dict[str, Any]:
        chave = (
            ncm, _normalizar_texto(descricao), _so_digitos(ex_tipi),
            contexto["regime"], contexto["operacao"], contexto["finalidade"], contexto["data_operacao"],
        )
        if chave not in self._cache_pis:
            self._cache_pis[chave] = PISCOFINSNacionalService.analisar(
                ncm,
                contexto={
                    "regime": contexto["regime"],
                    "operacao": contexto["operacao"],
                    "finalidade": contexto["finalidade"],
                    "data_operacao": contexto["data_operacao"],
                },
                descricao=descricao,
                ex_tipi=ex_tipi,
            )
        return self._cache_pis[chave]

    def _analisar_icms(
        self, ncm: str, descricao: str, contexto: Dict[str, Any]
    ) -> Dict[str, Any]:
        chave = (
            ncm, _normalizar_texto(descricao), contexto["uf_origem"], contexto["uf_destino"],
            contexto["data_operacao"], bool(contexto.get("consumidor_final")),
            bool(contexto.get("mercadoria_importada")), bool(contexto.get("contrato_fidelidade")),
            contexto.get("regime", ""),
        )
        if chave not in self._cache_icms:
            try:
                servico = (
                    ICMSMGLoteIntegracaoService
                    if contexto["uf_destino"] == "MG"
                    else ICMSUFService
                )
                self._cache_icms[chave] = servico.analisar(
                    ncm,
                    contexto={
                        "uf_origem": contexto["uf_origem"],
                        "uf_destino": contexto["uf_destino"],
                        "data_operacao": contexto["data_operacao"],
                        "consumidor_final": contexto.get("consumidor_final", False),
                        "mercadoria_importada": contexto.get("mercadoria_importada", False),
                        "contrato_fidelidade": contexto.get("contrato_fidelidade", False),
                        "regime": contexto.get("regime", ""),
                    },
                    descricao=descricao,
                )
            except ValueError as erro:
                self._cache_icms[chave] = {
                    "status": "ICMS NÃO ANALISADO",
                    "confirmado": False,
                    "exige_revisao": True,
                    "aliquota_operacao_confirmada": False,
                    "aliquota_interna_confirmada": False,
                    "fcp_confirmado": False,
                    "st_decisao_confirmada": False,
                    "st_confirmado": False,
                    "observacao": str(erro),
                }
        return self._cache_icms[chave]

    @staticmethod
    def _comparar_texto(divergencias: List[str], campo: str, atual: Any, esperado: Any) -> None:
        esperado_n = _normalizar_texto(esperado)
        if not esperado_n:
            return
        atual_n = _normalizar_texto(atual)
        if atual_n != esperado_n:
            divergencias.append(f"{campo}: atual '{str(atual or '-').strip() or '-'}' ≠ esperado '{esperado}'.")

    @staticmethod
    def _comparar_numero(
        divergencias: List[str], campo: str, atual: Any, esperado: Optional[float]
    ) -> None:
        if esperado is None:
            return
        if not _quase_igual(atual, esperado):
            divergencias.append(
                f"{campo}: atual {_numero(atual):.2f}% ≠ esperado {float(esperado):.2f}%."
            )

    @staticmethod
    def _opcional_numero(valor: Any) -> Optional[float]:
        if valor in (None, ""):
            return None
        try:
            return float(valor)
        except (TypeError, ValueError):
            return None


class ExportadorAnaliseTributariaLoteXLSX:
    """Gera planilha de conferência com resumo, itens, divergências e fontes."""

    CABECALHOS = (
        "Fonte", "Arquivo", "Documento", "Item", "Código", "Descrição", "NCM",
        "Descrição oficial", "CFOP", "UF origem", "UF destino", "Data", "Status",
        "Segurança %", "CST PIS atual", "CST PIS esperado", "PIS atual %", "PIS esperado %",
        "CST Cofins atual", "CST Cofins esperado", "Cofins atual %", "Cofins esperado %",
        "CST ICMS atual", "ICMS atual %", "ICMS esperado %", "CEST atual", "CEST esperado",
        "MVA original %", "MVA ajustada %", "MVA XML %", "MVA aplicada %", "Tipo MVA",
        "FCP-ST XML %", "FCP/FEM esperado %", "Status FCP/FEM", "Vigência/referência ST",
        "ST atual R$", "IPI atual %", "IPI TIPI ref. %", "Divergências", "Pendências", "Observação",
    )

    @classmethod
    def exportar(cls, resultado: ResultadoAnaliseLote, caminho: str | Path) -> str:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as erro:
            raise RuntimeError("Instale o pacote openpyxl para exportar a análise em lote.") from erro

        destino = Path(caminho)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws_resumo = wb.active
        ws_resumo.title = "Resumo"
        ws_itens = wb.create_sheet("Itens")
        ws_div = wb.create_sheet("Divergências")
        ws_fontes = wb.create_sheet("Fontes")

        azul = "1F4E78"
        azul_claro = "D9EAF7"
        amarelo = "FFF2CC"
        vermelho = "F4CCCC"
        verde = "D9EAD3"
        cinza = "E7E6E6"
        borda = Border(bottom=Side(style="thin", color="B7B7B7"))

        ws_resumo["A1"] = "FiscalPro — Análise Tributária em Lote"
        ws_resumo["A1"].font = Font(size=15, bold=True, color="FFFFFF")
        ws_resumo["A1"].fill = PatternFill("solid", fgColor=azul)
        ws_resumo.merge_cells("A1:D1")
        ws_resumo["A2"] = "Apoio à conferência. Resultados condicionais devem ser revisados antes da escrituração."
        ws_resumo.merge_cells("A2:D2")
        resumo = resultado.resumo()
        linhas_resumo = (
            ("Arquivos processados", resumo["arquivos_processados"]),
            ("Itens analisados", resumo["itens"]),
            ("NCMs únicos", resumo["ncm_unicos"]),
            ("Confirmados", resumo["confirmados"]),
            ("Divergências", resumo["divergencias"]),
            ("Revisar", resumo["revisar"]),
            ("Sem NCM", resumo["sem_ncm"]),
            ("Erros de arquivo", resumo["erros_arquivos"]),
        )
        for indice, (rotulo, valor) in enumerate(linhas_resumo, start=4):
            ws_resumo.cell(indice, 1, rotulo).font = Font(bold=True)
            ws_resumo.cell(indice, 2, valor)
            ws_resumo.cell(indice, 1).fill = PatternFill("solid", fgColor=azul_claro)
        ws_resumo["A14"] = "Contexto aplicado"
        ws_resumo["A14"].font = Font(bold=True, color="FFFFFF")
        ws_resumo["A14"].fill = PatternFill("solid", fgColor=azul)
        for linha, (chave, valor) in enumerate(resultado.contexto.items(), start=15):
            ws_resumo.cell(linha, 1, str(chave).replace("_", " ").title())
            ws_resumo.cell(linha, 2, str(valor))
        if resultado.erros:
            inicio = 15 + len(resultado.contexto) + 2
            ws_resumo.cell(inicio, 1, "Erros e avisos de leitura").font = Font(bold=True, color="FFFFFF")
            ws_resumo.cell(inicio, 1).fill = PatternFill("solid", fgColor=azul)
            for deslocamento, erro in enumerate(resultado.erros, start=1):
                ws_resumo.cell(inicio + deslocamento, 1, erro)
                ws_resumo.merge_cells(start_row=inicio + deslocamento, start_column=1, end_row=inicio + deslocamento, end_column=4)

        ws_itens.append(cls.CABECALHOS)
        for celula in ws_itens[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor=azul)
            celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for item in resultado.itens:
            ws_itens.append(cls._linha_item(item))
            linha = ws_itens.max_row
            cor = {
                "CONFIRMADO": verde,
                "DIVERGÊNCIA": vermelho,
                "REVISAR": amarelo,
                "SEM NCM": cinza,
                "ERRO": vermelho,
            }.get(item.status, amarelo)
            ws_itens.cell(linha, 13).fill = PatternFill("solid", fgColor=cor)
            for celula in ws_itens[linha]:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
                celula.border = borda
        ws_itens.freeze_panes = "A2"
        ws_itens.auto_filter.ref = ws_itens.dimensions

        ws_div.append(("Arquivo", "Documento", "Item", "NCM", "Descrição", "Status", "Divergência / pendência"))
        for celula in ws_div[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor=azul)
        for item in resultado.itens:
            mensagens = item.divergencias or item.pendencias
            for mensagem in mensagens:
                ws_div.append((item.arquivo, item.documento, item.numero_item, item.ncm, item.descricao, item.status, mensagem))
        ws_div.freeze_panes = "A2"
        ws_div.auto_filter.ref = ws_div.dimensions

        ws_fontes.append(("Tema", "Fundamento", "URL"))
        for celula in ws_fontes[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor=azul)
        fontes_vistas: set[Tuple[str, str, str]] = set()
        for item in resultado.itens:
            entradas = (
                ("PIS/Cofins", item.fundamento_piscofins, item.fonte_piscofins),
                ("ICMS/ICMS-ST", item.fundamento_icms, item.fonte_icms),
            )
            for entrada in entradas:
                if entrada[1] or entrada[2]:
                    fontes_vistas.add(entrada)
        for entrada in sorted(fontes_vistas):
            ws_fontes.append(entrada)
        ws_fontes.freeze_panes = "A2"

        larguras_itens = {
            1: 14, 2: 30, 3: 22, 4: 9, 5: 14, 6: 38, 7: 12, 8: 45, 9: 9,
            10: 10, 11: 10, 12: 12, 13: 14, 14: 12, 15: 12, 16: 15, 17: 12, 18: 14,
            19: 14, 20: 17, 21: 14, 22: 16, 23: 14, 24: 12, 25: 14, 26: 14, 27: 15,
            28: 14, 29: 14, 30: 13, 31: 14, 32: 18, 33: 13, 34: 16, 35: 38,
            36: 30, 37: 13, 38: 12, 39: 15, 40: 45, 41: 45, 42: 55,
        }
        for coluna, largura in larguras_itens.items():
            ws_itens.column_dimensions[get_column_letter(coluna)].width = largura
        for ws, larguras in (
            (ws_resumo, {1: 34, 2: 26, 3: 18, 4: 18}),
            (ws_div, {1: 30, 2: 22, 3: 10, 4: 12, 5: 38, 6: 14, 7: 70}),
            (ws_fontes, {1: 18, 2: 70, 3: 65}),
        ):
            for coluna, largura in larguras.items():
                ws.column_dimensions[get_column_letter(coluna)].width = largura
            for linha in ws.iter_rows():
                for celula in linha:
                    celula.alignment = Alignment(vertical="top", wrap_text=True)

        for coluna in (17, 18, 21, 22, 24, 25, 28, 29, 30, 31, 33, 34, 38, 39):
            for linha in range(2, ws_itens.max_row + 1):
                ws_itens.cell(linha, coluna).number_format = "0.00"
        for linha in range(2, ws_itens.max_row + 1):
            ws_itens.cell(linha, 37).number_format = 'R$ #,##0.00'

        wb.save(destino)
        return str(destino.resolve())

    @staticmethod
    def _linha_item(item: ResultadoItemLote) -> Tuple[Any, ...]:
        return (
            item.fonte_tipo, item.arquivo, item.documento, item.numero_item, item.codigo,
            item.descricao, item.ncm, item.ncm_oficial, item.cfop, item.uf_origem,
            item.uf_destino, item.data_operacao, item.status, item.confiabilidade,
            item.cst_pis_atual, item.cst_pis_esperado, item.aliquota_pis_atual,
            item.aliquota_pis_esperada, item.cst_cofins_atual, item.cst_cofins_esperado,
            item.aliquota_cofins_atual, item.aliquota_cofins_esperada, item.cst_icms_atual,
            item.aliquota_icms_atual, item.aliquota_icms_esperada, item.cest_atual,
            item.cest_esperado, item.mva_original, item.mva_ajustada, item.mva_st_atual,
            item.mva_aplicada, item.mva_tipo, item.aliquota_fcp_st_atual, item.fcp_esperado,
            item.fcp_status, item.st_vigencia_referencia, item.valor_icms_st_atual,
            item.aliquota_ipi_atual, item.aliquota_ipi_referencia,
            " | ".join(item.divergencias), " | ".join(item.pendencias), item.observacao,
        )


__all__ = [
    "AnaliseTributariaLoteService",
    "ExportadorAnaliseTributariaLoteXLSX",
    "ItemBrutoLote",
    "ResultadoAnaliseLote",
    "ResultadoItemLote",
]
