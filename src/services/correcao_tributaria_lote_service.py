"""Correção tributária assistida em lote — Sprint 17.3.0.

Converte somente divergências confirmadas da análise em lote em propostas
rastreáveis para arquivos SPED diretos. XMLs e membros internos de ZIP são
mantidos exclusivamente como fonte de conferência, pois não devem ser
reescritos pelo FiscalPro.

A aplicação sempre cria uma nova cópia do SPED, valida se a linha original não
mudou desde a análise, recalcula dependências documentais de PIS/Cofins quando
possível e gera relatórios TXT e Excel. O arquivo original nunca é alterado.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
import hashlib
import re

from src.services.analise_tributaria_lote_service import (
    ResultadoAnaliseLote,
    ResultadoItemLote,
)
from src.sped.estatisticas import CalculadorEstatisticasSPED
from src.sped.indice import IndiceSPED
from src.sped.leitor import LeitorSPEDUnificado
from src.sped.pre_validador import PreValidadorPVA
from src.sped.recalculador_piscofins import RecalculadorPISCOFINS


@dataclass(slots=True)
class PropostaCorrecaoLote:
    identificador: int
    caminho_origem: str
    numero_linha: int
    registro: str
    campo: str
    indice_campo: int
    valor_atual: str
    valor_sugerido: str
    confiabilidade: float
    documento: str = ""
    numero_item: str = ""
    codigo: str = ""
    ncm: str = ""
    origem_regra: str = ""
    fundamento: str = ""
    justificativa: str = ""
    tributo_recalculo: str = ""
    selecionada: bool = True
    conflito: bool = False

    @property
    def aplicavel(self) -> bool:
        return (
            not self.conflito
            and self.numero_linha > 0
            and bool(self.registro)
            and self.indice_campo >= 0
            and self.valor_sugerido != ""
            and self.valor_atual != self.valor_sugerido
        )


@dataclass(slots=True)
class ResultadoPreparacaoCorrecaoLote:
    propostas: List[PropostaCorrecaoLote] = field(default_factory=list)
    avisos: List[str] = field(default_factory=list)
    itens_xml_somente_leitura: int = 0
    itens_sped_sem_correcao_segura: int = 0
    arquivos_sped: int = 0

    @property
    def selecionadas(self) -> List[PropostaCorrecaoLote]:
        return [p for p in self.propostas if p.selecionada and p.aplicavel]

    @property
    def total(self) -> int:
        return len(self.propostas)

    @property
    def total_selecionadas(self) -> int:
        return len(self.selecionadas)

    @property
    def conflitos(self) -> List[PropostaCorrecaoLote]:
        return [p for p in self.propostas if p.conflito]

    def obter(self, identificador: int) -> PropostaCorrecaoLote:
        for proposta in self.propostas:
            if proposta.identificador == identificador:
                return proposta
        raise KeyError(f"Proposta {identificador} não encontrada.")


@dataclass(frozen=True, slots=True)
class AlteracaoCorrecaoLote:
    caminho_origem: str
    caminho_saida: str
    numero_linha: int
    registro: str
    campo: str
    valor_anterior: str
    valor_novo: str
    documento: str
    numero_item: str
    codigo: str
    ncm: str
    origem_regra: str
    fundamento: str
    justificativa: str
    automatica: bool = False


@dataclass(slots=True)
class ArquivoCorrigidoLote:
    caminho_origem: Path
    caminho_saida: Path
    caminho_relatorio_txt: Path
    tipo_sped: str
    total_confirmadas: int
    total_alteracoes: int
    total_recalculadas: int
    documentos_retotalizados: int
    erros_pre_pva: int
    avisos_pre_pva: int
    requer_reapuracao_bloco_m: bool
    hash_original_sha256: str
    hash_saida_sha256: str


@dataclass(slots=True)
class ResultadoAplicacaoCorrecaoLote:
    arquivos: List[ArquivoCorrigidoLote] = field(default_factory=list)
    alteracoes: List[AlteracaoCorrecaoLote] = field(default_factory=list)
    avisos: List[str] = field(default_factory=list)
    caminho_relatorio_excel: Optional[Path] = None

    @property
    def total_arquivos(self) -> int:
        return len(self.arquivos)

    @property
    def total_alteracoes(self) -> int:
        return len(self.alteracoes)


class CorrecaoTributariaLoteService:
    """Prepara e aplica correções confirmadas em cópias de arquivos SPED."""

    CAMPOS_C170 = {
        "CST PIS": (24, "codigo", "PIS"),
        "Base PIS": (25, "numero", "PIS"),
        "Alíquota PIS": (26, "numero", "PIS"),
        "CST COFINS": (30, "codigo", "COFINS"),
        "Base COFINS": (31, "numero", "COFINS"),
        "Alíquota COFINS": (32, "numero", "COFINS"),
        "Alíquota ICMS": (13, "numero", ""),
    }
    CAMPO_CEST_0200 = (12, "codigo")

    def preparar(self, analise: ResultadoAnaliseLote) -> ResultadoPreparacaoCorrecaoLote:
        resultado = ResultadoPreparacaoCorrecaoLote()
        caches: Dict[str, tuple[List[str], str, Dict[str, int]]] = {}
        chaves: Dict[tuple[str, int, str, int], PropostaCorrecaoLote] = {}
        arquivos_validos: set[str] = set()

        for item in analise.itens:
            if not item.fonte_tipo.upper().startswith("SPED"):
                if item.status == "DIVERGÊNCIA":
                    resultado.itens_xml_somente_leitura += 1
                continue
            if item.status != "DIVERGÊNCIA":
                continue
            caminho = self._caminho_sped_direto(item.arquivo)
            if caminho is None:
                resultado.itens_sped_sem_correcao_segura += 1
                resultado.avisos.append(
                    f"{item.arquivo}: divergência não preparada porque o SPED está dentro de ZIP "
                    "ou não existe mais como arquivo direto. Extraia o TXT e analise novamente."
                )
                continue
            chave_arquivo = str(caminho.resolve())
            if chave_arquivo not in caches:
                linhas, encoding = LeitorSPEDUnificado().ler(caminho)
                caches[chave_arquivo] = (linhas, encoding, self._mapear_0200(linhas))
            linhas, _encoding, mapa_0200 = caches[chave_arquivo]
            arquivos_validos.add(chave_arquivo)
            antes = len(chaves)
            self._propostas_item(item, caminho, linhas, mapa_0200, chaves, resultado.avisos)
            if len(chaves) == antes:
                resultado.itens_sped_sem_correcao_segura += 1

        propostas = sorted(
            chaves.values(),
            key=lambda p: (Path(p.caminho_origem).name.lower(), p.numero_linha, p.indice_campo),
        )
        for identificador, proposta in enumerate(propostas, start=1):
            proposta.identificador = identificador
            proposta.selecionada = proposta.aplicavel
        resultado.propostas = propostas
        resultado.arquivos_sped = len(arquivos_validos)
        if resultado.itens_xml_somente_leitura:
            resultado.avisos.append(
                f"{resultado.itens_xml_somente_leitura} item(ns) divergente(s) de XML foram mantidos "
                "somente para conferência; XML fiscal assinado não é reescrito."
            )
        if not propostas:
            resultado.avisos.append(
                "Nenhuma divergência com campo, valor esperado e confirmação suficientes foi encontrada em SPED direto."
            )
        return resultado

    def atualizar_selecao(
        self,
        preparacao: ResultadoPreparacaoCorrecaoLote,
        identificador: int,
        selecionada: bool,
    ) -> PropostaCorrecaoLote:
        proposta = preparacao.obter(identificador)
        if selecionada and not proposta.aplicavel:
            raise RuntimeError("Esta proposta possui conflito ou não é aplicável com segurança.")
        proposta.selecionada = bool(selecionada)
        return proposta

    def aplicar(
        self,
        preparacao: ResultadoPreparacaoCorrecaoLote,
        pasta_saida: str | Path,
    ) -> ResultadoAplicacaoCorrecaoLote:
        selecionadas = preparacao.selecionadas
        if not selecionadas:
            raise RuntimeError("Nenhuma correção foi marcada para aplicar.")
        destino = Path(pasta_saida)
        destino.mkdir(parents=True, exist_ok=True)
        if not destino.is_dir():
            raise RuntimeError(f"A pasta de saída não é válida: {destino}")

        agrupadas: Dict[str, List[PropostaCorrecaoLote]] = {}
        for proposta in selecionadas:
            agrupadas.setdefault(proposta.caminho_origem, []).append(proposta)

        resultado = ResultadoAplicacaoCorrecaoLote(avisos=list(preparacao.avisos))
        for caminho_origem, propostas in sorted(agrupadas.items()):
            arquivo_resultado, alteracoes, avisos = self._aplicar_arquivo(
                Path(caminho_origem), propostas, destino
            )
            resultado.arquivos.append(arquivo_resultado)
            resultado.alteracoes.extend(alteracoes)
            resultado.avisos.extend(avisos)

        resultado.caminho_relatorio_excel = self._exportar_excel(resultado, destino)
        return resultado

    def _propostas_item(
        self,
        item: ResultadoItemLote,
        caminho: Path,
        linhas: List[str],
        mapa_0200: Dict[str, int],
        chaves: Dict[tuple[str, int, str, int], PropostaCorrecaoLote],
        avisos: List[str],
    ) -> None:
        linha_c170 = int(item.numero_linha_fonte or 0)
        if linha_c170 <= 0 or linha_c170 > len(linhas):
            avisos.append(
                f"{caminho.name} — item {item.numero_item}: linha C170 não localizada com segurança."
            )
            return
        if self._codigo(linhas[linha_c170 - 1]) != "C170":
            avisos.append(
                f"{caminho.name} — linha {linha_c170}: o registro não é mais C170; proposta ignorada."
            )
            return

        if item.piscofins_confirmado:
            self._adicionar_c170(
                chaves, caminho, linhas, item, linha_c170, "CST PIS",
                item.cst_pis_esperado, item.confiabilidade_piscofins,
                "Motor Nacional de PIS/COFINS", item.fundamento_piscofins,
                item.piscofins_status,
            )
            if (
                str(item.cst_pis_esperado or "").zfill(2) in {"04", "05", "06", "07", "08", "09"}
                and item.aliquota_pis_esperada is not None
                and abs(float(item.aliquota_pis_esperada)) < 0.0001
            ):
                self._adicionar_c170(
                    chaves, caminho, linhas, item, linha_c170, "Base PIS", 0.0,
                    item.confiabilidade_piscofins, "Motor Nacional de PIS/COFINS",
                    item.fundamento_piscofins,
                    "Base zerada por CST confirmado sem incidência na operação analisada.",
                )
            self._adicionar_c170(
                chaves, caminho, linhas, item, linha_c170, "Alíquota PIS",
                item.aliquota_pis_esperada, item.confiabilidade_piscofins,
                "Motor Nacional de PIS/COFINS", item.fundamento_piscofins,
                item.piscofins_status,
            )
            self._adicionar_c170(
                chaves, caminho, linhas, item, linha_c170, "CST COFINS",
                item.cst_cofins_esperado, item.confiabilidade_piscofins,
                "Motor Nacional de PIS/COFINS", item.fundamento_piscofins,
                item.piscofins_status,
            )
            if (
                str(item.cst_cofins_esperado or "").zfill(2) in {"04", "05", "06", "07", "08", "09"}
                and item.aliquota_cofins_esperada is not None
                and abs(float(item.aliquota_cofins_esperada)) < 0.0001
            ):
                self._adicionar_c170(
                    chaves, caminho, linhas, item, linha_c170, "Base COFINS", 0.0,
                    item.confiabilidade_piscofins, "Motor Nacional de PIS/COFINS",
                    item.fundamento_piscofins,
                    "Base zerada por CST confirmado sem incidência na operação analisada.",
                )
            self._adicionar_c170(
                chaves, caminho, linhas, item, linha_c170, "Alíquota COFINS",
                item.aliquota_cofins_esperada, item.confiabilidade_piscofins,
                "Motor Nacional de PIS/COFINS", item.fundamento_piscofins,
                item.piscofins_status,
            )

        if item.aliquota_icms_confirmada:
            self._adicionar_c170(
                chaves, caminho, linhas, item, linha_c170, "Alíquota ICMS",
                item.aliquota_icms_esperada, item.confiabilidade_icms,
                "Motor ICMS por UF", item.fundamento_icms, item.icms_status,
            )

        if item.st_confirmado and item.cest_esperado:
            linha_0200 = mapa_0200.get(str(item.codigo or "").strip())
            if linha_0200 is None:
                avisos.append(
                    f"{caminho.name} — item {item.numero_item}: CEST confirmado, mas o 0200 do código "
                    f"'{item.codigo}' não foi localizado."
                )
            else:
                indice, tipo = self.CAMPO_CEST_0200
                atual = self._obter_campo(linhas[linha_0200 - 1], indice)
                sugerido = self._normalizar(item.cest_esperado, tipo)
                if self._diferente(atual, sugerido, tipo):
                    proposta = PropostaCorrecaoLote(
                        identificador=0,
                        caminho_origem=str(caminho.resolve()),
                        numero_linha=linha_0200,
                        registro="0200",
                        campo="CEST",
                        indice_campo=indice,
                        valor_atual=atual,
                        valor_sugerido=sugerido,
                        confiabilidade=float(item.confiabilidade_st or 100.0),
                        documento=item.documento,
                        numero_item=item.numero_item,
                        codigo=item.codigo,
                        ncm=item.ncm,
                        origem_regra="Motor ICMS-ST por UF",
                        fundamento=item.fundamento_icms,
                        justificativa=item.st_status,
                    )
                    self._consolidar(chaves, proposta, avisos)

    def _adicionar_c170(
        self,
        chaves: Dict[tuple[str, int, str, int], PropostaCorrecaoLote],
        caminho: Path,
        linhas: List[str],
        item: ResultadoItemLote,
        numero_linha: int,
        campo: str,
        esperado: Any,
        confiabilidade: float,
        origem_regra: str,
        fundamento: str,
        justificativa: str,
    ) -> None:
        if esperado is None or esperado == "":
            return
        indice, tipo, tributo = self.CAMPOS_C170[campo]
        atual = self._obter_campo(linhas[numero_linha - 1], indice)
        sugerido = self._normalizar(esperado, tipo)
        if not self._diferente(atual, sugerido, tipo):
            return
        proposta = PropostaCorrecaoLote(
            identificador=0,
            caminho_origem=str(caminho.resolve()),
            numero_linha=numero_linha,
            registro="C170",
            campo=campo,
            indice_campo=indice,
            valor_atual=atual,
            valor_sugerido=sugerido,
            confiabilidade=float(confiabilidade or 100.0),
            documento=item.documento,
            numero_item=item.numero_item,
            codigo=item.codigo,
            ncm=item.ncm,
            origem_regra=origem_regra,
            fundamento=fundamento,
            justificativa=justificativa,
            tributo_recalculo=tributo,
        )
        self._consolidar(chaves, proposta, [])

    @staticmethod
    def _consolidar(
        chaves: Dict[tuple[str, int, str, int], PropostaCorrecaoLote],
        proposta: PropostaCorrecaoLote,
        avisos: List[str],
    ) -> None:
        chave = (
            proposta.caminho_origem.lower(), proposta.numero_linha,
            proposta.registro, proposta.indice_campo,
        )
        existente = chaves.get(chave)
        if existente is None:
            chaves[chave] = proposta
            return
        if existente.valor_sugerido != proposta.valor_sugerido:
            existente.conflito = True
            existente.selecionada = False
            existente.justificativa = (
                f"Conflito entre valores confirmados: {existente.valor_sugerido} e {proposta.valor_sugerido}."
            )
            avisos.append(
                f"{Path(proposta.caminho_origem).name} — linha {proposta.numero_linha} / "
                f"{proposta.campo}: valores confirmados conflitantes; correção não será aplicada."
            )

    def _aplicar_arquivo(
        self,
        caminho: Path,
        propostas: List[PropostaCorrecaoLote],
        pasta_saida: Path,
    ) -> tuple[ArquivoCorrigidoLote, List[AlteracaoCorrecaoLote], List[str]]:
        linhas, encoding = LeitorSPEDUnificado().ler(caminho)
        novas = list(linhas)
        alteracoes: List[AlteracaoCorrecaoLote] = []
        avisos: List[str] = []
        afetadas_piscofins: set[tuple[int, str]] = set()

        for proposta in sorted(propostas, key=lambda p: (p.numero_linha, p.indice_campo)):
            indice_linha = proposta.numero_linha - 1
            if indice_linha < 0 or indice_linha >= len(novas):
                raise RuntimeError(
                    f"{caminho.name}: a linha {proposta.numero_linha} não existe mais. Analise o arquivo novamente."
                )
            if self._codigo(novas[indice_linha]) != proposta.registro:
                raise RuntimeError(
                    f"{caminho.name}: a linha {proposta.numero_linha} não corresponde mais ao registro "
                    f"{proposta.registro}. Analise o arquivo novamente."
                )
            atual_agora = self._obter_campo(novas[indice_linha], proposta.indice_campo)
            tipo = "numero" if "Alíquota" in proposta.campo else "codigo"
            if self._diferente(atual_agora, proposta.valor_atual, tipo):
                raise RuntimeError(
                    f"{caminho.name}: o campo {proposta.campo} da linha {proposta.numero_linha} mudou "
                    "desde a análise. Nenhuma cópia foi gerada para esse arquivo."
                )
            novas[indice_linha] = self._substituir_campo(
                novas[indice_linha], proposta.indice_campo, proposta.valor_sugerido
            )
            alteracoes.append(AlteracaoCorrecaoLote(
                caminho_origem=str(caminho),
                caminho_saida="",
                numero_linha=proposta.numero_linha,
                registro=proposta.registro,
                campo=proposta.campo,
                valor_anterior=atual_agora,
                valor_novo=proposta.valor_sugerido,
                documento=proposta.documento,
                numero_item=proposta.numero_item,
                codigo=proposta.codigo,
                ncm=proposta.ncm,
                origem_regra=proposta.origem_regra,
                fundamento=proposta.fundamento,
                justificativa=proposta.justificativa,
                automatica=False,
            ))
            if proposta.tributo_recalculo:
                afetadas_piscofins.add((proposta.numero_linha, proposta.tributo_recalculo))

        indice = IndiceSPED().construir(novas)
        estatisticas = CalculadorEstatisticasSPED().calcular(novas, indice)
        recalculo = RecalculadorPISCOFINS().recalcular(
            novas, afetadas_piscofins, estatisticas.tipo_sped
        )
        novas = recalculo.linhas
        avisos.extend(recalculo.avisos)

        caminho_saida = self._caminho_saida_unico(caminho, pasta_saida)
        for item in recalculo.alteracoes:
            alteracoes.append(AlteracaoCorrecaoLote(
                caminho_origem=str(caminho),
                caminho_saida=str(caminho_saida),
                numero_linha=item.numero_linha,
                registro=item.registro,
                campo=item.campo,
                valor_anterior=item.valor_anterior,
                valor_novo=item.valor_novo,
                documento="",
                numero_item="",
                codigo="",
                ncm="",
                origem_regra="Recálculo automático PIS/COFINS",
                fundamento="Dependência documental do C170/C100",
                justificativa=item.justificativa,
                automatica=True,
            ))

        pre_validacao = PreValidadorPVA().validar(novas, estatisticas.tipo_sped)
        self._gravar_linhas(caminho_saida, novas, encoding, caminho)
        alteracoes = [
            AlteracaoCorrecaoLote(**{
                **a.__dict__, "caminho_saida": str(caminho_saida)
            }) if hasattr(a, "__dict__") else AlteracaoCorrecaoLote(
                caminho_origem=a.caminho_origem,
                caminho_saida=str(caminho_saida),
                numero_linha=a.numero_linha,
                registro=a.registro,
                campo=a.campo,
                valor_anterior=a.valor_anterior,
                valor_novo=a.valor_novo,
                documento=a.documento,
                numero_item=a.numero_item,
                codigo=a.codigo,
                ncm=a.ncm,
                origem_regra=a.origem_regra,
                fundamento=a.fundamento,
                justificativa=a.justificativa,
                automatica=a.automatica,
            )
            for a in alteracoes
        ]
        caminho_relatorio = caminho_saida.with_name(
            f"{caminho_saida.stem}_RELATORIO.txt"
        )
        self._gravar_relatorio_txt(
            caminho_relatorio, caminho, caminho_saida, estatisticas.tipo_sped,
            propostas, alteracoes, recalculo.avisos, pre_validacao,
            recalculo.documentos_retotalizados, recalculo.requer_reapuracao_bloco_m,
        )
        arquivo_resultado = ArquivoCorrigidoLote(
            caminho_origem=caminho,
            caminho_saida=caminho_saida,
            caminho_relatorio_txt=caminho_relatorio,
            tipo_sped=estatisticas.tipo_sped,
            total_confirmadas=len(propostas),
            total_alteracoes=len(alteracoes),
            total_recalculadas=len(recalculo.alteracoes),
            documentos_retotalizados=recalculo.documentos_retotalizados,
            erros_pre_pva=len(pre_validacao.erros),
            avisos_pre_pva=len(pre_validacao.avisos),
            requer_reapuracao_bloco_m=recalculo.requer_reapuracao_bloco_m,
            hash_original_sha256=self._sha256(caminho),
            hash_saida_sha256=self._sha256(caminho_saida),
        )
        return arquivo_resultado, alteracoes, avisos

    @staticmethod
    def _caminho_sped_direto(valor: str) -> Optional[Path]:
        texto = str(valor or "").strip()
        if not texto or "::" in texto:
            return None
        caminho = Path(texto)
        if not caminho.is_file() or caminho.suffix.lower() != ".txt":
            return None
        return caminho

    @staticmethod
    def _mapear_0200(linhas: Iterable[str]) -> Dict[str, int]:
        mapa: Dict[str, int] = {}
        for numero, linha in enumerate(linhas, start=1):
            if CorrecaoTributariaLoteService._codigo(linha) != "0200":
                continue
            codigo = CorrecaoTributariaLoteService._obter_campo(linha, 1).strip()
            if codigo and codigo not in mapa:
                mapa[codigo] = numero
        return mapa

    @staticmethod
    def _codigo(linha: str) -> str:
        texto = linha.rstrip("\r\n")
        if not (texto.startswith("|") and texto.endswith("|")):
            return ""
        partes = texto.split("|")
        return partes[1].strip().upper() if len(partes) > 2 else ""

    @staticmethod
    def _obter_campo(linha: str, indice_campo: int) -> str:
        partes = linha.rstrip("\r\n").split("|")
        indice = indice_campo + 1
        return partes[indice] if indice < len(partes) else ""

    @staticmethod
    def _substituir_campo(linha: str, indice_campo: int, valor: str) -> str:
        if linha.endswith("\r\n"):
            texto, fim = linha[:-2], "\r\n"
        elif linha.endswith("\n"):
            texto, fim = linha[:-1], "\n"
        elif linha.endswith("\r"):
            texto, fim = linha[:-1], "\r"
        else:
            texto, fim = linha, ""
        partes = texto.split("|")
        indice = indice_campo + 1
        if not (texto.startswith("|") and texto.endswith("|")) or indice >= len(partes) - 1:
            raise RuntimeError(f"Campo {indice_campo} não localizado na linha SPED.")
        partes[indice] = valor
        return "|".join(partes) + fim

    @classmethod
    def _normalizar(cls, valor: Any, tipo: str) -> str:
        if tipo == "numero":
            numero = cls._decimal(valor)
            if numero is None:
                return ""
            numero = numero.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            return f"{numero:.2f}".replace(".", ",")
        return re.sub(r"\D", "", str(valor or "").strip()) if str(valor or "").strip() else ""

    @classmethod
    def _diferente(cls, atual: Any, esperado: Any, tipo: str) -> bool:
        if tipo == "numero":
            a = cls._decimal(atual)
            b = cls._decimal(esperado)
            if a is None and b is None:
                return False
            if a is None or b is None:
                return True
            return abs(a - b) > Decimal("0.005")
        return cls._normalizar(atual, tipo) != cls._normalizar(esperado, tipo)

    @staticmethod
    def _decimal(valor: Any) -> Optional[Decimal]:
        texto = str(valor if valor not in (None, "") else "").strip().replace("%", "")
        if not texto:
            return None
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        try:
            return Decimal(texto)
        except (InvalidOperation, ValueError):
            return None

    @staticmethod
    def _caminho_saida_unico(origem: Path, pasta: Path) -> Path:
        base = pasta / f"{origem.stem}_CORRIGIDO_FISCALPRO_17_3{origem.suffix}"
        if not base.exists():
            return base
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return pasta / f"{origem.stem}_CORRIGIDO_FISCALPRO_17_3_{timestamp}{origem.suffix}"

    @staticmethod
    def _gravar_linhas(caminho: Path, linhas: List[str], encoding: str, origem: Path | None = None) -> None:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        encoding_saida = encoding
        if encoding.lower().replace("_", "-") == "utf-8-sig" and origem is not None:
            try:
                possui_bom = origem.read_bytes().startswith(b"\xef\xbb\xbf")
            except OSError:
                possui_bom = False
            if not possui_bom:
                encoding_saida = "utf-8"
        with caminho.open("w", encoding=encoding_saida, newline="") as stream:
            stream.writelines(linhas)

    @staticmethod
    def _sha256(caminho: Path) -> str:
        digest = hashlib.sha256()
        with caminho.open("rb") as stream:
            for bloco in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(bloco)
        return digest.hexdigest()

    @staticmethod
    def _gravar_relatorio_txt(
        caminho: Path,
        origem: Path,
        saida: Path,
        tipo_sped: str,
        propostas: List[PropostaCorrecaoLote],
        alteracoes: List[AlteracaoCorrecaoLote],
        avisos_recalculo: List[str],
        pre_validacao: Any,
        documentos_retotalizados: int,
        requer_bloco_m: bool,
    ) -> None:
        linhas = [
            "FISCALPRO — RELATÓRIO DE CORREÇÃO TRIBUTÁRIA ASSISTIDA EM LOTE",
            "=" * 78,
            f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
            f"Arquivo original: {origem}",
            f"Cópia corrigida: {saida}",
            f"Tipo de SPED: {tipo_sped}",
            f"Correções confirmadas pela usuária: {len(propostas)}",
            f"Alterações totais, incluindo recálculos: {len(alteracoes)}",
            f"Documentos C100 retotalizados: {documentos_retotalizados}",
            f"Pré-Validador: {len(pre_validacao.erros)} erro(s) e {len(pre_validacao.avisos)} aviso(s)",
            "",
            "ALTERAÇÕES",
            "-" * 78,
        ]
        for item in alteracoes:
            linhas.extend([
                f"Linha {item.numero_linha} | {item.registro} | {item.campo}",
                f"  Antes: {item.valor_anterior or '-'}",
                f"  Depois: {item.valor_novo or '-'}",
                f"  Documento/item: {item.documento or '-'} / {item.numero_item or '-'}",
                f"  Produto: {item.codigo or '-'} | NCM {item.ncm or '-'}",
                f"  Origem: {item.origem_regra}",
                f"  Fundamento: {item.fundamento or '-'}",
                f"  Justificativa: {item.justificativa or '-'}",
                "",
            ])
        linhas.extend(["AVISOS", "-" * 78])
        if avisos_recalculo:
            linhas.extend(f"- {aviso}" for aviso in avisos_recalculo)
        else:
            linhas.append("- Nenhum aviso de recálculo.")
        if requer_bloco_m:
            linhas.append(
                "- O Bloco M não foi alterado automaticamente. Recalcule a apuração e valide a cópia no PVA oficial."
            )
        if pre_validacao.erros:
            linhas.append("- Erros do Pré-Validador:")
            linhas.extend(f"  • {item.mensagem}" for item in pre_validacao.erros[:100])
        if pre_validacao.avisos:
            linhas.append("- Avisos do Pré-Validador:")
            linhas.extend(f"  • {item.mensagem}" for item in pre_validacao.avisos[:100])
        linhas.extend([
            "",
            "SEGURANÇA",
            "-" * 78,
            "O arquivo original não foi alterado.",
            "Somente propostas marcadas e confirmadas foram aplicadas na cópia.",
            "XMLs permanecem somente para leitura e conferência.",
        ])
        caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")

    @staticmethod
    def _exportar_excel(resultado: ResultadoAplicacaoCorrecaoLote, pasta: Path) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
            from openpyxl.utils import get_column_letter
        except ImportError as erro:
            raise RuntimeError("A biblioteca openpyxl é necessária para gerar o relatório Excel.") from erro

        nome = f"Relatorio_Correcao_Tributaria_Lote_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        caminho = pasta / nome
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.append(["FiscalPro — Correção Tributária Assistida em Lote"])
        ws.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M:%S")])
        ws.append(["Arquivos corrigidos", resultado.total_arquivos])
        ws.append(["Alterações totais", resultado.total_alteracoes])
        ws.append([])
        ws.append([
            "Arquivo original", "Cópia corrigida", "Tipo SPED", "Confirmadas",
            "Alterações", "Recálculos", "C100 retotalizados", "Erros PVA",
            "Avisos PVA", "Reapurar Bloco M", "SHA-256 original", "SHA-256 saída",
        ])
        for arq in resultado.arquivos:
            ws.append([
                str(arq.caminho_origem), str(arq.caminho_saida), arq.tipo_sped,
                arq.total_confirmadas, arq.total_alteracoes, arq.total_recalculadas,
                arq.documentos_retotalizados, arq.erros_pre_pva, arq.avisos_pre_pva,
                "SIM" if arq.requer_reapuracao_bloco_m else "NÃO",
                arq.hash_original_sha256, arq.hash_saida_sha256,
            ])

        wa = wb.create_sheet("Alterações")
        wa.append([
            "Arquivo original", "Arquivo corrigido", "Linha", "Registro", "Campo",
            "Antes", "Depois", "Documento", "Item", "Código", "NCM",
            "Origem da regra", "Fundamento", "Justificativa", "Automática",
        ])
        for item in resultado.alteracoes:
            wa.append([
                item.caminho_origem, item.caminho_saida, item.numero_linha,
                item.registro, item.campo, item.valor_anterior, item.valor_novo,
                item.documento, item.numero_item, item.codigo, item.ncm,
                item.origem_regra, item.fundamento, item.justificativa,
                "SIM" if item.automatica else "NÃO",
            ])

        ww = wb.create_sheet("Avisos")
        ww.append(["Aviso"])
        for aviso in dict.fromkeys(resultado.avisos):
            ww.append([aviso])

        for planilha in wb.worksheets:
            planilha.freeze_panes = "A2"
            planilha.auto_filter.ref = planilha.dimensions
            for celula in planilha[1]:
                celula.font = Font(bold=True)
                celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for coluna in range(1, planilha.max_column + 1):
                largura = 12
                for linha in range(1, min(planilha.max_row, 200) + 1):
                    valor = planilha.cell(linha, coluna).value
                    largura = max(largura, min(60, len(str(valor or "")) + 2))
                planilha.column_dimensions[get_column_letter(coluna)].width = largura
        wb.save(caminho)
        return caminho


__all__ = [
    "PropostaCorrecaoLote",
    "ResultadoPreparacaoCorrecaoLote",
    "AlteracaoCorrecaoLote",
    "ArquivoCorrigidoLote",
    "ResultadoAplicacaoCorrecaoLote",
    "CorrecaoTributariaLoteService",
]
